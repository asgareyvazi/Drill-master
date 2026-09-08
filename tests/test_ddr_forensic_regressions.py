"""2026-09 forensic regressions. Domain tests are not desktop certification."""
import copy
import json
import math
from datetime import date
from pathlib import Path

import pytest

from core.database import (DatabaseManager, Company, Project, Well, Section,
                           DailyReport, BHAReport, SurveyPoint, MudReport, BulkMaterials,
                           BOPComponent, SafetyReport, ServiceCompanyPOB, AuditLog)
from core.domain_records import (bha_record, BHA_FIELDS, chemical_type, CHEMICAL_TYPES,
                                 material_route, optional_date, npt_record, collection_value)
from core.combo_identity import ComboOption, resolve_options
from core.ddr_import_service import DDRImportService
from core.survey_records import plot_series


@pytest.fixture
def db(tmp_path):
    manager = DatabaseManager()
    manager.db_path = str(tmp_path / "forensic.db")
    assert manager.initialize()
    with manager.session_scope() as session:
        # initialize() creates a real default company/project/well. Reuse
        # those identities so report and child contexts genuinely agree.
        company = session.get(Company, 1)
        company.name, company.code = "Operator", "OP"
        project = session.get(Project, 1)
        project.name, project.code = "Audit", "AUDIT"
        well = session.get(Well, 1)
        well.name, well.code = "Audit Well", "AW"
        section = Section(name="Audit Section", well_id=well.id)
        session.add(section)
        session.flush()
        report = DailyReport(well_id=well.id, section_id=section.id, report_number=1, report_date=date(2024, 10, 22))
        session.add(report)
    yield manager
    manager.close()


@pytest.mark.parametrize("value,expected", [("ID", "ID"), ("Long Label", "ID"), (" long-label ", "ID"), ("alias", "ID"), (42, "ID"), (0, None), ("invalid", None), ("", None), (None, None)])
def test_catalog_identity_never_qt_index(value, expected):
    result = resolve_options(value, [ComboOption("ID", "Long Label", "42", ("alias",))], field="test")
    assert result.identity == expected
    assert result.accepted == (expected is not None)


def test_catalog_ambiguous():
    result = resolve_options("alias", [ComboOption("a", "A", "1", ("alias",)), ComboOption("b", "B", "2", ("alias",))], field="test")
    assert result.method == "ambiguous" and result.identity is None


@pytest.mark.parametrize("status", ["NPT", "npt", "Non-Productive Time"])
def test_npt_activation_and_company(status):
    row, reviews = npt_record({"status": status, "npt_code": "F-PUMP", "contractor": "operator"}, ["Operator"])
    assert row["is_npt"] is True and row["main_code"] == "F-PUMP"
    assert row["contractor"] == "Operator" and not reviews


def test_npt_unknown_company_is_not_first_item():
    row, reviews = npt_record({"status": "NPT", "npt_code": "F-PUMP", "contractor": "Unknown Co"}, ["Operator"])
    assert row["is_npt"] and row["contractor"] is None
    assert reviews[0]["original_value"] == "Unknown Co"
    assert reviews[0]["status"] == "REVIEW_REQUIRED"


def test_npt_source_status_persists_and_creates_normal_npt_domain(db):
    service = DDRImportService(db, 1)
    result = service._save_time_logs(1, [{"time_from": "00:00", "time_to": "01:00", "duration": 1,
                                         "status": "NPT", "npt_code": "F-PUMP", "contractor": "Operator"}])
    assert result["valid"] == 1
    assert db.auto_update_from_daily_report(1)
    assert db.auto_update_from_daily_report(1)
    rows = db.get_npt_reports(report_id=1)
    assert len(rows) == 1 and rows[0]["responsible_party"] == "Operator"


@pytest.mark.parametrize("name,expected", [("Bentonite", "Viscosifier"), ("barite", "Weight Material"),
    ("Caustic Soda", "Alkalinity"), ("API Starch", "Filtration Control"), ("Bit Lube", "Lubricant"),
    ("KCl", "Shale Inhibitor"), ("Salt", None), ("Anti-Foam", None), ("Mystery blend", None), ("", None)])
def test_chemical_existing_taxonomy(name, expected):
    result = chemical_type(name)
    assert result.identity == expected
    assert result.identity is None or result.identity in CHEMICAL_TYPES


@pytest.mark.parametrize("name,route", [("Bentonite", "mud"), ("Barite", "mud"), ("Mud", "mud"),
                                      ("Diesel", "fuel_water"), ("Fresh Water", "fuel_water"), ("Unknown inventory", "review")])
def test_material_routing(name, route):
    assert material_route({"material_name": name}) == route


def test_mud_stock_type_and_inventory_destination(db):
    data = {"bulk_materials": [{"material_name": "Barite", "on_hand": 45, "received": 24, "unit": "1.5 MT/BB"},
                               {"material_name": "Fresh Water", "current_stock": 100},
                               {"material_name": "Diesel", "current_stock": 50}]}
    result = db.save_imported_multi_tab_data_atomic(1, 1, data)
    assert result["failed"] == 0
    with db.session_scope() as session:
        assert session.query(BulkMaterials).count() == 2
        chemicals = json.loads(session.query(MudReport).one().chemicals_json)
    assert chemicals[0]["type"] == "Weight Material"
    assert chemicals[0]["stock"] == 45 and chemicals[0]["used"] is None


def test_bha_named_projection_manual_and_import_roundtrip(db):
    source = {"component_name": '9-1/2" DC', "od": 9.5, "length": 18.04, "cum_length": 19.68,
              "_source_cells": {"bha.component_name": "R50C2"}, "_source_file": "C:/reports/DDR.xlsx"}
    row = bha_record(source)
    assert row["Component Name"] == source["component_name"] and row["Tool Type"] == "Drill Collar"
    assert row["Weight (kg)"] is None and row["Connection Type"] is None
    assert row["Cumulative Length (m)"] == 19.68
    db.save_imported_multi_tab_data_atomic(1, 1, {"bha_components": [source]})
    loaded = db.get_bha_report(1, report_id=1)["bha_configs"]
    assert db.save_bha_report(1, {"report_id": 1, "bha_data_json": loaded})
    again = db.get_bha_report(1, report_id=1)["bha_configs"]
    assert again == loaded
    assert again[0]["_provenance"]["source_record"]["_source_cells"] == source["_source_cells"]


@pytest.mark.parametrize("value", [{"cell": "R48C2"}, ["metadata"], "C:/reports/DDR.xlsx", "/home/reports/DDR.xlsx"])
def test_bha_metadata_never_in_any_display_field(value):
    row = bha_record({field: value for field in BHA_FIELDS})
    assert all(row[field] is None for field in BHA_FIELDS)
    assert row["_provenance"]


@pytest.mark.parametrize("name,expected", [('17-1/2" PDC Bit', "Bit"), ('9-1/2" Bit Sub(F.Valve, Baffel in)', "Sub (Bit Sub)"),
                                         ("S.Stab", "Stabilizer (String)"), ("XOS", "X-Over Sub"), ('5" HWDP', "HWDP")])
def test_bha_description_is_distinct_from_type(name, expected):
    row = bha_record({"component_name": name})
    assert row["Component Name"] == name and row["Tool Type"] == expected


@pytest.mark.parametrize("value,expected", [(None, None), ("", None), ("None", None),
    (date(2024, 10, 22), date(2024, 10, 22)), ("22-Oct-2024", date(2024, 10, 22)),
    ("2024/10/22", date(2024, 10, 22)), ("22/10/2024", date(2024, 10, 22))])
def test_optional_dates(value, expected):
    assert optional_date(value) == expected


def test_invalid_date():
    with pytest.raises(ValueError, match="Invalid date"):
        optional_date("not a date")


def test_pob_partial_failure_and_reload(db):
    result = db.save_imported_multi_tab_data_atomic(1, 1, {"pob_records": [
        {"company_name": "Valid", "date_in": "22-Oct-2024", "personnel_count": 3},
        {"company_name": "Missing date", "date_in": None, "personnel_count": 2},
        {"company_name": "Bad date", "date_in": "invalid", "personnel_count": 1}]})
    assert result["failed"] == 0 and result["review"] == 1
    rows = db.get_service_company_pob(1, report_id=1)
    assert len(rows) == 2
    assert next(r for r in rows if r["company_name"] == "Missing date")["date_in"] is None
    assert db.save_service_company_pob({"well_id": 1, "report_id": 1, "company_name": "Manual", "date_in": "22-Oct-2024"})


@pytest.mark.parametrize("kind", ["bop", "waste"])
def test_empty_and_populated_safety_collection_contract(db, kind):
    getter = db.get_bop_components if kind == "bop" else db.get_waste_records
    assert getter(report_id=1) == []
    with db.session_scope() as session:
        session.add(SafetyReport(well_id=1, report_id=1, report_date=date(2024, 10, 22)))
    report = db.get_safety_report(1, report_id=1)
    assert report["bop_stack_json"] == report["waste_history_json"] == []
    if kind == "bop":
        payload = {"bop_components": [{"name": "Annular", "type": "Annular", "working_pressure": 3000}]}
    else:
        payload = {"waste_records": [{"waste_type": "Cuttings", "volume": 12, "record_date": "2024-10-22"}]}
    result = db.save_imported_multi_tab_data_atomic(1, 1, payload)
    assert result["failed"] == 0
    assert len(getter(report_id=1)) == 1
    report = db.get_safety_report(1, report_id=1)
    assert len(report["bop_stack_json" if kind == "bop" else "waste_history_json"]) == 1


def test_collection_failure_is_not_empty(db, monkeypatch):
    def fail():
        raise RuntimeError("database unavailable")
    monkeypatch.setattr(db, "create_session", fail)
    for getter in (db.get_bop_components, db.get_waste_records):
        with pytest.raises(RuntimeError, match="database unavailable"):
            getter(report_id=1)
    assert collection_value(None, "test") == []
    with pytest.raises(ValueError):
        collection_value({"unexpected": 1}, "test")


def test_bad_bop_does_not_cancel_valid_survey(db):
    result = db.save_imported_multi_tab_data_atomic(1, 1, {
        "bop_components": [{"name": "Bad", "type": "Annular", "working_pressure": "invalid"}],
        "surveys": [{"md": 100, "inc": 0, "azi": 0}]})
    assert result["failed"] == 0 and result["surveys"] == 1
    assert result["review_rows"][0]["status"] == "INVALID_SOURCE"


def test_survey_calculations_isolation_reload_and_plot_data(db):
    sources = [{"md": "100 m", "inc": "10°", "azi": "45 deg"},
               {"md": "200 m", "inc": 15, "azi": 60},
               {"md": 300, "inc": 20, "azi": "INVALID"},
               {"md": 400, "inc": 20, "azi": None}]
    result = db.save_survey_records([dict(r, well_id=1, report_id=1) for r in sources])
    assert (result["accepted"], result["rejected"], result["calculated"]) == (3, 1, 2)
    rows = db.load_survey_points(report_id=1)
    assert rows[0]["tvd"] == pytest.approx(100 * math.cos(math.radians(10)))
    assert rows[1]["dls"] > 0
    assert rows[2]["azi"] is None and rows[2]["tvd"] is None
    series = plot_series(rows)
    assert series["md"] == [100, 200]
    assert series["tvd"] == [r["tvd"] for r in rows[:2]]
    assert result["review_items"][0]["original_value"]["azi"] == "INVALID"
    assert [r["azi"] for r in sources] == ["45 deg", 60, "INVALID", None]


def test_survey_update_inclination_and_no_zero_fallback(db):
    db.save_survey_points([{"well_id": 1, "report_id": 1, "md": 100, "inc": 2, "azi": 0}])
    db.save_survey_points([{"well_id": 1, "report_id": 1, "md": 100, "inc": 10, "azi": 0}])
    row = db.load_survey_points(report_id=1)[0]
    assert row["inc"] == 10 and row["tvd"] == pytest.approx(100 * math.cos(math.radians(10)))


def test_w13_paths_cwd_independent(tmp_path, monkeypatch):
    from core.runtime_config import drill_pipe_reference_paths
    before = drill_pipe_reference_paths()
    monkeypatch.chdir(tmp_path)
    assert drill_pipe_reference_paths() == before
    monkeypatch.setenv("DRILLMASTER_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("DRILLMASTER_DRILLPIPE_PATH", "vendor.xlsx")
    assert drill_pipe_reference_paths()[0] == tmp_path / "vendor.xlsx"


def test_partial_vs_complete_mud_validation():
    from core.validators import MudValidator
    partial = MudValidator.validate({"solid_percent": 7, "oil_percent": None, "water_percent": None})
    assert partial.is_valid and "Partial composition" in partial.summary()
    assert "expected ~100" not in partial.summary()
    invalid_total = MudValidator.validate({"solid_percent": 7, "oil_percent": 5, "water_percent": 0})
    assert "expected ~100" in invalid_total.summary()
    assert not MudValidator.validate({"solid_percent": 7, "oil_percent": 5, "water_percent": 88}).warnings


def test_real_workbook_actual_application_service_and_disk_reload(db):
    path = next(Path(__file__).resolve().parents[1].glob("*DDR*.xlsx"))
    service = DDRImportService(db, 1)
    report, extracted = service.extract_file(str(path))
    before = copy.deepcopy(extracted)
    result = service.import_records(extracted)
    assert result["failed"] == 0, result
    assert result["report_id"] and extracted == before
    db.engine.dispose()  # force fresh SQLite connections on reload
    rid = result["report_id"]
    surveys = db.load_survey_points(report_id=rid)
    assert [r["md"] for r in surveys] == [50, 108, 146]
    assert all(r["azi"] is None and r["tvd"] is None for r in surveys)
    bha = db.get_bha_report(service.well_id, report_id=rid)["bha_configs"]
    assert len(bha) == 9 and bha[0]["Component Name"] == '17-1/2" MT Bit'
    mud = db.get_mud_report(report_id=rid)
    chemicals = json.loads(mud["chemicals_json"])
    assert len(chemicals) == 34
    assert next(r for r in chemicals if r["product"] == "Barite")["type"] == "Weight Material"
    assert next(r for r in chemicals if r["product"] == "Barite")["stock"] == 45
    assert mud["kcl"] is None and mud["oil_percent"] is None
    with db.session_scope() as session:
        assert session.query(BulkMaterials).filter_by(report_id=rid).count() == 0
        before_counts = {model.__name__: session.query(model).count() for model in (DailyReport, BHAReport, SurveyPoint, ServiceCompanyPOB, BOPComponent)}
        audit = json.loads(session.query(AuditLog).filter_by(action="ddr_import").one().details)
        assert audit["source"]["metadata"]["raw_ir"]
        assert audit["result"]["review_items"]
    again = service.import_records(extracted)
    assert again["reimport"] and again["imported"] == 0 and again["report_id"] == rid
    with db.session_scope() as session:
        assert before_counts == {model.__name__: session.query(model).count() for model in (DailyReport, BHAReport, SurveyPoint, ServiceCompanyPOB, BOPComponent)}


def test_bha_manual_schema_and_mud_notification_static_wiring():
    root = Path(__file__).resolve().parents[1]
    downhole = (root / "tabs/w4_Downhole_Widget.py").read_text()
    mud = (root / "tabs/w3_drilling_report.py").read_text()
    assert '"Component Name"' in downhole
    assert '"bha_data": bha_data' in downhole
    assert 'self.show_warning(validation.summary())' not in mud
    assert 'QMessageBox.warning(self, "Mud validation", validation.summary())' in mud
    # Static wiring only; this assertion does not certify a rendered GUI.


def test_formula_only_rows_are_not_phantom_time_records():
    from core.excel_intelligence import DynamicTableExtractor
    class Merge:
        def get_value(self, r, c):
            return None, False
    extractor = DynamicTableExtractor({(1, 3): '=IF(A1="","",1)'}, Merge())
    columns = [{"col": 1, "field": "From", "canonical": "time_log.time_from"},
               {"col": 2, "field": "To", "canonical": "time_log.time_to"},
               {"col": 3, "field": "Hrs", "canonical": "time_log.duration"}]
    assert extractor._classify_row(1, columns) == "formula_only"
    result = extractor._build_result({"columns": columns}, "Variation", 1, 1, columns, 1, ["R1: formula_only"])
    assert result.records == [] and result.rejected_rows == 1


def test_manual_survey_append_uses_existing_trajectory_context(db):
    from core.engineering.engines.trajectory import TrajectoryCalculator
    source = [{"md": 100, "inc": 10, "azi": 45}, {"md": 200, "inc": 20, "azi": 60}]
    for row in source:
        db.save_survey_points([dict(row, well_id=1, report_id=1)])
    loaded = db.load_survey_points(report_id=1)
    expected = TrajectoryCalculator.calculate(source)
    assert loaded[-1]["north"] == pytest.approx(expected[-1]["north"])
    assert loaded[-1]["tvd"] == pytest.approx(expected[-1]["tvd"])


def test_manual_safety_repeated_save_no_child_duplicates(db):
    data = {"well_id": 1, "report_id": 1, "report_date": date(2024, 10, 22),
            "bop_stack_json": [{"Name": "Annular", "Type": "Annular", "WP (psi)": "3000", "Last Test": ""}],
            "waste_history_json": [{"Date": "2024-10-22", "Type": "Cuttings", "Volume (BBL)": "12", "pH": 7}]}
    assert db.save_safety_report(data)
    assert db.save_safety_report(data)
    assert len(db.get_bop_components(report_id=1)) == 1
    assert len(db.get_waste_records(report_id=1)) == 1
    assert db.get_safety_report(1, report_id=1)["bop_stack_json"][0]["Name"] == "Annular"


@pytest.mark.parametrize("padding,reorder", [(0, False), (3, True), (7, False)])
def test_generic_multiformat_table_pipeline(db, tmp_path, padding, reorder):
    from openpyxl import Workbook
    workbook = Workbook()
    header = workbook.active
    header.title = "Report context"
    for values in [("daily_report.report_date", date(2024, 10, 23)), ("daily_report.report_number", 2),
                   ("well_info.name", "Audit Well"), ("well_info.section_name", "Audit Section")]:
        header.append(values)
    survey = workbook.create_sheet("Measurements variation")
    for _ in range(padding):
        survey.append([])
    fields = ["Azimuth", "MD (m)", "Inclination"] if reorder else ["MD (m)", "Inclination", "Azimuth"]
    survey.append([None] * padding + fields)
    for values in ([45, 100, 10], [60, 200, 20]) if reorder else ([100, 10, 45], [200, 20, 60]):
        survey.append([None] * padding + list(values))
    survey.row_dimensions[padding + 2].hidden = True
    tools = workbook.create_sheet("Tools variation")
    tools.append(["Component Name", "OD (in)", "Length (m)"])
    tools.append(['9-1/2" DC', 9.5, 18.04])
    path = tmp_path / f"variation-{padding}.xlsx"
    workbook.save(path)
    service = DDRImportService(db, 1)
    extraction, payload = service.extract_file(str(path), template={})
    assert extraction.template_version == "generic-ir"
    assert len(payload["surveys"]) == 2
    assert len(payload["bha_components"]) == 1
    result = service.import_records(payload)
    assert result["failed"] == 0 and result["report_id"], result
    rows = db.load_survey_points(report_id=result["report_id"])
    assert len(rows) == 2 and rows[1]["dls"] > 0
    assert [r["md"] for r in rows] == [100, 200]


def test_well_identity_obeys_project_and_rejects_ambiguity(db):
    with db.session_scope() as session:
        project = Project(name="Other", code="OTHER", company_id=1)
        session.add(project)
        session.flush()
        other = Well(name="Audit Well", code="OTHER-WELL", project_id=project.id)
        session.add(other)
        session.flush()
        other_id = other.id
    assert DDRImportService(db, 1)._resolve_import_well({"name": "audit well"}) == 1
    assert DDRImportService(db, other_id)._resolve_import_well({"name": "Audit Well"}) == other_id
    with pytest.raises(ValueError, match="ambiguous"):
        DDRImportService(db)._resolve_import_well({"name": "Audit Well"})


def test_import_fingerprint_is_scoped_to_target_well(db):
    db.save_import_audit(1, "same-source", {}, {"well_id": 1, "report_id": 1})
    assert db.find_import_audit("same-source", well_id=1)
    assert db.find_import_audit("same-source", well_id=99999) is None


def test_derived_cleanup_keeps_manual_points_and_resets_code_usage(db):
    from datetime import datetime
    from core.database import ActivityCode, TimeLog24H, TimeDepthData
    with db.session_scope() as session:
        session.get(DailyReport, 1).depth_2400 = 100
        session.add(ActivityCode(well_id=1, main_phase="Drilling", main_code="F-PUMP", sub_code="F-PUMP", code_name="Pump"))
        session.add(TimeDepthData(well_id=1, report_id=1, timestamp=datetime(2024, 10, 22, 12), depth=50))
    DDRImportService(db, 1)._save_time_logs(1, [{"time_from": "00:00", "time_to": "01:00", "duration": 1,
                                             "status": "NPT", "npt_code": "F-PUMP", "contractor": "Operator"}])
    assert db.auto_update_from_daily_report(1)
    assert len(db.get_npt_reports(report_id=1)) == 1
    with db.session_scope() as session:
        code = session.query(ActivityCode).one()
        assert code.usage_count == 1 and code.last_used == date(2024, 10, 22)
        session.query(TimeLog24H).filter_by(report_id=1).delete()
        session.get(DailyReport, 1).depth_2400 = None
    assert db.auto_update_from_daily_report(1)
    assert db.get_npt_reports(report_id=1) == []
    with db.session_scope() as session:
        assert session.query(TimeDepthData).one().depth == 50
        code = session.query(ActivityCode).one()
        assert code.usage_count == 0 and code.total_hours == 0 and code.last_used is None


def test_survey_repository_uses_shared_engine(db):
    from core.repositories.report_repository import SurveyRepository
    repo = SurveyRepository(db)
    assert repo.save_points([{"well_id": 1, "report_id": 1, "md": 100, "inc": 10, "azi": 45},
                             {"well_id": 1, "report_id": 1, "md": 200, "inc": 20, "azi": 60}]) == 2
    assert db.load_survey_points(report_id=1)[1]["dls"] > 0
    assert repo.last_review_items == []


def test_import_does_not_activate_orm_measurement_defaults(db):
    from core.database import FuelWaterInventory, MaterialRequest, SevenDaysLookahead
    result = db.save_imported_multi_tab_data_atomic(1, 1, {
        "safety": {"days_without_lti": 468}, "fuel_water": {"fuel_stock": 100},
        "daily_report": {"material_request_detail": "Need valve"},
        "lookahead": [{"day": "2024-10-24", "activity": "Planned work", "date_start": "2024-10-24"}],
    })
    assert result["failed"] == 0
    with db.session_scope() as session:
        safety = session.query(SafetyReport).one()
        assert safety.waste_ph is None and safety.test_pressure is None and safety.lti_count is None
        fuel = session.query(FuelWaterInventory).one()
        assert fuel.fuel_stock == 100 and fuel.fuel_received is None and fuel.fuel_type is None
        request = session.query(MaterialRequest).one()
        assert request.requested_quantity is None and request.requested_unit is None
        plan = session.query(SevenDaysLookahead).one()
        assert plan.plan_date == date(2024, 10, 24) and plan.actual_start is None and plan.actual_end is None
