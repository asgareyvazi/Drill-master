"""Privacy regression check and DDR sanitizer (M43).

Positive cases are assembled at runtime so that this file -- which is itself
part of the scanned population -- never contains a literal that a rule
would reject.
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools" / "privacy"))

import privacy_check as pc  # noqa: E402
import sanitize_ddr_workbook as sd  # noqa: E402

FIXTURE = ROOT / "tests" / "fixtures" / "synthetic_ddr_golden_a.xlsx"
SYNTHETIC_PERSON = "Qx" + "zorvalli"          # an invented surname used as the protected token
SYNTHETIC_COORDINATE = "7," + "654,321"      # an invented coordinate in thousands notation


def _policy(**overrides):
    policy = pc.load_policy()
    policy = copy.deepcopy(policy)
    policy["denylist_sha256"] = [pc.token_digest(SYNTHETIC_PERSON), pc.token_digest(SYNTHETIC_COORDINATE)]
    policy.update(overrides)
    return policy


def _scanner(**overrides):
    return pc.Scanner(_policy(**overrides))


# -- text rules --------------------------------------------------------------

@pytest.mark.parametrize("line, rule", [
    ("set MINERU=C:" + "\\Users\\" + "jdoe" + "\\Desktop\\x.exe", "windows-user-profile-path"),
    ("-----BEGIN " + "RSA PRIVATE KEY-----", "private-key-block"),
    ("token=" + "ghp_" + "a" * 36, "github-token"),
    ("key " + "AKIA" + "ABCDEFGHIJKLMNOP", "aws-access-key"),
    ("mail " + "ops" + "@" + "realcorp.ir", "email-outside-allowed-domains"),
    ("Supervisor: A." + SYNTHETIC_PERSON, "denylisted-identity-token"),
    ("Northing " + SYNTHETIC_COORDINATE + " N", "denylisted-identity-token"),
    ("northing " + SYNTHETIC_COORDINATE.replace(",", ""), "denylisted-identity-token"),
    ("NAME " + SYNTHETIC_PERSON.upper(), "denylisted-identity-token"),
    # equipment serial numbers copied from a report (M43: bit serials survived in evidence JSON)
    ('"serial no": "' + "4" * 10 + ' | ' + "5" * 10 + '"', "equipment-serial-number"),
    ('"serial_number": "' + "8" * 8 + '"', "equipment-serial-number"),
    ("Serial Number = " + "7" * 7, "equipment-serial-number"),
    ("BIT SERIAL #: " + "3" * 9, "equipment-serial-number"),
])
def test_text_rules_fire_on_synthetic_positives(line, rule):
    rules = [f.rule for f in _scanner().scan_text(line, "doc.md")]
    assert rule in rules


@pytest.mark.parametrize("line", [
    "C:" + "\\Users\\runner\\work\\x.ps1",
    "C:" + "\\Users\\<you>\\Desktop",
    "%USERPROFILE%\\mineru-env",
    "author 123+bot" + "@" + "users.noreply.github.com",
    "demo" + "@" + "drillmaster.com",
    "diff \\n+" + "@pytest.mark.parametrize",
    "depth 3441 m, MW 71 pcf, Barite 45 sx",
    "Qxzorv is not the protected surname",
    "def test_serialized_provenance_is_not_domain_value(value):",
    '"equipment_id": "serial-A"',
    '"serial no": "[REDACTED:serial]"',
    "serial: 12345",
    "Serial Number: SN-A1-B2",
])
def test_text_rules_stay_quiet_on_negatives(line):
    assert _scanner().scan_text(line, "doc.md") == []


def test_findings_never_echo_the_protected_value():
    findings = _scanner().scan_text("Client Rep: S." + SYNTHETIC_PERSON, "x.md")
    rendered = "\n".join(f.render() for f in findings)
    assert findings and SYNTHETIC_PERSON.lower() not in rendered.lower()


def test_policy_contains_digests_not_plaintext():
    raw = (ROOT / "tools" / "privacy" / "policy.json").read_text(encoding="utf-8")
    policy = json.loads(raw)
    assert policy["denylist_sha256"], "the denylist must not be empty"
    assert all(len(d) == 64 for d in policy["denylist_sha256"])


def test_bad_policy_is_a_configuration_error(tmp_path):
    bad = tmp_path / "policy.json"
    bad.write_text(json.dumps({"denylist_sha256": ["not-a-digest"]}), encoding="utf-8")
    assert pc.main(["--policy", str(bad), "README.md"]) == 2


# -- file rules --------------------------------------------------------------

def _xlsx_bytes(text_cell: str) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Well Name:"
    wb.active["B1"] = text_cell
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_real_ddr_filename_is_prohibited_anywhere(tmp_path):
    name = "12-DDR " + "RIG-1 WELL-2 2024-Jan-01.xlsx"
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / name).write_bytes(_xlsx_bytes("SYN-1"))
    rules = [f.rule for f in _scanner().scan_file(tmp_path, f"sub/{name}")]
    assert "prohibited-filename" in rules and "data-file-not-allowlisted" in rules


def test_data_files_need_an_allowlist_entry_pinned_by_hash(tmp_path):
    data = _xlsx_bytes("SYN-1")
    (tmp_path / "fixture.xlsx").write_bytes(data)
    assert [f.rule for f in _scanner().scan_file(tmp_path, "fixture.xlsx")] == ["data-file-not-allowlisted"]
    allow = [{"path": "fixture.xlsx", "sha256": hashlib.sha256(data).hexdigest(), "provenance": "test"}]
    assert _scanner(allowlisted_data_files=allow).scan_file(tmp_path, "fixture.xlsx") == []
    (tmp_path / "fixture.xlsx").write_bytes(_xlsx_bytes("SYN-2"))  # same name, different bytes
    assert [f.rule for f in _scanner(allowlisted_data_files=allow).scan_file(tmp_path, "fixture.xlsx")] == [
        "allowlisted-data-file-changed"]


def test_identity_inside_an_allowlisted_workbook_is_still_found(tmp_path):
    data = _xlsx_bytes("P." + SYNTHETIC_PERSON)
    (tmp_path / "f.xlsx").write_bytes(data)
    allow = [{"path": "f.xlsx", "sha256": hashlib.sha256(data).hexdigest(), "provenance": "test"}]
    findings = _scanner(allowlisted_data_files=allow).scan_file(tmp_path, "f.xlsx")
    assert [f.rule for f in findings] == ["denylisted-identity-token"]
    assert findings[0].path.startswith("f.xlsx!xl/")


def test_spreadsheets_are_not_rejected_wholesale():
    policy = pc.load_policy()
    assert any(e["path"].endswith(".xlsx") for e in policy["allowlisted_data_files"])
    assert not any(p.endswith(r"\.xlsx$") for p in policy["prohibited_filename_patterns"])


def test_tracked_repository_population_is_clean():
    """The CI gate: the committed tree has no privacy findings."""
    result = subprocess.run([sys.executable, str(ROOT / "tools" / "privacy" / "privacy_check.py")],
                            cwd=ROOT, capture_output=True, text=True, timeout=300)
    assert result.returncode == 0, result.stdout[-4000:]


# -- sanitizer ---------------------------------------------------------------

def _source_workbook(tmp_path: Path) -> Path:
    """A small workbook with a formula + cached value, an image, printer settings and doc properties."""
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image

    png = tmp_path / "logo.png"
    Image.new("RGB", (20, 10), (200, 30, 30)).save(png)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DDR Remark"
    ws["A1"], ws["B1"] = "Client Rep:", "S." + SYNTHETIC_PERSON
    ws["A2"], ws["B2"] = "Northing (m):", 7654321
    ws["A3"], ws["B3"] = "MW (PCF):", 71
    ws["A4"] = '="Prepared by "&B1'
    ws.merge_cells("C1:D1")
    ws.add_image(XLImage(str(png)), "F1")
    wb.properties.creator = "Real Author"
    hidden = wb.create_sheet("setting")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Field list"
    path = tmp_path / "source.xlsx"
    wb.save(path)
    # emulate Excel: a cached formula result and a printer-settings part
    with zipfile.ZipFile(path) as zf:
        files = {n: zf.read(n) for n in zf.namelist()}
    sheet = files["xl/worksheets/sheet1.xml"].decode()
    sheet = sheet.replace("<f>\"Prepared by \"&amp;B1</f><v></v>", "")
    sheet = sheet.replace('<c r="A4" t="str"><f>"Prepared by "&amp;B1</f><v></v></c>',
                          f'<c r="A4" t="str"><f>"Prepared by "&amp;B1</f><v>Prepared by S.{SYNTHETIC_PERSON}</v></c>')
    if f"Prepared by S.{SYNTHETIC_PERSON}" not in sheet:
        sheet = sheet.replace('<c r="A4"', f'<c r="A4" t="str"><f>"Prepared by "&amp;B1</f><v>Prepared by S.'
                              f'{SYNTHETIC_PERSON}</v></c><c r="Z99"', 1)
    sheet = sheet.replace("</worksheet>", '<pageSetup r:id="rIdP1"/></worksheet>') if "<pageSetup" not in sheet \
        else sheet
    files["xl/worksheets/sheet1.xml"] = sheet.encode()
    rels = files.get("xl/worksheets/_rels/sheet1.xml.rels", b"").decode() or (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>')
    rels = rels.replace("</Relationships>", '<Relationship Id="rIdP1" Type="http://schemas.openxmlformats.org/'
                        'officeDocument/2006/relationships/printerSettings" '
                        'Target="../printerSettings/printerSettings1.bin"/></Relationships>')
    files["xl/worksheets/_rels/sheet1.xml.rels"] = rels.encode()
    files["xl/printerSettings/printerSettings1.bin"] = ("DEVICE " + SYNTHETIC_PERSON).encode("utf-16-le")
    with zipfile.ZipFile(path, "w") as zf:
        for name, data in files.items():
            zf.writestr(name, data)
    return path


def _spec(tmp_path: Path) -> dict:
    spec = {"text_substitutions": [["S." + SYNTHETIC_PERSON, "P.Synth01"]],
            "cell_values": {"DDR Remark!B2": 1000000},
            "verify_absent": [SYNTHETIC_PERSON, "Real Author", "7654321"]}
    path = tmp_path / "spec.json"
    path.write_text(json.dumps(spec), encoding="utf-8")
    return sd.load_spec(path)


def test_sanitizer_removes_identity_and_keeps_structure(tmp_path):
    import openpyxl
    src = _source_workbook(tmp_path)
    out = sd.sanitize(src.read_bytes(), _spec(tmp_path))
    assert sd.find_leaks(out, [SYNTHETIC_PERSON, "Real Author", "7654321"]) == []
    (tmp_path / "out.xlsx").write_bytes(out)
    values = openpyxl.load_workbook(tmp_path / "out.xlsx", data_only=True)
    formulas = openpyxl.load_workbook(tmp_path / "out.xlsx")
    ws = values["DDR Remark"]
    assert ws["B1"].value == "P.Synth01" and ws["B2"].value == 1000000 and ws["B3"].value == 71
    assert ws["A4"].value == "Prepared by P.Synth01"              # cached formula result substituted
    assert formulas["DDR Remark"]["A4"].value.startswith("=")       # formula preserved
    assert [str(r) for r in ws.merged_cells.ranges] == ["C1:D1"]
    assert values["setting"].sheet_state == "hidden"
    with zipfile.ZipFile(io.BytesIO(out)) as zf:
        names = zf.namelist()
        assert not any(n.startswith("xl/printerSettings/") for n in names)
        assert "printerSettings" not in zf.read("[Content_Types].xml").decode()
        media = [n for n in names if n.startswith("xl/media/")]
        assert media
        from PIL import Image
        with Image.open(io.BytesIO(zf.read(media[0]))) as img:
            assert img.size == (20, 10) and img.getextrema() == ((255, 255), (255, 255), (255, 255))


def test_sanitizer_is_deterministic(tmp_path):
    src = _source_workbook(tmp_path).read_bytes()
    assert sd.sanitize(src, _spec(tmp_path)) == sd.sanitize(src, _spec(tmp_path))


def test_sanitizer_fails_closed_when_a_token_survives(tmp_path):
    src = _source_workbook(tmp_path)
    spec = _spec(tmp_path)
    spec["text_substitutions"] = []          # nothing replaced -> leak check must refuse
    with pytest.raises(sd.SanitizationError) as err:
        sd.sanitize(src.read_bytes(), spec)
    assert SYNTHETIC_PERSON not in str(err.value)   # names parts, never the token


def test_sanitizer_refuses_to_overwrite_and_rejects_bad_specs(tmp_path):
    src = _source_workbook(tmp_path)
    spec = tmp_path / "spec.json"
    spec.write_text(json.dumps({"text_substitutions": [["a", "a-b"]]}), encoding="utf-8")
    with pytest.raises(sd.SanitizationError):
        sd.load_spec(spec)
    spec.write_text(json.dumps({"text_substitutions": [], "cell_values": {"B2": 1}}), encoding="utf-8")
    with pytest.raises(sd.SanitizationError):
        sd.load_spec(spec)
    spec.write_text(json.dumps({"text_substitutions": []}), encoding="utf-8")
    with pytest.raises(SystemExit):
        sd.main([str(src), str(src), "--spec", str(spec)])


def test_cell_dump_uses_the_same_spec(tmp_path):
    spec = _spec(tmp_path)
    dump = {"DDR Remark": {"B1": {"value": "S." + SYNTHETIC_PERSON, "row": 1, "col": 2, "col_letter": "B"},
                           "B2": {"value": 7654321, "row": 2, "col": 2, "col_letter": "B"},
                           "B3": {"value": 71, "row": 3, "col": 2, "col_letter": "B"}}}
    out = sd.sanitize_cell_dump(dump, spec)
    assert [out["DDR Remark"][k]["value"] for k in ("B1", "B2", "B3")] == ["P.Synth01", 1000000, 71]
    assert out["DDR Remark"]["B1"]["row"] == 1


# -- the committed fixture -----------------------------------------------------

def test_committed_fixture_keeps_the_golden_workbook_structure():
    import openpyxl
    wb = openpyxl.load_workbook(FIXTURE, data_only=True)
    assert wb.sheetnames == ["DDR Remark", "DDR Data", "Lookahead", "Service Company", "Activity Codes", "setting"]
    assert wb["setting"].sheet_state == "hidden"
    assert wb["DDR Remark"]["W3"].value == "SYNW-101"
    assert wb["DDR Remark"]["AI11"].value == 1000000 and wb["DDR Remark"]["AI12"].value == 500000
    with zipfile.ZipFile(FIXTURE) as zf:
        assert not any(n.startswith("xl/printerSettings/") for n in zf.namelist())
        assert b"<dc:creator>" not in zf.read("docProps/core.xml") or \
            b"<dc:creator></dc:creator>" in zf.read("docProps/core.xml")


def test_no_real_workbook_at_the_repository_root():
    assert not list(ROOT.glob("*DDR*.xlsx")), "operational workbooks must not be committed (see docs/privacy/)"
