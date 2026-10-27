"""M48 UI scenarios, run in an isolated subprocess by test_m48_analysis_ui.py.

Real ``AnalysisWidget`` + real ``SelectionManager`` signals on the synthetic
fixture (tests/m48_synthetic.py), Qt offscreen platform. Modal dialogs are
replaced by recorders so a scenario can never block. Each scenario prints one
``M48RESULT {json}`` line.
"""
import json
import os
import sys
import tempfile
import traceback
import warnings

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


def _blanket_ignores():
    return [f for f in warnings.filters
            if f[0] == "ignore" and f[1] is None and f[2] is Warning and f[3] is None]


BLANKET_BEFORE = len(_blanket_ignores())

import matplotlib  # noqa: E402

BACKEND_BEFORE = matplotlib.get_backend()

from PySide6.QtCore import QCoreApplication, QEvent, Qt  # noqa: E402
from PySide6.QtWidgets import QApplication, QMessageBox  # noqa: E402

app = QApplication.instance() or QApplication([])

from core.permissions import permissions  # noqa: E402

permissions.set_user({"id": 1, "username": "m48", "role": "engineer"})

import tabs.w12_Analysis as w12  # noqa: E402
from core import analysis_evidence as ev  # noqa: E402
from core.selection_manager import SelectionManager  # noqa: E402
from m48_synthetic import build  # noqa: E402

BLANKET_AFTER = len(_blanket_ignores())
BACKEND_AFTER = matplotlib.get_backend()

DIALOGS = []
for _name in ("warning", "information", "critical", "question"):
    setattr(QMessageBox, _name, staticmethod(
        lambda *a, _n=_name, **k: DIALOGS.append((_n, a[2] if len(a) > 2 else "")) or QMessageBox.Ok))

SCENARIOS = []


def scenario(fn):
    SCENARIOS.append(fn)
    return fn


class Ctx:
    def __init__(self):
        self.db, self.ids = build()
        self.sel = SelectionManager()
        self.sel.clear()
        self.tab = w12.AnalysisWidget(db_manager=self.db)
        self.tab.show()
        app.processEvents()   # the event loop must run (layout/paint call Qt virtuals)
        self.tmp = tempfile.mkdtemp(prefix="m48ui_")
        DIALOGS.clear()
        # Default: any file dialog is "cancelled" (a scenario that needs a file
        # sets it with save_as); a real modal dialog could never block a run.
        w12.QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: ("", ""))

    def close(self):
        self.tab.cleanup()
        self.sel.clear()
        self.tab.hide()
        self.tab.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)

    # helpers -------------------------------------------------------------
    def well(self, key):
        self.sel.select_well(self.ids[key], {"name": f"SYN-{key}"})
        app.processEvents()

    def bore(self, key, well="B"):
        self.sel.select_wellbore(self.ids[key], {"id": self.ids[key], "name": key,
                                                 "well_id": self.ids[well]})
        app.processEvents()

    def kpi(self, key):
        return next(c for c in self.tab.kpi_cards if c.metric_key == key).metric_value

    def perf(self, key):
        return next(c for c in self.tab.perf_kpi_cards if c.metric_key == key).metric_value

    def analysis(self, kind):
        self.tab.analytics_type.blockSignals(True)
        self.tab.analytics_type.setCurrentText(kind)
        self.tab.analytics_type.blockSignals(False)
        self.tab.run_advanced_analysis()
        return self.tab.results_text.toPlainText()

    def evidence(self, metric="rop", level=ev.SCOPE_WELL):
        t = self.tab
        t.evidence_metric.setCurrentIndex(t.evidence_metric.findData(metric))
        t.evidence_level.setCurrentIndex(t.evidence_level.findData(level))
        t.evidence_load_btn.click()
        return t._evidence_rows

    def save_as(self, path):
        w12.QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (path, ""))


def _all_kpis_are(ctx, state):
    return all(c.metric_value.state == state for c in ctx.tab.kpi_cards)


# ---------------------------------------------------------------------------
# Environment contracts
# ---------------------------------------------------------------------------

@scenario
def no_global_warning_suppression_and_backend_untouched(ctx):
    assert BLANKET_AFTER == BLANKET_BEFORE, (BLANKET_BEFORE, BLANKET_AFTER)
    assert BACKEND_AFTER == BACKEND_BEFORE, (BACKEND_BEFORE, BACKEND_AFTER)
    import inspect
    source = inspect.getsource(w12)
    assert "filterwarnings" not in source and "matplotlib.use(" not in source
    assert "import matplotlib.pyplot" not in source and "plt." not in source


@scenario
def npt_pie_canvas_uses_shared_backend_selector(ctx):
    from core.matplotlib_widgets import select_figure_canvas
    ctx.well("A")
    canvas_class, _is_qt = select_figure_canvas()
    canvases = ctx.tab.findChildren(canvas_class)
    assert canvases, "the NPT pie must be a select_figure_canvas() canvas"


@scenario
def page_paints_through_the_event_loop_after_selection(ctx):
    # Regression (found by workflow_m48 during M48): a Python attribute named
    # ``metric`` on a QFrame card shadowed the virtual QPaintDevice.metric()
    # and crashed the first event-loop pass. Paint the whole page for real.
    ctx.well("C")
    for index in range(ctx.tab.analysis_tabs.count()):
        ctx.tab.analysis_tabs.setCurrentIndex(index)
        app.processEvents()
    image = ctx.tab.grab()
    assert not image.isNull() and image.width() > 0
    # Readable cards: panel and card box rules are scoped by object name, so
    # no padding/border/margin cascades into the value labels (it clipped the
    # KPI numbers in the M48 screenshots). Size-independent box-model check.
    ctx.tab.analysis_tabs.setCurrentIndex(2)
    app.processEvents()
    cards = [*ctx.tab.kpi_cards, *ctx.tab.perf_kpi_cards, ctx.tab.npt_total_card]
    for card in cards:
        m = card.value_label.contentsMargins()
        assert max(m.left(), m.top(), m.right(), m.bottom()) <= 2, (getattr(card, "metric_key", "npt"), m)
    # Depth History: the container's padding must not cascade into the table
    # header (it clipped the column titles in the M48 screenshots).
    for widget in (ctx.tab.time_depth_table, ctx.tab.time_depth_table.horizontalHeader()):
        m = widget.contentsMargins()
        assert max(m.left(), m.top(), m.right(), m.bottom()) <= 2, (type(widget).__name__, m)
    # The header grid gives each widget its own cells (shared cells painted the
    # well/controls rows over the Data Quality and Plan-vs-Actual labels).
    grid = ctx.tab.status_label.parentWidget().layout()
    cells = {}
    for i in range(grid.count()):
        r, c, rs, cs = grid.getItemPosition(i)
        for cell in ((rr, cc) for rr in range(r, r + rs) for cc in range(c, c + cs)):
            assert cell not in cells, ("header cell shared", cell)
            cells[cell] = i
    assert ctx.tab.quality_label.isVisible() and ctx.tab.plan_variance_label.isVisible()
    assert "#2ecc71" in ctx.tab.status_label.styleSheet()      # "Updated" is not red
    ctx.sel.clear()
    app.processEvents()
    assert "#e74c3c" in ctx.tab.status_label.styleSheet()


@scenario
def style_sheets_never_use_ambiguous_8_digit_hex(ctx):
    # Qt reads #AARRGGBB; a CSS-style #RRGGBBAA turned the blue depth card green
    # and the grey help panel yellow. Every style sheet on the page is checked.
    import re
    from PySide6.QtWidgets import QWidget
    ctx.well("C")
    offenders = []
    for widget in [ctx.tab, *ctx.tab.findChildren(QWidget)]:
        sheet = widget.styleSheet()
        if re.search(r"#[0-9a-fA-F]{8}\b", sheet):
            offenders.append((type(widget).__name__, widget.objectName()))
    assert not offenders, offenders
    assert w12._qss_rgba("#3498DB", "30") == "rgba(52, 152, 219, 48)"   # by hand


@scenario
def card_attributes_never_shadow_the_qt_api(ctx):
    from PySide6.QtCore import Signal
    from PySide6.QtWidgets import QFrame
    cards = [*ctx.tab.kpi_cards, *ctx.tab.perf_kpi_cards, ctx.tab.npt_total_card,
             ctx.tab.npt_percent_card, ctx.tab.npt_daily_card]

    def check(stage):
        for card in cards:
            own = set(vars(card)) if hasattr(card, "__dict__") else set()
            # PySide caches bound signals in the instance dict under their own
            # names; those are Qt's, not a shadowing Python attribute.
            clashes = [name for name in own if hasattr(QFrame, name)
                       and not isinstance(getattr(QFrame, name), Signal)]
            assert not clashes, (stage, clashes)

    check("constructed")
    ctx.well("C")                 # card setters run on a real value ...
    check("value")
    ctx.sel.clear()               # ... and on the no-selection state
    app.processEvents()
    check("cleared")


# ---------------------------------------------------------------------------
# KPI states and selection lifecycle
# ---------------------------------------------------------------------------

@scenario
def no_well_shows_no_selection_everywhere(ctx):
    t = ctx.tab
    assert _all_kpis_are(ctx, ev.STATE_NO_SELECTION)
    assert all(c.metric_value.state == ev.STATE_NO_SELECTION for c in t.perf_kpi_cards)
    assert t.evidence_state_label.text() == "Evidence: select a well"
    assert not t.evidence_open_btn.isEnabled()


@scenario
def full_well_values_units_and_tooltips(ctx):
    ctx.well("C")
    # Well C: depth 700 m, 7 report dates, ROP mean (1+...+7)/7 = 4.0, NPT 7 x 1 h
    assert ctx.kpi("current_depth").text() == "700.0 m"
    assert ctx.kpi("rig_days").value == 7
    assert ctx.kpi("avg_rop").text() == "4.0 m/hr"
    assert ctx.kpi("total_npt").text() == "7.0 hrs"
    card = next(c for c in ctx.tab.kpi_cards if c.metric_key == "avg_rop")
    assert "[m/hr]" in card.toolTip() and "finite and >= 0" in card.toolTip()
    assert ctx.perf("best_rop").value == 7.0
    assert ctx.tab.status_label.text() == "🟢 Updated"


@scenario
def real_zero_rop_is_shown_as_zero(ctx):
    ctx.well("A")
    rows = ctx.evidence("rop")
    first = rows[0]
    assert (first.value, first.state) == (0.0, ev.STATE_VALUE)
    assert ctx.tab.evidence_table.item(0, 1).text() == "0.00"
    # canonical average over 0, 10, 8, 12, 6 = 36 / 5 = 7.2 (zero counted, None/-3 excluded)
    assert ctx.kpi("avg_rop").value == 7.2


@scenario
def well_without_reports_is_no_data_not_zero(ctx):
    from core.database import Well
    s = ctx.db.create_session()
    w = Well(name="SYN-EMPTY", code="SE", project_id=s.get(Well, ctx.ids["A"]).project_id)
    s.add(w)
    s.commit()
    wid = w.id
    s.close()
    ctx.sel.select_well(wid, {"name": "SYN-EMPTY"})
    assert _all_kpis_are(ctx, ev.STATE_MISSING)
    assert all("0.0" not in c.value_label.text() for c in ctx.tab.kpi_cards)
    assert ctx.perf("avg_rop").state == ev.STATE_MISSING


@scenario
def switch_a_to_b_leaves_nothing_stale(ctx):
    ctx.well("A")
    ctx.evidence("rop")
    ctx.well("B")
    assert ctx.kpi("current_depth").value == 700.0
    assert ctx.kpi("rig_days").value == 3
    rows = ctx.tab._evidence_rows
    assert rows and all(r.report_id in (None, *ctx.ids["B_reports"]) for r in rows)


@scenario
def deselect_clears_every_value(ctx):
    ctx.well("C")
    ctx.sel.select_well(None)
    assert _all_kpis_are(ctx, ev.STATE_NO_SELECTION)
    t = ctx.tab
    assert t.npt_table.rowCount() == 0 and t.evidence_table.rowCount() == 0
    assert not t.update_timer.isActive()
    ctx.well("C")
    ctx.sel.clear()
    assert _all_kpis_are(ctx, ev.STATE_NO_SELECTION)
    assert t.results_text.toPlainText() == "No well selected"


@scenario
def refresh_error_after_result_shows_error_and_retry_recovers(ctx):
    ctx.well("C")
    assert ctx.kpi("current_depth").value == 700.0
    original = w12.AnalysisWidget._canonical_scope_kpis

    def broken(self):
        raise RuntimeError("synthetic read failure")
    w12.AnalysisWidget._canonical_scope_kpis = broken
    try:
        ctx.tab.refresh_now()
        assert _all_kpis_are(ctx, ev.STATE_ERROR)
        assert ctx.kpi("current_depth").text() == "Read error"   # not 700, not "No data"
        assert ctx.tab.status_label.text().startswith("🔴 Error in: read kpis")
        assert "Refresh Now" in ctx.tab.status_label.text()
    finally:
        w12.AnalysisWidget._canonical_scope_kpis = original
    ctx.tab.refresh_now()
    assert ctx.kpi("current_depth").value == 700.0


@scenario
def canonical_error_dict_is_an_error_not_no_data(ctx):
    ctx.well("C")
    ctx.tab._canonical_scope_kpis = lambda: {"error": "OperationalError"}
    ctx.tab.refresh_now()
    assert _all_kpis_are(ctx, ev.STATE_ERROR)


@scenario
def well_bore_well_round_trip(ctx):
    ctx.well("B")
    assert ctx.kpi("rig_days").value == 3
    ctx.bore("OH")
    assert ctx.kpi("rig_days").value == 2 and ctx.kpi("current_depth").value == 600.0
    # NPT daily average: 2 h over 2 report dates = 1.0 h/day
    assert ctx.tab.npt_daily_card.metric_value.value == 1.0
    # W12's own scoped queries (not only the canonical KPIs): OH has 2 reports
    assert [r["depth"] for r in ctx.tab.chart_data["time_depth"]["data"]] == [500.0, 600.0]
    assert ctx.tab.recent_reports_table.rowCount() == 2
    ctx.bore("ST1")
    assert ctx.kpi("current_depth").value == 700.0
    # Back to the whole well: re-selecting the SAME well is a no-op in
    # SelectionManager by design; the whole-well view is select_wellbore(None).
    ctx.sel.select_wellbore(None)
    assert ctx.kpi("rig_days").value == 3 and ctx.tab.current_wellbore_id in (None, 0)
    ctx.bore("OH")
    ctx.well("C")                      # a different well clears the bore scope
    ctx.well("B")
    assert ctx.kpi("rig_days").value == 3 and not ctx.tab.current_wellbore_id


@scenario
def section_of_previous_well_never_scopes_the_new_well(ctx):
    ctx.well("B")
    ctx.tab.current_section_id = 12345       # a section selected under well B
    ctx.well("C")
    assert ctx.tab.current_section_id is None
    assert "section=None" in ctx.tab._cache_key("kpis")


@scenario
def stale_read_is_never_painted_over_a_newer_selection(ctx):
    ctx.well("C")
    original = w12.AnalysisWidget.calculate_kpis
    switched = []

    def slow(self, session):
        result = original(self, session)
        if not switched:
            switched.append(True)
            ctx.well("A")  # the user switches while well C is being read
        return result
    w12.AnalysisWidget.calculate_kpis = slow
    try:
        ctx.tab.refresh_now()
    finally:
        w12.AnalysisWidget.calculate_kpis = original
    assert ctx.tab.current_well_id == ctx.ids["A"]
    assert ctx.kpi("current_depth").value == 300.0   # A's depth, not C's 700


@scenario
def data_change_is_picked_up_by_auto_update_and_refresh(ctx):
    from core.database import DailyReport
    ctx.well("C")
    s = ctx.db.create_session()
    s.get(DailyReport, ctx.ids["C_reports"][-1]).depth_2400 = 750.0
    s.commit()
    s.close()
    ctx.tab.auto_update_data()        # timer path: same cache window, new revision
    assert ctx.kpi("current_depth").value == 750.0
    s = ctx.db.create_session()
    s.get(DailyReport, ctx.ids["C_reports"][-1]).depth_2400 = 760.0
    s.commit()
    s.close()
    ctx.tab.refresh()                 # MainWindow.refresh_all_tabs path
    assert ctx.kpi("current_depth").value == 760.0


# ---------------------------------------------------------------------------
# M49 - cache integrity: edits that change no count, id, sum or timestamp
# ---------------------------------------------------------------------------

def _raw_sql(ctx, statement, **params):
    """Write like an external tool would: plain SQL, no ORM onupdate hooks."""
    from sqlalchemy import text
    s = ctx.db.create_session()
    try:
        s.execute(text(statement), params)
        s.commit()
    finally:
        s.close()


def _npt_view(ctx):
    data = ctx.tab.chart_data.get("npt") or {}
    return dict(data.get("categories") or {}), ctx.tab.npt_category_card.value_label.text()


@scenario
def m49_currency_only_change_updates_cost_analysis_without_refresh(ctx):
    ctx.well("A")
    ctx.tab.auto_update_check.setChecked(True)
    text = ctx.analysis("Cost Analysis")
    assert "actual EUR 500.00" in text
    # Only the currency code changes: same row count, ids, amounts; raw SQL so
    # CostRecord.updated_at (an ORM onupdate) does not move either.
    _raw_sql(ctx, "UPDATE cost_records SET currency = 'GBP' WHERE well_id = :w AND currency = 'EUR'",
             w=ctx.ids["A"])
    ctx.tab.auto_update_data()                       # timer path, no Refresh / Run click
    text = ctx.tab.results_text.toPlainText()
    assert "actual GBP 500.00" in text and "EUR" not in text, text


@scenario
def m49_npt_code_only_change_updates_npt_view_without_refresh(ctx):
    ctx.well("C")
    ctx.tab.auto_update_check.setChecked(True)
    categories, card = _npt_view(ctx)
    assert categories == {"NPT-Z": 7.0} and "NPT-Z" in card, (categories, card)
    # The code changes, the duration does not (TimeLog24H has no updated_at).
    _raw_sql(ctx, "UPDATE time_logs_24h SET main_code = 'NPT-W' WHERE main_code = 'NPT-Z'")
    ctx.tab.auto_update_data()
    categories, card = _npt_view(ctx)
    assert categories == {"NPT-W": 7.0} and "NPT-W" in card and "NPT-Z" not in card, (categories, card)


@scenario
def m49_refresh_and_auto_update_agree_with_a_cold_read(ctx):
    ctx.well("C")
    ctx.tab.auto_update_check.setChecked(True)
    ctx.analysis("Cost Analysis")
    _raw_sql(ctx, "UPDATE time_logs_24h SET main_code = 'NPT-V' WHERE main_code = 'NPT-Z'")
    _raw_sql(ctx, "UPDATE cost_records SET currency = 'CAD' WHERE well_id = :w", w=ctx.ids["C"])
    ctx.tab.auto_update_data()
    auto = (_npt_view(ctx), ctx.tab.results_text.toPlainText())
    ctx.tab.refresh()
    forced = (_npt_view(ctx), ctx.tab.results_text.toPlainText())
    cold = w12.AnalysisWidget(db_manager=ctx.db)     # a fresh page: nothing cached
    try:
        cold.on_well_changed(ctx.ids["C"], {"name": "SYN-C"})
        cold.analytics_type.blockSignals(True)
        cold.analytics_type.setCurrentText("Cost Analysis")
        cold.analytics_type.blockSignals(False)
        cold.update_all_data()
        app.processEvents()
        data = cold.chart_data.get("npt") or {}
        reference = ((dict(data.get("categories") or {}), cold.npt_category_card.value_label.text()),
                     cold.results_text.toPlainText())
    finally:
        cold.cleanup()
        cold.deleteLater()
    assert auto[0] == forced[0] == reference[0] and auto[0][0] == {"NPT-V": 7.0}, (auto[0], forced[0])
    assert "CAD 7,000.00" in auto[1] and "CAD 7,000.00" in forced[1] and "CAD 7,000.00" in reference[1]


@scenario
def m49_cached_values_never_cross_into_a_new_scope(ctx):
    t = ctx.tab
    ctx.well("C")
    assert ctx.kpi("avg_rop").value == 4.0 and "NPT-Z" in _npt_view(ctx)[1]
    c_keys = set(t.data_cache)
    assert c_keys, "well C must have populated the cache"
    ctx.well("A")
    assert ctx.kpi("avg_rop").value == 7.2 and ctx.kpi("current_depth").value == 300.0
    categories, card = _npt_view(ctx)
    assert "NPT-Z" not in categories and "NPT-Z" not in card, (categories, card)
    scope_a = t._scope_key()
    assert all(scope_a in key for key in t.data_cache if key not in c_keys), list(t.data_cache)
    ctx.well("B")
    ctx.bore("ST1")
    assert ctx.kpi("avg_rop").value == 25.0                     # ST1 only: (20 + 30) / 2
    ctx.well("C")
    assert ctx.kpi("avg_rop").value == 4.0 and "NPT-Z" in _npt_view(ctx)[1]


@scenario
def m49_read_error_during_refresh_is_never_cached_or_shown_as_no_data(ctx):
    t = ctx.tab
    ctx.well("C")
    t.auto_update_check.setChecked(True)
    original = t.get_npt_data
    calls = []

    def failing(session):
        calls.append(1)
        raise RuntimeError("synthetic read failure")
    t.get_npt_data = failing
    try:
        t.refresh()
        card = t.npt_category_card.value_label.text()
        assert "Read error" in t.npt_total_card.value_label.text() or "Read error" in card, card
        assert "No data" not in card
        assert not any(key.startswith("npt|") for key in t.data_cache), list(t.data_cache)
        assert "Error" in t.status_label.text()
    finally:
        t.get_npt_data = original
    t.auto_update_data()                             # no forced refresh: the error was not cached
    assert _npt_view(ctx)[0] == {"NPT-Z": 7.0}
    assert calls == [1]

@scenario
def cache_key_covers_section_report_and_revision(ctx):
    ctx.well("C")
    t = ctx.tab
    base = t._cache_key("kpis")
    t.current_report_id = 99
    assert t._cache_key("kpis") != base
    t.current_report_id = None
    t.current_section_id = 5
    assert t._cache_key("kpis") != base
    t.current_section_id = None
    t._data_rev = ("other",)
    assert t._cache_key("kpis") != base


@scenario
def manual_auto_update_tick_starts_the_timer(ctx):
    # PySide6 6.8: stateChanged emits an int and 2 == Qt.Checked is False,
    # so the old `state == Qt.Checked` comparison never started the timer.
    assert (2 == Qt.Checked) is False
    ctx.well("C")
    t = ctx.tab
    t.auto_update_check.setChecked(False)
    assert not t.update_timer.isActive()
    t.auto_update_check.setChecked(True)   # emits stateChanged(int 2)
    assert t.update_timer.isActive()
    assert t.update_timer.interval() == t._auto_update_interval_ms()


# ---------------------------------------------------------------------------
# Projection / NPT / cost / risk text
# ---------------------------------------------------------------------------

@scenario
def projection_assessed_only_with_enough_dates(ctx):
    ctx.well("C")
    text = ctx.analysis("ROP Prediction")
    # C: ROP = day index on 7 consecutive dates -> slope exactly 1.000
    assert "Trend slope: 1.000 m/hr per day" in text and "NOT a validated forecast" in text
    assert "no cause is assessed" in text
    ctx.well("A")
    text = ctx.analysis("ROP Prediction")
    # A: valid dates 1, 2, 5, 6 (3 missing, 4 negative) = 4 < 5
    assert "NOT ASSESSED" in text and "Valid dated observations: 4" in text
    assert "1 x invalid: negative value" in text and "1 x missing: not recorded" in text


@scenario
def npt_forecast_uses_coverage_valid_days_including_zero(ctx):
    ctx.well("A")
    text = ctx.analysis("NPT Forecasting")
    # A: valid days 1 (0), 2 (2.0), 5 (1.5), 6 (0) = 4 -> NOT ASSESSED; 2 valid zeros
    assert "Coverage-valid report dates: 4 of 6 (2 with a valid zero" in text
    assert "NOT ASSESSED" in text
    assert ev.UNCLASSIFIED_NPT in text
    ctx.well("C")
    assert "Trend slope: 0.000 hours per day" in ctx.analysis("NPT Forecasting")


@scenario
def npt_coverage_label_names_partial_and_missing_days(ctx):
    ctx.well("A")
    text = ctx.tab.npt_coverage_label.text()
    assert "4 of 6" in text and "partial 1" in text and "missing 1" in text


@scenario
def cost_golden_single_currency(ctx):
    ctx.well("C")
    text = ctx.analysis("Cost Analysis")
    # actual 7000, planned 6000, variance planned - actual = -1000,
    # NPT cost 7000 x 7 / 168 = 291.67, productive 6708.33, per metre 7000 / 700 = 10.00
    for expected in ("Actual (recorded) total:  USD 7,000.00",
                     "Estimate (planned) total: USD 6,000.00",
                     "USD -1,000.00", "USD 291.67", "USD 6,708.33", "USD 10.00/m"):
        assert expected in text, expected
    assert "$" not in text


@scenario
def cost_mixed_currency_is_never_summed(ctx):
    ctx.well("A")
    text = ctx.analysis("Cost Analysis")
    assert "several currencies - totals NOT aggregated" in text
    assert "Actual total NOT ASSESSED" in text
    assert "1,500" not in text and "$" not in text
    # Each currency stays in its own currency (fixture: USD actual 1000 /
    # planned 1200, EUR actual 500); the total is "not summed", not "unknown".
    assert "Actual (recorded) total:  — (not summed: several currencies)" in text
    assert "actual USD 1,000.00; estimate USD 1,200.00" in text
    assert "actual EUR 500.00" in text
    assert "currency unknown" not in text


@scenario
def cost_is_never_attributed_to_a_wellbore(ctx):
    ctx.well("B")
    ctx.bore("OH")
    text = ctx.analysis("Cost Analysis")
    assert f"the selected wellbore #{ctx.ids['OH']} is NOT attributed a share" in text


@scenario
def risk_states(ctx):
    ctx.well("C")
    text = ctx.analysis("Risk Assessment")
    assert "Overall level: LOW" in text and "NOT a probability" in text
    assert "Equipment: NOT ASSESSED" in text and "Weather: NOT ASSESSED" in text
    ctx.well("A")
    text = ctx.analysis("Risk Assessment")
    assert "Overall level: NOT ASSESSED" in text and "LOW" not in text.split("Level rule")[0]


@scenario
def analysis_read_error_is_reported_not_blank(ctx):
    ctx.well("C")
    ctx.tab.analyze_risk = lambda session: (_ for _ in ()).throw(RuntimeError("x"))
    text = ctx.analysis("Risk Assessment")
    assert "read/processing error (RuntimeError)" in text


# ---------------------------------------------------------------------------
# Evidence workflow and exports
# ---------------------------------------------------------------------------

def _contrast(hex_a, hex_b):
    """WCAG 2.x contrast ratio of two #rrggbb colours (independent formula)."""
    def lum(h):
        chans = [int(h[i:i + 2], 16) / 255 for i in (1, 3, 5)]
        lin = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in chans]
        return 0.2126 * lin[0] + 0.7152 * lin[1] + 0.0722 * lin[2]
    hi, lo = sorted((lum(hex_a), lum(hex_b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


@scenario
def scope_label_is_readable_on_the_header_in_every_state(ctx):
    import re
    header_bg = "#2c3e50"
    seen = set()
    for step in (lambda: ctx.well("A"), lambda: ctx.well("B"), lambda: ctx.bore("OH")):
        step()
        colour = re.search(r"color:\s*(#[0-9a-fA-F]{6})", ctx.tab.scope_label.styleSheet()).group(1)
        assert _contrast(colour, header_bg) >= 3.0, (ctx.tab.scope_label.text(), colour)
        seen.add(ctx.tab.scope_label.text())
    assert len(seen) == 3          # whole well / whole-well aggregate / wellbore


@scenario
def analytics_chart_title_states_the_outcome_and_never_goes_stale(ctx):
    def title():
        # setTitle(None) hides the label but keeps its text: hidden = no title.
        label = ctx.tab.analytics_plot.getPlotItem().titleLabel
        return label.text if label.isVisible() else ""
    ctx.well("A")                                   # 4 valid ROP dates < 5
    ctx.analysis("ROP Prediction")
    assert "NOT ASSESSED" in title()
    ctx.analysis("Cost Analysis")                   # next analysis: no stale title
    assert "Projection" not in title() and "NOT ASSESSED" not in title()
    assert title() == "Cost: NOT CHARTED - several currencies are never summed (no FX conversion)"
    ctx.well("C")                                   # 7 valid dates -> assessed
    ctx.analysis("ROP Prediction")
    assert title().startswith("Projection (linear fit")
    assert ctx.tab.analytics_plot.getPlotItem().getAxis("left").labelText == "ROP (m/hr)"
    ctx.analysis("Risk Assessment")
    assert title().endswith("(not a probability)")
    assert ctx.tab.analytics_plot.getPlotItem().getAxis("left").labelText == "Rule score (0-10)"
    ticks = ctx.tab.analytics_plot.getPlotItem().getAxis("bottom")._tickLevels
    assert [name for _, name in ticks[0]] == ["NPT Risk", "Safety"]
    ctx.analysis("ROP Prediction")                  # custom ticks never leak onward
    assert ctx.tab.analytics_plot.getPlotItem().getAxis("bottom")._tickLevels is None
    assert ctx.tab.analytics_plot.getPlotItem().getAxis("left").labelText == "ROP (m/hr)"
    ctx.analysis("Risk Assessment")
    ctx.sel.clear()                                 # deselect: nothing of SYN-C remains
    app.processEvents()
    assert title() == ""
    assert ctx.tab.analytics_plot.getPlotItem().getAxis("bottom")._tickLevels is None
    assert ctx.tab.analytics_plot.getPlotItem().getAxis("left").labelText == "Value"
    assert not ctx.tab.analytics_plot.getPlotItem().items


@scenario
def evidence_selection_rings_its_point_and_point_click_selects_the_row(ctx):
    from datetime import datetime as _dt
    ctx.well("C")
    rows = ctx.evidence("rop")
    ctx.tab.select_evidence_row(2)                     # 2026-02-03, 3.0 m/hr
    ring = ctx.tab._evidence_highlight
    assert ring is not None and ring in ctx.tab.evidence_plot.getPlotItem().items
    (x,), (y,) = ring.getData()
    assert y == rows[2].value == 3.0
    assert x == _dt(2026, 2, 3).timestamp()
    ctx.tab.select_evidence_row(4)                     # moving the selection moves the ring
    assert ring not in ctx.tab.evidence_plot.getPlotItem().items
    assert ctx.tab._evidence_highlight.getData()[1][0] == 5.0
    # point -> row: the trend scatter carries each point's row index
    scatter = next(i for i in ctx.tab.evidence_plot.getPlotItem().items
                   if isinstance(i, w12.pg.ScatterPlotItem) and i is not ctx.tab._evidence_highlight)
    ctx.tab._on_evidence_point_clicked(scatter, [scatter.points()[1]])
    assert ctx.tab._selected_evidence() is rows[1]
    # an ambiguous row (no value) gets no ring
    ctx.well("B")
    rows = ctx.evidence("rop")
    ctx.tab.select_evidence_row(next(i for i, r in enumerate(rows) if r.state == ev.STATE_AMBIGUOUS))
    assert ctx.tab._evidence_highlight is None


@scenario
def legacy_row_never_double_counts_a_date_in_whole_well_cards(ctx):
    # SYN-B: legacy 99.0 (no report_id) on the date of linked report B1.
    ctx.well("B")
    assert ctx.kpi("avg_rop").value == 15.5          # (5+7+20+30)/4, by hand
    assert ctx.perf("best_rop").value == 30.0        # never the legacy 99.0
    rows = ctx.evidence("rop")
    legacy = [r for r in rows if r.reason == ev.LEGACY_DUPLICATE_REASON]
    assert len(legacy) == 1 and legacy[0].drilldown_report_id is None


@scenario
def evidence_ambiguous_point_disables_drilldown(ctx):
    ctx.well("B")
    rows = ctx.evidence("rop")
    index = next(i for i, r in enumerate(rows) if r.state == ev.STATE_AMBIGUOUS)
    ctx.tab.select_evidence_row(index)
    assert not ctx.tab.evidence_open_btn.isEnabled()
    assert ctx.tab.evidence_open_btn.toolTip().startswith("Ambiguous ownership")
    assert ctx.tab.open_evidence_report() is False
    # Depth: two reports (OH, ST1) share 2026-01-02; each row keeps its report_id
    # but neither may be opened as "the" source report.
    rows = ctx.evidence("depth")
    index = next(i for i, r in enumerate(rows) if r.state == ev.STATE_AMBIGUOUS)
    assert rows[index].report_id is not None
    ctx.tab.select_evidence_row(index)
    assert not ctx.tab.evidence_open_btn.isEnabled()
    assert ctx.tab.open_evidence_report() is False
    # The disabled state must be visible, not only functional: the rendered
    # background differs from the enabled "Show Evidence" button next to it.
    def _bg(button):
        image = button.grab().toImage()
        return image.pixelColor(6, image.height() // 2).name()
    assert ctx.tab.evidence_load_btn.isEnabled()
    assert _bg(ctx.tab.evidence_open_btn) != _bg(ctx.tab.evidence_load_btn), \
        (_bg(ctx.tab.evidence_open_btn), _bg(ctx.tab.evidence_load_btn))


@scenario
def evidence_scope_levels_and_date_range(ctx):
    from PySide6.QtCore import QDate
    ctx.well("C")
    t = ctx.tab
    assert len(ctx.evidence("rop", ev.SCOPE_REPORT)) == 0
    assert "NOT ASSESSED - no daily report selected" in t.evidence_state_label.text()
    t.evidence_range_check.setChecked(True)
    t.evidence_from.setDate(QDate(2026, 2, 2))
    t.evidence_to.setDate(QDate(2026, 2, 4))
    rows = ctx.evidence("rop", ev.SCOPE_WELL)
    assert [r.value for r in rows] == [2.0, 3.0, 4.0]
    assert "2026-02-02 to 2026-02-04" in t.evidence_state_label.text()


@scenario
def html_report_escapes_text_and_uses_metric_values(ctx):
    from core.database import Well
    s = ctx.db.create_session()
    s.get(Well, ctx.ids["C"]).name = "SYN-<b>C</b>"
    s.commit()
    s.close()
    ctx.sel.select_well(ctx.ids["C"], {"name": "SYN-<b>C</b>"})
    html = ctx.tab._generate_report_html()
    assert "SYN-&lt;b&gt;C&lt;/b&gt;" in html and "SYN-<b>C" not in html
    assert "700.0 m" in html and "4.0 m/hr" in html


@scenario
def export_all_data_is_typed_from_datasets(ctx):
    from openpyxl import load_workbook
    ctx.well("A")
    path = os.path.join(ctx.tmp, "all.xlsx")
    assert ctx.tab.export_all_data(path) == [path]
    book = load_workbook(path)
    perf = book["performance"]
    header = [c.value for c in perf[1]]
    rop_col = header.index("avg_rop_m_hr")
    values = [row[rop_col].value for row in perf.iter_rows(min_row=2)]
    # A: 0.0 (real zero, numeric), 10.0, None (empty), -3.0 (stored as recorded), 8, 12, 6
    assert values[0] == 0 and perf.cell(2, rop_col + 1).data_type == "n"
    assert values[2] is None
    assert "metadata" in book.sheetnames


@scenario
def export_all_charts_pdf_honours_the_chosen_filename(ctx):
    ctx.well("C")
    path = os.path.join(ctx.tmp, "chosen_name.pdf")
    ctx.tab.export_all_charts(path)
    assert os.path.exists(path) and os.path.getsize(path) > 0


@scenario
def chart_image_export_has_exporters(ctx):
    import pyqtgraph as pg
    assert hasattr(pg, "exporters") and hasattr(pg.exporters, "ImageExporter")


# ---------------------------------------------------------------------------
# Three end-to-end user flows (real widgets, button clicks)
# ---------------------------------------------------------------------------

@scenario
def e2e_performance_report_export_switch_well(ctx):
    import csv
    t = ctx.tab
    ctx.well("C")
    t.analysis_tabs.setCurrentWidget(t.evidence_tab)
    rows = ctx.evidence("rop")
    assert len(rows) == 7
    t.select_evidence_row(2)                     # the 2026-02-03 point (value 3.0)
    assert t.evidence_open_btn.isEnabled()
    t.evidence_open_btn.click()
    target = ctx.ids["C_reports"][2]
    assert ctx.sel.current_report_id == target and t.current_report_id == target
    # The opened report gets its data-quality summary; the tooltip lists one
    # metric per line (it used to show a literal backslash-n).
    app.processEvents()
    assert t.quality_label.text().startswith("Data Quality: ") and "—" not in t.quality_label.text()
    tip = t.quality_label.toolTip()
    assert "\\n" not in tip and len(tip.splitlines()) > 1, tip
    path = os.path.join(ctx.tmp, "flow1.csv")
    ctx.save_as(path)
    t.evidence_csv_btn.click()
    with open(path, encoding="utf-8", newline="") as fh:
        lines = list(csv.reader(fh))
    assert any(r[:2] == ["# scope", f"Well #{ctx.ids['C']} (whole well)"] for r in lines)
    body = lines[next(i for i, r in enumerate(lines) if r and r[0] == "metric") + 1:]
    assert [r[2] for r in body] == ["1.0", "2.0", "3.0", "4.0", "5.0", "6.0", "7.0"]
    ctx.well("A")
    assert t.current_report_id is None
    assert all(r.report_id in ctx.ids["A_reports"] for r in t._evidence_rows)


@scenario
def e2e_npt_coverage_forecast_not_assessed_export(ctx):
    from openpyxl import load_workbook
    t = ctx.tab
    ctx.well("A")
    t.analysis_tabs.setCurrentIndex(t.analysis_tabs.indexOf(t.evidence_tab))
    assert "partial 1" in t.npt_coverage_label.text()
    assert "NOT ASSESSED" in ctx.analysis("NPT Forecasting")
    rows = ctx.evidence("npt")
    states = [r.state for r in rows]
    assert states == [ev.STATE_VALUE, ev.STATE_VALUE, ev.STATE_PARTIAL, ev.STATE_MISSING,
                      ev.STATE_VALUE, ev.STATE_VALUE]
    path = os.path.join(ctx.tmp, "flow2.xlsx")
    ctx.save_as(path)
    t.evidence_xlsx_btn.click()
    sheet = load_workbook(path)["evidence"]
    got = [(r[2].value, r[4].value) for r in sheet.iter_rows(min_row=2)]
    assert got[0] == (0, "value") and got[2] == (None, "partial") and got[3] == (None, "missing")


@scenario
def e2e_cost_and_risk(ctx):
    t = ctx.tab
    ctx.well("A")
    assert "totals NOT aggregated" in ctx.analysis("Cost Analysis")
    assert "Overall level: NOT ASSESSED" in ctx.analysis("Risk Assessment")
    ctx.well("C")
    assert "USD 7,000.00" in ctx.analysis("Cost Analysis")
    assert "Overall level: LOW" in ctx.analysis("Risk Assessment")
    ctx.sel.select_well(None)
    assert t.results_text.toPlainText() == "No well selected"


def main(selected=None):
    for fn in SCENARIOS:
        if selected and fn.__name__ not in selected:
            continue
        ctx = None
        try:
            ctx = Ctx()
            fn(ctx)
            result = {"name": fn.__name__, "ok": True}
        except Exception:
            result = {"name": fn.__name__, "ok": False, "error": traceback.format_exc()[-3000:]}
        finally:
            if ctx is not None:
                try:
                    ctx.close()
                except Exception:
                    pass
        print("M48RESULT " + json.dumps(result), flush=True)


if __name__ == "__main__":
    main(set(sys.argv[1:]) or None)
