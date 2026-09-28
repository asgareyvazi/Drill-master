"""Mission 31 adversarial scenarios.

Every scenario drives real engine/service/database/packaging paths. Nothing here
mocks a database, a report renderer or a service: the point is to observe the
actual behaviour of the delivered code for the matrix Mission 31 requires
(plan/actual, currency, safety, scope, ownership, snapshot, import, export,
backup, permissions).
"""
from __future__ import annotations

import hashlib
import os
from datetime import date, time

import pytest

from core.actual_vs_plan import ActualVsPlanEngine, activity_plan_totals, compare
from core.cost_semantics import format_money, normalize_currency, summarize_costs
from core.operational_time import summarize_time_logs
from core.safety_semantics import safety_kpis
from core.scope_attribution import (
    ALREADY,
    AMBIGUOUS,
    INVALID,
    RESOLVED,
    ScopeAttributionService,
)
from test_scope_attribution import _base, _get_report, _mgr, _report, _section, _well, _wellbore


@pytest.fixture
def env():
    db = _mgr()
    project = _base(db)
    a, b = _well(db, project, "A"), _well(db, project, "B")
    ba, bb = _wellbore(db, a, "Same"), _wellbore(db, b, "Same")
    sa, sb = _section(db, a, "Same", wellbore_id=ba), _section(db, b, "Same", wellbore_id=bb)
    ra = _report(db, a, 1, wellbore_id=ba, section_id=sa)
    rb = _report(db, b, 1, wellbore_id=bb, section_id=sb)
    yield db, dict(a=a, b=b, ba=ba, bb=bb, sa=sa, sb=sb, ra=ra, rb=rb, project=project)
    db.close()


# --------------------------------------------------------------------------
# S8: Plan / Actual / Forecast / Target matrix
# --------------------------------------------------------------------------
@pytest.mark.parametrize(
    "planned,actual,expected",
    [
        (None, None, ("unavailable", None, None)),
        (None, 0, ("unavailable", None, None)),
        (0, None, ("unavailable", None, None)),
        (0, 0, ("on-track", 0.0, 0.0)),
        (100, None, ("unavailable", None, None)),
        (None, 100, ("unavailable", None, None)),
        (100, 80, ("behind", -20.0, -20.0)),
        (80, 100, ("ahead", 20.0, 25.0)),
        (100, 100, ("on-track", 0.0, 0.0)),
    ],
)
def test_plan_actual_matrix_all_required_pairs(planned, actual, expected):
    result = compare("Depth", planned, actual)
    assert (result.status, result.variance, result.variance_pct) == expected


def test_plan_only_and_actual_only_never_fabricate_counterparts():
    planned_only = ActualVsPlanEngine.compare_metrics({"depth_m": 100}, {"hours": 5})
    assert planned_only.success is False or planned_only.validation_status == "missing_input"
    assert not planned_only.values
    unavailable = planned_only.metadata["unavailable_metrics"]
    assert set(unavailable) == {"depth_m", "hours", "rop_m_per_hr", "npt_hours", "cost"}
    for metric in unavailable.values():
        assert metric["planned"] is None or metric["actual"] is None


def test_activity_totals_require_complete_rows_and_keep_explicit_zero():
    complete = [{"planned_duration_hours": 0, "planned_depth_to": 0},
                {"planned_duration_hours": 4, "planned_depth_to": 10}]
    assert activity_plan_totals(complete) == {"hours": 4, "depth": 10}
    # an unknown duration invalidates the hours total but not an independent max
    incomplete_hours = complete + [{"planned_duration_hours": None, "planned_depth_to": 20}]
    assert activity_plan_totals(incomplete_hours) == {"hours": None, "depth": 20}
    # an unknown depth invalidates the depth total but not the hours sum
    incomplete_depth = complete + [{"planned_duration_hours": 2, "planned_depth_to": None}]
    assert activity_plan_totals(incomplete_depth) == {"hours": 6, "depth": None}


def test_persisted_plan_actual_variance_matches_engine(env):
    db, x = env
    from datetime import date as _date

    db.save_cost_record({"well_id": x["a"], "category": "Drilling", "planned_cost": 100.0,
                         "actual_cost": 80.0, "currency": "USD", "cost_date": _date(2026, 1, 1)})
    summary = summarize_costs(db.get_cost_records(x["a"]))
    assert summary["status"] == "single-currency"
    assert summary["total_planned"] == 100.0 and summary["total_actual"] == 80.0
    assert compare("Cost", summary["total_planned"], summary["total_actual"]).variance_pct == -20.0


# --------------------------------------------------------------------------
# S10/S11: currency
# --------------------------------------------------------------------------
def _cost(planned, actual, currency):
    return {"category": "C", "planned_cost": planned, "actual_cost": actual, "currency": currency}


def test_multi_currency_rows_have_no_single_total_but_keep_groups():
    summary = summarize_costs([_cost(100, 50, "USD"), _cost(200, 100, "EUR")])
    assert summary["status"] == "multi-currency"
    assert summary["currency"] is None
    assert summary["total_planned"] is None and summary["total_actual"] is None
    groups = {(g["currency"], g["planned"], g["actual"]) for g in summary["groups"]}
    assert groups == {("USD", 100.0, 50.0), ("EUR", 200.0, 100.0)}


@pytest.mark.parametrize("marker", ["", "unknown", "N/A", "not recorded", "—", None])
def test_unknown_currency_never_authorizes_a_money_total(marker):
    summary = summarize_costs([_cost(100, 50, marker)])
    assert summary["status"] == "unknown-currency"
    assert summary["total_planned"] is None and summary["total_actual"] is None
    assert normalize_currency(marker) is None
    assert format_money(None, marker) == "— (currency unknown)"


def test_null_amount_invalidates_only_its_own_side():
    summary = summarize_costs([_cost(100, None, "USD"), _cost(50, 25, "USD")])
    assert summary["total_planned"] == 150.0
    assert summary["total_actual"] is None
    assert summary["variance"] is None


# --------------------------------------------------------------------------
# S12: safety
# --------------------------------------------------------------------------
class _Safety:
    def __init__(self, report_date, days_without_lti=None, lti_count=None, near_miss_count=None):
        self.report_date = report_date
        self.days_without_lti = days_without_lti
        self.lti_count = lti_count
        self.near_miss_count = near_miss_count


def test_missing_safety_record_is_not_zero_incidents():
    kpis = safety_kpis([])
    assert kpis == {"days_without_lti": None, "total_lti": None, "total_near_miss": None}


def test_explicit_zero_events_survive_as_zero():
    kpis = safety_kpis([_Safety(date(2026, 1, 1), days_without_lti=12, lti_count=0, near_miss_count=0)])
    assert kpis == {"days_without_lti": 12, "total_lti": 0, "total_near_miss": 0}


def test_one_and_multiple_events_are_counted():
    one = safety_kpis([_Safety(date(2026, 1, 1), days_without_lti=3, lti_count=1, near_miss_count=2)])
    assert (one["total_lti"], one["total_near_miss"]) == (1, 2)
    many = safety_kpis([_Safety(date(2026, 1, 1), lti_count=1, near_miss_count=0),
                        _Safety(date(2026, 1, 2), lti_count=2, near_miss_count=None)])
    assert many["total_lti"] == 3
    assert many["total_near_miss"] is None  # one unrecorded field unknownises that total only


def test_incomplete_assessment_and_ambiguous_latest_day_are_unknown():
    assert safety_kpis([_Safety(date(2026, 1, 1), lti_count=None)])["total_lti"] is None
    ambiguous = [_Safety(date(2026, 1, 5), days_without_lti=1), _Safety(date(2026, 1, 5), days_without_lti=2)]
    assert safety_kpis(ambiguous)["days_without_lti"] is None


def test_time_logs_missing_vs_zero_and_npt_classification():
    from types import SimpleNamespace

    assert summarize_time_logs([])["total_hours"] is None
    zeros = summarize_time_logs([SimpleNamespace(duration=0, is_npt=False)])
    assert zeros["total_hours"] == 0 and zeros["npt_hours"] == 0 and zeros["productive_hours"] == 0
    unknown_flag = summarize_time_logs([SimpleNamespace(duration=5, is_npt=None)])
    assert unknown_flag["total_hours"] == 5 and unknown_flag["npt_hours"] is None


# --------------------------------------------------------------------------
# S1-S5: scope, ownership, attribution
# --------------------------------------------------------------------------
def test_same_named_scope_entities_do_not_cross_between_wells(env):
    db, x = env
    from core.database import TimeLog24H

    with db.session_scope() as session:
        session.add(TimeLog24H(report_id=x["ra"], time_from=time(0), time_to=time(10), duration=10,
                               activity_description="A only", is_npt=False))
        session.add(TimeLog24H(report_id=x["rb"], time_from=time(0), time_to=time(1), duration=1,
                               activity_description="B only", is_npt=True))
    from core.data_quality import DataQualityService

    quality_a = DataQualityService(db).for_report(x["ra"])
    quality_b = DataQualityService(db).for_report(x["rb"])
    assert quality_a is not None and quality_b is not None
    from core.actual_vs_plan import section_actual_days

    with db.session_scope() as session:
        assert section_actual_days(session, x["sa"]) == pytest.approx(10 / 24)
        assert section_actual_days(session, x["sb"]) == pytest.approx(1 / 24)


def test_null_scope_with_single_candidate_resolves_once_then_reports_already(env):
    db, x = env
    lone = _wellbore(db, x["a"], "Lone")
    sec = _section(db, x["a"], "LoneSection", wellbore_id=lone)
    orphan = _report(db, x["a"], 7, wellbore_id=None, section_id=sec)
    service = ScopeAttributionService(db)
    analysis = service.analyze(well_id=x["a"])
    resolved = [o for o in analysis.outcomes if o.report_id == orphan and o.dimension == "wellbore"]
    assert [o.status for o in resolved] == [RESOLVED]
    assert resolved[0].target_id == lone
    assert _get_report(db, orphan)[0] is None  # analyze never mutates
    first = service.resolve(well_id=x["a"])
    assert first.applied >= 1
    assert _get_report(db, orphan)[0] == lone
    second = service.resolve(well_id=x["a"])
    assert second.applied == 0
    assert second.wellbore[ALREADY] >= 1


def test_two_bores_one_section_is_ambiguous_and_never_overwritten(env):
    """A bore may never be guessed; a unique section may be resolved (documented rule)."""
    db, x = env
    _wellbore(db, x["a"], "Second")
    ambiguous = _report(db, x["a"], 9, wellbore_id=None, section_id=None)
    service = ScopeAttributionService(db)
    analysis = service.analyze(well_id=x["a"])
    outcomes = [o for o in analysis.outcomes if o.report_id == ambiguous and o.dimension == "wellbore"]
    assert [o.status for o in outcomes] == [AMBIGUOUS]
    assert outcomes[0].target_id is None
    resolved = service.resolve(well_id=x["a"])
    assert resolved.wellbore[AMBIGUOUS] == 1
    applied = _get_report(db, ambiguous)
    assert applied[0] is None, "two candidate bores must never be reduced to one silently"
    if applied[1] is not None:
        # only the unique-section rule may have been applied, and only to the section column
        outcome = [o for o in resolved.outcomes if o.report_id == ambiguous and o.dimension == "section"][0]
        assert outcome.status == RESOLVED and outcome.method == "via_unique_section_in_well"


def test_cross_well_attribution_is_reported_invalid(env):
    db, x = env
    service = ScopeAttributionService(db)
    scoped = service.analyze(well_id=x["a"])
    assert scoped.total_reports == 1  # the well filter excludes well B entirely
    assert {o.report_id for o in scoped.outcomes} == {x["ra"]}
    everything = service.analyze()
    assert everything.total_reports == 2
    assert not any(o.status == INVALID for o in everything.outcomes)  # coherent fixture stays coherent

    # contradict the ownership chain through the raw boundary, then re-analyze
    from core.database import DailyReport

    with db.engine.connect() as connection:
        connection.execute(
            DailyReport.__table__.update().where(DailyReport.id == x["ra"]).values(wellbore_id=x["bb"])
        )
        connection.commit()
    contradicted = service.analyze(well_id=x["a"])
    statuses = {o.status for o in contradicted.outcomes if o.report_id == x["ra"]}
    assert statuses == {INVALID}
    repaired = service.resolve(well_id=x["a"])
    assert repaired.applied == 0
    assert _get_report(db, x["ra"])[0] == x["bb"]  # reported, never silently rewritten


def test_ownership_conflict_is_rejected_before_commit(env):
    db, x = env
    from core.database import DailyReport, OwnershipIntegrityError

    with pytest.raises(OwnershipIntegrityError):
        with db.session_scope() as session:
            report = session.get(DailyReport, x["ra"])
            report.wellbore_id = x["bb"]  # bore belongs to the other well
            session.flush()


def test_raw_sql_bypass_is_not_silently_repaired(env):
    """SQLite has no cross-table CHECK; the invariant is ORM-enforced.

    Mission 31 records the actual enforcement level: a raw UPDATE can write a
    contradictory row, and the ORM must then NOT silently rewrite it (silent
    authoritative repair is worse than an explicit invalidity).
    """
    db, x = env
    from core.database import DailyReport

    engine = db.engine
    with engine.connect() as connection:
        connection.execute(
            DailyReport.__table__.update()
            .where(DailyReport.id == x["ra"])
            .values(wellbore_id=x["bb"])
        )
        connection.commit()
    assert _get_report(db, x["ra"])[0] == x["bb"]  # read back verbatim: no auto-repair
    service = ScopeAttributionService(db)
    report = service.analyze(well_id=x["a"])
    statuses = {o.status for o in report.outcomes if o.report_id == x["ra"]}
    assert statuses == {INVALID}, "contradictory legacy/external scope must be reported, not repaired"


# --------------------------------------------------------------------------
# S17: snapshot integrity
# --------------------------------------------------------------------------
def test_report_snapshot_is_not_aliased_to_live_master_rows(env):
    db, x = env
    from core.database import Base, DailyReport

    models = {cls.__name__: cls for cls in Base.registry._class_registry.values() if isinstance(cls, type)}
    with db.session_scope() as session:
        report = session.get(DailyReport, x["ra"])
        report.summary = "captured"
        report.header_snapshot = {"well": "A", "nested": {"value": 1}}
        snapshot = None
        from core.report_snapshot import build_report_snapshot, snapshot_is_complete

        snapshot = build_report_snapshot(session, report, models)
        assert snapshot_is_complete(snapshot)
        report.summary = "mutated after snapshot"
        report.header_snapshot["nested"]["value"] = 999
        report.depth_2400 = 777
    assert snapshot["report"]["summary"] == "captured"
    assert snapshot["report"]["depth_2400"] == 100
    assert snapshot["report"]["header_snapshot"]["nested"]["value"] == 1


# --------------------------------------------------------------------------
# S15/S16: duplicate and malformed import
# --------------------------------------------------------------------------
def test_duplicate_import_is_idempotent(env):
    db, x = env
    from core.database import BulkMaterials

    payload = {"bulk_materials": [dict(material_name="Diesel", unit="L", initial_stock=100,
                                       received=10, used=5, report_date=date(2026, 1, 1))]}
    first = db.save_imported_multi_tab_data_atomic(x["a"], x["ra"], payload)
    second = db.save_imported_multi_tab_data_atomic(x["a"], x["ra"], payload)
    assert first["imported"] >= 1 and second["imported"] >= 1
    with db.session_scope() as session:
        rows = session.query(BulkMaterials).filter_by(report_id=x["ra"]).all()
        assert len(rows) == 1, "re-import must upsert the rows this report owns"
        assert rows[0].initial_stock == 100 and rows[0].received == 10 and rows[0].used == 5
    assert second["failed"] == 0


def test_malformed_import_rolls_back_without_partial_rows(env):
    db, x = env
    from core.database import BulkMaterials

    good = {"bulk_materials": [dict(material_name="Diesel", unit="L", initial_stock=100,
                                    received=1, used=1, report_date=date(2026, 1, 1))]}
    assert db.save_imported_multi_tab_data_atomic(x["a"], x["ra"], good)["imported"] >= 1
    with db.session_scope() as session:
        before = session.query(BulkMaterials).filter_by(report_id=x["ra"]).count()
    bad = {"bulk_materials": [dict(material_name="Diesel", unit="L", initial_stock=10,
                                   received="broken", report_date=date(2026, 1, 1))]}
    from core.import_diagnostics import PersistenceError

    with pytest.raises(PersistenceError):
        db.save_imported_multi_tab_data_atomic(x["a"], x["ra"], bad)
    with db.session_scope() as session:
        rows = session.query(BulkMaterials).filter_by(report_id=x["ra"]).all()
        assert len(rows) == before, "a malformed row must not leave a partial import behind"
        assert all(r.initial_stock == 100 for r in rows), "the previous good row must survive intact"


# --------------------------------------------------------------------------
# S20: export parity for critical fields
# --------------------------------------------------------------------------
def test_report_export_html_and_excel_keep_unknowns_unfabricated(tmp_path):
    db = _mgr()
    project = _base(db)
    well = _well(db, project, "ExportWell")
    bore = _wellbore(db, well, "Original")
    section = _section(db, well, "12.25in", wellbore_id=bore)
    report_id = _report(db, well, 1, wellbore_id=bore, section_id=section, depth=1234)
    db.save_cost_record({"well_id": well, "category": "Drilling", "planned_cost": 1000.0,
                         "actual_cost": None, "currency": None, "cost_date": date(2026, 1, 1)})
    from core.report_engine import CostReportEngine, DDRReportEngine, EOWRReportEngine

    ddr_path = tmp_path / "ddr.html"
    assert DDRReportEngine(db).generate(report_id, str(ddr_path), format="html") is True
    ddr = ddr_path.read_text(encoding="utf-8")
    assert "ExportWell" in ddr
    assert "1234.0" in ddr               # recorded depth is printed
    assert ">0.0 m<" not in ddr          # an unreported depth is never printed as a measured zero
    assert "—" in ddr                    # unknown values use the explicit unknown marker

    cost = CostReportEngine(db)
    cost_html, cost_excel = tmp_path / "cost.html", tmp_path / "cost.xlsx"
    assert cost.generate(well, str(cost_html), format="html") is True
    assert cost.generate(well, str(cost_excel), format="excel") is True
    html = cost_html.read_text(encoding="utf-8")
    assert "— (currency unknown)" in html
    assert "No currency conversion is performed" in html
    from openpyxl import load_workbook

    workbook = load_workbook(cost_excel)
    summary = {row[0]: row[1] for row in workbook["Cost Summary"].iter_rows(values_only=True)}
    assert summary["Actual currency status"] == "unknown-currency"
    source = [row for row in workbook["Source lines"].iter_rows(values_only=True)]
    assert ("Drilling", None, 1000, None) in source, source

    eowr_path = tmp_path / "eowr.html"
    assert EOWRReportEngine(db).generate(well, str(eowr_path), format="html") is True
    eowr = eowr_path.read_text(encoding="utf-8")
    assert "9. Safety Summary" in eowr
    assert "Not recorded: no safety data was submitted for this scope." in eowr
    assert "Not recorded: no trajectory/survey data was submitted for this scope." in eowr


def test_unknown_numbers_are_not_rendered_as_measured_zeros():
    from core.text_utils import fmt_num

    assert fmt_num(None, 1) == "—"
    assert fmt_num(None, 2, default=0.0) == "0.00"   # explicit "no movement" still available
    assert fmt_num(0, 1) == "0.0"                    # explicit zero stays a fact
    assert fmt_num("not-a-number", 1, default=None) == "—"


# --------------------------------------------------------------------------
# S18/S19: backup and permissions
# --------------------------------------------------------------------------
def _file_db(tmp_path) -> "object":
    from core.database import DatabaseManager

    manager = DatabaseManager()
    manager.db_path = str(tmp_path / "live.db")
    manager.initialize()
    return manager


def test_failed_backup_preserves_the_previous_good_backup(tmp_path, monkeypatch):
    db = _file_db(tmp_path)
    destination = tmp_path / "backup.db"
    assert db.backup_to(destination) == str(destination)
    digest_before = hashlib.sha256(destination.read_bytes()).hexdigest()
    # source vanishes: a naive implementation would publish an empty database
    os.remove(db.db_path)
    assert db.backup_to(destination) is None
    assert hashlib.sha256(destination.read_bytes()).hexdigest() == digest_before
    db.close()


def test_permission_denied_destination_reports_failure(tmp_path):
    db = _file_db(tmp_path)
    locked = tmp_path / "locked"
    locked.mkdir()
    os.chmod(locked, 0o500)
    try:
        assert db.backup_to(locked / "backup.db") is None
        assert db.backup_to(db.db_path) is None  # never overwrite the live database
    finally:
        os.chmod(locked, 0o700)
        db.close()


def test_missing_or_locked_destination_is_not_a_silent_success(tmp_path):
    db = _file_db(tmp_path)
    empty_name = tmp_path / "not-a-dir"
    empty_name.write_text("")
    assert db.backup_to(empty_name / "backup.db") is None
    db.close()


# --------------------------------------------------------------------------
# S22: well-control kill sheet must never compute from absent kick data
# --------------------------------------------------------------------------
def _kill(**overrides):
    from core.engineering.well_control_kill_sheet import (
        build_canonical_kill_sheet_inputs,
        compute_kill_sheet,
    )

    raw = dict(tvd_m=3000.0, md_m=3100.0, shoe_tvd_m=2500.0, hole_size_in=12.25,
               casing_id_in=8.835, mw_pcf=100.0, frac_gradient_psi_ft=0.75,
               sidpp_psi=500.0, sicp_psi=700.0, pit_gain_bbl=20.0, scr1_psi=800.0,
               scr1_spm=30.0, scr2_psi=850.0, scr2_spm=60.0, pump_output_bbl_stk=0.1,
               method="Driller", well_type="development", pipes_m=[])
    raw.update(overrides)
    return compute_kill_sheet(build_canonical_kill_sheet_inputs(**raw))


@pytest.mark.parametrize("absent", ["sidpp_psi", "scr1_psi", "pit_gain_bbl",
                                    "pump_output_bbl_stk", "md_m", "sicp_psi"])
def test_kill_sheet_refuses_to_compute_from_a_missing_kick_input(absent):
    """Absence used to be silently coerced to 0.0 and still returned success."""
    result = _kill(**{absent: None})
    assert result.success is False
    assert "INPUT_INVALID" in result.error
    assert absent in result.error


def test_kill_sheet_still_computes_when_the_zero_is_recorded():
    complete = _kill()
    assert complete.success is True
    zero_sidpp = _kill(sidpp_psi=0.0)
    assert zero_sidpp.success is True
    # a recorded zero SIDPP is a real (smaller) overpressure, not a missing one
    assert zero_sidpp.kill_mw_ppg != complete.kill_mw_ppg


# --------------------------------------------------------------------------
# S20b: PDF/Excel exports keep absence distinct from measured zero
# --------------------------------------------------------------------------
def test_ddr_pdf_html_does_not_invent_zeros_for_missing_values():
    from core.ddr_pdf_export import DDRPDFExporter

    html = DDRPDFExporter(None)._build_html(
        {"report_date": "2026-01-01", "depth_0000": None, "depth_2400": 0.0},
        {"name": "W"}, [], {"avg_rop": None, "wob_min": None, "wob_max": 10},
        {"mw": None, "pv": 22.0},
    )
    assert "1234" not in html
    assert "—" in html                       # the unknown marker is present
    assert "0.0 m" in html                   # a recorded zero is still printed
    assert "None" not in html


def test_ddr_excel_leaves_unrecorded_measurements_empty(tmp_path):
    from openpyxl import load_workbook

    from core.report_engine import DDRReportEngine

    db = _mgr()
    project = _base(db)
    well = _well(db, project, "ExcelWell")
    bore = _wellbore(db, well, "Original")
    section = _section(db, well, "12.25in", wellbore_id=bore)
    report_id = _report(db, well, 1, wellbore_id=bore, section_id=section, depth=2500)
    target = tmp_path / "ddr.xlsx"
    assert DDRReportEngine(db).generate(report_id, str(target), format="excel") is True
    rows = {row[0]: row[1] for row in load_workbook(target)["DDR Header"].iter_rows(values_only=True)}
    assert float(rows["Depth 24:00"]) == 2500    # recorded value survives
    assert rows["Depth 00:00"] in (None, "")     # unrecorded stays empty, not 0
    assert rows["Depth 00:00"] != 0
    db.close()


# --------------------------------------------------------------------------
# S23: equipment hours — an unknown duration is not a measured zero
# --------------------------------------------------------------------------
def test_imported_solid_control_hours_absent_or_malformed_stay_unknown(env):
    db, ids = env
    result = db.save_imported_multi_tab_data_atomic(ids["a"], ids["ra"], {"solid_control": [
        {"equipment": "Shaker #1"},
        {"equipment": "Shaker #2", "daily_hrs": "bad"},
        {"equipment": "Shaker #3", "daily_hrs": 12},
    ]})
    assert result["status"] != "FAILED"
    from core.database import EquipmentLog

    session = db.create_session()
    try:
        rows = {row.equipment_name: row.hours_worked
                for row in session.query(EquipmentLog).filter_by(report_id=ids["ra"]).all()}
    finally:
        session.close()
    assert rows == {"Shaker #1": None, "Shaker #2": None, "Shaker #3": 12.0}
    assert result["review"] >= 1                   # the malformed cell is tracked, not silent


def test_equipment_summary_treats_unknown_hours_as_unknown(env):
    db, ids = env
    assert db.save_equipment_records(ids["a"], ids["ra"], "Pump", [
        {"equipment_name": "A", "hours_worked": 5},
        {"equipment_name": "B"},
    ])
    summary = db.get_equipment_summary(report_id=ids["ra"])
    assert summary["hours_not_recorded"] == 1
    assert summary["total_hours"] is None         # 5 + unknown is not 5
    assert summary["by_type"]["Pump"]["total_hours"] is None


def test_save_equipment_log_without_hours_persists_null(env):
    db, ids = env
    identity = db.save_equipment_log({"well_id": ids["a"], "report_id": ids["ra"],
                                      "equipment_type": "Pump", "equipment_name": "C"})
    rows = {row["id"]: row for row in db.get_equipment_logs(report_id=ids["ra"])}
    assert rows[identity]["hours_worked"] is None


# --------------------------------------------------------------------------
# S24: persons-on-board — an unknown headcount is not a smaller safe total
# --------------------------------------------------------------------------
def test_unknown_headcount_never_produces_a_pob_total(env):
    db, ids = env
    db.save_service_company_pob({"well_id": ids["a"], "report_id": ids["ra"],
                                 "company_name": "Crew", "personnel_count": 5})
    db.save_service_company_pob({"well_id": ids["a"], "report_id": ids["ra"],
                                 "company_name": "Unknown Co"})
    rows = {row["company_name"]: row["personnel_count"]
            for row in db.get_service_company_pob(report_id=ids["ra"])}
    assert rows == {"Crew": 5, "Unknown Co": None}
    # not 5, not 0: an incomplete roster has no total
    assert db.calculate_total_pob(report_id=ids["ra"]) is None


def test_pob_total_is_exact_only_when_every_headcount_is_recorded(env):
    db, ids = env
    for name, count in (("A", 5), ("B", 2), ("C", 0)):
        db.save_service_company_pob({"well_id": ids["a"], "report_id": ids["ra"],
                                     "company_name": name, "personnel_count": count})
    assert db.calculate_total_pob(report_id=ids["ra"]) == 7   # recorded zero stays 0


def test_service_company_without_headcount_persists_null(env):
    db, ids = env
    identity = db.save_service_company({"well_id": ids["a"], "report_id": ids["ra"],
                                        "company_name": "Rigless"})
    rows = {row["id"]: row for row in db.get_service_companies(report_id=ids["ra"])}
    assert rows[identity]["personnel_count"] is None          # not silently 1


# --------------------------------------------------------------------------
# S25: aggregates/derived series — an unrecorded quantity is not a zero
# --------------------------------------------------------------------------
def _orm(db, model, **fields):
    session = db.create_session()
    try:
        row = model(**fields)
        session.add(row)
        session.flush()
        row_id = row.id
        session.commit()
        return row_id
    finally:
        session.close()


def test_stock_and_request_totals_are_unknown_when_a_quantity_is_unrecorded(env):
    from core.database import BulkMaterials

    db, ids = env
    complete = _orm(db, BulkMaterials, well_id=ids["a"], report_id=ids["ra"],
                    report_date=date(2026, 1, 1),
                    material_name="Diesel", initial_stock=100.0, received=20.0, used=5.0)
    totals = db.calculate_bulk_totals(well_id=ids["a"])
    assert (totals["total_initial_stock"], totals["total_received"], totals["total_used"]) == (100.0, 20.0, 5.0)

    _orm(db, BulkMaterials, well_id=ids["a"], report_id=ids["ra"],
         report_date=date(2026, 1, 1), material_name="Water")
    totals = db.calculate_bulk_totals(well_id=ids["a"])
    # The documented trichotomy: an unreported opening stock is UNKNOWN, so the
    # total is unknown (100 + missing is not 100). received/used follow the
    # documented "absent = no movement (0.0)" daily-report convention instead.
    assert totals["total_initial_stock"] is None
    assert totals["total_received"] == 20.0 and totals["total_used"] == 5.0
    assert totals["material_count"] == 2

    assert complete            # the complete row keeps its own values
    assert db.calculate_bulk_totals(report_id=ids["ra"])["total_initial_stock"] is None


def test_material_request_quantities_are_null_when_not_supplied(env):
    db, ids = env
    identity = db.save_material_request({"well_id": ids["a"], "report_id": ids["ra"],
                                        "request_date": date(2026, 1, 1),
                                        "requested_items": "Needle valve"})
    rows = {row["id"]: row for row in db.get_material_requests(report_id=ids["ra"])}
    assert rows[identity]["requested_quantity"] is None      # not 0.0
    balance = db.calculate_material_balance(report_id=ids["ra"])
    assert balance["total_requested"] is None
    assert balance["balance"] is None                        # no partial arithmetic


def test_transport_passenger_counts_are_null_when_not_supplied(env):
    from core.database import TransportLog

    db, ids = env
    identity = _orm(db, TransportLog, well_id=ids["a"], report_id=ids["ra"],
                    log_date=date(2026, 1, 1), vehicle_type="Crew boat", vehicle_name="MV-1")
    rows = {row["id"]: row for row in db.get_transport_logs(report_id=ids["ra"])}
    assert rows[identity]["passengers_in"] is None and rows[identity]["passengers_out"] is None


def test_chart_data_skips_unrecorded_points_instead_of_plotting_zeros(env):
    from core.database import DailyReport

    db, ids = env
    _orm(db, DailyReport, well_id=ids["a"], report_date=date(2026, 1, 1),
         depth_2400=1000.0, rop_meter=10.0)
    _orm(db, DailyReport, well_id=ids["a"], report_date=date(2026, 1, 2),
         depth_2400=None, rop_meter=12.0)
    _orm(db, DailyReport, well_id=ids["a"], report_date=date(2026, 1, 3),
         depth_2400=1200.0, rop_meter=None)
    chart = db.generate_time_depth_chart_data(ids["a"])
    assert chart["data_points"] == 1                       # only the fully known day
    assert chart["depths"] == [1000.0] and chart["rop"] == [10.0]


# --------------------------------------------------------------------------
# S26: dialogs, exports and derived aggregates must not invent measurements
# --------------------------------------------------------------------------
def _qt_app():
    """The process-wide application, which must be able to host widgets.

    A bare ``QCoreApplication`` left alive by another test cannot be replaced,
    and building a widget on it aborts the interpreter, so fail loudly here
    instead of crashing the whole session.
    """
    from PySide6.QtWidgets import QApplication

    app = QApplication.instance()
    if app is None:
        return QApplication([])
    assert isinstance(app, QApplication), (
        f"the session is running on a bare {type(app).__name__}; widget tests need a QApplication")
    return app


def test_bit_record_dialog_keeps_unrecorded_measurements_blank():
    """A bit record loaded without depths/hours must not become measured zeros."""
    _qt_app()  # a QApplication must exist before widgets are built
    from dialogs.drilling_report_dialogs import AddBitRecordDialog

    dialog = AddBitRecordDialog(None, edit_data={"Bit No": "1"})
    dialog._save()
    result = dialog.result
    for field in ("Depth In (m)", "Depth Out (m)", "Hours", "ROP (m/hr)",
                  "Size (in)", "WOB Min (klb)", "Rot. Max", "MW (pcf)", "TFA (in²)"):
        assert result[field] == "", f"{field} fabricated {result[field]!r} for an absent measurement"
    assert result["Metres Drilled"] == ""

    # Recorded values (including a genuine zero) still survive verbatim.
    dialog2 = AddBitRecordDialog(None, edit_data={
        "Bit No": "2", "Depth In (m)": 1000, "Depth Out (m)": 1100, "Hours": 10,
        "Size (in)": 8.5, "WOB Min (klb)": 0})
    dialog2._save()
    assert dialog2.result["Depth In (m)"] == "1000.0"
    assert dialog2.result["Metres Drilled"] == "100.0"
    assert dialog2.result["ROP (m/hr)"] == "10.0"
    assert dialog2.result["WOB Min (klb)"] == "0.0"      # explicit zero preserved


def test_bha_component_dialog_keeps_unrecorded_dimensions_blank():
    _qt_app()
    from dialogs.drilling_report_dialogs import AddBHAComponentDialog

    dialog = AddBHAComponentDialog(None)
    dialog._save()
    for field in ("OD (in)", "ID (in)", "Length (m)", "Weight (kg)", "Make-up Torque (ft-lb)"):
        assert dialog.result[field] == "", f"{field} fabricated {dialog.result[field]!r}"
    dialog2 = AddBHAComponentDialog(None)
    dialog2.od.setValue(6.75)
    dialog2.length.setValue(9.5)
    dialog2._save()
    assert dialog2.result["OD (in)"] == "6.75" and dialog2.result["Length (m)"] == "9.5"


def test_professional_export_leaves_unrecorded_duration_blank(env):
    """The exported spreadsheet states 0 h only when 0 h was recorded."""
    from core.database import TimeLog24H
    from core.professional_export import ProfessionalExcelExport
    from openpyxl import load_workbook
    import os
    import tempfile

    db, ids = env
    session = db.create_session()
    try:
        session.add(TimeLog24H(report_id=ids["ra"], time_from=time(0, 0),
                               time_to=time(6, 0), duration=None, main_code="DRL",
                               activity_description="no duration recorded"))
        session.add(TimeLog24H(report_id=ids["ra"], time_from=time(6, 0),
                               time_to=time(12, 0), duration=0.0, main_code="DRL",
                               activity_description="recorded zero"))
        session.commit()
    finally:
        session.close()
    with tempfile.TemporaryDirectory() as directory:
        target = os.path.join(directory, "export.xlsx")
        assert ProfessionalExcelExport(db).export(ids["a"], target, report_id=ids["ra"]) is True
        sheet = load_workbook(target)["Time Logs"]
        durations = [sheet.cell(row=row, column=3).value for row in (2, 3)]
    # Openpyxl stores an empty cell as None; the recorded 0.0 must stay 0.0.
    assert durations[0] is None, f"unrecorded duration exported as {durations[0]!r}"
    assert durations[1] == 0.0


def test_activity_code_usage_hours_are_unknown_when_a_duration_is_missing(env):
    """Derived activity totals must not present a subtotal as the total."""
    from core.database import ActivityCode, TimeLog24H

    db, ids = env
    session = db.create_session()
    try:
        session.add(ActivityCode(well_id=ids["a"], main_phase="drilling", main_code="DRL",
                                 sub_code="DRL-1", code_name="drilling"))
        session.add(ActivityCode(well_id=ids["a"], main_phase="drilling", main_code="DRL",
                                 sub_code="DRL-2", code_name="drilling"))
        session.commit()
        session.add(TimeLog24H(report_id=ids["ra"], time_from=time(0, 0),
                               time_to=time(6, 0), duration=6.0, main_code="DRL-1",
                               activity_description="recorded"))
        session.add(TimeLog24H(report_id=ids["ra"], time_from=time(6, 0),
                               time_to=time(12, 0), duration=None, main_code="DRL-2",
                               activity_description="unrecorded"))
        session.commit()
    finally:
        session.close()
    db.auto_update_from_daily_report(ids["ra"])
    session = db.create_session()
    try:
        hours = {row.sub_code: row.total_hours for row in session.query(ActivityCode).all()}
    finally:
        session.close()
    assert hours["DRL-1"] == 6.0
    assert hours["DRL-2"] is None, "an unknown duration produced a numeric hours total"


def test_full_hierarchy_orders_unknown_section_depth_last(env):
    from core.database import Section

    from sqlalchemy import text

    db, ids = env
    session = db.create_session()
    try:
        session.add(Section(well_id=ids["a"], name="unknown", depth_from=None))
        session.add(Section(well_id=ids["a"], name="deep", depth_from=2000.0))
        session.commit()
        # The ORM default writes 0.0 for an unset depth, so emulate the legacy row
        # this sorting rule exists for: a section whose depth is genuinely NULL.
        session.execute(text("UPDATE sections SET depth_from = NULL WHERE name = 'unknown'"))
        session.commit()
        stored = session.execute(
            text("SELECT name, depth_from FROM sections WHERE section_well_scope IS NULL"
                 ) if False else
            text("SELECT name, depth_from FROM sections WHERE well_id = :w AND name = 'unknown'"),
            {"w": ids["a"]}).fetchall()
        assert stored == [("unknown", None)], stored
    finally:
        session.close()
    hierarchy = db.get_full_hierarchy()
    well = next(w for company in hierarchy for project in company["projects"]
                for w in project["wells"] if w["id"] == ids["a"])
    unassigned = [row["name"] for row in well["unassigned_sections"]]
    assert unassigned == ["deep", "unknown"], unassigned


def test_legacy_time_log_validator_reports_unrecorded_duration():
    from core.validators import TimeLogValidator

    report = TimeLogValidator.validate_logs([
        {"duration": None, "main_code": "DRL"},
        {"duration": 12.0, "main_code": "DRL"},
    ])
    messages = [issue["message"] for issue in report.warnings]
    assert any("no recorded duration" in message for message in messages), messages
    assert not any("expected ~24h" in message for message in messages), \
        "an incomplete total must not be validated as if the missing hours were zero"


def test_derived_npt_skips_logs_without_a_recorded_duration(env):
    """An NPT row cannot claim a duration the operator never recorded."""
    from core.database import NPTReport, TimeLog24H

    db, ids = env
    session = db.create_session()
    try:
        session.add(TimeLog24H(report_id=ids["ra"], time_from=time(0, 0), time_to=time(3, 0),
                               duration=3.0, main_code="NPT", is_npt=True,
                               activity_description="recorded npt"))
        session.add(TimeLog24H(report_id=ids["ra"], time_from=time(3, 0), time_to=time(6, 0),
                               duration=None, main_code="NPT", is_npt=True,
                               activity_description="unrecorded npt"))
        session.commit()
    finally:
        session.close()
    assert db.auto_update_from_daily_report(ids["ra"]) is not False
    session = db.create_session()
    try:
        npt = session.query(NPTReport).filter(NPTReport.report_id == ids["ra"]).all()
    finally:
        session.close()
    assert len(npt) == 1, f"a fabricated 0 h NPT row was derived ({len(npt)} rows)"
    assert npt[0].duration_hours == 3.0


def test_failed_engine_calculations_are_not_reported_as_zero():
    """A failed/absent engineering calculation is unknown, not a measured 0."""
    from core.managers import DrillingManager

    assert DrillingManager.calculate_tfa([]) is None
    assert DrillingManager.calculate_rop(1000.0, 1200.0, 0.0) in (None, 0.0)  # no hours -> engine decides
    bad = DrillingManager.calculate_rop("not-a-number", 1200.0, 10.0)
    assert bad is None, f"a failed ROP calculation returned {bad!r}"
    assert DrillingManager.calculate_hsi("x", "y", "z") is None


def test_drilling_tab_persists_unknown_for_uncomputed_values():
    """The real drilling-report payload must not turn an uncomputed value into 0.0."""
    _qt_app()
    from tabs.w3_drilling_report import (_calc_spin, _calc_value, _set_calc,
                                    DrillingParametersTab)

    tab = DrillingParametersTab()
    _calc_value_keys = {"tfa", "avg_rop", "hsi"}
    # Nothing has been computed yet: the tab reports unknown, not measured zero.
    tab.calculate_tfa()
    tab.calculate_rop()
    tab.calculate_hsi()
    payload = tab.collect_data()
    assert payload["tfa"] is None, f"uncomputed TFA persisted as {payload['tfa']!r}"
    assert payload["avg_rop"] is None, f"uncomputed ROP persisted as {payload['avg_rop']!r}"
    assert payload["hsi"] is None, f"uncomputed HSI persisted as {payload['hsi']!r}"
    assert all(key in payload for key in _calc_value_keys)

    # A real engine result is stored verbatim, including a genuine zero.
    tab.nozzle_table.setRowCount(1)
    from PySide6.QtWidgets import QDoubleSpinBox, QSpinBox
    size = QDoubleSpinBox()
    size.setValue(16)
    qty = QSpinBox()
    qty.setValue(3)
    tab.nozzle_table.setCellWidget(0, 1, size)
    tab.nozzle_table.setCellWidget(0, 2, qty)
    tab.calculate_tfa()
    payload = tab.collect_data()
    assert payload["tfa"] and payload["tfa"] > 0

    # A genuine 0.0 result stays a measured zero, not "not computed". The engines
    # return no value at all for an invalid nozzle (shown as unknown), so the
    # zero-vs-unknown contract is asserted on the display helper itself.
    from PySide6.QtWidgets import QDoubleSpinBox
    probe = _calc_spin(QDoubleSpinBox())
    _set_calc(probe, 0.0)
    assert payload_is_zero(_calc_value(probe)), "a real 0.0 must not read back as unknown"

    # Loading a stored report that recorded no average ROP keeps it unknown.
    tab.load_from_dict({"avg_rop": None, "hsi": None, "tfa": None})
    loaded = tab.collect_data()
    assert loaded["avg_rop"] is None, f"a stored NULL reloaded as {loaded['avg_rop']!r}"
    assert loaded["hsi"] is None and loaded["tfa"] is None

    tab.clear_form()
    assert tab.collect_data()["avg_rop"] is None


def payload_is_zero(value):
    return value == 0.0 and value is not None


def test_engine_result_without_a_rop_value_is_unknown(monkeypatch):
    """A successful run that yields no ROP must not be reported as 0 m/hr."""
    from types import SimpleNamespace
    from core.managers import DrillingManager
    from core.engineering.engines import bit_performance

    monkeypatch.setattr(bit_performance.BitPerformanceEngine, "from_run",
                        staticmethod(lambda **kwargs: SimpleNamespace(success=True, values={})))
    assert DrillingManager.calculate_rop(1000.0, 1100.0, 10.0) is None


# --------------------------------------------------------------------------
# S27: plan / actual export matrix must not supply numbers the planner omitted
# --------------------------------------------------------------------------
def _plan_env(db, well_id):
    from datetime import date as _date
    from core.database import PlannedActivity, WellPlan

    with db.session_scope() as session:
        plan = WellPlan(well_id=well_id, plan_name="P", plan_version=1, is_active=True)
        session.add(plan)
        session.flush()
        plan_id = plan.id
        session.add(PlannedActivity(well_id=well_id, plan_id=plan_id, phase_code="DRL",
                                    activity_name="full", planned_depth_from=1000.0,
                                    planned_depth_to=1200.0, planned_duration_hours=10.0,
                                    planned_start=_date(2026, 1, 1), planned_end=_date(2026, 1, 2)))
        session.add(PlannedActivity(well_id=well_id, plan_id=plan_id, phase_code="DRL",
                                    activity_name="partial", planned_depth_from=None,
                                    planned_depth_to=None, planned_duration_hours=None,
                                    planned_start=_date(2026, 1, 3), planned_end=_date(2026, 1, 4)))
    return plan_id


def test_plan_excel_export_leaves_unplanned_fields_blank(env):
    """A plan row with no planned depth/duration must export empty, not 0 m / 0 hrs."""
    from core.report_engine import PlanReportEngine
    from openpyxl import load_workbook
    import os
    import tempfile

    db, ids = env
    _plan_env(db, ids["a"])
    with tempfile.TemporaryDirectory() as directory:
        target = os.path.join(directory, "plan.xlsx")
        assert PlanReportEngine(db).generate(ids["a"], target, format="excel") is True
        sheet = load_workbook(target)["Activities"]
        rows = {sheet.cell(row=r, column=2).value: r for r in (2, 3)}
        blank = rows["partial"]
        assert sheet.cell(row=blank, column=3).value is None, \
            f"unplanned depth_from exported as {sheet.cell(row=blank, column=3).value!r}"
        assert sheet.cell(row=blank, column=4).value is None
        assert sheet.cell(row=blank, column=5).value is None, \
            f"unplanned duration exported as {sheet.cell(row=blank, column=5).value!r}"
        complete = rows["full"]
        assert sheet.cell(row=complete, column=3).value == 1000.0
        assert sheet.cell(row=complete, column=5).value == 10.0


def test_plan_excel_export_keeps_a_planned_zero(env):
    """A depth/duration the planner recorded as 0 stays 0 in the export."""
    from core.database import PlannedActivity
    from core.report_engine import PlanReportEngine
    from openpyxl import load_workbook
    import os
    import tempfile

    db, ids = env
    plan_id = _plan_env(db, ids["a"])
    with db.session_scope() as session:
        session.add(PlannedActivity(well_id=ids["a"], plan_id=plan_id, phase_code="DRL",
                                    activity_name="zeroed", planned_depth_from=0.0,
                                    planned_depth_to=0.0, planned_duration_hours=0.0,
                                    planned_start=date(2026, 1, 5), planned_end=date(2026, 1, 6)))
    with tempfile.TemporaryDirectory() as directory:
        target = os.path.join(directory, "plan.xlsx")
        assert PlanReportEngine(db).generate(ids["a"], target, format="excel") is True
        sheet = load_workbook(target)["Activities"]
        row = {sheet.cell(row=r, column=2).value: r for r in range(2, 5)}["zeroed"]
        assert sheet.cell(row=row, column=3).value == 0.0
        assert sheet.cell(row=row, column=5).value == 0.0


# --------------------------------------------------------------------------
# S28: ownership invariant guards must be covered by regression tests
# (mutation controls O-OWNERSHIP-TYPE / O-OWNERSHIP-CROSS survived until these)
# --------------------------------------------------------------------------
def _wellbore_raise(db, **kwargs):
    """Create a wellbore through the guarded write path and return the error."""
    from core.database import Wellbore

    with pytest.raises(Exception) as caught:
        with db.session_scope() as session:
            session.add(Wellbore(**kwargs))
            session.flush()
    return caught.value


def test_invalid_wellbore_type_is_rejected(env):
    from core.database import OwnershipIntegrityError

    db, ids = env
    error = _wellbore_raise(db, well_id=ids["a"], name="bogus", wellbore_type="horizontal")
    assert isinstance(error, OwnershipIntegrityError), f"got {type(error).__name__}: {error}"
    assert "wellbore_type" in str(error)


def test_cross_well_sidetrack_lineage_is_rejected(env):
    from core.database import OwnershipIntegrityError

    db, ids = env
    # ids["ba"] belongs to well A; the sidetrack is created under well B.
    error = _wellbore_raise(db, well_id=ids["b"], name="crossed", wellbore_type="sidetrack",
                            parent_wellbore_id=ids["ba"])
    assert isinstance(error, OwnershipIntegrityError), f"got {type(error).__name__}: {error}"
    assert "crosses wells" in str(error)


def test_non_sidetrack_carrying_lineage_is_rejected(env):
    from core.database import OwnershipIntegrityError

    db, ids = env
    error = _wellbore_raise(db, well_id=ids["a"], name="lineaged", wellbore_type="original",
                            parent_wellbore_id=ids["ba"])
    assert isinstance(error, OwnershipIntegrityError), f"got {type(error).__name__}: {error}"
    assert "sidetrack" in str(error)


def test_valid_sidetrack_lineage_is_accepted(env):
    """The guards must reject only real violations, not the legitimate case."""
    from test_scope_attribution import _wellbore

    db, ids = env
    sidetrack = _wellbore(db, ids["a"], "legit-st", wtype="sidetrack", parent=ids["ba"])
    assert isinstance(sidetrack, int) and sidetrack > 0


# --------------------------------------------------------------------------
# S29: data-quality metrics must not score unrecorded data as zero
# (M33 Part L / E4 - core/data_quality.py)
# --------------------------------------------------------------------------
def _quality_env(db, well_id, report_id, durations):
    from core.database import TimeLog24H

    with db.session_scope() as session:
        for idx, duration in enumerate(durations):
            session.add(TimeLog24H(report_id=report_id, time_from=time(idx, 0), time_to=time(idx + 1, 0),
                                   duration=duration, main_code="DRL", activity_description="q"))


def _coverage_metric(service, report_id):
    return next(m for m in service.for_report(report_id) if m.name == "24h time coverage")


def test_time_coverage_is_unknown_when_a_duration_is_missing(env):
    from core.data_quality import DataQualityService

    db, ids = env
    _quality_env(db, ids["a"], ids["ra"], [6.0, None])
    service = DataQualityService(db)
    metric = _coverage_metric(service, ids["ra"])
    assert metric.value is None, f"unrecorded hours were scored as {metric.value}"
    assert metric.status == "unknown" and metric.confidence == 0.0
    assert metric.evidence["unrecorded_entries"] == 1
    assert metric.evidence["total_hours"] is None


def test_time_coverage_is_computed_when_every_duration_is_recorded(env):
    from core.data_quality import DataQualityService

    db, ids = env
    _quality_env(db, ids["a"], ids["ra"], [12.0, 12.0])
    service = DataQualityService(db)
    metric = _coverage_metric(service, ids["ra"])
    assert metric.value == 100.0 and metric.status == "good"
    assert metric.evidence["total_hours"] == 24.0


def test_recorded_zero_duration_is_not_treated_as_missing(env):
    from core.data_quality import DataQualityService

    db, ids = env
    _quality_env(db, ids["a"], ids["ra"], [0.0, 12.0])
    service = DataQualityService(db)
    metric = _coverage_metric(service, ids["ra"])
    assert metric.value == 50.0, "an explicitly recorded 0 h must stay a measured zero"


def test_summary_excludes_unknown_metrics_and_reports_confidence(env):
    from core.data_quality import DataQualityService

    db, ids = env
    _quality_env(db, ids["a"], ids["ra"], [6.0, None])
    summary = DataQualityService(db).summary(ids["ra"])
    assert "24h time coverage" in summary["unknown_metrics"]
    assert summary["confidence"] < 1.0
    assert summary["evidence"]["unknown_metric_count"] >= 1
    known = [m for m in summary["metrics"] if m["value"] is not None]
    assert summary["score"] == round(sum(m["value"] for m in known) / len(known), 1)


def test_derived_avg_rop_display_is_not_silently_clamped():
    """A real computed ROP must survive the derived-value widget.

    Regression for a Qt default: an unset ``QDoubleSpinBox`` maximum is 99.99,
    so a legitimate 150 m/hr ROP was silently truncated in the read-only display
    AND in the payload read from it (and a stored 150 came back truncated on
    reload). The documented domain range for ``avg_rop`` is 0-500 m/hr
    (``core/validators.py``), so 150 is inside the contract.
    """
    _qt_app()
    from core.managers import DrillingManager
    from tabs.w3_drilling_report import DrillingParametersTab

    expected = DrillingManager.calculate_rop(1000.0, 1150.0, 1.0)
    assert expected is not None and expected > 99.99, (
        f"the scenario must exceed the Qt default maximum (engine returned {expected!r})"
    )

    tab = DrillingParametersTab()
    tab.depth_in.setValue(1000.0)
    tab.depth_out.setValue(1150.0)
    tab.hours_on_bottom.setValue(1.0)
    tab.calculate_rop()
    assert tab.avg_rop.value() == pytest.approx(expected), (
        f"derived ROP was truncated to {tab.avg_rop.value()!r} (engine returned {expected!r})"
    )
    assert tab.collect_data()["avg_rop"] == pytest.approx(expected), (
        "the payload must carry the engine's number, not the widget's clamped one"
    )

    # A reloaded stored value must not be truncated on the way back in either.
    reloaded = DrillingParametersTab()
    reloaded.load_from_dict({"avg_rop": expected, "hsi": None, "tfa": None})
    assert reloaded.collect_data()["avg_rop"] == pytest.approx(expected)

    # Beyond the documented range the validator owns plausibility (it warns
    # "verify unit"); the display must still carry the engine's number instead of
    # truncating it to the bound.
    far = DrillingManager.calculate_rop(100.0, 1000.0, 1.0)
    assert far is not None and far > 500.0
    beyond = DrillingParametersTab()
    beyond.depth_in.setValue(100.0)
    beyond.depth_out.setValue(1000.0)
    beyond.hours_on_bottom.setValue(1.0)
    beyond.calculate_rop()
    assert beyond.avg_rop.value() == pytest.approx(far)
    assert beyond.collect_data()["avg_rop"] == pytest.approx(far)

    # The validator is what flags that out-of-range value for the user.
    from core.validators import DrillingParamsValidator
    warnings = DrillingParamsValidator.validate(beyond.collect_data()).warnings
    assert any("avg_rop" in str(w.get("field")) for w in warnings), warnings


def test_calc_spin_domain_maximum_is_explicit():
    """``_calc_spin`` must take the field's documented bound, not Qt's 99.99."""
    _qt_app()
    from PySide6.QtWidgets import QDoubleSpinBox
    from tabs.w3_drilling_report import _calc_spin

    plain = _calc_spin(QDoubleSpinBox())
    assert plain.maximum() == 99.99, "Qt's default is still the fallback for unbounded fields"
    bounded = _calc_spin(QDoubleSpinBox(), 500)
    assert bounded.maximum() == 500, "avg_rop's domain bound must reach the widget"
    assert bounded.minimum() == -1 and bounded.value() == -1, "sentinel contract unchanged"

    # ...and the real tab must configure that bound at the call site, so the
    # field's declared range is the documented domain range and not Qt's 99.99.
    from tabs.w3_drilling_report import DrillingParametersTab

    tab = DrillingParametersTab()
    assert tab.avg_rop.maximum() == 500, (
        f"the avg_rop display declares the Qt default {tab.avg_rop.maximum()} instead of "
        "the validator's 0-500 domain range"
    )
    assert tab.avg_rop.isReadOnly(), "derived displays stay read-only"

