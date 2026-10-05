"""Focused M38 regressions for units, unknown data, and export scope."""
from __future__ import annotations

from datetime import date

import pytest


@pytest.fixture
def m38_db(tmp_path):
    from core.database import (
        Company, CostRecord, DailyReport, DatabaseManager, Project, Section, Well,
    )

    db = DatabaseManager()
    db.db_path = str(tmp_path / "m38.sqlite")
    assert db.initialize(), db.last_diagnostic
    with db.session_scope() as session:
        company = Company(name="M38 Company", code="M38")
        session.add(company)
        session.flush()
        project = Project(company_id=company.id, name="M38 Project", code="M38-P")
        session.add(project)
        session.flush()
        well = Well(project_id=project.id, name="M38 Well", code="M38-W")
        session.add(well)
        session.flush()
        section = Section(well_id=well.id, name="8-1/2 in")
        session.add(section)
        session.flush()
        first = DailyReport(
            well_id=well.id, section_id=section.id, report_date=date(2026, 10, 1),
            report_number=1, depth_2400=1000.0, mw_pcf=70.0,
        )
        second = DailyReport(
            well_id=well.id, section_id=section.id, report_date=date(2026, 10, 2),
            report_number=2, depth_2400=1100.0, mw_pcf=72.0,
        )
        session.add_all([first, second])
        session.flush()
        ids = {"well": well.id, "section": section.id, "first": first.id, "second": second.id}
        session.add_all([
            CostRecord(well_id=well.id, category="Mud", description="day one", cost_date=date(2026, 10, 1), actual_cost=10),
            CostRecord(well_id=well.id, category="Fuel", description="day two", cost_date=date(2026, 10, 2), actual_cost=20),
        ])
    yield db, ids
    db.close()


def test_canonical_density_is_pcf_and_header_measurement_is_separate():
    from core.canonical_schema import FIELD_SPECS
    from core.database import CostRecord, DailyReport, MudReport

    assert FIELD_SPECS["mud_report.mw"].unit == "pcf"
    assert FIELD_SPECS["daily_report.mw_pcf"].unit == "pcf"
    assert FIELD_SPECS["daily_report.mw_pcf"].max_val == pytest.approx(187.01298639122953)
    assert "mw_pcf" in DailyReport.__table__.columns
    assert "mw" in MudReport.__table__.columns
    assert "report_id" not in CostRecord.__table__.columns


def test_v3_database_gets_additive_mw_pcf_column_without_backfill(m38_db):
    import sqlite3
    from core.database import DatabaseManager

    db, ids = m38_db
    from core.database import MudReport
    with db.session_scope() as session:
        session.add(MudReport(
            well_id=ids["well"], report_id=ids["first"],
            report_date=date(2026, 10, 1), mw=71.0,
        ))
    path = db.db_path
    db.close()
    with sqlite3.connect(path) as connection:
        connection.execute("PRAGMA foreign_keys=OFF")
        connection.execute("ALTER TABLE daily_reports DROP COLUMN mw_pcf")
        connection.execute("UPDATE schema_version SET version=3")
        connection.commit()

    upgraded = DatabaseManager()
    upgraded.db_path = path
    assert upgraded.initialize(), upgraded.last_diagnostic
    assert upgraded.schema_version == 4
    report = upgraded.get_daily_report_by_id(ids["first"])
    assert report["depth_2400"] == 1000 and report["mw_pcf"] is None
    assert upgraded.get_mud_report(report_id=ids["first"])["mw"] == 71.0
    upgraded.close()

    reopened = DatabaseManager()
    reopened.db_path = path
    assert reopened.initialize(), reopened.last_diagnostic
    assert reopened.schema_version == 4
    assert reopened.get_daily_report_by_id(ids["first"])["mw_pcf"] is None
    assert reopened.get_mud_report(report_id=ids["first"])["mw"] == 71.0
    reopened.close()


def test_mud_density_normalization_requires_explicit_unit_and_preserves_lineage():
    from core.mud_records import mud_density_pcf
    from core.unit_manager import UnitManager

    for value, unit in [(71, "pcf"), (9.5, "ppg"), (1.2, "sg"), (1000, "kg/m3")]:
        normalized, lineage = mud_density_pcf({"mw": value, "mw_unit": unit})
        expected = UnitManager.convert(value, "density", unit, "pcf")
        assert normalized["mw"] == pytest.approx(expected)
        assert normalized["mw_unit"] == "PCF"
        assert lineage["source_unit"] == unit
        assert lineage["normalized_unit"] == "pcf"
        assert lineage["conversion_rule"]

    with pytest.raises(ValueError, match="explicit source unit"):
        mud_density_pcf({"mw": 71})
    equivalent, _ = mud_density_pcf({"mw": "9 ppg", "mw_unit": "lb/gal"})
    assert equivalent["mw"] == pytest.approx(UnitManager.convert(9, "density", "ppg", "pcf"))
    with pytest.raises(ValueError, match="Conflicting"):
        mud_density_pcf({"mw": "71 pcf", "mw_unit": "ppg"})
    with pytest.raises(ValueError, match="outside canonical bounds"):
        mud_density_pcf({"mw": 30, "mw_unit": "ppg"})
    with pytest.raises(ValueError, match="greater than zero"):
        mud_density_pcf({"mw": 0, "mw_unit": "pcf"})


def test_mineru_density_uses_same_pcf_boundary_and_unknown_is_reviewable():
    from core.mineru_engine import DocumentNormalizer

    normalizer = DocumentNormalizer()
    sg = normalizer._normalize_document_value("1.20 SG", "mud_report.mw", "Mud Weight")
    pcf = normalizer._normalize_document_value(71, "mud_report.mw", "Mud Weight (PCF)")
    bare = normalizer._normalize_document_value(71, "mud_report.mw", "Mud Weight")
    assert sg.ok and sg.value == pytest.approx(1.2 * 8.34540445 * (1 / 0.133680556))
    assert sg.source_unit == "sg" and sg.unit == "pcf" and sg.conversion_rule
    assert pcf.ok and pcf.value == 71 and pcf.source_unit == "pcf"
    assert bare.needs_review and bare.value is None
    zero = normalizer._normalize_document_value(0, "mud_report.mw", "Mud Weight (PCF)")
    assert zero.needs_review and zero.value is None
    assert "greater than zero" in zero.review_reason


def test_data_quality_distinguishes_missing_report_empty_logs_and_explicit_zero(m38_db):
    from core.data_quality import DataQualityService
    from core.database import MudReport

    db, ids = m38_db
    service = DataQualityService(db)
    missing = service.summary(999999)
    assert missing["score"] is None and missing["status"] == "unknown"
    assert all(name in missing["unknown_metrics"] for name in (
        "Report completeness (required)", "24h time coverage", "Unit consistency",
    ))

    empty_logs = next(metric for metric in service.for_report(ids["first"])
                      if metric.name == "24h time coverage")
    assert empty_logs.value is None and empty_logs.status == "unknown"
    assert empty_logs.evidence["query_completed"] is True
    assert empty_logs.evidence["entries"] == 0

    with db.session_scope() as session:
        session.add(MudReport(
            well_id=ids["well"], report_id=ids["first"],
            report_date=date(2026, 10, 1), mw=0.0,
        ))
    unit_metric = next(metric for metric in service.for_report(ids["first"])
                       if metric.name == "Unit consistency")
    assert unit_metric.value == 80 and unit_metric.status == "warning"
    assert "mud_report.mw" in unit_metric.evidence["assessed_fields"]


def test_data_quality_report_query_failure_is_not_empty_success():
    from core.data_quality import DataQualityService

    class QueryFailure:
        @staticmethod
        def get_daily_report_by_id(_report_id):
            raise RuntimeError("database query unavailable")

    with pytest.raises(RuntimeError, match="database query unavailable"):
        DataQualityService(QueryFailure()).summary(123)


def test_mud_report_upsert_and_caller_owned_transaction_preserve_unknown_and_zero(m38_db):
    from sqlalchemy.exc import IntegrityError
    from core.database import MudReport

    db, ids = m38_db
    values = {
        "well_id": ids["well"], "report_id": ids["first"],
        "report_date": date(2026, 10, 1), "mw": None, "pv": 0.0,
    }
    record_id = db.save_mud_report(values)
    assert record_id
    loaded = db.get_mud_report(report_id=ids["first"])
    assert loaded["id"] == record_id and loaded["mw"] is None and loaded["pv"] == 0.0

    updated = {**values, "mw": 71.0, "pv": None}
    assert db.save_mud_report(updated) == record_id
    loaded = db.get_mud_report(report_id=ids["first"])
    assert loaded["id"] == record_id and loaded["mw"] == 71.0 and loaded["pv"] is None

    session = db.create_session()
    try:
        caller_update = {**updated, "mw": None, "pv": 0.0}
        assert db.save_mud_report(caller_update, session=session) == record_id
        session.flush()
        # Caller-owned session is not committed by the helper; rollback owns outcome.
        session.rollback()
    finally:
        session.close()
    loaded = db.get_mud_report(report_id=ids["first"])
    assert loaded["id"] == record_id and loaded["mw"] == 71.0 and loaded["pv"] is None

    failing = db.create_session()
    try:
        with pytest.raises(IntegrityError):
            db.save_mud_report({
                "well_id": ids["well"], "report_id": 999999,
                "report_date": date(2026, 10, 1), "mw": 65.0,
            }, session=failing)
        failing.rollback()
    finally:
        failing.close()
    with db.session_scope() as session:
        rows = session.query(MudReport).filter_by(report_id=ids["first"]).all()
        assert len(rows) == 1 and rows[0].id == record_id


def test_professional_export_scope_labels_costs_and_report_mw(m38_db, tmp_path):
    from openpyxl import load_workbook
    from core.professional_export import ProfessionalExcelExport, ProfessionalExportMetadata
    from core.database import MudReport
    from core.report_engine import DDRReportEngine, EOWRReportEngine

    db, ids = m38_db
    with db.session_scope() as session:
        session.add(MudReport(
            well_id=ids["well"], report_id=ids["first"],
            report_date=date(2026, 10, 1), mw=71.0,
        ))
    report_data = DDRReportEngine(db)._collect_data(ids["first"])
    report_html = DDRReportEngine(db)._build_html(report_data)
    assert "DDR Header MW (PCF)" in report_html and "70.0" in report_html
    assert "Mud Properties MW (PCF)" in report_html and "71.0" in report_html
    eowr_engine = EOWRReportEngine(db)
    eowr_data = eowr_engine._collect_data(ids["well"])
    eowr_html = eowr_engine._build_html(eowr_data)
    assert "DDR Header MW (PCF)" in eowr_html and "70.0" in eowr_html
    assert "Mud Sample MW (PCF)" in eowr_html and "71.0" in eowr_html
    eowr_path = tmp_path / "eowr.xlsx"
    assert eowr_engine._save_excel(eowr_data, str(eowr_path))
    eowr_workbook = load_workbook(eowr_path, data_only=True)
    eowr_daily_headers = [cell.value for cell in eowr_workbook["DailyReports"][1]]
    eowr_mud_headers = [cell.value for cell in eowr_workbook["MudReports"][1]]
    assert "DDR Header MW (PCF)" in eowr_daily_headers
    assert "Mud Sample MW (PCF)" in eowr_mud_headers
    metadata = ProfessionalExportMetadata.build(db, ids["well"], report_id=ids["first"])
    assert metadata["Section"] == "8-1/2 in"
    assert "Selected daily report" in metadata["Record Scope"]
    whole_well_metadata = ProfessionalExportMetadata.build(db, ids["well"])
    assert "Not assessed" in whole_well_metadata["Data Quality"]
    assert "Whole well" in whole_well_metadata["Record Scope"]
    with pytest.raises(ValueError, match="section"):
        ProfessionalExportMetadata.build(db, ids["well"], section_id=999, report_id=ids["first"])

    report_path = tmp_path / "report-scope.xlsx"
    assert ProfessionalExcelExport(db).export(ids["well"], str(report_path), report_id=ids["first"])
    workbook = load_workbook(report_path, data_only=True)
    daily = workbook["Daily Report"]
    daily_headers = [cell.value for cell in daily[1]]
    assert "mw_pcf (PCF)" in daily_headers
    assert daily.cell(row=2, column=daily_headers.index("mw_pcf (PCF)") + 1).value == 70
    costs = workbook["Cost"]
    headers = [cell.value for cell in costs[1]]
    assert "Export Scope" in headers
    scope_col = headers.index("Export Scope") + 1
    descriptions_col = headers.index("description") + 1
    assert costs.max_row == 3
    assert {costs.cell(row=r, column=descriptions_col).value for r in (2, 3)} == {"day one", "day two"}
    assert all("Whole well" in costs.cell(row=r, column=scope_col).value for r in (2, 3))
    raw_data = workbook["Raw Data"]
    raw_headers = [cell.value for cell in raw_data[2]]
    entity_col = raw_headers.index("Entity") + 1
    raw_scope_col = raw_headers.index("Scope") + 1
    raw_cost_rows = [r for r in range(2, raw_data.max_row + 1)
                     if raw_data.cell(row=r, column=entity_col).value == "CostRecord"]
    assert len(raw_cost_rows) == 2
    assert all("whole well" in raw_data.cell(row=r, column=raw_scope_col).value
               for r in raw_cost_rows)
    quality = workbook["Data Quality"]
    assert quality.cell(row=1, column=1).value == "Scope"
    assert quality.cell(row=2, column=1).value == "Report"

    well_path = tmp_path / "well-scope.xlsx"
    assert ProfessionalExcelExport(db).export(ids["well"], str(well_path))
    well_workbook = load_workbook(well_path, data_only=True)
    assert well_workbook["Cost"].max_row == 3
    assert well_workbook["Data Quality"].cell(row=2, column=1).value == "Well"
