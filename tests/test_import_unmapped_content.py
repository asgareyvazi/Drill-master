"""Unmapped populated content diagnostic (M43 closure, Phase 3).

Synthetic workbooks only. The diagnostic must name sheet and range of populated
cells that no extraction step used, never carry their values, never propose a
mapping, never change the import status, and be deterministic.
"""

from __future__ import annotations

import json
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from core.ddr_import_service import DDRImportService
from core.import_coverage import summary_line, unmapped_content_diagnostic
from core.import_ir import RawCell, RawDocument, SourceLocation

SECRET_TEXT = "SYNTHETIC-PRIVATE-NOTE-7781"
SECRET_NUMBER = 98765.4321


def _workbook(path, *, with_unknown_table=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["MD", "Inc", "Azi"])  # recognised survey signature (rows 1-4, cols A-C)
    for md, inc, azi in ((1000.0, 1.5, 120.0), (1100.0, 2.0, 121.0), (1200.0, 2.5, 122.0)):
        ws.append([md, inc, azi])
    if with_unknown_table:
        # A table whose headers match no canonical signature (F8:H11).
        ws["F8"], ws["G8"], ws["H8"] = "Zeta", "Omega", "Kappa"
        for offset, row in enumerate(((SECRET_NUMBER, 2, 3), (4, 5, 6), (7, 8, 9))):
            for col, value in zip("FGH", row):
                ws[f"{col}{9 + offset}"] = value
    ws["A20"] = SECRET_TEXT  # an isolated note
    wb.save(path)
    return path


@pytest.fixture
def service():
    return DDRImportService(db_manager=None)


def test_unrecognised_table_is_reported_with_location_and_counts_only(tmp_path, service):
    _rep, extracted = service.extract_file(_workbook(tmp_path / "synthetic.xlsx"), template=None)
    assert extracted.get("surveys"), "the recognised survey table must still be extracted"
    diagnostic = extracted["metadata"]["unmapped_content"]
    [sheet] = diagnostic["sheets"]
    assert sheet["sheet"] == "Data"
    ranges = {region["range"]: region for region in sheet["regions"]}
    assert ranges["F8:H11"] == {"range": "F8:H11", "cells": 12, "numeric_cells": 9, "text_cells": 3,
                                "rows": 4, "columns": 3, "kind": "table-like"}
    assert ranges["A20:A20"]["kind"] == "isolated" and ranges["A20:A20"]["numeric_cells"] == 0
    assert not any(r.startswith(("A1:", "B1:", "C1:")) for r in ranges), "survey span is accounted for"
    assert diagnostic["unmapped_cells"] == 13 and diagnostic["unmapped_numeric_cells"] == 9
    assert diagnostic["table_like_regions"] == 1 and diagnostic["informational"] is True
    text = json.dumps(diagnostic)
    assert SECRET_TEXT not in text and "98765" not in text and "Zeta" not in text


def test_diagnostic_is_deterministic_and_does_not_touch_the_fingerprint_identity(tmp_path, service):
    path = _workbook(tmp_path / "synthetic.xlsx")
    first = service.extract_file(path, template=None)[1]["metadata"]
    second = service.extract_file(path, template=None)[1]["metadata"]
    assert json.dumps(first["unmapped_content"], sort_keys=True) == json.dumps(second["unmapped_content"], sort_keys=True)
    from core.ddr_import_service import _source_fingerprint
    assert _source_fingerprint({"metadata": first}) == _source_fingerprint({"metadata": second})
    stripped = {k: v for k, v in first.items() if k != "unmapped_content"}
    assert _source_fingerprint({"metadata": stripped}) == _source_fingerprint({"metadata": first})


def test_fully_mapped_workbook_reports_nothing(tmp_path, service):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["MD", "Inc", "Azi"])
    ws.append([1000.0, 1.5, 120.0])
    ws.append([1100.0, 2.0, 121.0])
    wb.save(tmp_path / "clean.xlsx")
    diagnostic = service.extract_file(tmp_path / "clean.xlsx", template=None)[1]["metadata"]["unmapped_content"]
    assert diagnostic["unmapped_cells"] == 0 and diagnostic["sheets"][0]["regions"] == []
    assert summary_line(diagnostic) == ""


def test_regions_are_bounded_and_counted():
    cells = [RawCell(value=i, location=SourceLocation(sheet="S", row=1 + 3 * i, column=1)) for i in range(40)]
    report = SimpleNamespace(raw_document=RawDocument(source_file="x", cells=cells), field_results=[], table_results=[])
    diagnostic = unmapped_content_diagnostic(report, max_regions_per_sheet=5)
    [sheet] = diagnostic["sheets"]
    assert len(sheet["regions"]) == 5 and sheet["regions_total"] == 40 and sheet["regions_truncated"] == 35
    assert diagnostic["unmapped_cells"] == 40  # totals are never truncated


def test_field_value_and_its_label_are_accounted_for():
    cells = [RawCell(value="Depth", location=SourceLocation(sheet="S", row=2, column=1)),
             RawCell(value=1234.0, location=SourceLocation(sheet="S", row=2, column=2)),
             RawCell(value="   ", location=SourceLocation(sheet="S", row=5, column=5))]  # blank is not populated
    result = SimpleNamespace(sheet="S", row=2, col=2,
                             candidates=[{"row": 2, "col": 2, "sheet": "S", "label_row": 2, "label_col": 1}])
    report = SimpleNamespace(raw_document=RawDocument(source_file="x", cells=cells),
                             field_results=[result], table_results=[])
    diagnostic = unmapped_content_diagnostic(report)
    assert diagnostic["populated_cells"] == 2 and diagnostic["unmapped_cells"] == 0


def test_summary_line_carries_counts_but_no_values(tmp_path, service):
    diagnostic = service.extract_file(_workbook(tmp_path / "s.xlsx"), template=None)[1]["metadata"]["unmapped_content"]
    line = summary_line(diagnostic)
    assert "13 populated source cell(s)" in line and "9 numeric" in line and "1 table-like" in line
    assert SECRET_TEXT not in line


def test_import_result_and_lineage_carry_the_diagnostic_without_changing_status(tmp_path, monkeypatch):
    """End-to-end through the transactional import with a temp synthetic database."""
    from core.database import DatabaseManager

    data_dir = tmp_path / "data"
    data_dir.mkdir()
    monkeypatch.setenv("DRILLMASTER_ENV", "test")
    monkeypatch.setenv("DRILLMASTER_DATA_DIR", str(data_dir))
    db = DatabaseManager()
    db.initialize()
    with db.session_scope() as session:
        from core.database import Company, Project, Well
        company = Company(name="Synthetic Co", code="SC")
        session.add(company)
        session.flush()
        project = Project(name="Synthetic Project", code="SP", company_id=company.id)
        session.add(project)
        session.flush()
        well = Well(name="Synthetic Well", code="SW", project_id=project.id)
        session.add(well)
        session.flush()
        well_id = well.id
    results = {}
    for name, unknown in (("with.xlsx", True), ("without.xlsx", False)):
        service = DDRImportService(db_manager=db, well_id=well_id)
        _rep, extracted = service.extract_file(_workbook(tmp_path / name, with_unknown_table=unknown), template=None)
        extracted.setdefault("daily_report", {})["report_date"] = "2026-01-0" + ("1" if unknown else "2")
        results[name] = service._execute_import(extracted, refresh_ui=False)
    with_table, without_table = results["with.xlsx"], results["without.xlsx"]
    assert with_table["status"] == without_table["status"], "informational: status is unchanged"
    assert with_table["unmapped_content"]["unmapped_numeric_cells"] == 9
    assert any("13 populated source cell(s)" in line for line in with_table["details"])
    assert with_table["report_id"], "the import persisted a report"
    audit = db.find_import_audit(_fingerprint_of(service, tmp_path / "with.xlsx"), well_id=well_id)
    assert audit["result"]["unmapped_content"]["unmapped_numeric_cells"] == 9  # lineage
    assert SECRET_TEXT not in json.dumps(audit["result"]["unmapped_content"])
    db.close()


def _fingerprint_of(service, path):
    from core.ddr_import_service import _source_fingerprint
    _rep, extracted = service.extract_file(path, template=None)
    return _source_fingerprint({"metadata": extracted["metadata"]})
