from openpyxl import load_workbook

from core.database import CostRecord
from core.report_engine import EOWRReportEngine
from test_scope_attribution import _base, _mgr, _report, _well


class _FailCostQuerySession:
    """Delegate normal ORM work but fault only the EOWR cost query."""

    def __init__(self, session):
        self._session = session

    def query(self, model, *args, **kwargs):
        if model is CostRecord:
            raise RuntimeError("injected cost database failure")
        return self._session.query(model, *args, **kwargs)

    def __getattr__(self, name):
        return getattr(self._session, name)


def test_eowr_distinguishes_empty_cost_result_from_failed_cost_query(tmp_path, monkeypatch):
    db = _mgr()
    well_id = _well(db, _base(db), "A")
    _report(db, well_id, 1, depth=100)
    engine = EOWRReportEngine(db)

    empty = engine._collect_data(well_id)
    assert empty["cost_records"] == []
    assert empty["cost_query_status"] == "available"
    assert empty["summary"]["cost_currency_status"] == "no-records"
    empty_html = engine._build_html(empty)
    assert "No cost records were returned" in empty_html
    assert "this is not a reported zero" in empty_html

    original_create_session = db.create_session
    monkeypatch.setattr(
        db,
        "create_session",
        lambda: _FailCostQuerySession(original_create_session()),
    )
    unavailable = engine._collect_data(well_id)
    assert unavailable["cost_records"] is None
    assert unavailable["cost_query_status"] == "query-failed"
    assert unavailable["summary"]["total_cost"] is None
    assert unavailable["summary"]["cost_currency_status"] == "query-failed"

    html = engine._build_html(unavailable)
    assert "Cost data unavailable" in html
    assert "cost-record query failed" in html
    assert "No cost records were returned" not in html

    path = tmp_path / "eowr-cost-query-failed.xlsx"
    assert engine._save_excel(unavailable, str(path))
    workbook = load_workbook(path, data_only=True)
    try:
        summary = workbook["Summary"]
        rows = {summary.cell(row, 1).value: summary.cell(row, 2).value
                for row in range(5, summary.max_row + 1)}
        assert rows["Total Cost"] is None
        assert rows["Cost Data Status"] == "Unavailable — query failed"
    finally:
        workbook.close()
        db.close()
