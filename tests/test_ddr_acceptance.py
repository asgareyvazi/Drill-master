"""Environment-gated DDR acceptance tests.

These tests are intentionally not fixtures disguised as real acceptance.  CI
skips with an explicit reason unless the caller supplies the real DDR paths and
an externally managed MinerU installation:

    DRILLMASTER_TEST_DDR_XLSX=/real/DDR.xlsx
    DRILLMASTER_TEST_DDR_PDF=/real/DDR.pdf

The Excel test exercises the canonical extractor and atomic database boundary.
The PDF test exercises the external MinerU adapter, common IR, normalization,
review/provenance, and the same atomic boundary.  No path or date is invented
when a real document is unavailable.
"""

from __future__ import annotations

import json
import os
from datetime import date
from pathlib import Path

import pytest


REPO = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = REPO / "templates" / "OEOC_DDR_v3.json"


def _env_file(name: str) -> Path:
    value = os.getenv(name)
    if not value:
        pytest.skip(f"{name} is not set; real DDR acceptance is opt-in")
    path = Path(value).expanduser()
    if not path.is_file():
        pytest.skip(f"{name} does not point to a file: {path}")
    return path


def _db_manager():
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool
    from core.database import Base, DatabaseManager

    manager = DatabaseManager()
    manager.engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(manager.engine)
    manager.Session = sessionmaker(bind=manager.engine, autoflush=False, autocommit=False)
    return manager


def _seed_report(manager):
    from core.database import Company, DailyReport, Project, Section, Well

    session = manager.create_session()
    try:
        company = Company(name="Acceptance Company", code="ACCEPTANCE")
        session.add(company)
        session.flush()
        project = Project(name="Acceptance Project", code="ACCEPTANCE", company_id=company.id)
        session.add(project)
        session.flush()
        well = Well(name="Acceptance Well", code="ACCEPTANCE-WELL", project_id=project.id)
        session.add(well)
        session.flush()
        section = Section(name="Imported Section", well_id=well.id)
        session.add(section)
        session.flush()
        report = DailyReport(
            well_id=well.id,
            section_id=section.id,
            report_number=1,
            report_date=date(2024, 1, 1),
            status="Draft",
        )
        session.add(report)
        session.commit()
        return well.id, report.id
    finally:
        session.close()


def _assert_ir_details(raw_document):
    payload = raw_document.to_dict(include_cells=True)
    assert payload["ir_version"] >= 2
    assert payload["source_file"]
    assert "tables_detail" in payload
    assert "cell_values" in payload
    assert "section_title_values" in payload
    # A source may legitimately have no formal title/unit, but the contract
    # must be serializable and deserializable without losing them when present.
    from core.import_ir import RawDocument

    restored = RawDocument.from_dict(payload)
    assert restored.source_file == raw_document.source_file
    assert len(restored.cells) == len(raw_document.cells)
    assert len(restored.tables) == len(raw_document.tables)
    for cell in restored.cells[:25]:
        assert cell.location.file == raw_document.source_file
        assert cell.original_value == cell.value


def _assert_review_contract(rows):
    from core.import_quality import ImportReviewMatrix, ReviewItem

    matrix = ImportReviewMatrix.from_rows(rows)
    assert len(matrix.items) == len(rows)
    for item in matrix.items:
        payload = item.to_dict()
        restored = ReviewItem.from_dict(payload)
        assert restored.target_field == item.target_field
        assert restored.original_value == item.original_value
        assert restored.normalized_value == item.normalized_value
        assert restored.file == item.file


@pytest.mark.integration
def test_real_ddr_excel_canonical_ir_review_and_atomic_db():
    source = _env_file("DRILLMASTER_TEST_DDR_XLSX")
    if not TEMPLATE_PATH.is_file():
        pytest.fail(f"Canonical OEOC template is missing: {TEMPLATE_PATH}")

    from openpyxl import load_workbook
    from core.excel_intelligence import ExcelIntelligence

    workbook = load_workbook(source, data_only=False, read_only=False)
    try:
        template = json.loads(TEMPLATE_PATH.read_text(encoding="utf-8"))
        expected_sheets = {
            part.replace("_", " ").lower()
            for key in template
            if key.startswith("sheet_")
            for part in [key.split("_", 2)[-1]]
        }
        actual_sheets = {sheet.title.lower() for sheet in workbook.worksheets}
        if not expected_sheets.intersection(actual_sheets):
            pytest.fail(
                "The supplied DDR workbook does not match the canonical OEOC "
                f"template; sheets={sorted(actual_sheets)}"
            )
        report = ExcelIntelligence(workbook, template).extract()
    finally:
        workbook.close()

    assert report.raw_document is not None
    _assert_ir_details(report.raw_document)
    assert report.raw_document.metadata.get("engine") == "Excel"
    assert report.canonical_json, "Real Excel produced no canonical data"

    review_rows = []
    for result in report.field_results:
        if result.status != "OK" or result.certainty == "LOW" or result.canonical_field in report.source_tokens:
            source_token = report.source_tokens.get(result.canonical_field, {})
            review_rows.append(
                {
                    "file": source.name,
                    "sheet": result.sheet,
                    "source_cell": result.cell,
                    "original_value": source_token.get("original_value", result.original_value),
                    "normalized_value": None if source_token else result.normalized_value,
                    "target_field": result.canonical_field,
                    "confidence": result.confidence,
                    "decision": "REVIEW",
                    "validation_state": result.validation or "unvalidated",
                    "review_state": "unreviewed",
                }
            )
    _assert_review_contract(review_rows)

    manager = _db_manager()
    well_id, report_id = _seed_report(manager)
    persisted = manager.save_imported_multi_tab_data_atomic(
        well_id, report_id, dict(report.canonical_json)
    )
    assert persisted.get("failed") == 0, persisted
    assert persisted.get("imported", 0) > 0, persisted

    # The known historical crash token must never become a float conversion
    # exception.  If it exists in source, the extractor must preserve it as a
    # review token/NULL, not invent zero.
    for token in report.source_tokens.values():
        assert token.get("normalized_value") is None


@pytest.mark.integration
def test_real_ddr_pdf_mineru_common_ir_review_and_atomic_db():
    source = _env_file("DRILLMASTER_TEST_DDR_PDF")
    from core.mineru_engine import DocumentNormalizer, MinerUAdapter

    adapter = MinerUAdapter()
    health = adapter.health_check()
    if not health.available:
        pytest.skip("MinerU is unavailable in this environment: " + str(health.error or health.to_dict()))
    result = adapter.parse_file(source)
    assert result.success, result.error
    assert result.document is not None
    assert result.document.raw_files, "MinerU produced no inspectable markdown/JSON/assets"

    normalized = DocumentNormalizer().normalize(result.document)
    assert normalized.raw_document is not None
    _assert_ir_details(normalized.raw_document)
    assert normalized.provenance, "Real PDF produced no canonical provenance"
    assert all(item.get("source_file") == str(source.resolve()) for item in normalized.provenance)
    assert normalized.validation.valid, normalized.validation.errors

    review_rows = []
    for warning in normalized.warnings:
        source_info = warning.get("source") or {}
        review_rows.append(
            {
                "file": source.name,
                "sheet": source_info.get("source_sheet", ""),
                "page": source_info.get("source_page"),
                "detected_table": "MinerU document",
                "source_cell": (
                    f"page {source_info['source_page']}:row {source_info.get('source_row')}:column "
                    f"{source_info.get('source_column')}"
                    if source_info.get("source_page") is not None
                    else ""
                ),
                "coordinates": source_info.get("bounding_box"),
                "original_value": warning.get("value"),
                "normalized_value": warning.get("normalized_value"),
                "target_field": warning.get("field", ""),
                "confidence": source_info.get("confidence"),
                "decision": "REVIEW",
                "validation_state": "needs_review",
                "review_state": "unreviewed",
            }
        )
    _assert_review_contract(review_rows)

    # Persist the same canonical payload when the real DDR contains the
    # required report date.  Missing source dates remain a real acceptance
    # failure rather than being replaced with today's date.
    report_data = normalized.canonical_data.get("daily_report") or {}
    if not report_data.get("report_date"):
        pytest.fail("Real DDR PDF has no canonical daily_report.report_date; refusing to invent one")
    manager = _db_manager()
    well_id, report_id = _seed_report(manager)
    persisted = manager.save_imported_multi_tab_data_atomic(
        well_id, report_id, dict(normalized.canonical_data)
    )
    assert persisted.get("failed") == 0, persisted
    assert persisted.get("imported", 0) > 0, persisted

    # This is the regression guard for the old numeric conversion crash.
    for table in normalized.raw_document.tables:
        for row in table.rows:
            for cell in row:
                if str(cell.original_value).strip().lower() == "drilling data":
                    assert cell.validation_state in {"unvalidated", "needs_review", "valid"}
                    assert cell.review_state in {"unreviewed", "review", "accepted"}
