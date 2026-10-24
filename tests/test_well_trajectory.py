"""M44: whole-well trajectory composition (core/well_trajectory.py).

Golden values are closed-form circular-arc geometry, independent of the
minimum-curvature engine: a constant-azimuth build at a constant rate is an
exact circular arc of radius R = 30 m / rate(rad), so

    dTVD = R * sin(inc_end),  dH = R * (1 - cos(inc_end))       (from vertical)

and a straight hold of length L at inclination i adds L*cos(i), L*sin(i).
"""
import math
from datetime import date

import pytest

from core import well_trajectory as wt
from core.well_trajectory import (
    ERROR, INSUFFICIENT_DATA, NO_DATA, READY, REVIEW_REQUIRED, compose_well_trajectory, dls_for_display,
    load_well_trajectory, to_display_units,
)


def _rows(points, report_id=None, start_id=1):
    return [{"id": start_id + n, "md": md, "inc": inc, "azi": azi, "report_id": report_id}
            for n, (md, inc, azi) in enumerate(points)]


def _j_profile(azimuth=90.0):
    """Vertical to 600 m, build 3 deg/30 m to 30 deg at 900 m, hold to 1200 m."""
    points = [(0.0, 0.0, azimuth), (300.0, 0.0, azimuth), (600.0, 0.0, azimuth)]
    points += [(600.0 + 30.0 * k, 3.0 * k, azimuth) for k in range(1, 11)]
    points += [(900.0 + 30.0 * k, 30.0, azimuth) for k in range(1, 11)]
    return points


def _codes(trajectory):
    return [issue.code for issue in trajectory.issues]


# ---------------------------------------------------------------- goldens
def test_j_profile_matches_closed_form_arc_and_hold():
    radius = 30.0 / math.radians(3.0)                       # 572.958 m
    build_tvd, build_h = radius * math.sin(math.radians(30)), radius * (1 - math.cos(math.radians(30)))
    assert build_tvd == pytest.approx(286.479, abs=1e-3) and build_h == pytest.approx(76.762, abs=1e-3)

    t = compose_well_trajectory(_rows(_j_profile()))
    assert t.status == READY and t.drawable and len(t.stations) == 23
    at = {round(s.md): s for s in t.stations}
    assert at[600].tvd == pytest.approx(600.0, abs=1e-6) and at[600].east == pytest.approx(0.0, abs=1e-9)
    assert at[900].tvd == pytest.approx(600.0 + build_tvd, abs=1e-3)
    assert at[900].east == pytest.approx(build_h, abs=1e-3)
    end = at[1200]
    assert end.tvd == pytest.approx(600.0 + build_tvd + 300.0 * math.cos(math.radians(30)), abs=1e-3)  # 1146.287
    assert end.east == pytest.approx(build_h + 150.0, abs=1e-3)                                          # 226.762
    assert end.north == pytest.approx(0.0, abs=1e-6)
    assert end.hd == pytest.approx(end.east, abs=1e-6)
    # closure azimuth 90 deg -> VS equals the horizontal displacement
    assert t.vs_azimuth_source == "closure" and t.vs_azimuth == pytest.approx(90.0, abs=1e-6)
    assert end.vs == pytest.approx(end.east, abs=1e-6)


def test_dogleg_severity_is_build_rate_in_build_and_zero_in_hold():
    t = compose_well_trajectory(_rows(_j_profile()))
    at = {round(s.md): s for s in t.stations}
    for md in range(630, 901, 30):
        assert at[md].dls == pytest.approx(3.0, abs=1e-6), md
    for md in range(930, 1201, 30):
        assert at[md].dls == pytest.approx(0.0, abs=1e-6), md


def test_user_vs_azimuth_projects_the_horizontal_displacement():
    t = compose_well_trajectory(_rows(_j_profile()), vs_azimuth=150.0)
    end = t.stations[-1]
    assert t.vs_azimuth_source == "user" and t.vs_azimuth == pytest.approx(150.0)
    # VS = N cos(a) + E sin(a); N = 0 here so VS = E sin(150) = E / 2
    assert end.vs == pytest.approx(end.east * 0.5, abs=1e-6)
    negative = compose_well_trajectory(_rows(_j_profile()), vs_azimuth=-30.0)
    assert negative.vs_azimuth == pytest.approx(330.0)
    assert negative.stations[-1].vs == pytest.approx(end.east * math.sin(math.radians(330)), abs=1e-6)


def test_vertical_well_has_default_north_vs_azimuth_with_notice():
    t = compose_well_trajectory(_rows([(0, 0, 0), (500, 0, 0), (1000, 0, 0)]))
    assert t.status == READY
    assert t.vs_azimuth_source == "default_north" and t.vs_azimuth == 0.0
    assert t.closure_azimuth is None and t.closure_distance == pytest.approx(0.0)
    assert "VS_AZIMUTH_UNDEFINED" in _codes(t)
    assert t.stations[-1].tvd == pytest.approx(1000.0)


# -------------------------------------------- composition across reports
def test_repeated_station_across_reports_is_merged_with_provenance():
    first = _rows([(0, 0, 45), (100, 1, 45), (200, 2, 45)], report_id=11, start_id=1)
    second = _rows([(200, 2, 45), (300, 3, 45)], report_id=12, start_id=10)
    t = compose_well_trajectory(first + second)
    assert t.status == READY and t.input_rows == 5 and t.excluded_rows == 0
    assert [s.md for s in t.stations] == [0, 100, 200, 300]
    shared = t.stations[2]
    assert shared.survey_ids == (3, 10) and shared.report_ids == (11, 12)
    assert t.stations[3].report_ids == (12,)


def test_frames_are_continuous_not_restarted_per_report():
    """Second report's stations continue from the first report's position."""
    first = _rows(_j_profile()[:13], report_id=1)                 # to 900 m
    second = _rows(_j_profile()[12:], report_id=2, start_id=100)  # 900 m repeated, then hold
    whole = compose_well_trajectory(_rows(_j_profile()))
    split = compose_well_trajectory(first + second)
    assert split.status == READY
    assert split.stations[-1].tvd == pytest.approx(whole.stations[-1].tvd, abs=1e-9)
    assert split.stations[-1].east == pytest.approx(whole.stations[-1].east, abs=1e-9)
    # Restarting the second report at its own tie-on (N = E = 0, TVD = MD cos inc)
    # loses the build section: E 150 instead of 226.76, TVD 1039.23 instead of 1146.29.
    restarted = compose_well_trajectory(second)
    assert restarted.stations[-1].east == pytest.approx(150.0, abs=1e-6)
    assert restarted.stations[-1].tvd == pytest.approx(1200 * math.cos(math.radians(30)), abs=1e-6)
    assert split.stations[-1].east - restarted.stations[-1].east == pytest.approx(76.762, abs=1e-3)


def test_conflicting_readings_at_one_md_are_all_excluded_for_review():
    rows = _rows([(0, 0, 10), (100, 0.5, 10), (200, 1.0, 10)], report_id=1)
    rows += _rows([(200, 3.0, 10)], report_id=2, start_id=50)
    t = compose_well_trajectory(rows)
    assert t.status == REVIEW_REQUIRED and t.drawable
    assert [s.md for s in t.stations] == [0, 100]
    conflict = next(i for i in t.issues if i.code == "CONFLICTING_DUPLICATE")
    assert conflict.severity == "error" and conflict.md == 200 and set(conflict.survey_ids) == {3, 50}
    assert t.excluded_rows == 2


def test_md_and_angle_tolerances_and_azimuth_wrap():
    near = compose_well_trajectory(_rows([(0, 0, 0), (100, 1, 359.995), (100.005, 1, 0.0), (200, 2, 0)]))
    assert near.status == READY and [round(s.md, 3) for s in near.stations] == [0, 100, 200]
    apart = compose_well_trajectory(_rows([(0, 0, 0), (100, 1, 0), (100.02, 1, 0), (200, 2, 0)]))
    assert len(apart.stations) == 4
    disagree = compose_well_trajectory(_rows([(0, 0, 0), (100, 1, 0), (100, 1.02, 0), (200, 2, 0)]))
    assert "CONFLICTING_DUPLICATE" in _codes(disagree) and disagree.status == REVIEW_REQUIRED


# ------------------------------------------------ invalid / missing data
@pytest.mark.parametrize("row, code", [
    ({"md": None, "inc": 1, "azi": 1}, "INVALID_MD"),
    ({"md": -5, "inc": 1, "azi": 1}, "INVALID_MD"),
    ({"md": "abc", "inc": 1, "azi": 1}, "INVALID_MD"),
    ({"md": float("nan"), "inc": 1, "azi": 1}, "INVALID_MD"),
    ({"md": True, "inc": 1, "azi": 1}, "INVALID_MD"),
    ({"md": 150, "inc": None, "azi": 1}, "MISSING_ANGLE"),
    ({"md": 150, "inc": 1, "azi": None}, "MISSING_ANGLE"),
    ({"md": 150, "inc": 181, "azi": 1}, "INVALID_ANGLE"),
    ({"md": 150, "inc": -1, "azi": 1}, "INVALID_ANGLE"),
    ({"md": 150, "inc": 1, "azi": float("inf")}, "INVALID_ANGLE"),
    ({"md": 150, "inc": "x", "azi": 1}, "INVALID_ANGLE"),
])
def test_invalid_rows_are_excluded_and_reported(row, code):
    good = _rows([(0, 0, 0), (100, 1, 0), (200, 2, 0)])
    t = compose_well_trajectory(good + [dict(row, id=99)])
    assert t.status == REVIEW_REQUIRED
    assert [s.md for s in t.stations] == [0, 100, 200]
    issue = next(i for i in t.issues if i.code == code)
    assert issue.severity == "warning" and issue.survey_ids == (99,)
    assert t.excluded_rows == 1 and t.input_rows == 4


def test_no_rows_is_no_data_and_not_drawable():
    for rows in (None, []):
        t = compose_well_trajectory(rows)
        assert t.status == NO_DATA and not t.drawable and t.stations == []


def test_single_station_or_all_invalid_is_insufficient_and_not_drawable():
    single = compose_well_trajectory(_rows([(100, 1, 1)]))
    assert single.status == INSUFFICIENT_DATA and not single.drawable and len(single.stations) == 1
    invalid = compose_well_trajectory([{"id": 1, "md": None, "inc": 1, "azi": 1},
                                       {"id": 2, "md": 10, "inc": None, "azi": 1}])
    assert invalid.status == INSUFFICIENT_DATA and not invalid.drawable and invalid.stations == []
    assert invalid.excluded_rows == 2


def test_large_station_gap_is_flagged_but_does_not_block():
    t = compose_well_trajectory(_rows([(0, 0, 0), (100, 0, 0), (400, 0, 0)]))
    gap = next(i for i in t.issues if i.code == "LARGE_STATION_GAP")
    assert gap.severity == "warning" and gap.md == 400
    assert t.status == READY
    assert "LARGE_STATION_GAP" not in _codes(compose_well_trajectory(_rows([(0, 0, 0), (150, 0, 0)])))


def test_tie_on_assumption_is_always_stated():
    t = compose_well_trajectory(_rows([(500, 10, 45), (530, 11, 45)]))
    tie = next(i for i in t.issues if i.code == "TIE_ON_ASSUMED")
    assert tie.md == 500 and tie.severity == "info"
    assert t.stations[0].north == 0 and t.stations[0].east == 0
    assert t.stations[0].tvd == pytest.approx(500 * math.cos(math.radians(10)), abs=1e-9)


def test_engine_refusal_is_error_without_stations(monkeypatch):
    from core.engineering.engines.trajectory import TrajectoryCalculator
    from core.engineering.result import EngineeringError

    def refuse(*_args, **_kwargs):
        raise EngineeringError("synthetic refusal")
    monkeypatch.setattr(TrajectoryCalculator, "calculate", staticmethod(refuse))
    t = compose_well_trajectory(_rows([(0, 0, 0), (100, 1, 0)]))
    assert t.status == ERROR and not t.drawable and t.stations == []
    assert "ENGINE_REFUSED" in _codes(t)


def test_orm_like_objects_are_accepted():
    class Row:
        def __init__(self, id, md, inc, azi):
            self.id, self.md, self.inc, self.azi, self.report_id, self.tool = id, md, inc, azi, 7, "MWD"
    t = compose_well_trajectory([Row(1, 0, 0, 0), Row(2, 100, 2, 30)])
    assert t.status == READY and t.stations[1].tool == "MWD" and t.stations[1].report_ids == (7,)


# ------------------------------------------------------------------ units
def test_display_unit_conversions():
    assert to_display_units(1.0, "m") == 1.0
    assert to_display_units(0.3048, "ft") == pytest.approx(1.0, abs=1e-12)
    assert to_display_units(100.0, "ft") == pytest.approx(328.083989501, abs=1e-9)
    assert dls_for_display(3.0, "deg/30m") == 3.0
    assert dls_for_display(3.0, "deg/100ft") == pytest.approx(3.0 * 30.48 / 30.0, abs=1e-12)
    assert to_display_units(None, "ft") is None and dls_for_display(None, "deg/100ft") is None
    with pytest.raises(ValueError):
        to_display_units(1.0, "yd")
    with pytest.raises(ValueError):
        dls_for_display(1.0, "deg/10m")


# ------------------------------------------- synthetic DB + consumers
@pytest.fixture(scope="module")
def synthetic(tmp_path_factory):
    from synthetic_operations import build_synthetic_operations, make_file_database
    db = make_file_database(tmp_path_factory.mktemp("traj") / "syn.db")
    ids = build_synthetic_operations(db, end_date=date(2026, 1, 15))
    yield db, ids
    db.engine.dispose()


def _dev_expected():
    radius = 30.0 / math.radians(2.0)                       # 859.437 m
    tvd = 600 + radius * math.sin(math.radians(40)) + 600 * math.cos(math.radians(40))
    h = radius * (1 - math.cos(math.radians(40))) + 600 * math.sin(math.radians(40))
    return tvd, h                                            # 1612.07, 586.74


def test_synthetic_directional_well_composes_across_six_reports(synthetic):
    db, ids = synthetic
    tvd, h = _dev_expected()
    assert tvd == pytest.approx(1612.07, abs=0.01) and h == pytest.approx(586.74, abs=0.01)
    t = load_well_trajectory(db, ids["dev"])
    assert t.status == READY and t.input_rows == 66 and len(t.stations) == 61
    end = t.stations[-1]
    assert end.md == pytest.approx(1800.0)
    assert end.tvd == pytest.approx(tvd, abs=0.05) and end.hd == pytest.approx(h, abs=0.05)
    assert end.north == pytest.approx(-h * math.cos(math.radians(45)), abs=0.05)
    assert end.east == pytest.approx(h * math.sin(math.radians(45)), abs=0.05)
    assert t.vs_azimuth_source == "closure" and t.vs_azimuth == pytest.approx(135.0, abs=1e-6)
    # The stored derived columns of the last report are a per-report frame
    # (restarted at N = E = 0) and must not be what consumers present.
    stored = db.load_survey_points(well_id=ids["dev"])
    last_stored = max(stored, key=lambda r: (r["md"], r.get("report_id") or 0))
    assert last_stored["tvd"] < tvd - 200


def test_report_scope_and_review_well_via_database(synthetic):
    db, ids = synthetic
    one_report = load_well_trajectory(db, ids["dev"], report_id=ids["reports"][1])
    assert one_report.status == READY and one_report.stations[-1].md == pytest.approx(300.0)
    qa = load_well_trajectory(db, ids["qa"])
    assert qa.status == REVIEW_REQUIRED and [s.md for s in qa.stations] == [0, 100]
    assert {"CONFLICTING_DUPLICATE", "MISSING_ANGLE", "TIE_ON_ASSUMED"} <= set(_codes(qa))
    assert load_well_trajectory(db, ids["vert"]).status == NO_DATA
    assert load_well_trajectory(db, None).status == NO_DATA


def test_eowr_and_exports_use_the_composed_trajectory(synthetic, tmp_path):
    from openpyxl import load_workbook

    from core.api.rest_api import DrillMasterAPI
    from core.professional_export import ProfessionalExcelExport
    from core.report_engine import EOWRReportEngine

    db, ids = synthetic
    tvd, _h = _dev_expected()
    engine = EOWRReportEngine(db)
    data = engine._collect_data(ids["dev"])
    assert data["surveys"][-1].tvd == pytest.approx(tvd, abs=0.05)       # not the 1378.88 frame value
    assert data["survey_trajectory"].status == READY
    html = engine._build_html(data)
    assert "Recomputed by minimum curvature from 61 distinct stations" in html
    assert "No survey uncertainty model" in html

    api_rows = DrillMasterAPI(db).get_trajectory(ids["dev"])
    assert len(api_rows) == 61 and api_rows[-1]["tvd"] == pytest.approx(tvd, abs=0.05)

    path = tmp_path / "pro.xlsx"
    assert ProfessionalExcelExport(db).export(ids["dev"], str(path))
    sheet = load_workbook(path)["Survey"]
    header = [c.value for c in sheet[1]]
    tvd_col = header.index("TVD (m)")
    values = [row[tvd_col] for row in sheet.iter_rows(min_row=2, values_only=True)
              if isinstance(row[tvd_col], (int, float))]
    assert len(values) == 61 and values[-1] == pytest.approx(tvd, abs=0.05)
    notes = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True) if isinstance(row[0], str)]
    assert any("Recomputed by minimum curvature" in note for note in notes)


def test_module_constants_are_exact_unit_definitions():
    assert wt.FT_PER_M == pytest.approx(3.280839895013123, abs=1e-15)
    assert wt.DLS_30M_TO_100FT == pytest.approx(30.48 / 30.0, abs=1e-15)
