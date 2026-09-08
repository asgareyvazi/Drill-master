"""Follow-up contracts and real service/edit lifecycle; not desktop certification."""
import copy
import json
from datetime import date
from pathlib import Path

import pytest

from core.database import DatabaseManager, DailyReport, Section
from core.domain_records import collection_value, bha_records, is_metadata_value
from core.save_outcome import SaveOutcome, SaveIssue, save_all, public_status
from core.mud_records import mud_density_pcf, preserve_widget_values
from core.ddr_import_service import DDRImportService


@pytest.fixture
def db(tmp_path):
    manager = DatabaseManager()
    manager.db_path = str(tmp_path / "lifecycle.db")
    assert manager.initialize()
    with manager.session_scope() as session:
        section = Section(well_id=1, name="Lifecycle")
        session.add(section)
        session.flush()
        session.add(DailyReport(well_id=1, section_id=section.id, report_number=1, report_date=date(2024, 10, 22)))
    yield manager
    manager.close()


@pytest.mark.parametrize("value", [None, "null", " null ", "", "  ", [], "[]"])
def test_empty_collection_contract(value):
    assert collection_value(value, "records") == []


@pytest.mark.parametrize("value", ["{bad json", {}, "{}", [None], "[1]"])
def test_invalid_collection_not_silently_empty(value):
    with pytest.raises((ValueError, TypeError)):
        collection_value(value, "records")


def test_save_all_partial_success_and_exact_diagnostics():
    def bad():
        raise ValueError("POB date_out 'nonsense' is not a date")
    result = save_all([("Mud", lambda: True), ("POB", bad), ("BHA", lambda: False)])
    assert not result and result.saved == 1 and result.status == "SYSTEM_ERROR"
    assert result.issues[0].status == "INVALID_SOURCE"
    assert result.issues[0].section == "POB" and "date_out" in result.summary()
    assert result.issues[0].corrective_action and result.issues[0].traceback
    assert result.issues[1].section == "BHA"


def test_save_all_detailed_no_stale_child_outcome():
    class Owner:
        last_save_outcome = SaveOutcome(issues=[SaveIssue("old", "stale")])
        def save(self):
            return True
    assert save_all([("current", Owner().save)])
    result = save_all([("Survey", lambda: SaveOutcome(saved=2, issues=[SaveIssue("Survey", "Missing azimuth", status="REVIEW_REQUIRED", row=3, field="azi")]))])
    assert result.status == "REVIEW_REQUIRED" and "row 3 [azi]" in result.summary()


@pytest.mark.parametrize("legacy,new", [("ACCEPT", "SUCCESS"), ("VALIDATION_ERROR", "INVALID_SOURCE"), ("PERSISTENCE_ERROR", "SYSTEM_ERROR"), ("UNSUPPORTED", "UNSUPPORTED"), ("REVIEW_REQUIRED", "REVIEW_REQUIRED")])
def test_status_bridge(legacy, new):
    assert public_status(legacy) == new


def stations():
    return [{"well_id": 1, "report_id": 1, "md": 100, "inc": 10, "azi": 45},
            {"well_id": 1, "report_id": 1, "md": 200, "inc": 20, "azi": 60}]


def test_survey_edit_md_delete_and_recalculate(db):
    db.save_survey_records(stations())
    rows = db.load_survey_points(report_id=1)
    first_id = rows[0]["id"]
    rows[0]["md"] = 110
    db.save_survey_records(rows, replace_scope=(1, 1))
    actual = db.load_survey_points(report_id=1)
    assert len(actual) == 2 and actual[0]["id"] == first_id and actual[0]["md"] == 110
    result = db.save_survey_records(actual[:1], replace_scope=(1, 1))
    assert result["calculated"] == 1 and len(db.load_survey_points(report_id=1)) == 1
    db.save_survey_records([], replace_scope=(1, 1))
    assert db.load_survey_points(report_id=1) == []


def test_invalid_snapshot_never_deletes_existing_stations(db):
    db.save_survey_records(stations())
    rows = db.load_survey_points(report_id=1)
    rows[0]["inc"] = "BAD"
    result = db.save_survey_records(rows[:1], replace_scope=(1, 1))
    assert result["status"] == "INVALID_SOURCE" and result["rejected"] == 1
    assert len(db.load_survey_points(report_id=1)) == 2


def test_survey_cross_context_and_collision_are_atomic(db):
    db.save_survey_records(stations())
    rows = db.load_survey_points(report_id=1)
    rows[0]["md"] = rows[1]["md"]
    with pytest.raises(ValueError, match="collides"):
        db.save_survey_records(rows[:1])
    with pytest.raises(ValueError, match="context"):
        db.save_survey_records(stations(), replace_scope=(1, 999))
    assert [r["md"] for r in db.load_survey_points(report_id=1)] == [100, 200]


def test_safety_invalid_edit_keeps_children_and_valid_clear_deletes(db):
    data = {"well_id": 1, "report_id": 1, "report_date": date(2024, 10, 22),
            "bop_stack_json": [{"Name": "Annular", "Type": "RS", "WP (psi)": 3000}],
            "waste_history_json": [{"Date": "2024-10-22", "Type": "Cuttings", "Volume (BBL)": 5}]}
    assert db.save_safety_report(data)
    before = db.get_safety_report(1, report_id=1)
    data["bop_stack_json"][0]["WP (psi)"] = "invalid"
    data["waste_history_json"][0]["Date"] = "None"
    assert db.save_safety_report(data)
    assert db.last_safety_review
    after = db.get_safety_report(1, report_id=1)
    assert before["bop_stack_json"] == after["bop_stack_json"]
    assert before["waste_history_json"] == after["waste_history_json"]
    data.update(bop_stack_json=[], waste_history_json=[])
    assert db.save_safety_report(data)
    after = db.get_safety_report(1, report_id=1)
    assert after["bop_stack_json"] == [] and after["waste_history_json"] == []


def test_safety_latest_report_does_not_mix_well_children(db):
    with db.session_scope() as session:
        session.add(DailyReport(well_id=1, section_id=1, report_number=2, report_date=date(2024, 10, 23)))
    for rid, name in [(1, "A"), (2, "B")]:
        assert db.save_safety_report({"well_id": 1, "report_id": rid, "report_date": date(2024, 10, 21 + rid),
            "bop_stack_json": [{"Name": name, "Type": "RS", "WP (psi)": 3000}]})
    assert [r["Name"] for r in db.get_safety_report(1)["bop_stack_json"]] == ["B"]


@pytest.mark.parametrize("value", [{"source": "A1"}, "{'bha.component_name': 'A1'}", "/tmp/report.xlsx", r"C:\data\report.xlsx"])
def test_serialized_provenance_is_not_domain_value(value):
    assert is_metadata_value(value)
    record = bha_records([{"component_name": value, "weight": value, "connection": value}])[0]
    assert record["Component Name"] is None and record["Weight (kg)"] is None and record["Connection Type"] is None
    assert record["_provenance"]


def test_bha_repository_edit_delete_and_cumulative(db):
    from core.repositories.report_repository import BHARepository
    repo = BHARepository(db)
    rows = [{"component_name": "Bit", "length": 1}, {"component_name": "Collar", "length": 10}]
    rid = repo.save(1, {"report_id": 1, "bha_data": rows})
    data = db.get_bha_report(1, report_id=1)["bha_configs"]
    assert data[1]["Cumulative Length (m)"] == 11
    provenance = copy.deepcopy(data[1]["_provenance"])
    data[0]["Length (m)"] = 2
    assert repo.save(1, {"report_id": 1, "bha_data": data}) == rid
    loaded = db.get_bha_report(1, report_id=1)["bha_configs"]
    assert loaded[1]["Cumulative Length (m)"] == 12 and loaded[1]["_provenance"] == provenance
    repo.save(1, {"report_id": 1, "bha_data": []})
    assert db.get_bha_report(1, report_id=1)["bha_configs"] == []


@pytest.mark.parametrize("value,unit", [(71, "pcf"), (9.5, "ppg"), (1.2, "sg"), ("9.5 ppg", None)])
def test_mud_density_reaches_pcf_domain(value, unit):
    from core.unit_manager import UnitManager
    result, lineage = mud_density_pcf({"mw": value, "mw_unit": unit})
    magnitude = 9.5 if isinstance(value, str) else value
    assert result["mw"] == pytest.approx(UnitManager.convert(magnitude, "density", unit or "ppg", "pcf"))
    assert lineage["domain_unit"] == "pcf"


def test_untouched_widgets_preserve_missingness_and_explicit_zero():
    original = {"kcl": None, "water_percent": None, "ph": None}
    displayed = {"kcl": 0, "water_percent": 0, "ph": 9.5}
    assert preserve_widget_values(original, displayed, displayed) == original
    assert preserve_widget_values(original, displayed, displayed, {"kcl"})["kcl"] == 0
    edited = {**displayed, "water_percent": 88}
    assert preserve_widget_values(original, displayed, edited)["water_percent"] == 88


def test_semantic_scalar_label_does_not_become_measurement():
    from core.excel_intelligence import FieldExtractor, LabelDetector, MergeCellAnalyzer
    from openpyxl import Workbook
    w = Workbook(); sheet = w.active
    sheet.append(["Wind Speed (km/h)", None, "Material Type", 999])
    cells = {(1, c): sheet.cell(1, c).value for c in range(1, 5)}
    extractor = FieldExtractor(cells, MergeCellAnalyzer(sheet), LabelDetector(cells))
    result = extractor.extract({"row": 1, "col": 1, "field": "Wind Speed"}, "safety.wind_speed")
    assert result.value is None


def test_real_golden_source_reviews_and_domain_roundtrip(db):
    path = Path(__file__).resolve().parents[1] / "08-DDR OEOC-208 AZNS-207 2024-Oct-22.xlsx"
    service = DDRImportService(db, 1)
    extraction, payload = service.extract_file(str(path))
    # Known source density is not discarded by the ReviewItem boundary.
    assert extraction.source_tokens["mud_report.mw"]["normalized_value"] == 71
    assert not any(r["canonical_field"] == "mud_report.mw" for r in payload["metadata"]["review_matrix"])
    result = service.import_records(payload)
    assert result["failed"] == 0
    rid, wid = result["report_id"], result["well_id"]
    mud = db.get_mud_report(report_id=rid)
    assert mud["mw"] == 71 and mud["water_percent"] is None and mud["kcl"] is None
    original = json.loads(mud["chemicals_json"])
    # Manual named edit does not change the imported measurements or metadata.
    mud["summary"] = "Reviewed without inventing measurements"
    assert db.save_mud_report(mud)
    reloaded = db.get_mud_report(report_id=rid)
    assert json.loads(reloaded["chemicals_json"]) == original
    surveys = db.load_survey_points(report_id=rid)
    assert db.save_survey_records(surveys, replace_scope=(wid, rid))["calculated"] == 0


def test_planning_snapshot_keeps_dates_actuals_ids_and_deletes(db):
    from core.database import SevenDaysLookahead
    # More than seven imported plan rows must not disappear on manual load/save.
    records = [{"plan_date": date(2024, 10, 22), "day_number": i + 1, "activity": f"Work {i}"} for i in range(11)]
    assert db.save_lookahead_records(1, 1, records)
    rows = db.get_seven_days_lookahead(report_id=1)
    assert len(rows) == 11
    original_id = rows[0]["id"]
    rows[0]["activity"] = "Edited work"
    assert db.save_lookahead_records(1, 1, rows)
    loaded = db.get_seven_days_lookahead(report_id=1)
    assert len(loaded) == 11 and loaded[0]["id"] == original_id
    assert all(row["plan_date"] == date(2024, 10, 22) and row["actual_start"] is None for row in loaded)
    loaded[0]["plan_date"] = "nonsense"
    result = db.save_lookahead_records(1, 1, loaded[:1])
    assert result.status == "INVALID_SOURCE" and len(db.get_seven_days_lookahead(report_id=1)) == 11
    assert db.save_lookahead_records(1, 1, [])
    with db.session_scope() as session:
        assert session.query(SevenDaysLookahead).count() == 0


def test_real_3d_draw_boundary_and_no_data(tmp_path):
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.figure import Figure
    from core.trajectory_plot import draw_trajectory_3d
    from core.survey_records import prepare_surveys
    figure = Figure()
    axes = figure.add_subplot(111, projection="3d")
    assert draw_trajectory_3d(axes, [])["status"] == "NO_DATA"
    assert draw_trajectory_3d(axes, [{"md": 100, "inc": 10, "azi": None}])["status"] == "INSUFFICIENT_DATA"
    complete, _, _ = prepare_surveys(stations())
    result = draw_trajectory_3d(axes, complete)
    assert result["status"] == "READY" and len(axes.lines) == 1
    east, north, tvd = axes.lines[0].get_data_3d()
    assert list(east) == result["series"]["east"] and list(north) == result["series"]["north"]
    assert list(tvd) == result["series"]["tvd"]
    figure.savefig(tmp_path / "trajectory.png")
    assert (tmp_path / "trajectory.png").stat().st_size > 1000
    complete[0]["north"] = float("nan")
    assert len(draw_trajectory_3d(axes, complete)["series"]["md"]) == 1


def test_actual_main_window_save_all_method_no_false_success():
    """Execute actual coordinator method with widget protocol fakes, not Qt."""
    import ast
    from types import SimpleNamespace
    source = (Path(__file__).resolve().parents[1] / "main_window.py").read_text()
    tree = ast.parse(source)
    method = next(n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef) and n.name == "save_all_tabs")
    namespace = {}
    exec(compile(ast.Module(body=[method], type_ignores=[]), "main_window.py", "exec"), namespace)
    messages = []
    tabs = [SimpleNamespace(save_data=lambda: True), SimpleNamespace(save_data=lambda: False)]
    owner = SimpleNamespace(tab_widget=SimpleNamespace(count=lambda: 2, widget=lambda i: tabs[i], tabText=lambda i: ["Mud", "POB"][i]),
        status_manager=SimpleNamespace(show_success=lambda *args: messages.append(("success", args)),
                                       show_error=lambda *args: messages.append(("error", args))),
        _invalidate_hierarchy_cache=lambda: None)
    assert namespace["save_all_tabs"](owner) is False
    assert messages[0][0] == "error" and "POB" in messages[0][1][1]


def test_whitespace_does_not_defeat_formula_only_classification():
    from openpyxl import Workbook
    from core.excel_intelligence import DynamicTableExtractor, MergeCellAnalyzer
    workbook = Workbook(); sheet = workbook.active
    sheet.append(["=1+1", "\n"])
    cells = {(1, 1): "=1+1", (1, 2): "\n"}
    extractor = DynamicTableExtractor(cells, MergeCellAnalyzer(sheet))
    assert extractor._classify_row(1, [{"col": 1}, {"col": 2}]) == "formula_only"


def test_chemical_closing_balance_never_invents_opening_or_overwrites_stock():
    from core.mud_records import chemical_balance
    source = {"product": "Barite", "received": 0, "used": 2, "stock": 45}
    assert chemical_balance(source) == {"status": "REVIEW_REQUIRED", "field": "opening_stock", "closing_stock": None}
    assert source["stock"] == 45
    source.update(opening_stock=47, returned=0, adjusted=0)
    assert chemical_balance(source)["closing_stock"] == 45
    source["opening_stock"] = "N.C"
    assert chemical_balance(source)["status"] == "REVIEW_REQUIRED"
    source["opening_stock"] = "invalid quantity"
    assert chemical_balance(source)["status"] == "INVALID_SOURCE"


def test_equipment_identity_repeat_and_actionable_hours(db):
    row = {"well_id": 1, "report_id": 1, "equipment_type": "Pump", "equipment_name": "Pump A", "equipment_id": "serial-A", "hours_worked": 12}
    identity = db.save_equipment_log(row)
    assert db.save_equipment_log(dict(row, hours_worked=13)) == identity
    assert len(db.get_equipment_logs(report_id=1)) == 1
    with pytest.raises(ValueError, match="hours_worked"):
        db.save_equipment_log(dict(row, hours_worked="NaN"))
    assert db.get_equipment_logs(report_id=1)[0]["hours_worked"] == 13


def test_pob_partial_update_preserves_dates_and_invalid_count(db):
    row = {"well_id": 1, "report_id": 1, "company_name": "Crew", "personnel_count": 3, "date_in": date(2024, 1, 1)}
    identity = db.save_service_company_pob(row)
    db.save_service_company_pob({"id": identity, "remarks": "Edited only remarks"})
    assert db.get_service_company_pob(report_id=1)[0]["date_in"] == date(2024, 1, 1)
    with pytest.raises(ValueError, match="personnel_count"):
        db.save_service_company_pob(dict(row, personnel_count=-1))


def test_validator_dictionary_issues_propagate_through_save_all():
    from core.validators import MudValidator
    from core.save_outcome import validation_outcome
    validation = MudValidator.validate({"solid_percent": 7, "oil_percent": None, "water_percent": None})
    result = validation_outcome("Mud", validation, saved=1)
    assert result.status == "REVIEW_REQUIRED" and result.saved == 1
    assert result.issues[0].field and result.issues[0].reason


@pytest.mark.integration
def test_actual_golden_lifecycle_save_all_and_clean_replay(tmp_path):
    from tools.certify_ddr_lifecycle import certify_lifecycle
    workbook = Path(__file__).resolve().parents[1] / "08-DDR OEOC-208 AZNS-207 2024-Oct-22.xlsx"
    if not workbook.is_file():
        pytest.skip("Actual Golden A workbook is unavailable")
    evidence = certify_lifecycle(workbook, tmp_path / "real-a-lifecycle")
    assert evidence["save_all"]["status"] == "REVIEW_REQUIRED"
    assert evidence["edit_reload_equal"] and evidence["clean_source_replay_equal"]
    assert evidence["survey_calculation"]["calculated"] == 0
    assert evidence["rendered_3d_agg"]["status"] == "INSUFFICIENT_DATA"


def test_equipment_snapshot_deletion_clear_and_atomic_invalid_row(db):
    rows = [{"equipment_name": "A", "hours_worked": 1}, {"equipment_name": "B", "hours_worked": 2}]
    assert db.save_equipment_records(1, 1, "Pump", rows)
    rows[0]["hours_worked"] = 9
    rows[1]["hours_worked"] = "bad"
    with pytest.raises(ValueError):
        db.save_equipment_records(1, 1, "Pump", rows)
    assert db.get_equipment_logs(report_id=1)[0]["hours_worked"] == 1
    assert db.save_equipment_records(1, 1, "Pump", rows[:1])
    assert len(db.get_equipment_logs(report_id=1)) == 1
    assert db.save_equipment_records(1, 1, "Pump", [])
    assert db.get_equipment_logs(report_id=1) == []
