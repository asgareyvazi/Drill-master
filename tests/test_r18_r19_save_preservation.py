"""R18/R19: actual BaseTab/Downhole methods + production SQLite, not native Qt.

The protocol supplies only widget APIs. The production bindings, dirty tracker,
coordinator, parsers, DB services/repositories and Excel exporter are unchanged.
"""

import ast
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from datetime import date
import json

import pytest
from sqlalchemy import event

from core import editor_bindings
from core.database import BHAReport, DownholeEquipment, FormationReport
from core.editor_state import EditSection
from core.legacy_bha import LegacyBHAReadOnlyError
from core.repositories.base import BaseRepository
from tests.test_real_user_acceptance_regressions import (
    production_db as production_db,
    downhole_editor,
    actual_class,
    Table,
)


def snapshot(*roots):
    def read(widget):
        if isinstance(widget, Table):
            return tuple(tuple(item.text() if item else None for item in row) for row in widget.rows)
        return widget.text()

    return tuple(read(root) for root in roots)


@pytest.fixture
def loaded(production_db, monkeypatch):
    db = production_db
    obj, _ = downhole_editor(db)
    db.save_bha_report(
        obj.current_well,
        {
            "report_id": obj.current_report_id,
            "bha_name": "Current BHA",
            "bha_data": [{"component_name": "Bit", "length": 0.3, "od": 8.5, "weight": 0, "id": None}],
        },
    )
    obj.load_all_data_from_db()
    monkeypatch.setattr(editor_bindings, "editor_snapshot", snapshot)
    # No Qt validator exists in this protocol; numeric table validation remains
    # the actual BHAManager/domain_records validation, exercised below.
    monkeypatch.setattr(editor_bindings, "validate_editor_inputs", lambda *args: None)
    obj.configure_save_tracking()
    return obj


def stored(obj):
    return {
        model.__name__: obj.db.generic_get_list(model, {"report_id": obj.current_report_id})
        for model in (BHAReport, DownholeEquipment, FormationReport)
    }


def edit_length(obj, value="0.45"):
    col = next(i for i, header in enumerate(obj.bha_table.headers) if header.text() == "Length (m)")
    obj.bha_table.item(0, col).setText(value)


def edit_equipment(obj):
    obj.equipment_manager.load_data([{"equipment_name": "MWD", "rotation_hours": 0, "install_date": None}])


def main_method(name):
    path = Path(__file__).resolve().parents[1] / "main_window.py"
    tree = ast.parse(path.read_text())
    cls = next(node for node in tree.body if isinstance(node, ast.ClassDef) and node.name == "MainWindow")
    method = next(node for node in cls.body if isinstance(node, ast.FunctionDef) and node.name == name)
    namespace = {}
    exec(compile(ast.Module(body=[method], type_ignores=[]), str(path), "exec"), namespace)
    return namespace[name]


def global_save(obj, *others):
    tabs = (obj, *others)
    window = SimpleNamespace(
        tab_widget=SimpleNamespace(
            count=lambda: len(tabs), widget=lambda i: tabs[i], tabText=lambda i: type(tabs[i]).__name__
        ),
        status_manager=SimpleNamespace(show_success=lambda *args: None, show_error=lambda *args: None),
        _invalidate_hierarchy_cache=lambda: None,
    )
    main_method("save_all_tabs")(window)
    return window.last_save_outcome


def auto_save(obj):
    manager = actual_class("core/managers.py", "AutoSaveManager")
    return manager.save_widget("Downhole", obj)


@pytest.mark.parametrize("route", [global_save, auto_save])
def test_clean_never_invokes_saver_or_mutates_db(loaded, route):
    before = stored(loaded)
    for section in loaded._edit_sections.values():
        section.writer = lambda: pytest.fail("Clean saver invoked")
    result = route(loaded)
    assert result.disposition == "NO_CHANGES" and result.saved == 0
    assert "Nothing to save" in result.summary()
    assert stored(loaded) == before


@pytest.mark.parametrize("route", [global_save, auto_save])
def test_one_dirty_valid_saves_only_it_and_clears(loaded, route):
    before = stored(loaded)
    loaded._edit_sections["Formation"].writer = lambda: pytest.fail("Clean formation saved")
    loaded._edit_sections["Downhole Equipment"].writer = lambda: pytest.fail("Clean equipment saved")
    edit_length(loaded)
    assert loaded.is_dirty
    result = route(loaded)
    assert result.disposition == "SAVED" and result.saved == 1 and not loaded.is_dirty
    assert stored(loaded)["BHAReport"] != before["BHAReport"]
    assert stored(loaded)["DownholeEquipment"] == before["DownholeEquipment"]
    assert route(loaded).disposition == "NO_CHANGES"


def test_two_valid_dirty_sections(loaded):
    edit_length(loaded)
    edit_equipment(loaded)
    result = global_save(loaded)
    assert result.saved == 2 and result.counts["SAVED"] == 2
    assert not loaded.is_dirty
    assert len(stored(loaded)["BHAReport"]) == len(stored(loaded)["DownholeEquipment"]) == 1
    assert stored(loaded)["FormationReport"] == []


def test_invalid_and_valid_are_honest_partial_success(loaded):
    before = stored(loaded)
    edit_length(loaded, "invalid-length")
    edit_equipment(loaded)
    result = global_save(loaded)
    assert not result and result.disposition == "VALIDATION_ERROR" and result.saved == 1
    assert "Length" in result.summary()
    assert loaded._edit_sections["BHA"].dirty
    assert not loaded._edit_sections["Downhole Equipment"].dirty
    assert stored(loaded)["BHAReport"] == before["BHAReport"]
    assert stored(loaded)["DownholeEquipment"]


def test_cancelled_dialog_no_mutation_no_false_dirty(loaded):
    from types import SimpleNamespace
    import sys

    before = stored(loaded)
    # Actual add_bha_tool method must not copy cancelled dialog values.
    previous = sys.modules.get("dialogs.drilling_report_dialogs")
    sys.modules["dialogs.drilling_report_dialogs"] = SimpleNamespace(
        AddBHAComponentDialog=lambda *_: SimpleNamespace(exec=lambda: False)
    )
    try:
        loaded.add_bha_tool()
    finally:
        if previous is None:
            del sys.modules["dialogs.drilling_report_dialogs"]
        else:
            sys.modules["dialogs.drilling_report_dialogs"] = previous
    assert not loaded.is_dirty and global_save(loaded).saved == 0
    assert stored(loaded) == before


def test_failed_child_context_explicitly_blocks_old_report_write(loaded):
    db = loaded.db
    old = stored(loaded)
    edit_length(loaded)
    new = db.save_daily_report(
        {
            "well_id": loaded.current_well,
            "section_id": loaded.current_section,
            "report_date": date(2026, 9, 9),
            "report_number": 2,
        }
    )["id"]
    loaded._on_report_changed_internal(new, {})

    def failed(*args):
        raise RuntimeError("Injected child report load failure")

    loaded.on_report_changed = failed
    assert not loaded._apply_pending_context()
    for section in loaded._edit_sections.values():
        section.writer = lambda: pytest.fail("Stale context saver invoked")
    result = global_save(loaded)
    assert result.disposition == "CONTEXT_BLOCKED" and result.saved == 0
    assert "Downhole" in result.summary() and loaded.is_dirty
    assert all(not rows for rows in stored(loaded).values())
    assert db.generic_get_list(BHAReport, {"report_id": old["BHAReport"][0]["report_id"]}) == old["BHAReport"]


def test_hidden_clean_is_no_change(loaded):
    loaded.visible = False
    before = stored(loaded)
    assert global_save(loaded).disposition == "NO_CHANGES"
    assert stored(loaded) == before


def test_read_only_callback_existence_is_not_dirty(loaded):
    loaded._edit_sections = {}
    loaded.save_data = lambda: pytest.fail("Read-only callback invoked")
    before = stored(loaded)
    assert global_save(loaded).disposition == "NO_CHANGES" and not loaded.is_dirty
    assert stored(loaded) == before


def test_failed_save_retains_dirty_then_success_acknowledges(loaded):
    before = stored(loaded)
    edit_length(loaded)
    section = loaded._edit_sections["BHA"]
    writer = section.writer

    def failed():
        raise RuntimeError("Injected unavailable persistence")

    section.writer = failed
    result = global_save(loaded)
    assert result.disposition == "SYSTEM_ERROR" and section.dirty
    assert stored(loaded) == before
    section.writer = writer
    assert global_save(loaded).saved == 1 and not section.dirty


def test_save_reopen_exact_null_zero_and_source_provenance(loaded):
    edit_length(loaded)
    from core.domain_records import bha_records

    expected = bha_records(loaded.bha_manager.get_all_data())
    assert global_save(loaded).saved == 1
    loaded.db.close()
    assert loaded.db.initialize()
    loaded.load_all_data_from_db()
    assert loaded.bha_manager.get_all_data() == expected
    rows = loaded.db.get_bha_report(loaded.current_well, loaded.current_report_id)["bha_configs"]
    assert rows[0]["ID (in)"] is None and rows[0]["Weight (kg)"] == 0
    assert not loaded.is_dirty and global_save(loaded).saved == 0


def test_invalid_never_becomes_null(loaded):
    before = stored(loaded)
    edit_length(loaded, "not a number")
    result = auto_save(loaded)
    assert result.disposition == "VALIDATION_ERROR" and loaded.is_dirty
    assert stored(loaded) == before


def test_no_dml_for_all_clean_with_sql_observer(loaded):
    writes = []

    def observe(conn, cursor, statement, params, context, many):
        if statement.lstrip().split()[0].upper() in {"INSERT", "UPDATE", "DELETE"}:
            writes.append(statement)

    event.listen(loaded.db.engine, "before_cursor_execute", observe)
    try:
        assert global_save(loaded).saved == auto_save(loaded).saved == 0
    finally:
        event.remove(loaded.db.engine, "before_cursor_execute", observe)
    assert writes == []


LEGACY = {
    "Assembly A – drilling": [
        {
            "component_name": "Bit",
            "length": 0.3,
            "od": 8.5,
            "id": None,
            "weight": 0,
            "_provenance": {"source_record": {"configuration": "Assembly A – drilling", "line": 1}},
        },
        {"component_name": "Drill collar", "length": 9.5, "od": 6.5, "id": 2.25, "weight": None},
    ],
    "Assembly B – survey": [
        {
            "component_name": "MWD",
            "length": 8.8,
            "od": 6.75,
            "id": 0,
            "weight": None,
            "_provenance": {"source_record": {"configuration": "Assembly B – survey", "line": 1}},
        },
        {"component_name": "Bit", "length": 0.4, "od": 8.375, "id": None, "weight": 0},
    ],
}


@pytest.fixture(params=[False, True], ids=["json-map", "historical-json-string"])
def legacy(production_db, request):
    db = production_db
    obj, _ = downhole_editor(db)
    original = json.dumps(LEGACY, ensure_ascii=False) if request.param else deepcopy(LEGACY)
    identity = db.generic_save(
        BHAReport,
        {
            "well_id": obj.current_well,
            "report_id": obj.current_report_id,
            "bha_name": "Original report BHA",
            "bha_data_json": original,
        },
    )
    obj.load_all_data_from_db()
    return obj, identity, original


def test_legacy_list_select_inspect_read_only_restart(legacy):
    obj, identity, original = legacy
    for _ in range(2):
        assert [key for label, key in obj.bha_selector.items if key is not None] == list(LEGACY)
        assert obj.bha_name_input.read_only and obj.bha_table.edit_triggers == 0
        assert not obj.save_bha_btn.enabled and not obj.delete_bha_btn.enabled
        for name, source in LEGACY.items():
            obj.bha_selector.setCurrentText(name)
            obj.load_bha_config()
            assert obj.bha_name_input.text() == name
            actual = obj.bha_manager.get_all_data()
            assert [row["Length (m)"] for row in actual] == [row["length"] for row in source]
            assert [row["OD (in)"] for row in actual] == [row["od"] for row in source]
            assert [row["ID (in)"] for row in actual] == [row["id"] for row in source]
            assert [row["Weight (kg)"] for row in actual] == [row["weight"] for row in source]
            assert not obj.add_bha_tool() and not obj.remove_bha_tool()
            assert not obj.save_bha_config() and not obj.delete_bha_config()
            assert obj.db.get_bha_report(obj.current_well, obj.current_report_id)["bha_configs"] == original
        obj.db.close()
        assert obj.db.initialize()
        obj.load_all_data_from_db()


@pytest.mark.parametrize(
    "route",
    [
        "service",
        "generic-update",
        "repository-update",
        "generic-delete",
        "repository-delete",
        "generic-insert",
        "repository-insert",
    ],
)
def test_legacy_destructive_bypass_blocked(legacy, route):
    obj, identity, original = legacy
    db = obj.db
    values = {"id": identity, "bha_name": "Replacement", "bha_data_json": []}
    insert = {"well_id": obj.current_well, "report_id": obj.current_report_id, "bha_data_json": []}
    actions = {
        "service": lambda: db.save_bha_report(obj.current_well, {"report_id": obj.current_report_id, "bha_data": []}),
        "generic-update": lambda: db.generic_save(BHAReport, values),
        "repository-update": lambda: BaseRepository(db).save(BHAReport, values),
        "generic-delete": lambda: db.generic_delete(BHAReport, identity),
        "repository-delete": lambda: BaseRepository(db).delete(BHAReport, identity),
        "generic-insert": lambda: db.generic_save(BHAReport, insert),
        "repository-insert": lambda: BaseRepository(db).save(BHAReport, insert),
    }
    with pytest.raises(LegacyBHAReadOnlyError, match="read-only"):
        actions[route]()
    assert db.get_bha_report(obj.current_well, obj.current_report_id)["bha_configs"] == original
    assert len(db.generic_get_list(BHAReport, {"report_id": obj.current_report_id})) == 1


@pytest.mark.parametrize(
    "payload", [{"bha_report": {"bha_data": []}}, {"bha_components": [{"component_name": "Bit", "length": 0.5}]}]
)
def test_legacy_import_replacement_rolls_back(legacy, payload):
    obj, _, original = legacy
    from core.import_diagnostics import PersistenceError

    with pytest.raises(PersistenceError, match="read-only"):
        obj.db.save_imported_multi_tab_data_atomic(obj.current_well, obj.current_report_id, payload)
    assert obj.db.get_bha_report(obj.current_well, obj.current_report_id)["bha_configs"] == original


def test_legacy_excel_contains_both_names_and_lossless_raw_archive(legacy, tmp_path):
    from core.professional_export import ProfessionalExcelExport
    from openpyxl import load_workbook

    obj, identity, original = legacy
    target = tmp_path / "legacy.xlsx"
    assert ProfessionalExcelExport(obj.db).export(obj.current_well, str(target), report_id=obj.current_report_id)
    wb = load_workbook(target)
    assert [row[0] for row in wb["BHA"].iter_rows(min_row=2, values_only=True)] == [
        name for name, rows in LEGACY.items() for row in rows
    ]
    chunks = [
        row[2]
        for row in wb["Raw Data"].iter_rows(min_row=3, values_only=True)
        if row[0] == "BHAReport" and row[1] == identity
    ]
    assert json.loads("".join(chunks))["bha_data_json"] == original
    wb.close()
    assert obj.db.get_bha_report(obj.current_well, obj.current_report_id)["bha_configs"] == original


def test_successful_local_save_is_not_repeated_by_global_or_timer(loaded):
    edit_length(loaded)
    assert loaded.save_bha_config()
    before = stored(loaded)
    loaded._edit_sections["BHA"].writer = lambda: pytest.fail("Already persisted local save repeated")
    assert global_save(loaded).saved == auto_save(loaded).saved == 0
    assert stored(loaded) == before


def test_reverted_edit_is_clean_without_write(loaded):
    before = stored(loaded)
    edit_length(loaded)
    assert loaded.is_dirty
    edit_length(loaded, "0.3")
    assert not loaded.is_dirty and global_save(loaded).saved == 0
    assert stored(loaded) == before


def test_null_survives_unrelated_dirty_field(loaded):
    edit_length(loaded)
    assert global_save(loaded).saved == 1
    rows = loaded.db.get_bha_report(loaded.current_well, loaded.current_report_id)["bha_configs"]
    assert rows[0]["ID (in)"] is None
    loaded.load_all_data_from_db()
    assert loaded.bha_manager.get_all_data()[0]["ID (in)"] is None


def test_valid_explicit_zero_is_not_missing(loaded):
    weight = next(i for i, h in enumerate(loaded.bha_table.headers) if h.text() == "Weight (kg)")
    loaded.bha_table.item(0, weight).setText("0")
    assert global_save(loaded).saved == 1
    loaded.load_all_data_from_db()
    assert loaded.bha_manager.get_all_data()[0]["Weight (kg)"] == 0


def test_successful_parent_with_failed_child_snapshot_stays_blocked(loaded):
    before = stored(loaded)
    edit_length(loaded)
    section = loaded._edit_sections["BHA"]
    section.loaded_context = (loaded.current_well, -1)
    # Parent IDs are current; a child that did not load cannot borrow its readiness.
    assert loaded.save_context_ready()
    section.writer = lambda: pytest.fail("Unloaded child saver invoked")
    result = global_save(loaded)
    assert result.disposition == "CONTEXT_BLOCKED" and "BHA" in result.summary()
    assert stored(loaded) == before and section.dirty


def test_same_context_failed_reload_does_not_acknowledge_edits(loaded):
    from core.editor_state import editor_loaded

    before = stored(loaded)
    edit_length(loaded)

    @editor_loaded("BHA")
    def failed_load(self):
        raise RuntimeError("Reload unavailable")

    with pytest.raises(RuntimeError):
        failed_load(loaded)
    assert loaded._edit_sections["BHA"].dirty
    assert global_save(loaded).disposition == "CONTEXT_BLOCKED"
    assert stored(loaded) == before


def test_well_header_save_preserves_context_and_partial_invalid_sibling(loaded):
    from tests.test_real_user_acceptance_regressions import editor

    base = type(editor())
    cls = actual_class("tabs/w1_well_info.py", "WellInfoTab", DrillTabBase=base, Signal=lambda: None, QComboBox=object)
    well = cls.__new__(cls)
    well.__dict__.update(editor().__dict__)
    well.widget_name = "Well"
    well.db = loaded.db
    well.current_well = well.db.get_well_by_id(loaded.current_well)
    for kind in ("well", "section", "report"):
        identity = getattr(loaded, f"current_{kind}_id")
        setattr(well, f"current_{kind}_id", identity)
        setattr(well, f"_loaded_{kind}_id", identity)
    values = deepcopy(well.current_well)
    well._loaded_form_values = deepcopy(values)
    well.get_form_data = lambda: deepcopy(values)

    def forbidden(*args):
        pytest.fail("Saving existing Well changed selection/reloaded sibling editors")

    well.main_window = SimpleNamespace(populate_hierarchy=forbidden, sel_manager=SimpleNamespace(select_well=forbidden))
    well.data_saved = SimpleNamespace(emit=lambda: None)
    section = EditSection("Well", lambda: values, well.save_data)
    well._tracked_edit_sections = [section]
    well._edit_sections = {"Well": section}
    values["name"] = "Changed well header"
    edit_length(loaded, "bad length")
    before_bha = stored(loaded)
    result = global_save(well, loaded)
    assert result.saved == 1 and result.disposition == "VALIDATION_ERROR"
    assert well.db.get_well_by_id(loaded.current_well)["name"] == "Changed well header"
    assert stored(loaded) == before_bha and loaded.is_dirty
    assert well.current_report_id == loaded.current_report_id


def test_legacy_selection_never_causes_dirty_save_but_equipment_can_save(legacy, monkeypatch):
    obj, _, original = legacy
    monkeypatch.setattr(editor_bindings, "editor_snapshot", snapshot)
    monkeypatch.setattr(editor_bindings, "validate_editor_inputs", lambda *args: None)
    obj.configure_save_tracking()
    for name in LEGACY:
        obj.bha_selector.setCurrentText(name)
        obj.load_bha_config()
        assert not obj.is_dirty and global_save(obj).saved == 0
    edit_equipment(obj)
    assert global_save(obj).saved == 1
    assert obj.db.get_bha_report(obj.current_well, obj.current_report_id)["bha_configs"] == original


def test_legacy_json_archive_all_alternatives_and_cancel(legacy, tmp_path):
    obj, _, original = legacy
    path = tmp_path / "all-configurations.json"
    namespace = obj.export_bha_data.__globals__
    namespace["QFileDialog"] = SimpleNamespace(getSaveFileName=lambda *args: ("", ""))
    assert obj.export_bha_data() is False and not path.exists()
    namespace["QFileDialog"] = SimpleNamespace(getSaveFileName=lambda *args: (str(path), "JSON"))
    assert obj.export_bha_data()
    assert json.loads(path.read_text())["bha_configs"] == original
    assert obj.db.get_bha_report(obj.current_well, obj.current_report_id)["bha_configs"] == original


def test_unsupported_dirty_editor_does_not_claim_saved_or_clear(loaded):
    before = stored(loaded)
    edit_length(loaded)
    loaded._edit_sections["BHA"].writer = lambda: editor_bindings.unsupported_editor(
        "No lossless persistence path; edit retained"
    )
    result = global_save(loaded)
    assert result.disposition == "UNSUPPORTED" and result.saved == 0 and not result
    assert loaded.is_dirty and stored(loaded) == before
