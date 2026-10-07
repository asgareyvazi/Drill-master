"""Deterministic file-backed release scenario using a generated synthetic workbook."""
from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from core.database import AuditLog, Company, DailyReport, DatabaseManager, Project, Section, Well, Wellbore
from core.ddr_import_service import DDRImportService
from core.engineering.engines.mse import MSEEngine
from core.professional_export import ProfessionalExcelExport
from core.repositories.mse_repository import MSECalculationRepository


MSE_INPUTS = {
    "wob_lbf": 21000.0,
    "rpm": 95.0,
    "torque_ft_lbf": 6400.0,
    "rop_ft_hr": 22.0,
    "bit_diameter_in": 8.5,
}


def _make_synthetic_workbook(path: Path) -> None:
    workbook = Workbook()
    context = workbook.active
    context.title = "Report context"
    for pair in (
        ("daily_report.report_date", date(2026, 10, 7)),
        ("daily_report.report_number", 42),
        ("well_info.name", "M42 Synthetic Well"),
        ("well_info.wellbore_name", "Original"),
        ("well_info.section_name", "12-1/4 in"),
        ("well_info.project_name", "M42 Temporary Project"),
    ):
        context.append(pair)
    context.append(["Persian review note", "گزارش آزمایشی"])

    survey = workbook.create_sheet("Measurements")
    survey.append(["MD (m)", "Inclination", "Azimuth"])
    survey.append([100, 5, 15])
    survey.append([200, 10, 30])
    survey.row_dimensions[3].hidden = True
    survey["F1"] = "unrelated separated cell"
    tools = workbook.create_sheet("Tools")
    tools.append(["Component Name", "OD (in)", "Length (m)"])
    tools.append(["Bit", 8.5, 0.3])
    workbook.save(path)
    workbook.close()


def test_disposable_hierarchy_synthetic_import_engineering_export_and_restore(tmp_path, monkeypatch):
    data_root = tmp_path / "first-user-data"
    db_path = data_root / "drillmaster.sqlite"
    monkeypatch.setenv("DRILLMASTER_ENV", "test")
    monkeypatch.delenv("DRILLMASTER_ENVIRONMENT", raising=False)
    monkeypatch.setenv("DRILLMASTER_DATA_DIR", str(data_root))
    monkeypatch.setenv("DRILLMASTER_DB_PATH", str(db_path))
    monkeypatch.setenv("DRILLMASTER_AI_IMPORT", "0")

    manager = DatabaseManager()
    assert manager.initialize(), manager.last_diagnostic
    with manager.session_scope() as session:
        company = Company(name="M42 Temporary Company", code="M42-C")
        session.add(company)
        session.flush()
        project = Project(company_id=company.id, name="M42 Temporary Project", code="M42-P")
        session.add(project)
        session.flush()
        well = Well(project_id=project.id, name="M42 Synthetic Well", code="M42-W")
        session.add(well)
        session.flush()
        original = Wellbore(well_id=well.id, name="Original", wellbore_type="original")
        session.add(original)
        session.flush()
        section = Section(well_id=well.id, wellbore_id=original.id, name="12-1/4 in", depth_from=0, depth_to=1000)
        session.add(section)
        session.flush()
        well_id, original_bore_id, section_id = well.id, original.id, section.id

    workbook_path = tmp_path / "synthetic-ddr.xlsx"
    _make_synthetic_workbook(workbook_path)
    importer = DDRImportService(manager, well_id)
    extraction, proposed = importer.extract_file(str(workbook_path), template={})
    assert extraction.raw_document is not None
    assert extraction.raw_document.source_file == str(workbook_path)
    assert extraction.template_version == "generic-ir"
    assert proposed["metadata"]["raw_ir"]
    assert proposed["well_info"]["wellbore_name"] == "Original"
    assert proposed["metadata"]["review_matrix"], "generic extraction must go through explicit review"
    review_order = [
        (row["target_field"], row.get("source_location"), row.get("original_value"))
        for row in proposed["metadata"]["review_matrix"]
    ]
    repeated_extraction, repeated = importer.extract_file(str(workbook_path), template={})
    assert [
        (row["target_field"], row.get("source_location"), row.get("original_value"))
        for row in repeated["metadata"]["review_matrix"]
    ] == review_order
    assert repeated_extraction.raw_document.to_dict(include_cells=True) == extraction.raw_document.to_dict(include_cells=True)
    for review_item in proposed["metadata"]["review_matrix"]:
        review_item["decision"] = "ACCEPT"  # explicit operator-confirmation simulation

    result = importer.import_records(proposed)
    assert result["failed"] == 0 and result["report_id"]
    report_id = result["report_id"]
    report = manager.get_daily_report_by_id(report_id)
    assert report["report_number"] == 42
    assert report["report_date"] == date(2026, 10, 7)
    surveys = manager.load_survey_points(report_id=report_id)
    assert [row["md"] for row in surveys] == [100, 200]
    assert len(manager.get_bha_report(well_id, report_id=report_id)["bha_configs"]) == 1

    # Add a same-named sidetrack section after import and prove the persisted
    # report remains on its original bore instead of being re-attributed by name.
    with manager.session_scope() as session:
        sidetrack = Wellbore(well_id=well_id, name="Sidetrack", wellbore_type="sidetrack",
                             parent_wellbore_id=original_bore_id, kickoff_md=600)
        session.add(sidetrack)
        session.flush()
        session.add(Section(well_id=well_id, wellbore_id=sidetrack.id, name="12-1/4 in"))
        session.flush()
        same_named_sections = session.query(Section).filter_by(well_id=well_id, name="12-1/4 in").all()
        assert len(same_named_sections) >= 2
        assert any(item.id == section_id and item.wellbore_id == original_bore_id for item in same_named_sections)
        assert any(item.wellbore_id == sidetrack.id for item in same_named_sections)
        stored_report = session.get(DailyReport, report_id)
        assert stored_report.section_id is not None
        assert session.get(Section, stored_report.section_id).wellbore_id == original_bore_id
        assert stored_report.wellbore_id == original_bore_id

    mse = MSEEngine.calculate(**MSE_INPUTS)
    assert mse.success
    history_id = MSECalculationRepository(manager).save_run(
        inputs=MSE_INPUTS, result_values=mse.values, method=MSEEngine.METHOD,
        label="M42 synthetic release scenario", well_id=well_id,
    )
    export_path = tmp_path / "synthetic-report.xlsx"
    assert ProfessionalExcelExport(manager).export(well_id, str(export_path), report_id=report_id)
    exported = load_workbook(export_path, data_only=True, read_only=True)
    try:
        assert {"Executive Summary", "Daily Report", "Audit"} <= set(exported.sheetnames)
        daily = exported["Daily Report"]
        headers = {cell.value: cell.column for cell in daily[1]}
        assert "report_number" in headers
        assert daily.cell(2, headers["report_number"]).value == 42
    finally:
        exported.close()

    backup_path = tmp_path / "external" / "release-scenario.db"
    assert manager.backup_to(backup_path) == str(backup_path)
    manager.close()
    restore_root = tmp_path / "restored-user-data"
    restore_root.mkdir()
    restored_path = restore_root / "drillmaster.sqlite"
    shutil.copy2(backup_path, restored_path)
    monkeypatch.setenv("DRILLMASTER_DATA_DIR", str(restore_root))
    monkeypatch.setenv("DRILLMASTER_DB_PATH", str(restored_path))
    restored = DatabaseManager()
    assert restored.initialize(), restored.last_diagnostic
    try:
        with restored.session_scope() as session:
            assert session.query(Company).filter_by(name="M42 Temporary Company").count() == 1
            assert session.query(Project).filter_by(name="M42 Temporary Project").count() == 1
            assert session.query(Well).filter_by(id=well_id).count() == 1
            restored_report = session.get(DailyReport, report_id)
            assert restored_report.wellbore_id == original_bore_id
            assert session.query(AuditLog).filter_by(action="ddr_import", entity_id=report_id).count() >= 1
        assert restored.load_survey_points(report_id=report_id) == surveys
        saved = MSECalculationRepository(restored).get(history_id)
        assert saved is not None and saved.result["mse_psi"] == mse.values["mse_psi"]
        assert restored.authenticate_user("admin", "admin123")
    finally:
        restored.close()
