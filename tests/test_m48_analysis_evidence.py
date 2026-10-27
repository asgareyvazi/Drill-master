"""M48 - core contracts of core.analysis_evidence (Qt-free).

Golden values are derived by hand in the comments (or by an independent
implementation such as numpy.polyfit) - never by calling the function under
test to produce its own expectation.
"""
import csv
import math
import random
from datetime import date, timedelta

import numpy as np
import pytest

from core import analysis_evidence as ev

D = date(2026, 1, 1)


def day(n):
    return D + timedelta(days=n - 1)


# ---------------------------------------------------------------------------
# Value states: 0.0 is a value, None is missing, NaN/inf/negative are invalid
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("raw, state, value", [
    (0, ev.STATE_VALUE, 0.0),
    (0.0, ev.STATE_VALUE, 0.0),
    (-0.0, ev.STATE_VALUE, 0.0),
    ("7.5", ev.STATE_VALUE, 7.5),
    (None, ev.STATE_MISSING, None),
    (float("nan"), ev.STATE_INVALID, None),
    (float("inf"), ev.STATE_INVALID, None),
    (float("-inf"), ev.STATE_INVALID, None),
    (-0.001, ev.STATE_INVALID, None),
    ("abc", ev.STATE_INVALID, None),
    (True, ev.STATE_INVALID, None),
])
def test_classify_number_boundaries(raw, state, value):
    got_state, got_value, reason = ev.classify_number(raw)
    assert (got_state, got_value) == (state, value)
    assert bool(reason) == (state != ev.STATE_VALUE)


def test_metric_text_never_turns_missing_into_zero_and_keeps_units():
    assert ev.metric(0.0, "m").text() == "0.0 m"
    assert ev.metric(None, "m").text() == "No data"           # no unit on a state word
    assert ev.metric(float("nan"), "m").state == ev.STATE_INVALID
    assert ev.metric(-2.0, "m").text() == "-2.0 m"             # allowed when the quantity may be negative
    assert ev.MetricValue(ev.STATE_ERROR).text() != ev.MetricValue(ev.STATE_MISSING).text()
    assert ev.MetricValue(ev.STATE_PARTIAL, 3.0, "h").text() == "3.0 (partial) h"


def test_every_kpi_definition_has_title_unit_and_definition():
    for key, (title, unit, definition) in ev.KPI_DEFINITIONS.items():
        assert title and unit and len(definition) > 20, key
        assert f"[{unit}]" in ev.kpi_tooltip(key)


# ---------------------------------------------------------------------------
# ROP
# ---------------------------------------------------------------------------

def test_summarize_rop_golden():
    # valid rows: 0.0 and 5.0 -> average 2.5, best 5.0;
    # None -> missing (1); NaN, -2, inf, "x" -> invalid (4)
    out = ev.summarize_rop([0.0, 5.0, None, float("nan"), -2.0, float("inf"), "x"])
    assert out == {"average": 2.5, "best": 5.0, "valid_count": 2,
                   "missing_count": 1, "invalid_count": 4}


def test_summarize_rop_all_zero_is_zero_not_missing():
    out = ev.summarize_rop([0.0, 0.0])
    assert out["average"] == 0.0 and out["best"] == 0.0


def test_summarize_rop_nothing_valid_is_none():
    assert ev.summarize_rop([None, float("nan")])["average"] is None


def test_midpoint_series_golden():
    # (10, 20) -> 15; (0, 0) -> 0 (real zero); missing / NaN / negative bound -> excluded
    pairs = [(10, 20), (0, 0), (None, 5), (float("nan"), 3), (-1, 4)]
    assert ev.midpoint_series(pairs) == [15.0, 0.0]


def _row(n, rop, report_id, bore=None, hours=None, source_id=None):
    return {"date": day(n), "rop": rop, "report_id": report_id, "wellbore_id": bore,
            "hours": hours, "source_id": source_id or report_id}


def test_rop_daily_series_same_date_hours_weighted_golden():
    # 10 m/hr for 2 h and 20 m/hr for 6 h -> (20 + 120) / 8 = 17.5 m/hr
    series, excluded = ev.rop_daily_series([_row(1, 10.0, 7, 1, 2.0, 1), _row(1, 20.0, 7, 1, 6.0, 2)])
    assert [(o.date, o.value, o.report_id) for o in series] == [(day(1), 17.5, 7)]
    assert excluded == []


def test_rop_daily_series_same_date_without_hours_is_ambiguous_not_double_counted():
    series, excluded = ev.rop_daily_series([_row(1, 10.0, 7, 1, None, 1), _row(1, 20.0, 8, 1, None, 2)])
    assert series == []
    assert len(excluded) == 1 and excluded[0].state == ev.STATE_AMBIGUOUS
    assert excluded[0].drilldown_report_id is None


def test_legacy_row_is_only_an_unambiguous_fallback():
    # legacy 99 on a date with a linked 5 -> excluded; legacy 8 alone -> used
    series, excluded = ev.rop_daily_series([
        _row(1, 5.0, 7), _row(1, 99.0, None, source_id=50), _row(2, 8.0, None, source_id=51)])
    assert [(o.date.day, o.value, o.report_id) for o in series] == [(1, 5.0, 7), (2, 8.0, None)]
    assert [(o.date.day, o.reason) for o in excluded] == [(1, ev.LEGACY_DUPLICATE_REASON)]
    # a linked row with an invalid rate still supersedes the legacy row
    series, excluded = ev.rop_daily_series([_row(3, -1.0, 9), _row(3, 4.0, None, source_id=52)])
    assert series == [] and {o.state for o in excluded} == {ev.STATE_INVALID, ev.STATE_AMBIGUOUS}


def test_canonical_average_rop_never_double_counts_a_legacy_date(syn):
    from core.operations_intelligence import OperationsIntelligenceService
    db, ids = syn
    # SYN-B rows: OH 5.0, 7.0; ST1 20.0, 30.0; legacy 99.0 on 2026-01-01 (a
    # date with linked report B1) -> excluded. Hand golden (5+7+20+30)/4.
    kpis = OperationsIntelligenceService(db).analyze_well(ids["B"])["kpis"]
    assert kpis["average_rop"] == 15.5


def test_canonical_average_rop_keeps_an_unambiguous_legacy_row():
    from datetime import date as _date
    from core.database import DrillingParameters
    from core.operations_intelligence import OperationsIntelligenceService
    from m48_synthetic import build
    db, ids = build()
    session = db.create_session()
    try:
        # SYN-C has no row on 2026-02-08: a legacy 8.0 there is unambiguous.
        session.add(DrillingParameters(well_id=ids["C"], report_id=None,
                                       report_date=_date(2026, 2, 8), avg_rop=8.0))
        session.commit()
    finally:
        session.close()
    kpis = OperationsIntelligenceService(db).analyze_well(ids["C"])["kpis"]
    assert kpis["average_rop"] == 4.5          # (1+2+...+7+8)/8 = 36/8


def test_rop_daily_series_two_wellbores_same_date_is_ambiguous_even_with_hours():
    series, excluded = ev.rop_daily_series([_row(1, 10.0, 7, 1, 2.0), _row(1, 20.0, 8, 2, 2.0)])
    assert series == [] and excluded[0].reason == "rows from several wellbores share this date"


def test_rop_daily_series_rejects_nan_inf_negative_with_reasons():
    rows = [_row(1, float("nan"), 1), _row(2, float("inf"), 2), _row(3, -1.0, 3),
            _row(4, None, 4), _row(5, 0.0, 5)]
    series, excluded = ev.rop_daily_series(rows)
    assert [(o.date, o.value) for o in series] == [(day(5), 0.0)]
    reasons = {o.date: (o.state, o.reason) for o in excluded}
    assert reasons[day(1)] == (ev.STATE_INVALID, "non-finite value (NaN/inf)")
    assert reasons[day(2)] == (ev.STATE_INVALID, "non-finite value (NaN/inf)")
    assert reasons[day(3)] == (ev.STATE_INVALID, "negative value")
    assert reasons[day(4)] == (ev.STATE_MISSING, "not recorded")


def test_ambiguous_observation_never_drills_down_even_with_a_report_id():
    # load_depth_observations keeps each report's id on an ambiguous date
    row = ev.Observation("depth", day(1), 5.0, "m", ev.STATE_AMBIGUOUS, "2 reports", report_id=5)
    assert row.drilldown_report_id is None
    assert row.drilldown_reason().startswith("Ambiguous ownership")


def test_observation_drilldown_only_through_explicit_report_id():
    linked = ev.Observation("rop", day(1), 1.0, "m/hr", ev.STATE_VALUE, report_id=9)
    legacy = ev.Observation("rop", day(1), 1.0, "m/hr", ev.STATE_VALUE, report_id=None)
    assert linked.drilldown_report_id == 9 and linked.drilldown_reason() == ""
    assert legacy.drilldown_report_id is None and "report_id" in legacy.drilldown_reason()


# ---------------------------------------------------------------------------
# NPT coverage
# ---------------------------------------------------------------------------

def _rep(rid, n, bore=None):
    return {"report_id": rid, "date": day(n), "wellbore_id": bore, "section_id": None}


def _log(rid, hours, npt=False, code=None):
    return {"report_id": rid, "duration": hours, "is_npt": npt, "code": code}


def test_npt_daily_states_golden():
    reports = [_rep(1, 1), _rep(2, 2), _rep(3, 3), _rep(4, 4), _rep(5, 5), _rep(6, 6),
               _rep(7, 7), _rep(8, 7), _rep(9, 8), _rep(10, 9)]
    logs = [
        _log(1, 24.0),                                  # full coverage, no NPT -> valid zero
        _log(2, 2.0, True, "NPT-X"), _log(2, 22.0),     # 2.0 h
        _log(3, 12.0),                                  # 50 % -> partial, NOT a zero
        # report 4: no rows -> missing
        _log(5, 1.0, True), _log(5, -1.0),              # negative duration -> invalid
        _log(6, 20.0), _log(6, 5.0),                    # 25 h -> over-recorded, invalid
        _log(7, 24.0), _log(8, 24.0),                   # two reports one date -> ambiguous
        _log(9, 1.5, True, None), _log(9, 22.5),        # unclassified code
        _log(10, 23.0, None),                           # neither NPT nor productive -> invalid
    ]
    days = {d.date: d for d in ev.npt_daily_series(reports, logs)}
    assert (days[day(1)].state, days[day(1)].npt_hours) == (ev.STATE_VALUE, 0.0)
    assert (days[day(2)].state, days[day(2)].npt_hours) == (ev.STATE_VALUE, 2.0)
    assert days[day(3)].state == ev.STATE_PARTIAL and days[day(3)].coverage_percent == 50.0
    assert days[day(4)].state == ev.STATE_MISSING
    assert days[day(5)].state == ev.STATE_INVALID
    assert days[day(6)].state == ev.STATE_INVALID and "exceeds 24 h" in days[day(6)].reason
    assert days[day(7)].state == ev.STATE_AMBIGUOUS and days[day(7)].report_ids == [7, 8]
    assert days[day(8)].categories == {ev.UNCLASSIFIED_NPT: 1.5}
    assert days[day(9)].state == ev.STATE_INVALID
    cov = ev.npt_coverage_summary(list(days.values()))
    assert cov["complete"] is False and cov["covered_days"] == 3 and cov["days"] == 9


@pytest.mark.parametrize("recorded, state", [
    (22.7, ev.STATE_PARTIAL),   # 94.58 % < 95 %
    (23.0, ev.STATE_VALUE),     # 95.83 %
    (24.1, ev.STATE_VALUE),     # inside the 6-minute rounding tolerance
    (24.2, ev.STATE_INVALID),   # over-recorded
])
def test_npt_coverage_boundaries(recorded, state):
    (only,) = ev.npt_daily_series([_rep(1, 1)], [_log(1, recorded)])
    assert only.state == state


def test_npt_coverage_summary_empty_is_not_complete():
    assert ev.npt_coverage_summary([])["complete"] is False


# ---------------------------------------------------------------------------
# Projection
# ---------------------------------------------------------------------------

def test_projection_requires_five_distinct_dates():
    four = [(day(i), float(i)) for i in range(1, 5)]
    result = ev.linear_projection(four, horizon=3)
    assert result["status"] == "not_assessed" and result["projection"] == []
    assert "at least 5" in result["reason"]
    assert "NOT ASSESSED" in ev.describe_projection(result, "m/hr", "T")


def test_projection_rejects_duplicate_dates():
    pts = [(day(1), 1.0)] * 2 + [(day(i), 1.0) for i in range(2, 6)]
    assert ev.linear_projection(pts, horizon=1)["status"] == "not_assessed"


def test_projection_hand_golden_with_backtest():
    # x = 1..5, y = 2,4,5,4,5. mean x 3, mean y 4; Sxy = 4+0+0+0+2 = 6; Sxx = 10
    # slope 0.6, intercept 4 - 1.8 = 2.2 -> day 6: 5.8, day 7: 6.4
    pts = [(day(i), y) for i, y in zip(range(1, 6), (2, 4, 5, 4, 5))]
    r = ev.linear_projection(pts, horizon=2)
    assert r["status"] == "assessed"
    assert r["slope"] == pytest.approx(0.6) and r["intercept"] == pytest.approx(2.2)
    assert [p["day"] for p in r["projection"]] == [6, 7]
    assert [p["value"] for p in r["projection"]] == pytest.approx([5.8, 6.4])
    # Backtest, origin k=3: fit (1,2),(2,4),(3,5): slope 1.5, intercept 2/3 -> x4 = 6.667,
    #   |6.667-4| = 2.667; baseline mean 11/3 -> |3.667-4| = 0.333
    # origin k=4: fit 1..4: slope 0.7, intercept 2.0 -> x5 = 5.5, |5.5-5| = 0.5;
    #   baseline 3.75 -> 1.25
    bt = r["backtest"]
    assert bt["origins"] == 2
    assert bt["mae_model"] == pytest.approx((8 / 3 + 0.5) / 2)
    assert bt["mae_baseline"] == pytest.approx((1 / 3 + 1.25) / 2)
    assert bt["model_better"] is False
    text = ev.describe_projection(r, "m/hr", "T")
    assert "does NOT beat" in text and "NOT a validated forecast" in text
    assert "confidence interval" in text  # stated as NOT claimed (limits)


def test_projection_uses_calendar_days_and_clips_at_zero():
    # Days 1,2,3,4,5 values 10,8,6,4,2 -> slope -2, intercept 12;
    # day 6 -> 0 (not clipped), day 7 -> -2 clipped to 0
    pts = [(day(i), 12 - 2 * i) for i in range(1, 6)]
    r = ev.linear_projection(pts, horizon=2)
    assert [(p["value"], p["clipped"]) for p in r["projection"]] == [(0.0, False), (0.0, True)]
    gaps = [(day(n), float(n)) for n in (1, 3, 6, 8, 10)]
    assert ev.linear_projection(gaps, horizon=1)["x"] == [1, 3, 6, 8, 10]


def test_projection_property_matches_numpy_polyfit():
    rng = random.Random(4801)
    for _ in range(200):
        n = rng.randint(5, 30)
        offsets = sorted(rng.sample(range(0, 120), n))
        ys = [rng.uniform(0, 50) for _ in offsets]
        r = ev.linear_projection([(D + timedelta(days=o), y) for o, y in zip(offsets, ys)], 1)
        xs = [o - offsets[0] + 1 for o in offsets]
        slope, intercept = np.polyfit(xs, ys, 1)
        assert r["slope"] == pytest.approx(slope, rel=1e-9, abs=1e-9)
        assert r["intercept"] == pytest.approx(intercept, rel=1e-9, abs=1e-7)
        assert r["projection"][0]["value"] >= 0.0


# ---------------------------------------------------------------------------
# Property: ROP validity and one-observation-per-date invariants
# ---------------------------------------------------------------------------

def _independent_valid(v):
    if v is None or isinstance(v, bool):
        return False
    try:
        f = float(v)
    except (TypeError, ValueError):
        return False
    return f == f and f not in (float("inf"), float("-inf")) and f >= 0


def test_rop_validity_property():
    rng = random.Random(4802)
    pool = [None, float("nan"), float("inf"), -float("inf"), -5.0, 0.0, "x", True]
    for _ in range(300):
        values = [rng.choice(pool) if rng.random() < 0.4 else rng.uniform(0, 80)
                  for _ in range(rng.randint(0, 20))]
        expected = [float(v) for v in values if _independent_valid(v)]
        out = ev.summarize_rop(values)
        if expected:
            assert out["average"] == pytest.approx(sum(expected) / len(expected))
            assert out["best"] == max(expected)
        else:
            assert out["average"] is None and out["best"] is None
        assert out["valid_count"] + out["missing_count"] + out["invalid_count"] == len(values)


def test_rop_daily_series_property_one_observation_per_date():
    rng = random.Random(4803)
    for _ in range(300):
        rows = []
        for i in range(rng.randint(1, 15)):
            rows.append(_row(rng.randint(1, 5),
                             rng.choice([None, float("nan"), -1.0, 0.0, rng.uniform(0, 40)]),
                             rng.choice([None, 1, 2, 3]), rng.choice([None, 10, 11]),
                             rng.choice([None, 0.0, 2.0, 5.0]), source_id=i))
        series, excluded = ev.rop_daily_series(rows)
        dates = [o.date for o in series]
        assert len(dates) == len(set(dates))
        # Date-level ambiguity only: a superseded legacy row is ambiguous as a
        # row, but its date is decided by the report-linked rows.
        ambiguous = {o.date for o in excluded if o.state == ev.STATE_AMBIGUOUS
                     and o.reason != ev.LEGACY_DUPLICATE_REASON}
        assert not ambiguous & set(dates)
        # Independent oracle of the legacy rule: on a date with any linked row
        # only linked rows are candidates; otherwise the legacy rows are.
        valid_dates = set()
        for d in {r["date"] for r in rows}:
            same = [r for r in rows if r["date"] == d]
            linked = [r for r in same if r["report_id"] is not None]
            if any(_independent_valid(r["rop"]) for r in (linked or same)):
                valid_dates.add(d)
        assert valid_dates == set(dates) | ambiguous
        for o in excluded:
            if o.reason == ev.LEGACY_DUPLICATE_REASON:
                assert o.report_id is None and o.drilldown_report_id is None
        for o in series:
            assert math.isfinite(o.value) and o.value >= 0


def test_npt_valid_zero_only_with_full_coverage_property():
    rng = random.Random(4804)
    for _ in range(300):
        logs = [_log(1, rng.choice([None, -1.0, 0.0, rng.uniform(0, 12)]),
                     rng.choice([True, False, None]), rng.choice([None, "C1"]))
                for _ in range(rng.randint(0, 4))]
        (d,) = ev.npt_daily_series([_rep(1, 1)], logs)
        if d.state == ev.STATE_VALUE and d.npt_hours == 0.0:
            durations = [x["duration"] for x in logs]
            assert all(_independent_valid(h) for h in durations)
            assert sum(durations) >= 0.95 * 24 and sum(durations) <= 24.1
            assert not any(x["is_npt"] for x in logs)


# ---------------------------------------------------------------------------
# Risk (rule-based, not a probability)
# ---------------------------------------------------------------------------

FULL = {"complete": True, "covered_days": 3, "days": 3}


@pytest.mark.parametrize("npt, score", [(0, 3), (10, 3), (10.01, 5), (20, 5), (20.5, 7),
                                        (30, 7), (30.1, 9)])
def test_risk_npt_rule_boundaries(npt, score):
    assert ev.risk_scores(npt, 100, True)["NPT Risk"] == score


@pytest.mark.parametrize("days, score", [(91, 2), (90, 5), (31, 5), (30, 8), (0, 8)])
def test_risk_safety_rule_boundaries(days, score):
    assert ev.risk_scores(5, days, True)["Safety"] == score


def test_missing_safety_never_gives_low():
    a = ev.risk_assessment(2.0, FULL, None)
    assert a["overall"]["assessed"] is False and a["overall"]["level"] is None
    assert "Safety" in a["overall"]["reason"]


def test_incomplete_npt_coverage_is_not_assessed():
    a = ev.risk_assessment(2.0, {"complete": False, "covered_days": 1, "days": 3}, 100)
    npt = next(c for c in a["criteria"] if c["name"] == "NPT Risk")
    assert npt["assessed"] is False and "1 of 3" in npt["reason"]
    assert a["overall"]["assessed"] is False


def test_risk_overall_is_worst_required_criterion_golden():
    # NPT 5 % -> 3 (LOW); 100 days without LTI -> 2 -> worst 3 -> LOW
    assert ev.risk_assessment(5.0, FULL, 100)["overall"]["level"] == "LOW"
    # NPT 25 % -> 7; 10 days -> 8 -> worst 8 -> HIGH
    a = ev.risk_assessment(25.0, FULL, 10)
    assert (a["overall"]["level"], a["overall"]["score"]) == ("HIGH", 8)
    by_name = {c["name"]: c for c in a["criteria"]}
    assert by_name["Equipment"]["assessed"] is False and by_name["Weather"]["assessed"] is False
    assert by_name["Well Control"]["proxy"] is True and by_name["Well Control"]["required"] is False


# ---------------------------------------------------------------------------
# Typed export
# ---------------------------------------------------------------------------

def _rows():
    return [ev.Observation("rop", day(1), 0.0, "m/hr", ev.STATE_VALUE, "", 11, 5, None, "P#1"),
            ev.Observation("rop", day(2), None, "m/hr", ev.STATE_MISSING, "not recorded", 12, 5,
                           None, "P#2"),
            ev.Observation("rop", day(3), None, "m/hr", ev.STATE_INVALID, "=HYPERLINK(1)", None,
                           None, None, "P#3")]


def test_export_csv_keeps_zero_empty_none_and_neutralises_formulas(tmp_path):
    path = tmp_path / "e.csv"
    meta = ev.export_metadata("Well #1", "rop", "m/hr", "def", "all dates", {"well_id": 1})
    assert ev.export_csv(str(path), _rows(), meta) == 3
    with path.open(encoding="utf-8", newline="") as fh:
        lines = list(csv.reader(fh))
    header_at = next(i for i, r in enumerate(lines) if r and r[0] == "metric")
    assert any(r[0] == "# null_policy" for r in lines[:header_at])
    body = lines[header_at + 1:]
    assert body[0][2] == "0.0" and body[1][2] == "" and body[1][4] == "missing"
    assert body[2][5] == "'=HYPERLINK(1)"


def test_export_xlsx_cells_are_typed(tmp_path):
    from openpyxl import load_workbook
    path = tmp_path / "e.xlsx"
    ev.export_xlsx(str(path), _rows(), ev.export_metadata("S", "rop", "m/hr", "d", "r"))
    book = load_workbook(path)
    sheet = book["evidence"]
    assert [c.value for c in sheet[1]] == list(ev.EXPORT_COLUMNS)
    # a numeric zero cell (openpyxl reads a stored 0.0 back as int 0), not text or empty
    assert sheet["C2"].value == 0 and sheet["C2"].data_type == "n"
    assert sheet["C3"].value is None
    assert sheet["G2"].value == 11 and isinstance(sheet["G2"].value, int)
    assert "null_policy" in [r[0].value for r in book["metadata"].iter_rows()]


def test_export_tables_csv_and_xlsx(tmp_path):
    from openpyxl import load_workbook
    tables = {"perf": (("date", "rop_m_hr"), [(day(1), 0.0), (day(2), None), (day(3), "-1")])}
    meta = {"scope": "S"}
    (written,) = ev.export_tables_csv(str(tmp_path / "x.csv"), tables, meta)
    with open(written, encoding="utf-8", newline="") as fh:
        rows = list(csv.reader(fh))
    assert rows[-3:] == [["2026-01-01", "0.0"], ["2026-01-02", ""], ["2026-01-03", "'-1"]]
    ev.export_tables_xlsx(str(tmp_path / "x.xlsx"), tables, meta)
    sheet = load_workbook(tmp_path / "x.xlsx")["perf"]
    assert sheet["B2"].value == 0 and sheet["B2"].data_type == "n" and sheet["B3"].value is None


def test_filter_range_inclusive_and_drops_undated():
    rows = _rows() + [ev.Observation("rop", None, 1.0, "m/hr", ev.STATE_VALUE)]
    assert [r.date for r in ev.filter_range(rows, day(2), day(3))] == [day(2), day(3)]
    assert len(ev.filter_range(rows, None, None)) == 3


# ---------------------------------------------------------------------------
# Scope key / problems
# ---------------------------------------------------------------------------

def test_scope_key_covers_every_dimension():
    base = ev.AnalysisScope(1, 2, 3, 4, ev.SCOPE_WELL)
    variants = [ev.AnalysisScope(9, 2, 3, 4), ev.AnalysisScope(1, 9, 3, 4),
                ev.AnalysisScope(1, 2, 9, 4), ev.AnalysisScope(1, 2, 3, 9),
                ev.AnalysisScope(1, 2, 3, 4, ev.SCOPE_SECTION)]
    keys = {base.key()} | {v.key() for v in variants}
    assert len(keys) == 6


def test_scope_problems():
    assert ev.AnalysisScope(None).problem() == "no well selected"
    assert ev.AnalysisScope(1, level=ev.SCOPE_SECTION).problem() == "no section selected"
    assert ev.AnalysisScope(1, level=ev.SCOPE_REPORT).problem() == "no daily report selected"
    assert ev.AnalysisScope(1).problem() == ""


# ---------------------------------------------------------------------------
# d-exponent unit contract (formula unchanged; independent hand golden)
# ---------------------------------------------------------------------------

def test_d_exponent_independent_golden():
    from core.engineering.engines.bit_performance import BitPerformanceEngine as E
    # 60 ft/hr, 100 rpm, 1000 lbf, 12 in: log10(60/6000) = -2; log10(12000/12e6) = -3 -> 2/3
    assert E.d_exponent(60.0, 100.0, 1000.0, 12.0).value == pytest.approx(0.6667, abs=1e-4)
    # 100 ft/hr, 50 rpm, 10 000 lbf, 8.5 in:
    # log10(100/3000) = -1.4771213; log10(120000/8.5e6) = -1.8502377 -> 0.798341
    assert E.d_exponent(100.0, 50.0, 10000.0, 8.5).value == pytest.approx(0.7983, abs=1e-4)
    # W12 converts m/hr -> ft/hr (x 3.28084) and klb -> lbf (x 1000):
    # 18.288 m/hr = 60.0 ft/hr; 1 klb = 1000 lbf -> the first golden again
    assert E.d_exponent(18.288 * 3.28084, 100.0, 1 * 1000.0, 12.0).value == pytest.approx(
        0.6667, abs=1e-4)
    assert E.d_exponent(0.0, 100.0, 1000.0, 12.0).value is None


# ---------------------------------------------------------------------------
# DB loaders on the synthetic multi-well / multi-bore fixture
# ---------------------------------------------------------------------------

@pytest.fixture
def syn():
    from m48_synthetic import build
    return build()


def _evidence(db, scope, metric):
    s = db.create_session()
    try:
        return ev.evidence_for_metric(s, scope, metric)
    finally:
        s.close()


def test_whole_well_rop_evidence_states(syn):
    db, ids = syn
    rows = _evidence(db, ev.AnalysisScope(ids["A"]), "rop")
    got = [(r.date.day, r.state, r.value) for r in rows]
    # A6: (12 x 2 + 6 x 4) / 6 = 8.0 hours-weighted
    assert got == [(1, ev.STATE_VALUE, 0.0), (2, ev.STATE_VALUE, 10.0),
                   (3, ev.STATE_MISSING, None), (4, ev.STATE_INVALID, None),
                   (5, ev.STATE_VALUE, 8.0), (6, ev.STATE_VALUE, 8.0)]
    assert [r.report_id for r in rows] == ids["A_reports"]


def test_wellbore_filter_and_legacy_rows(syn):
    db, ids = syn
    whole = _evidence(db, ev.AnalysisScope(ids["B"]), "rop")
    # day 1: the legacy 99.0 row (no report_id) duplicates OH report B1, so the
    # linked 5.0 is that date's value and the legacy row is listed as excluded;
    # day 2: OH + ST1 share the date -> ambiguous; day 3: ST1 30.0.
    by_day = {}
    for r in whole:
        by_day.setdefault(r.date.day, []).append(r)
    assert [(r.state, r.value) for r in by_day[1] if r.state == ev.STATE_VALUE] == [(ev.STATE_VALUE, 5.0)]
    legacy = [r for r in by_day[1] if r.state != ev.STATE_VALUE]
    assert len(legacy) == 1 and legacy[0].reason == ev.LEGACY_DUPLICATE_REASON
    assert legacy[0].drilldown_report_id is None
    assert {r.state for r in by_day[2]} == {ev.STATE_AMBIGUOUS}
    assert [(r.state, r.value) for r in by_day[3]] == [(ev.STATE_VALUE, 30.0)]
    oh = _evidence(db, ev.AnalysisScope(ids["B"], ids["OH"]), "rop")
    assert [(r.date.day, r.value) for r in oh] == [(1, 5.0), (2, 7.0)]
    st = _evidence(db, ev.AnalysisScope(ids["B"], ids["ST1"]), "rop")
    assert [(r.date.day, r.value) for r in st] == [(2, 20.0), (3, 30.0)]
    assert all(r.value != 99.0 for r in oh + st)  # legacy row never attributed to a bore


def test_report_and_section_scope_loaders(syn):
    db, ids = syn
    rid = ids["C_reports"][2]
    rows = _evidence(db, ev.AnalysisScope(ids["C"], report_id=rid, level=ev.SCOPE_REPORT), "rop")
    assert [(r.report_id, r.value) for r in rows] == [(rid, 3.0)]
    npt = _evidence(db, ev.AnalysisScope(ids["B"], ids["OH"]), "npt")
    assert [(r.date.day, r.state, r.value) for r in npt] == [(1, ev.STATE_VALUE, 1.0),
                                                             (2, ev.STATE_VALUE, 1.0)]
    depth = _evidence(db, ev.AnalysisScope(ids["B"]), "depth")
    assert [r.state for r in depth if r.date.day == 2] == [ev.STATE_AMBIGUOUS] * 2


def test_data_revision_changes_only_with_data(syn):
    from core.database import DailyReport
    db, ids = syn
    s = db.create_session()
    try:
        first = ev.data_revision(s, ids["A"])
        assert ev.data_revision(s, ids["A"]) == first
        s.get(DailyReport, ids["A_reports"][0]).depth_2400 = 111.0
        s.commit()
        assert ev.data_revision(s, ids["A"]) != first
    finally:
        s.close()


def test_unknown_evidence_metric_raises(syn):
    db, ids = syn
    with pytest.raises(ValueError):
        _evidence(db, ev.AnalysisScope(ids["A"]), "nope")
