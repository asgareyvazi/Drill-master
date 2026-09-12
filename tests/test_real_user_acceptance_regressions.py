"""Acceptance remediations: actual Python methods with table protocols, NOT native Qt.

The Qt substitutes exercise nullable/persistence and pending-event contracts only.
Production fixtures use generated credentials and an explicitly created hierarchy.
"""
import ast
import copy
from datetime import date
import json
import logging
from pathlib import Path
import re
import secrets
from types import SimpleNamespace

import pytest

from core.database import DatabaseManager, Company, Project, Well, Section
from core.domain_records import collection_value, restore_named_text
from core.repositories.base import BaseRepository
from core.repositories.well_repository import WellRepository
from core.editor_state import editor_loaded, editor_saved
from core.save_outcome import save_all
from core.text_utils import safe_str

ROOT = Path(__file__).resolve().parents[1]


class Item:
    def __init__(self, text):
        self.value, self.roles, self.background = text, {}, None

    def text(self):
        return self.value

    def setText(self, text):
        self.value = text

    def setData(self, role, value):
        self.roles[role] = copy.deepcopy(value)

    def data(self, role):
        return copy.deepcopy(self.roles.get(role))

    def setTextAlignment(self, _):
        pass

    def setBackground(self, color):
        self.background = color


class Table:
    def __init__(self):
        self.headers, self.rows = [], []

    def setColumnCount(self, count):
        self.headers = [None] * count

    def setHorizontalHeaderLabels(self, labels):
        self.headers = [Item(label) for label in labels]

    def setEditTriggers(self, value):
        self.edit_triggers = value

    def setColumnWidth(self, *_):
        pass

    def columnCount(self):
        return len(self.headers)

    def rowCount(self):
        return len(self.rows)

    def setRowCount(self, count):
        assert count == 0
        self.rows = []

    def insertRow(self, row):
        self.rows.insert(row, [None] * self.columnCount())

    def setItem(self, row, col, item):
        self.rows[row][col] = item

    def item(self, row, col):
        return self.rows[row][col]

    def horizontalHeaderItem(self, col):
        return self.headers[col]


class Color:
    """Protocol only; QColor rendering/validity itself is not certified here."""
    def __init__(self, value=None):
        self.value = value

    def isValid(self):
        return isinstance(self.value, str) and bool(re.fullmatch(r"#[0-9a-fA-F]{6}", self.value))


class Widget:
    def showEvent(self, _):
        pass


def actual_class(path, name, **namespace):
    tree = ast.parse((ROOT / path).read_text())
    node = next(n for n in tree.body if isinstance(n, ast.ClassDef) and n.name == name)
    env = dict(editor_loaded=editor_loaded, editor_saved=editor_saved,
               QAbstractItemView=SimpleNamespace(NoEditTriggers=0, DoubleClicked=1, EditKeyPressed=2),
               QTableWidgetItem=Item, TableManager=lambda _: None,
               Qt=SimpleNamespace(UserRole=256, AlignRight=1, AlignVCenter=2),
               QColor=Color, QWidget=Widget, logger=logging.getLogger(name),
               safe_str=safe_str, restore_named_text=restore_named_text,
               collection_value=collection_value, json=json)
    env.update(namespace)
    exec(compile(ast.Module(body=[node], type_ignores=[]), str(path), "exec"), env)
    return env[name]


def manager(name):
    table = Table()
    cls = actual_class("tabs/w4_Downhole_Widget.py", name)
    result = cls(table)
    result.setup_table()
    return result


@pytest.fixture
def production_db(monkeypatch, tmp_path):
    monkeypatch.setenv("DRILLMASTER_ENV", "production")
    monkeypatch.setenv("DRILLMASTER_DATA_DIR", str(tmp_path))
    db = DatabaseManager(bootstrap_passwords={"admin": secrets.token_urlsafe(24)})
    db.db_path = str(tmp_path / "acceptance.db")
    assert db.initialize()
    assert db.generic_get_list(Well) == []
    company = db.generic_save(Company, {"name": "Acceptance operator", "code": "QA"})
    project = db.generic_save(Project, {"name": "Production Acceptance Test", "code": "PAT", "company_id": company})
    well = db.generic_save(Well, {"name": "TEST-WELL-001", "code": "TEST-WELL-001", "project_id": project})
    section = db.generic_save(Section, {"name": "Test section", "well_id": well})
    report = db.save_daily_report({"well_id": well, "section_id": section, "report_number": 1, "report_date": date(2026, 9, 8)})
    db.acceptance_ids = (company, project, well, section, report["id"])
    yield db
    db.close()


@pytest.mark.parametrize("color", [None, "", "invalid", 0, {"unsupported": True}, "#123456"])
def test_optional_color_and_untouched_values(color):
    formation = manager("FormationManager")
    source = {"formation_name": "Sandstone", "top_md": 0, "base_md": 100.5, "color": color}
    formation.load_data([source])
    row = formation.get_all_data()[0]
    assert row["Top MD (m)"] == 0 and row["Base MD (m)"] == 100.5
    assert row["Thickness (m)"] is None
    assert row["_provenance"]["source_record"] == source
    formation.table.item(0, 8).setText("Explicit description edit")
    assert formation.get_all_data()[0]["Description"] == "Explicit description edit"
    formation.load_data(None)
    assert formation.table.rowCount() == 0


@pytest.mark.parametrize("name", ["BHAManager", "DownholeEquipmentManager", "FormationManager"])
def test_collection_error_is_not_silently_empty(name):
    with pytest.raises(ValueError):
        manager(name).load_data('{"not": "records"}')
    with pytest.raises(ValueError):
        manager(name).load_data("not-json")


def test_downhole_zero_null_and_persistence_reopen(production_db):
    db = production_db
    _, _, well, _, report = db.acceptance_ids
    equipment = manager("DownholeEquipmentManager")
    equipment.load_data([{"equipment_name": "MWD", "rotation_hours": 0, "install_date": None}])
    row = equipment.get_all_data()[0]
    assert row["Rotation Hours"] == 0 and row["Install Date"] is None
    db.save_downhole_equipment(well, {"report_id": report, "equipment_data_json": [row]})
    formations = manager("FormationManager")
    formations.load_data([{"name": "Sandstone", "top_md": 0}])
    db.save_formation_report(well, {"report_id": report, "formations": formations.get_all_data()})
    db.close()
    reopened = DatabaseManager()
    reopened.db_path = db.db_path
    assert reopened.initialize()
    equipment.load_data(reopened.get_downhole_equipment(well, report)["equipment_data"])
    formations.load_data(reopened.get_formation_report(well, report)["formations"])
    assert equipment.get_all_data()[0]["Rotation Hours"] == 0
    assert equipment.get_all_data()[0]["Install Date"] is None
    assert formations.get_all_data()[0]["Color"] is None
    assert formations.get_all_data()[0]["Top MD (m)"] == 0
    reopened.close()


def editor():
    cls = actual_class("core/base_tab.py", "DrillTabBase")
    obj = cls.__new__(cls)
    obj.widget_name = "Acceptance editor"
    obj.visible = False
    obj.isVisible = lambda: obj.visible
    obj.calls = []
    for kind in ("well", "section", "report"):
        setattr(obj, f"current_{kind}_id", None)
        setattr(obj, f"current_{kind}_data", {})
        setattr(obj, f"_loaded_{kind}_id", None)
        setattr(obj, f"_pending_{kind}", None)
        setattr(obj, f"on_{kind}_changed", lambda identity, data, kind=kind: obj.calls.append((kind, identity)))
    return obj


def test_pending_failed_parent_is_retained_and_save_blocked():
    obj = editor()
    obj._on_well_changed_internal(11, {"name": "Well"})
    obj._on_section_changed_internal(22, {})
    obj._on_report_changed_internal(33, {})
    def fail(*_):
        raise RuntimeError("load failed")
    obj.on_well_changed = fail
    assert not obj._apply_pending_context()
    assert obj._pending_well[0] == 11 and obj._loaded_well_id is None
    assert not obj.save_context_ready() and obj.calls == []
    obj.on_well_changed = lambda identity, data: obj.calls.append(("well", identity))
    obj.showEvent(None)
    assert obj.calls == [("well", 11), ("section", 22), ("report", 33)]
    assert obj.save_context_ready()


@pytest.mark.parametrize("visible", [False, True])
def test_force_refresh_preserves_report_and_section(visible):
    obj = editor()
    obj.visible = visible
    obj._on_well_changed_internal(11, {"name": "Well"})
    obj._on_section_changed_internal(22, {"name": "Section"})
    obj._on_report_changed_internal(33, {"summary": "Report"})
    obj._apply_pending_context()
    obj.calls = []
    obj.force_refresh()
    assert (obj.current_well_id, obj.current_section_id, obj.current_report_id) == (11, 22, 33)
    obj._apply_pending_context()
    assert obj.calls == [("well", 11), ("section", 22), ("report", 33)]
    assert obj.current_report_data == {"summary": "Report"}


def test_save_all_does_not_erase_new_context_with_stale_empty_view(production_db):
    db = production_db
    _, _, well, _, report = db.acceptance_ids
    db.save_bha_report(well, {"report_id": report, "bha_name": "BHA 1", "bha_data": [{"component_name": "Bit", "length": 0.3}]})
    class StaleEditor:
        def save_context_ready(self):
            return False

        def save_data(self):
            return db.save_bha_report(well, {"report_id": report, "bha_data": []})
    result = save_all([("BHA", StaleEditor().save_data), ("independent valid", lambda: True)])
    assert result.status == "REVIEW_REQUIRED" and result.saved == 1
    assert len(db.get_bha_report(well, report)["bha_configs"]) == 1


def test_well_invalid_date_does_not_clear_saved_value_or_mutate_input(production_db):
    db = production_db
    well = db.acceptance_ids[2]
    value = {"id": well, "spud_date": "2026-09-01"}
    assert db.save_well(value)
    assert value["spud_date"] == "2026-09-01"
    with pytest.raises(ValueError):
        db.save_well({"id": well, "spud_date": "not-a-date"})
    assert db.get_well_by_id(well)["spud_date"] == date(2026, 9, 1)


def test_repository_context_and_deleted_identity(production_db):
    db = production_db
    company, project, well, _, _ = db.acceptance_ids
    other = db.generic_save(Project, {"name": "Other", "code": "OTHER", "company_id": company})
    repo = WellRepository(db)
    with pytest.raises(ValueError, match="project"):
        repo.resolve_identity({"name": "Unscoped well"})
    second = repo.resolve_identity({"name": "TEST-WELL-001", "project_id": other})
    assert second != well
    assert repo.resolve_identity({"name": "TEST-WELL-001", "project_id": project}) == well
    with pytest.raises(ValueError, match="Ambiguous"):
        repo.resolve_identity({"name": "TEST-WELL-001"})
    assert db.delete_well(second)
    with pytest.raises(ValueError, match="no longer exists"):
        BaseRepository(db).save(Well, {"id": second, "name": "Resurrected", "project_id": other})
    assert not repo.get_by_name_or_code("Resurrected")


@pytest.mark.parametrize("engine_name", ["DDRReportEngine", "NPTReportEngine", "CostReportEngine", "PlanReportEngine", "EOWRReportEngine"])
def test_report_engines_import_without_qt(engine_name):
    from core import report_engine
    assert getattr(report_engine, engine_name)


class Text:
    def __init__(self): self.value = ""
    def setReadOnly(self, value): self.read_only = value
    def setEnabled(self, value): self.enabled = value
    def text(self): return self.value
    def setText(self, value): self.value = value
    def clear(self): self.value = ""


class Combo:
    def __init__(self): self.items, self.selected = [], None
    def blockSignals(self, _): pass
    def clear(self): self.items, self.selected = [], None
    def addItem(self, text, identity=None): self.items.append((text, identity))
    def setCurrentText(self, text): self.selected = next(identity for label, identity in self.items if label == text)
    def currentData(self): return self.selected


def downhole_editor(db):
    base = type(editor())
    message = SimpleNamespace(Yes=1, No=2, question=lambda *_: 1)
    cls = actual_class("tabs/w4_Downhole_Widget.py", "DownholeWidget", DrillTabBase=base, QMessageBox=message)
    obj = cls.__new__(cls)
    obj.__dict__.update(editor().__dict__)
    obj.db = db
    _, _, well, section, report = db.acceptance_ids
    for kind, identity in (("well", well), ("section", section), ("report", report)):
        setattr(obj, f"current_{kind}_id", identity)
        setattr(obj, f"_loaded_{kind}_id", identity)
    obj.current_well, obj.current_section = well, section
    obj.bha_name_input, obj.bha_selector, obj.well_label = Text(), Combo(), Text()
    for name in ("add_bha_tool_btn", "remove_bha_tool_btn", "save_bha_btn", "delete_bha_btn", "export_bha_btn"):
        setattr(obj, name, Text())
    for kind, name in (("bha", "BHAManager"), ("equipment", "DownholeEquipmentManager"), ("formation", "FormationManager")):
        value = manager(name)
        setattr(obj, f"{kind}_manager", value)
        setattr(obj, f"{kind}_table", value.table)
    obj.messages = []
    for method in ("show_success", "show_error", "show_warning", "show_message"):
        setattr(obj, method, lambda text: obj.messages.append(text))
    return obj, message


def test_bha_button_save_cancel_delete_and_restart(production_db):
    db = production_db
    obj, message = downhole_editor(db)
    obj.load_all_data_from_db()
    obj.bha_manager.load_data([{"component_name": "Drill collar", "length": 9, "od": 8}])
    obj.bha_name_input.setText("BHA 1")
    assert obj.save_bha_config()
    db.close()
    assert db.initialize()
    obj.load_all_data_from_db()
    assert obj.bha_manager.get_all_data()[0]["Length (m)"] == 9
    message.question = lambda *_: message.No
    assert not obj.delete_bha_config()
    assert len(db.get_bha_report(obj.current_well, obj.current_report_id)["bha_configs"]) == 1
    message.question = lambda *_: message.Yes
    assert obj.delete_bha_config()
    db.close()
    assert db.initialize()
    assert db.get_bha_report(obj.current_well, obj.current_report_id) is None
    obj.load_all_data_from_db()
    assert obj.bha_table.rowCount() == 0


def test_downhole_null_empty_corrupt_switch_clear(production_db):
    from core.database import BHAReport
    db = production_db
    obj, _ = downhole_editor(db)
    db.save_bha_report(obj.current_well, {"report_id": obj.current_report_id, "bha_name": "BHA", "bha_data": [{"component_name": "Bit", "length": .3}]})
    obj.load_all_data_from_db()
    assert obj.bha_table.rowCount() == 1
    other = db.save_daily_report({"well_id": obj.current_well, "section_id": obj.current_section,
        "report_number": 2, "report_date": date(2026, 9, 9)})["id"]
    obj.current_report_id = other
    obj.load_all_data_from_db()
    assert obj.bha_table.rowCount() == 0
    identity = db.generic_save(BHAReport, {"well_id": obj.current_well, "report_id": other, "bha_name": "NULL BHA", "bha_data_json": None})
    obj.load_all_data_from_db()
    assert obj.bha_table.rowCount() == 0
    db.generic_save(BHAReport, {"id": identity, "bha_data_json": "malformed JSON"})
    with pytest.raises(ValueError): obj.load_all_data_from_db()
    assert obj._downhole_loaded_context is None and not obj.save_context_ready()
    obj.on_selection_cleared()
    assert obj.bha_table.rowCount() == obj.formation_table.rowCount() == obj.equipment_table.rowCount() == 0


def test_invalid_bha_numeric_edit_is_not_silently_erased(production_db):
    db = production_db
    _, _, well, _, report = db.acceptance_ids
    db.save_bha_report(well, {"report_id": report, "bha_name": "BHA", "bha_data": [{"component_name": "Bit", "length": .3}]})
    with pytest.raises(ValueError, match="Length"):
        db.save_bha_report(well, {"report_id": report, "bha_data": [{"component_name": "Bit", "length": "invalid"}]})
    assert db.get_bha_report(well, report)["bha_configs"][0]["Length (m)"] == .3


def test_eowr_incomplete_survey_and_optional_plan_export(production_db, tmp_path):
    from core.report_engine import EOWRReportEngine, DDRReportEngine
    from core.database import SurveyPoint, MudReport
    db = production_db
    _, _, well, _, report = db.acceptance_ids
    db.generic_save(SurveyPoint, {"well_id": well, "report_id": report, "md": 100, "inc": 10, "azi": None})
    db.generic_save(MudReport, {"well_id": well, "report_id": report, "report_date": date(2026, 9, 8), "mw": 71})
    for fmt, ext in (("html", "html"), ("excel", "xlsx")):
        assert EOWRReportEngine(db).generate(well, str(tmp_path / f"eowr.{ext}"), format=fmt)
        assert DDRReportEngine(db).generate(report, str(tmp_path / f"ddr.{ext}"), format=fmt)
    text = (tmp_path / "eowr.html").read_text()
    assert "—" in text and "<td>None</td>" not in text
    assert "pcf</td>" in (tmp_path / "ddr.html").read_text()


def test_professional_export_scope_nulls_lineage_and_failures(production_db, tmp_path, monkeypatch):
    from core.database import SurveyPoint
    from core.professional_export import ProfessionalExcelExport
    from openpyxl import load_workbook
    db = production_db
    _, _, well, section, report = db.acceptance_ids
    another = db.save_daily_report({"well_id": well, "section_id": section, "report_date": date(2026, 9, 9)})["id"]
    db.generic_save(SurveyPoint, {"well_id": well, "report_id": report, "md": 100, "inc": 10, "azi": None})
    db.generic_save(SurveyPoint, {"well_id": well, "report_id": another, "md": 200, "inc": 20, "azi": 20})
    db.save_bha_report(well, {"report_id": report, "bha_name": "BHA", "bha_data": [{"component_name": "Bit", "length": .3, "_provenance": {"source_record": {"note": "x" * 40000}}}]})
    target = tmp_path / "آزمون.xlsx"
    exporter = ProfessionalExcelExport(db)
    assert exporter.export(well, str(target), report_id=report)
    wb = load_workbook(target)
    assert wb["Survey"].max_row == 2 and wb["Survey"].cell(2, 3).value is None
    assert wb["BHA"].max_row == 2
    assert all("_provenance" not in str(c.value) for row in wb["BHA"] for c in row)
    chunks = [row[2] for row in wb["Raw Data"].iter_rows(min_row=3, values_only=True) if row[0] == "BHAReport"]
    raw = json.loads("".join(chunks))
    assert len(raw["bha_data_json"][0]["_provenance"]["source_record"]["note"]) == 40000
    wb.close()
    def fail(*_, **__): raise RuntimeError("Injected unavailable query")
    monkeypatch.setattr(db, "get_mud_report", fail)
    assert not exporter.export(well, str(tmp_path / "failed.xlsx"), report_id=report)
    assert not (tmp_path / "failed.xlsx").exists()
    assert not exporter.export(999999, str(tmp_path / "wrong-scope.xlsx"), report_id=report)


def test_csv_export_real_file_cancel_and_write_failure(tmp_path, monkeypatch):
    import sys
    import types
    qt = types.ModuleType("PySide6.QtWidgets")
    target = tmp_path / "آزمون.csv"
    qt.QFileDialog = SimpleNamespace(getSaveFileName=lambda *_: (str(target), "CSV"))
    errors = []
    qt.QMessageBox = SimpleNamespace(critical=lambda *args: errors.append(args))
    monkeypatch.setitem(sys.modules, "PySide6.QtWidgets", qt)
    export = actual_class("core/managers.py", "ExportManager")()
    table = manager("FormationManager").table
    table.cellWidget = lambda *_: None
    assert export.export_table_with_dialog(table, "QA") == str(target)
    assert "Formation Name" in target.read_text()
    qt.QFileDialog.getSaveFileName = lambda *_: ("", "CSV")
    assert export.export_table_with_dialog(table, "QA") is None
    qt.QFileDialog.getSaveFileName = lambda *_: (str(tmp_path), "CSV")
    assert export.export_table_with_dialog(table, "QA") is False and errors


def test_autosave_uses_same_pending_context_guard():
    cls = actual_class("core/managers.py", "AutoSaveManager")
    class Pending:
        def save_context_ready(self): return False
        def save_data(self): raise AssertionError("Stale widget must not be saved")
    assert cls.save_widget("pending", Pending()).status == "REVIEW_REQUIRED"


@pytest.mark.parametrize("method,data", [
    ("save_bha_report", {"bha_name": "Wrong scope", "bha_data": []}),
    ("save_downhole_equipment", {"equipment_data_json": []}),
    ("save_formation_report", {"formations": []})])
def test_downhole_services_reject_cross_well_report(production_db, method, data):
    db = production_db
    _, project, well, _, report = db.acceptance_ids
    other = db.generic_save(Well, {"name": "Other well", "project_id": project})
    with pytest.raises(ValueError, match="does not belong"):
        getattr(db, method)(other, {"report_id": report, **data})
    assert db.get_bha_report(well, report) is None
    assert db.get_downhole_equipment(well, report) is None
    assert db.get_formation_report(well, report) is None


def test_downhole_totals_do_not_fabricate_missing_measurements(production_db):
    bha = manager("BHAManager")
    bha.load_data([{"component_name": "Bit", "length": .3, "weight": None}])
    assert bha.calculate_totals() == (.3, None)
    obj, _ = downhole_editor(production_db)
    obj.bha_manager.load_data([{"component_name": "Incomplete bit"}])
    obj.calculate_bha_totals()
    assert obj.messages[-1].startswith("REVIEW_REQUIRED")
    equipment = manager("DownholeEquipmentManager")
    equipment.load_data([{"equipment_name": "MWD", "rotation_hours": 0}])
    assert equipment.calculate_hours() == {"sliding": None, "rotation": 0, "pumping": None, "total": None}
    equipment.table.item(0, 6).setText("invalid")
    with pytest.raises(ValueError): equipment.calculate_hours()


def test_unknown_service_date_does_not_mean_up_to_date():
    equipment = manager("DownholeEquipmentManager")
    equipment.check_service_due.__func__.__globals__["date"] = date
    equipment.load_data([{"equipment_name": "Unknown", "next_service": None},
                         {"equipment_name": "Invalid", "next_service": "bad date"},
                         {"equipment_name": "Overdue", "next_service": "2000-01-01"}])
    assert equipment.check_service_due() == [{"name": "Overdue", "row": 2}]
    assert equipment.service_review == ["Unknown", "Invalid"]


def _analysis_probe(db, well_id):
    """A minimal stand-in exposing the two AnalysisWidget query methods.

    The methods only read ``self.db`` / ``self.current_well_id``; exercising them
    on a SimpleNamespace avoids constructing the full Qt widget while still
    running the real production SQL.
    """
    from tabs.w12_Analysis import AnalysisWidget
    probe = SimpleNamespace(db=db, current_well_id=well_id)
    session = db.create_session()
    try:
        return (
            AnalysisWidget.calculate_kpis(probe, session),
            AnalysisWidget.get_npt_data(probe, session),
        )
    finally:
        session.close()


def test_analysis_kpis_report_unknown_not_zero(production_db):
    """A well with a DDR but no drilling parameters / time logs has unknown
    ROP / NPT / efficiency — never a fabricated 0.0 m/hr, 0% NPT, 100% efficiency
    (R-4 NULL->zero fabrication). This also matches the canonical
    OperationsIntelligenceService, which returns None for the same well."""
    from core.text_utils import fmt_num
    _, _, well, _, _ = production_db.acceptance_ids
    kpis, npt = _analysis_probe(production_db, well)
    assert kpis["avg_rop"] is None
    assert kpis["best_rop"] is None
    assert kpis["npt_percentage"] is None
    assert kpis["efficiency"] is None
    assert npt["total_npt"] is None
    assert npt["npt_percentage"] is None
    # Rendered as the repository's unknown marker, not a fabricated number.
    assert fmt_num(kpis["avg_rop"], 1, default=None) == "—"
    assert fmt_num(kpis["npt_percentage"], 1, default=None) == "—"
    assert fmt_num(npt["npt_percentage"], 1, default=None) == "—"

    # Canonical intelligence layer must agree that the metrics are unknown.
    from core.operations_intelligence import OperationsIntelligenceService
    canonical = OperationsIntelligenceService(production_db).analyze_well(well)["kpis"]
    assert canonical["average_rop"] is None
    assert canonical["npt_percent"] is None


def test_analysis_kpis_real_zero_npt_stays_zero(production_db):
    """A well with recorded time and no NPT rows has a genuine 0.0h / 0% NPT and
    100% efficiency — an explicit recorded zero must not be turned into "—"."""
    from datetime import time
    from core.database import TimeLog24H, DrillingParameters
    _, _, well, _, report = production_db.acceptance_ids
    production_db.generic_save(TimeLog24H, {
        "report_id": report, "time_from": time(0, 0), "time_to": time(23, 59),
        "duration": 24.0, "is_npt": False, "main_code": "DRILL",
    })
    production_db.generic_save(DrillingParameters, {
        "well_id": well, "report_id": report, "report_date": date(2026, 9, 8),
        "avg_rop": 12.5,
    })
    kpis, npt = _analysis_probe(production_db, well)
    assert kpis["avg_rop"] == 12.5
    assert kpis["npt_percentage"] == 0.0
    assert kpis["efficiency"] == 100.0
    assert npt["total_npt"] == 0.0
    assert npt["npt_percentage"] == 0.0


def test_planning_npt_unknown_not_zero_then_real_zero(production_db):
    """w10 Planning shares the NPT KPI: unknown when no time is recorded, a real
    0% once time exists with no NPT rows."""
    from datetime import time
    from tabs import w10_Planning_Widget
    from core.database import TimeLog24H

    # NPTReportTab is wrapped by @make_scrollable; recover the undecorated class
    # (captured in the wrapper's closure) so the pure query logic can be exercised
    # without constructing the full Qt scroll-area widget.
    scrollable = w10_Planning_Widget.NPTReportTab
    inner_cls = next(
        cell.cell_contents for cell in scrollable.__init__.__closure__
        if isinstance(cell.cell_contents, type)
        and hasattr(cell.cell_contents, "get_npt_data")
    )

    def npt(well_id):
        probe = SimpleNamespace(db=production_db, current_well_id=well_id,
                                current_report_id=None, current_section_id=None)
        session = production_db.create_session()
        try:
            return inner_cls.get_npt_data(probe, session)
        finally:
            session.close()

    _, _, well, _, report = production_db.acceptance_ids
    unknown = npt(well)
    assert unknown["total_npt"] is None and unknown["npt_percentage"] is None
    production_db.generic_save(TimeLog24H, {
        "report_id": report, "time_from": time(0, 0), "time_to": time(23, 59),
        "duration": 24.0, "is_npt": False, "main_code": "DRILL",
    })
    real = npt(well)
    assert real["total_npt"] == 0.0 and real["npt_percentage"] == 0.0


def test_eowr_and_ddr_time_analysis_no_false_full_productivity(production_db, tmp_path):
    """Report exports must not claim 100% productive / 0% NPT for a well/day that
    has no recorded time (report_engine R-4)."""
    from core.report_engine import EOWRReportEngine, DDRReportEngine
    _, _, well, _, report = production_db.acceptance_ids
    assert DDRReportEngine(production_db).generate(report, str(tmp_path / "ddr.html"), format="html")
    ddr = (tmp_path / "ddr.html").read_text()
    # No time logged -> Time Analysis percentages are unknown, not 100%/0%.
    assert "(100%)" not in ddr
    assert "<b>100%</b>" not in ddr
    assert "—" in ddr
    assert EOWRReportEngine(production_db).generate(well, str(tmp_path / "eowr.html"), format="html")
    eowr = (tmp_path / "eowr.html").read_text()
    assert "<td>None</td>" not in eowr
    assert "—" in eowr
