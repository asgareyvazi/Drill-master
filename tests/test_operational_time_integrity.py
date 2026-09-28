"""Missing time must survive the engine, W12 and export boundaries."""
from datetime import time
from types import SimpleNamespace

import pytest

from core.database import DailyReport, TimeLog24H
from core.operational_time import summarize_time_logs
from core.operations_intelligence import OperationsIntelligenceService
from test_scope_attribution import _mgr, _base, _well, _wellbore, _section, _report


@pytest.mark.parametrize("durations, expected", [
    ([], (None, None, None, None)),
    ([None], (None, None, None, None)),
    ([0], (0, 0, 0, None)),
    ([2, None], (None, None, None, None)),
    ([2, 0], (2, 2, 0, 100)),
])
def test_time_summary_and_all_service_scopes(durations, expected):
    rows = [SimpleNamespace(duration=d, is_npt=True) for d in durations]
    fields = ("total_hours", "npt_hours", "productive_hours", "npt_percent")
    assert tuple(summarize_time_logs(rows)[f] for f in fields) == expected
    m = _mgr()
    w = _well(m, _base(m), "A")
    b = _wellbore(m, w, "A")
    sec = _section(m, w, "S", wellbore_id=b)
    r = _report(m, w, 1, wellbore_id=b, section_id=sec)
    with m.create_session() as s:
        for duration in durations:
            s.add(TimeLog24H(report_id=r, time_from=time(0), time_to=time(1),
                            duration=duration, is_npt=True))
        s.commit()
    service = OperationsIntelligenceService(m)
    for method, key in [(service.analyze_well, w), (service.analyze_wellbore, b),
                        (service.analyze_section, sec)]:
        kpis = method(key)["kpis"]
        assert tuple(kpis[f] for f in fields) == expected


def test_known_npt_is_not_a_complete_percentage_with_unknown_productive_time():
    result = summarize_time_logs([SimpleNamespace(duration=2, is_npt=True),
                                  SimpleNamespace(duration=None, is_npt=False)])
    assert result["npt_hours"] == 2
    assert result["total_hours"] is None
    assert result["npt_percent"] is None


def test_progress_requires_two_consecutive_known_depths():
    m = _mgr()
    w = _well(m, _base(m), "A")
    _report(m, w, 1, depth=100)
    service = OperationsIntelligenceService(m)
    assert service.analyze_well(w)["kpis"]["daily_progress"] is None
    r = _report(m, w, 2, depth=None)
    assert service.analyze_well(w)["kpis"]["daily_progress"] is None
    with m.create_session() as s:
        s.get(DailyReport, r).depth_2400 = 100
        s.commit()
    assert service.analyze_well(w)["kpis"]["daily_progress"] == 0


@pytest.mark.parametrize("duration", [None, 0.0, 2.0])
def test_npt_html_excel_keep_unknown_and_zero(tmp_path, duration):
    from core.report_engine import NPTReportEngine, EOWRReportEngine
    from openpyxl import load_workbook
    m = _mgr()
    w = _well(m, _base(m), "A")
    r = _report(m, w, 1)
    with m.create_session() as s:
        s.add(TimeLog24H(report_id=r, time_from=time(0), time_to=time(1),
                        duration=duration, is_npt=True))
        s.commit()
    engine = NPTReportEngine(m)
    data = engine._collect_data(w)
    assert data["total_npt"] == duration
    assert data["entries"][0]["hours"] == duration
    assert data["by_main_code"]["Unknown"] == duration
    html = engine._build_html(data)
    assert "NPT SUMMARY" in html
    if duration is None:
        assert "—" in html
        assert data["npt_pct"] is None
    path = tmp_path / "npt.xlsx"
    assert engine._save_excel(data, str(path))
    wb = load_workbook(path)
    assert wb["NPT Events"].cell(2, 4).value == duration
    wb.close()
    eowr = EOWRReportEngine(m)._collect_data(w)
    assert eowr["summary"]["total_npt"] == duration


def test_npt_empty_time_with_cost_does_not_fail_or_invent_zero():
    from core.database import CostRecord
    from core.report_engine import NPTReportEngine
    m = _mgr()
    w = _well(m, _base(m), "A")
    _report(m, w, 1)
    with m.create_session() as s:
        s.add(CostRecord(well_id=w, category="Rig", currency="USD", actual_cost=200))
        s.commit()
    engine = NPTReportEngine(m)
    data = engine._collect_data(w)
    assert data is not None
    assert data["total_npt"] is None
    assert data["npt_cost"] is None
    assert "NPT SUMMARY" in engine._build_html(data)


def test_npt_period_uses_period_report_count_and_refuses_whole_well_cost():
    from datetime import date
    from core.database import CostRecord
    from core.report_engine import NPTReportEngine
    m = _mgr()
    w = _well(m, _base(m), "A")
    r = _report(m, w, 1)
    _report(m, w, 2)
    with m.create_session() as s:
        s.add(TimeLog24H(report_id=r, time_from=time(0), time_to=time(2), duration=2, is_npt=True))
        s.add(CostRecord(well_id=w, category="Rig", currency="USD", actual_cost=200))
        s.commit()
    engine = NPTReportEngine(m)
    data = engine._collect_data(w, from_date=date(2026, 1, 1), to_date=date(2026, 1, 1))
    assert data["report_count"] == 1
    assert data["daily_avg"] == 2
    assert data["npt_cost"] is None
    assert engine._collect_data(w)["npt_cost"] == 200
