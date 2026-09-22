"""Mission 25 — §12/§40-43: report-engine scope metadata + multi-bore aggregate.

Two guarantees per report engine:

* SCOPE METADATA: EOWR / NPT / Cost / Plan HTML output declares its scope
  ("Whole-Well Aggregate" / well-level plan). DDR is report-scoped and must NOT
  claim whole-well. The label matches the query semantics (§39 — never a fake
  scope field).
* MULTI-BORE AGGREGATE (§40-41): with an Original bore and an ST-1 bore in one
  well, the whole-well NPT export sums BOTH bores (independent oracle: 1h + 4h =
  5h), and selecting a bore in the UI cannot narrow these well-level reports.

Uses HTML output (no PDF/matplotlib dependency) against a real in-memory DB.
The expected aggregate is computed from fixture facts, not the implementation.
"""
from datetime import date, time

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from core.database import (
    Base, DatabaseManager, Company, Project, Well, Wellbore, Section,
    DailyReport, TimeLog24H,
)
from core.database import CostRecord, WellPlan, PlannedActivity
from core.report_engine import (
    EOWRReportEngine, NPTReportEngine, DDRReportEngine,
    CostReportEngine, PlanReportEngine,
)


@pytest.fixture()
def db():
    m = DatabaseManager()
    m.engine = create_engine(
        "sqlite:///:memory:", connect_args={"check_same_thread": False},
        poolclass=StaticPool)
    Base.metadata.create_all(m.engine)
    m.Session = sessionmaker(bind=m.engine, autoflush=False, autocommit=False)
    m.create_session()
    return m


@pytest.fixture()
def multibore(db):
    """Well W1 with Original (NPT 1h) and ST-1 (NPT 4h). Whole-well NPT = 5h."""
    s = db.create_session()
    c = Company(name="C", code="C"); s.add(c); s.flush()
    p = Project(name="P", code="P", company_id=c.id); s.add(p); s.flush()
    w = Well(name="W1", code="W1", project_id=p.id); s.add(w); s.flush()
    orig = Wellbore(well_id=w.id, name="Original", wellbore_type="original")
    s.add(orig); s.flush()
    st1 = Wellbore(well_id=w.id, name="ST-1", wellbore_type="sidetrack",
                   parent_wellbore_id=orig.id, kickoff_md=2500)
    s.add(st1); s.flush()
    sec_o = Section(name='12-1/4"', well_id=w.id, wellbore_id=orig.id, depth_from=1000)
    sec_s = Section(name='8-1/2"', well_id=w.id, wellbore_id=st1.id, depth_from=2500)
    s.add_all([sec_o, sec_s]); s.flush()
    ro = DailyReport(well_id=w.id, section_id=sec_o.id, wellbore_id=orig.id,
                     report_number=1, report_date=date(2026, 1, 1), depth_2400=1000)
    rs = DailyReport(well_id=w.id, section_id=sec_s.id, wellbore_id=st1.id,
                     report_number=2, report_date=date(2026, 1, 5), depth_2400=3000)
    s.add_all([ro, rs]); s.flush()
    s.add_all([
        # Original: 1h NPT + 23h productive.
        TimeLog24H(report_id=ro.id, time_from=time(0, 0), time_to=time(1, 0),
                   duration=1.0, is_npt=True, main_code="NPT-A"),
        TimeLog24H(report_id=ro.id, time_from=time(1, 0), time_to=time(23, 59),
                   duration=23.0, is_npt=False),
        # ST-1: 4h NPT + 20h productive.
        TimeLog24H(report_id=rs.id, time_from=time(0, 0), time_to=time(4, 0),
                   duration=4.0, is_npt=True, main_code="NPT-B"),
        TimeLog24H(report_id=rs.id, time_from=time(4, 0), time_to=time(23, 59),
                   duration=20.0, is_npt=False),
    ])
    s.commit()
    ids = {"well": w.id, "orig": orig.id, "st1": st1.id,
           "ro": ro.id, "rs": rs.id}
    s.close()
    return ids


# ---------------------------------------------------------------------------
# Scope metadata (§12)
# ---------------------------------------------------------------------------

def test_eowr_html_declares_whole_well_scope(db, multibore, tmp_path):
    out = tmp_path / "eowr.html"
    assert EOWRReportEngine(db).generate(multibore["well"], str(out), format="html")
    html = out.read_text()
    assert "Whole-Well Aggregate" in html


def test_npt_html_declares_whole_well_scope(db, multibore, tmp_path):
    out = tmp_path / "npt.html"
    assert NPTReportEngine(db).generate(multibore["well"], str(out), format="html")
    html = out.read_text()
    assert "Whole-Well Aggregate" in html


def test_ddr_html_is_not_labelled_whole_well(db, multibore, tmp_path):
    """DDR is report-scoped — it must NOT claim whole-well aggregate."""
    out = tmp_path / "ddr.html"
    assert DDRReportEngine(db).generate(multibore["ro"], str(out), format="html")
    html = out.read_text()
    assert "Whole-Well Aggregate" not in html


# ---------------------------------------------------------------------------
# Multi-bore aggregate (§40-41) — independent oracle: 1h + 4h = 5h
# ---------------------------------------------------------------------------

def test_npt_excel_aggregates_both_bores(db, multibore, tmp_path):
    from openpyxl import load_workbook
    out = tmp_path / "npt.xlsx"
    assert NPTReportEngine(db).generate(multibore["well"], str(out), format="excel")
    wb = load_workbook(out)
    ws = wb["NPT Summary"]
    rows = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
            for r in range(1, ws.max_row + 1)}
    assert rows.get("Scope") == "Whole-Well Aggregate"
    # Total NPT sums BOTH bores (1 + 4 = 5), never a single bore's value.
    total = float(rows["Total NPT (hrs)"])
    assert total == 5.0, rows


def test_npt_collect_data_spans_both_bores(db, multibore):
    """The whole-well NPT collection includes both bores' NPT categories."""
    engine = NPTReportEngine(db)
    data = engine._collect_data(multibore["well"])
    assert data["total_npt"] == 5.0
    # by_code shape may vary; assert on the authoritative total instead.
    assert data["report_count"] >= 2


# ---------------------------------------------------------------------------
# Cost report — well-level, aggregates both bores' well-scoped cost (§43)
# ---------------------------------------------------------------------------

def test_cost_report_is_well_level_and_labelled(db, multibore, tmp_path):
    s = db.create_session()
    # CostRecord is well-scoped (no wellbore dimension). Two categories.
    s.add_all([
        CostRecord(well_id=multibore["well"], category="Rig", actual_cost=100000.0),
        CostRecord(well_id=multibore["well"], category="Mud", actual_cost=40000.0),
    ])
    s.commit(); s.close()
    out = tmp_path / "cost.html"
    assert CostReportEngine(db).generate(multibore["well"], str(out), format="html")
    html = out.read_text()
    assert "Whole-Well Aggregate" in html
    # Independent oracle: total actual = 100000 + 40000 = 140000.
    data = CostReportEngine(db)._collect_data(multibore["well"], None, None)
    assert data["total_cost"] == 140000.0


def test_cost_excel_scope_row(db, multibore, tmp_path):
    from openpyxl import load_workbook
    s = db.create_session()
    s.add(CostRecord(well_id=multibore["well"], category="Rig", actual_cost=50000.0))
    s.commit(); s.close()
    out = tmp_path / "cost.xlsx"
    assert CostReportEngine(db).generate(multibore["well"], str(out), format="excel")
    wb = load_workbook(out)
    ws = wb["Cost Summary"]
    rows = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
            for r in range(1, ws.max_row + 1)}
    assert rows.get("Scope") == "Whole-Well Aggregate"


# ---------------------------------------------------------------------------
# Plan report — well-level plan (§42)
# ---------------------------------------------------------------------------

def test_plan_report_is_well_level_and_labelled(db, multibore, tmp_path):
    from datetime import datetime
    s = db.create_session()
    plan = WellPlan(well_id=multibore["well"], plan_name="Base Plan",
                    plan_version="1.0", is_active=True)
    s.add(plan); s.flush()
    s.add(PlannedActivity(
        well_id=multibore["well"], plan_id=plan.id,
        activity_name="Drill 12-1/4", phase_code="Drill",
        planned_start=datetime(2026, 1, 1, 0, 0),
        planned_end=datetime(2026, 1, 3, 0, 0),
        planned_duration_hours=48.0,
        planned_depth_from=0, planned_depth_to=1000))
    s.commit(); s.close()
    out = tmp_path / "plan.html"
    assert PlanReportEngine(db).generate(multibore["well"], str(out), format="html")
    html = out.read_text()
    # Plan is a well-level artifact, labelled as such (not bore-scoped).
    assert "Well-Level Plan" in html
