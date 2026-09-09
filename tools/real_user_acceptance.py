"""Isolated Production service acceptance, NOT desktop/Windows certification.

Creates a NEW evidence directory and generated admin credential; never resets an
existing DB. The original workbook is hashed, read only, and never rewritten.
Usage: python tools/real_user_acceptance.py WORKBOOK --output NEW_DIRECTORY
"""
from __future__ import annotations

import argparse
import copy
from datetime import date, datetime, time
import hashlib
import json
import logging
import os
from pathlib import Path
import platform
import secrets
import sqlite3
import sys

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core import database as models
from core.database import DatabaseManager, Company, Project, Well, Section, User, Base
from core.ddr_import_service import DDRImportService
from core.save_outcome import save_all
from tools.certify_ddr import snapshot


def integrity(db):
    with sqlite3.connect(db.db_path) as con:
        return {"integrity_check": con.execute("PRAGMA integrity_check").fetchall(),
                "foreign_key_check": con.execute("PRAGMA foreign_key_check").fetchall(),
                "report_well_mismatches": {table: con.execute(
                    f'SELECT count(*) FROM "{table}" c JOIN daily_reports r ON c.report_id=r.id WHERE c.well_id != r.well_id'
                ).fetchone()[0] for table in Base.metadata.tables
                    if {"well_id", "report_id"}.issubset(Base.metadata.tables[table].columns.keys())}}


def open_db(path, bootstrap=None):
    db = DatabaseManager(bootstrap_passwords=bootstrap)
    db.db_path = str(path)
    if not db.initialize():
        raise RuntimeError(db.last_diagnostic)
    return db


def schema_crud_probe(db, ids):
    """Mechanical ORM/API lifecycle for all models, NOT domain validation proof.

    Required non-FK fields get explicitly synthetic QA values; optional domain
    measurements remain absent. Credentials and the system audit log are not fabricated; history-table probes
    are mechanical only and do not certify approval workflows.
    Existing hierarchy is explicit; new parent IDs come from actual insert results.
    """
    from sqlalchemy import Boolean, Date, DateTime, Float, Integer, JSON, Time
    repo = {m.local_table.name: m.class_ for m in Base.registry.mappers}
    excluded = {"users", "audit_logs", "companies", "projects", "wells", "sections", "daily_reports"}
    inserted, result = [], []
    by_table = dict(ids)
    for table in Base.metadata.sorted_tables:
        model = repo.get(table.name)
        if model is None or table.name in excluded:
            continue
        values = {}
        for column in table.columns:
            if column.primary_key:
                continue
            if column.foreign_keys:
                fk = next(iter(column.foreign_keys))
                target = by_table.get(fk.column.table.name)
                if target is not None:
                    values[column.name] = target
                continue
            if column.nullable or column.default is not None or column.server_default is not None:
                continue
            typ = column.type
            if isinstance(typ, DateTime): value = datetime(2026, 9, 8, 8)
            elif isinstance(typ, Date): value = date(2026, 9, 8)
            elif isinstance(typ, Time): value = time(8)
            elif isinstance(typ, Boolean): value = False
            elif isinstance(typ, Integer): value = 1
            elif isinstance(typ, Float): value = 1.0
            elif isinstance(typ, JSON): value = []
            else: value = ("QA " + model.__name__)[:getattr(typ, "length", None)]
            values[column.name] = value
        entry = {"model": model.__name__, "scope": "synthetic schema CRUD; not engineering input acceptance"}
        try:
            identity = db.generic_save(model, values)
            by_table[table.name] = identity
            inserted.append((model, identity, entry))
            entry["create"] = True
            editable = next((c.name for c in table.columns if c.name in ("remarks", "description", "notes", "summary")), None)
            edit = {"id": identity}
            if editable:
                edit[editable] = "Explicit QA edit — ویرایش"
            assert db.generic_save(model, edit) == identity
            row = db.generic_get_list(model, {"id": identity})[0]
            assert not editable or row[editable] == edit[editable]
            entry["edit_reload"] = True
        except Exception as exc:
            entry["error"] = f"{type(exc).__name__}: {exc}"
        result.append(entry)
    # Actual new engine/connection, not session identity-map caching.
    db.close()
    assert db.initialize()
    for model, identity, entry in inserted:
        entry["reopen"] = bool(db.generic_get_list(model, {"id": identity}))
    for model, identity, entry in reversed(inserted):
        try:
            entry["delete"] = db.generic_delete(model, identity)
            entry["deleted_reload"] = not db.generic_get_list(model, {"id": identity})
        except Exception as exc:
            entry["delete_error"] = f"{type(exc).__name__}: {exc}"
    return result


def run(workbook, output):
    workbook, output = Path(workbook).resolve(), Path(output).resolve()
    output.mkdir(parents=True, exist_ok=False)
    os.environ["DRILLMASTER_ENV"] = "production"
    os.environ["DRILLMASTER_DATA_DIR"] = str(output / "profile")
    logging.basicConfig(filename=output / "runtime.log", level=logging.INFO, force=True)
    digest = hashlib.sha256(workbook.read_bytes()).hexdigest()
    secret = secrets.token_urlsafe(24)
    db = open_db(output / "acceptance.db", {"admin": secret})
    evidence = {"platform": platform.platform(), "python": sys.version,
                "mode": "production", "database": str(db.db_path), "source_sha256": digest,
                "native_gui": "BLOCKED: missing system Qt/GL libraries; NOT VERIFIED", "windows_python312": "NOT VERIFIED",
                "golden_b": "NOT AVAILABLE / NOT VERIFIED"}
    with db.session_scope() as session:
        evidence["first_run"] = {"users": session.query(User).count(), "wells": session.query(Well).count(), "companies": session.query(Company).count()}
    assert evidence["first_run"] == {"users": 1, "wells": 0, "companies": 0}
    assert db.authenticate_user("admin", secret)
    assert not db.authenticate_user("admin", secrets.token_urlsafe(24))
    assert not db.authenticate_user("not-an-account", secret)
    db.close()
    db = open_db(output / "acceptance.db")
    assert db.authenticate_user("admin", secret)
    del secret
    evidence["authentication"] = "correct/wrong/unknown/reopen without bootstrap PASS; service level"
    company = db.generic_save(Company, {"name": "Acceptance operator", "code": "PAT"})
    project = db.generic_save(Project, {"name": "Production Acceptance Test", "code": "PAT", "company_id": company})
    assert db.save_well({"name": "TEST-WELL-001", "code": "TEST-WELL-001", "project_id": project, "spud_date": "2026-09-01"})
    well = db.generic_get_list(Well, {"code": "TEST-WELL-001"})[0]["id"]
    section = db.generic_save(Section, {"name": "12.25 in test section", "well_id": well, "hole_size": 12.25})
    report = db.save_daily_report({"well_id": well, "section_id": section, "report_number": 1,
                                  "report_date": date(2026, 9, 8), "summary": "Production Acceptance Test — آزمون تولید"})["id"]
    evidence["manual_context"] = {"company": company, "project": project, "well": well, "section": section, "report": report}
    assert db.save_well({"id": well, "rig_name": "Acceptance rig"})
    assert db.get_well_by_id(well)["rig_name"] == "Acceptance rig"
    try:
        db.save_well({"id": well, "spud_date": "invalid date"})
        raise AssertionError("Invalid date was accepted")
    except ValueError:
        pass
    assert db.get_well_by_id(well)["spud_date"] == date(2026, 9, 1)
    manual = {"formations": [{"formation_name": "Sandstone", "top_md": 0, "base_md": 100, "color": None}],
              "bha_components": [{"component_name": "PDC bit", "od": 12.25, "length": 0.3}, {"component_name": "Drill collar", "od": 8, "id": 2.8, "length": 9}],
              "downhole_equipment": [{"equipment_name": "MWD", "rotation_hours": 0, "install_date": None}],
              "surveys": [{"md": 0, "inc": 0, "azi": 0}, {"md": 100, "inc": 10, "azi": 45}, {"md": 200, "inc": 20, "azi": 60}, {"md": 300, "inc": 20, "azi": None}]}
    evidence["manual_operational_save"] = db.save_imported_multi_tab_data_atomic(well, report, manual)
    assert not evidence["manual_operational_save"]["failed"]
    db.save_formation_report(well, {"report_id": report, "formations": manual["formations"]})
    assert db.get_formation_report(well, report)["formations"][0]["Color"] is None
    evidence["schema_crud"] = schema_crud_probe(db, {"companies": company, "projects": project, "wells": well,
                                                     "sections": section, "daily_reports": report, "users": 1})
    service = DDRImportService(db, well)
    extraction, payload = service.extract_file(str(workbook))
    result = service.import_records(payload)
    assert result["report_id"] and result["failed"] == 0, result
    imported_well, imported_report = result["well_id"], result["report_id"]
    evidence["golden_a"] = result
    evidence["extraction"] = {"template_version": extraction.template_version}
    before = snapshot(db, imported_report)
    db.close()
    db = open_db(output / "acceptance.db")
    assert snapshot(db, imported_report) == before
    replay = DDRImportService(db, imported_well).import_records(payload)
    assert not replay["failed"] and snapshot(db, imported_report) == before
    evidence["import_reopen_replay_equal"] = True
    evidence["imported_counts"] = {name: len(rows) for name, rows in before.items()}
    # Explicit service-prepared edits. No widget-interaction claim.
    mud = copy.deepcopy(db.generic_get_list(models.MudReport, {"report_id": imported_report})[0])
    mud["summary"] = "Acceptance edit — ویرایش"
    chemicals = json.loads(mud["chemicals_json"])
    original = copy.deepcopy(chemicals)
    chemicals[0]["received"] = 2
    mud["chemicals_json"] = json.dumps(chemicals)
    bha = copy.deepcopy(db.generic_get_list(models.BHAReport, {"report_id": imported_report})[0])
    bha["bha_data_json"][0]["Remarks"] = "Acceptance edit"
    formation = copy.deepcopy(db.generic_get_list(models.FormationReport, {"report_id": imported_report})[0])
    formation["formations"] = formation.pop("formations_json")
    formation["formations"][0]["Description"] = "Acceptance edit"
    outcome = save_all([("Mud", lambda: db.save_mud_report(mud)),
                        ("BHA", lambda: db.save_bha_report(imported_well, bha)),
                        ("Formation", lambda: db.save_formation_report(imported_well, formation))])
    evidence["save_all_valid"] = outcome.to_dict()
    assert outcome, outcome.summary()
    edited = snapshot(db, imported_report)
    saved_chemicals = json.loads(edited["MudReport"][0]["chemicals_json"])
    assert saved_chemicals[0]["received"] == 2 and saved_chemicals[0]["stock"] == original[0]["stock"]
    assert saved_chemicals[1:] == original[1:]
    assert edited["MudReport"][0]["water_percent"] is None
    db.close()
    db = open_db(output / "acceptance.db")
    assert snapshot(db, imported_report) == edited
    evidence["edit_reopen_equal"] = True
    def invalid(): raise ValueError("Explicit invalid QA input")
    def system_failure(): raise RuntimeError("Injected QA failure")
    evidence["save_all_partial"] = save_all([("valid", lambda: True), ("invalid", invalid), ("system fault", system_failure)]).to_dict()
    evidence["save_all_no_callbacks"] = save_all([]).to_dict()
    from core.report_engine import DDRReportEngine, EOWRReportEngine, NPTReportEngine, CostReportEngine, PlanReportEngine
    from core.professional_export import ProfessionalExcelExport
    exports = {}
    for cls in (DDRReportEngine, EOWRReportEngine, NPTReportEngine, CostReportEngine, PlanReportEngine):
        for fmt, ext in (("html", "html"), ("excel", "xlsx")):
            destination = output / f"{cls.__name__}.{ext}"
            ok = cls(db).generate(imported_report if cls is DDRReportEngine else imported_well, str(destination), format=fmt)
            exports[destination.name] = {"returned_success": ok, "bytes": destination.stat().st_size if destination.exists() else 0}
    # A missing plan is expected absence, not a fabricated drilling schedule.
    plan = db.generic_save(models.WellPlan, {"well_id": imported_well, "plan_name": "Explicit acceptance plan",
        "is_active": True, "planned_final_depth": 1500, "planned_total_days": 3})
    db.generic_save(models.PlannedActivity, {"well_id": imported_well, "plan_id": plan,
        "activity_name": "Drill the next interval", "phase_code": "Drilling",
        "planned_start": datetime(2026, 9, 9, 0), "planned_end": datetime(2026, 9, 9, 12),
        "planned_duration_hours": 12, "planned_depth_from": 1000, "planned_depth_to": 1500})
    for fmt, ext in (("html", "html"), ("excel", "xlsx")):
        destination = output / f"Plan-with-data.{ext}"
        ok = PlanReportEngine(db).generate(imported_well, str(destination), format=fmt)
        assert ok, "Valid planned-activity export failed"
        exports[destination.name] = {"returned_success": ok, "bytes": destination.stat().st_size}
    destination = output / "Professional.xlsx"
    ok = ProfessionalExcelExport(db).export(imported_well, str(destination), report_id=imported_report)
    exports[destination.name] = {"returned_success": ok, "bytes": destination.stat().st_size if destination.exists() else 0}
    from openpyxl import load_workbook
    for filename, info in exports.items():
        if filename.endswith("xlsx") and info["bytes"]:
            wb = load_workbook(output / filename, read_only=True)
            info["sheets"] = wb.sheetnames
            info["rows"] = {ws.title: ws.max_row for ws in wb.worksheets}
            wb.close()
    evidence["exports"] = exports
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.figure import Figure
    from core.trajectory_plot import draw_trajectory_3d
    chart_results = {}
    for label, rows in (("empty", []), ("golden-a", db.load_survey_points(report_id=imported_report)), ("manual-valid-and-missing", db.load_survey_points(report_id=report))):
        fig = Figure()
        chart_results[label] = draw_trajectory_3d(fig.add_subplot(111, projection="3d"), rows)
        fig.savefig(output / f"trajectory-{label}.png")
    evidence["agg_charts_not_interactive"] = chart_results
    evidence["integrity_before_delete"] = integrity(db)
    evidence["report_delete"] = db.delete_daily_report(imported_report)
    assert not db.get_daily_report_by_id(imported_report)
    evidence["report_children_after_delete"] = {name: len(rows) for name, rows in snapshot(db, imported_report).items()}
    evidence["well_delete"] = db.delete_well(imported_well)
    evidence["manual_well_survives"] = bool(db.get_well_by_id(well))
    evidence["integrity_after_delete"] = integrity(db)
    db.close()
    assert hashlib.sha256(workbook.read_bytes()).hexdigest() == digest
    evidence["original_workbook_unchanged"] = True
    (output / "acceptance.json").write_text(json.dumps(evidence, default=str, indent=2), encoding="utf-8")
    print(json.dumps({key: evidence[key] for key in ("first_run", "authentication", "imported_counts", "import_reopen_replay_equal", "edit_reopen_equal", "exports", "well_delete", "integrity_after_delete")}, indent=2))
    return evidence


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    run(args.workbook, args.output)
