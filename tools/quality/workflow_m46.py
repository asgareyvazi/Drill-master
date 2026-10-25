#!/usr/bin/env python3
"""M46 end-to-end workflows through the real services, SYNTHETIC data only.

Document flow (14 steps): import -> extract -> review -> approve/reject ->
audit -> Persian search -> highlight -> link well/report -> open from the
well / report -> change -> remove (document stays) -> reindex -> duplicate
-> no-FTS5 fallback.

Survey flow (10 steps): stored per-report rows -> whole-well composition ->
an INDEPENDENT minimum-curvature reproduction (written here, not imported
from the application) station by station -> analytic J-profile end point ->
REST API, EOWR engine and professional Excel export compared with the
independent reference -> feet display -> invalid survey well -> survey
editor scope (W1).

OCR: a deterministic stand-in engine returns fixed SYNTHETIC text so the
reviewable OCR path runs without OCR model weights.  It is labelled as such
in the output and says nothing about real OCR accuracy.

Usage::

    QT_QPA_PLATFORM=offscreen python tools/quality/workflow_m46.py [--json OUT]

Exit status 0 only when every step passes.
"""
from __future__ import annotations

import argparse
import json
import math
import os
import sys
import tempfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
ARABIC = "\u06af\u0632\u0627\u0631\u0634 \u0644\u0648\u0644\u0647 \u062c\u062f\u0627\u0631\u064a SYN-DEV-01 \u0643\u064a\u0633\u064a\u0646\u06af"


class StandInOcr:
    """Deterministic stand-in for an OCR engine (NOT an OCR model)."""
    name = "stand-in (synthetic, no model)"

    def recognize(self, image_bytes):
        return "SYNTHETIC scanned line SYN-DEV-01 pump pressure 2500 psi", 0.42


def _independent_min_curvature(rows):
    """Textbook minimum curvature (ratio factor RF = 2/b * tan(b/2)), tie-on at
    the first station with N = E = 0 and TVD = MD * cos(inc).  Written
    independently of core.well_trajectory."""
    out = []
    for i, (md, inc, azi) in enumerate(rows):
        if i == 0:
            out.append((md, md * math.cos(math.radians(inc)), 0.0, 0.0))
            continue
        md0, inc0, azi0 = rows[i - 1]
        i1, i2, a1, a2 = map(math.radians, (inc0, inc, azi0, azi))
        cos_b = math.cos(i2 - i1) - math.sin(i1) * math.sin(i2) * (1 - math.cos(a2 - a1))
        b = math.acos(max(-1.0, min(1.0, cos_b)))
        rf = 1.0 if b < 1e-9 else 2.0 / b * math.tan(b / 2.0)
        dmd = md - md0
        _, tvd, n, e = out[-1]
        out.append((md,
                    tvd + dmd / 2 * (math.cos(i1) + math.cos(i2)) * rf,
                    n + dmd / 2 * (math.sin(i1) * math.cos(a1) + math.sin(i2) * math.cos(a2)) * rf,
                    e + dmd / 2 * (math.sin(i1) * math.sin(a1) + math.sin(i2) * math.sin(a2)) * rf))
    return out


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--json")
    args = parser.parse_args(argv)
    work = Path(tempfile.mkdtemp(prefix="drillmaster-m46-workflow-"))
    os.environ.update({"DRILLMASTER_ENV": "test", "DRILLMASTER_DATA_DIR": str(work),
                       "DRILLMASTER_DB_PATH": str(work / "synthetic.db"), "DRILLMASTER_LOG_DIR": str(work / "logs"),
                       "DRILLMASTER_BACKUP_DIR": str(work / "backups"), "DRILLMASTER_AI_IMPORT": "0"})
    sys.path[:0] = [str(ROOT), str(ROOT / "tests")]
    from sqlalchemy import text

    from core import search_normalization as sn
    from core.document_library import DocumentLibrary, PageDecision, extract_document
    from core.document_links import DocumentLinkService
    from core.permissions import permissions, session_user
    from synthetic_operations import build_synthetic_operations, make_file_database

    steps = []

    def step(flow, name, ok, detail=""):
        steps.append({"flow": flow, "step": len([s for s in steps if s["flow"] == flow]) + 1, "name": name,
                      "status": "PASS" if ok else "FAIL", "detail": detail})
        print(f"{flow:8s} {steps[-1]['step']:2d} {'PASS' if ok else 'FAIL'}  {name}  {detail}")

    def rows(sql, **params):
        with db.session_scope() as s:
            return s.execute(text(sql), params).fetchall()

    db = make_file_database(work / "synthetic.db")
    ids = build_synthetic_operations(db, end_date=date(2026, 1, 15))
    permissions.set_user(session_user(db.authenticate_user("engineer", "Synthetic-Eng-1!")))
    db.enforce_authorization(permissions)

    # ------------------------------------------------------------- documents
    folder = work / "docs"
    folder.mkdir()
    (folder / "syn_fa_note.txt").write_bytes(ARABIC.encode("cp1256"))
    (folder / "syn_scan.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"SYNTHETIC-not-a-real-image" * 4)
    (folder / "syn_scan_rejected.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"SYNTHETIC-rejected" * 4)
    library = DocumentLibrary(db)
    fa = extract_document(folder / "syn_fa_note.txt")
    scan = extract_document(folder / "syn_scan.png", ocr_engine=StandInOcr())
    rejected = extract_document(folder / "syn_scan_rejected.png", ocr_engine=StandInOcr())
    step("document", "import: files extracted", fa.importable and scan.importable,
         f"text={fa.status}, png={scan.status}")
    step("document", "extract: text page ok, OCR page needs review",
         fa.pages[0].status == "ok" and scan.pages[0].status == "needs_review" and scan.pages[0].method == "ocr")
    from core.document_library import default_decisions
    step("document", "review: OCR unapproved until a decision",
         default_decisions(scan)[1].approved is False and default_decisions(fa)[1].approved is True)
    out_fa = library.import_reviewed(fa)
    edited = "SYNTHETIC scanned line SYN-DEV-01 pump pressure 2500 psi (verified)"
    out_scan = library.import_reviewed(scan, {1: PageDecision(approved=True, text=edited)})
    out_rej = library.import_reviewed(rejected, {1: PageDecision(approved=False)})
    stored = rows("SELECT text, review_status FROM document_pages WHERE document_id=:d", d=out_scan.document_id)
    step("document", "approve/reject: approved OCR stored as edited, rejected file stores nothing",
         out_fa.status == out_scan.status == "imported" and out_rej.status == "nothing_approved"
         and stored == [(edited, "edited")])
    audit = rows("SELECT details FROM audit_logs WHERE action='document_import' AND entity_id=:d", d=out_scan.document_id)
    details = json.loads(audit[0][0]) if audit else {}
    step("document", "audit: import audited with OCR approval and edit",
         details.get("ocr_pages_approved") == [1] and details.get("edited_pages") == [1])
    persian_query = "\u0644\u0648\u0644\u0647 \u062c\u062f\u0627\u0631\u06cc \u06a9\u06cc\u0633\u06cc\u0646\u06af"   # Persian yeh/keheh
    result = library.search(persian_query)
    step("document", "Persian search finds Arabic-letter text",
         [h["document_id"] for h in result["hits"]] == [out_fa.document_id], result["method"].split(" (")[0])
    original = rows("SELECT text FROM document_pages WHERE document_id=:d", d=out_fa.document_id)[0][0]
    spans = sn.find_spans(original, sn.query_terms(persian_query))
    step("document", "highlight: spans cover the ORIGINAL (unnormalised) words",
         len(spans) == 3 and original[spans[0][0]:spans[0][1]] == "\u0644\u0648\u0644\u0647"
         and "\u064a" in original, f"spans={len(spans)}")
    links = DocumentLinkService(db)
    suggestions = links.suggest_wells(out_fa.document_id)
    before_links = rows("SELECT COUNT(*) FROM document_links")[0][0]
    link_id = links.create_link(out_fa.document_id, ids["dev"], ids["reports"][6], page_no=1)
    step("document", "link: suggestion only, link after explicit create",
         [s.well_id for s in suggestions][:1] == [ids["dev"]] and before_links == 0 and link_id)
    for_well = links.links_for_well(ids["dev"])
    for_report = links.links_for_report(ids["reports"][6])
    step("document", "open from Well Info / Daily Report: listed with page anchor",
         [(r["document_id"], r["page_no"]) for r in for_well] == [(out_fa.document_id, 1)]
         and [r["document_id"] for r in for_report] == [out_fa.document_id])
    links.update_link(link_id, ids["dev"], ids["reports"][5], page_no=1)
    change = rows("SELECT details FROM audit_logs WHERE action='document_link_change'")
    step("document", "change link: audited before/after",
         len(change) == 1 and json.loads(change[0][0])["to"]["report_id"] == ids["reports"][5])
    links.remove_link(link_id)
    step("document", "remove link: document and pages stay",
         rows("SELECT COUNT(*) FROM document_links")[0][0] == 0 and library.document(out_fa.document_id) is not None
         and rows("SELECT text FROM document_pages WHERE document_id=:d", d=out_fa.document_id)[0][0] == original)
    rebuilt = library.rebuild_index()
    again = library.search(persian_query)
    step("document", "reindex: same hits, original text untouched",
         [h["document_id"] for h in again["hits"]] == [out_fa.document_id]
         and rows("SELECT text FROM document_pages WHERE document_id=:d", d=out_fa.document_id)[0][0] == original,
         f"normalisation v{rebuilt.get('normalization_version', sn.NORMALIZATION_VERSION)}")
    dup = library.import_reviewed(extract_document(folder / "syn_fa_note.txt"))
    step("document", "duplicate: same content refused, original kept",
         dup.status == "duplicate" and dup.document_id == out_fa.document_id)
    library._ensure_index = lambda: None                     # as on a SQLite build without FTS5
    fallback = library.search(persian_query)
    library.close()
    step("document", "no FTS5: fallback scan gives the same hit",
         fallback["method"].startswith("Substring scan")
         and [h["document_id"] for h in fallback["hits"]] == [out_fa.document_id])

    # ---------------------------------------------------------------- survey
    from core.well_trajectory import FT_PER_M, READY, load_well_trajectory
    stored_rows = db.load_survey_points(well_id=ids["dev"])
    reports = sorted({r["report_id"] for r in stored_rows})
    step("survey", "stored rows span six reports", len(stored_rows) == 66 and len(reports) == 6,
         f"rows={len(stored_rows)}")
    trajectory = load_well_trajectory(db, ids["dev"])
    step("survey", "composition: distinct stations, READY", trajectory.status == READY and len(trajectory.stations) == 61)
    distinct = sorted({r["md"]: (r["md"], r["inc"], r["azi"]) for r in stored_rows}.values())
    reference = _independent_min_curvature(distinct)
    worst = max(max(abs(s.tvd - r[1]), abs(s.north - r[2]), abs(s.east - r[3]))
                for s, r in zip(trajectory.stations, reference))
    step("survey", "independent minimum curvature matches every station",
         len(reference) == 61 and worst < 1e-6, f"max |diff| = {worst:.2e} m")
    radius = 30.0 / math.radians(2.0)
    tvd_ref = 600 + radius * math.sin(math.radians(40)) + 600 * math.cos(math.radians(40))
    hd_ref = radius * (1 - math.cos(math.radians(40))) + 600 * math.sin(math.radians(40))
    end = trajectory.stations[-1]
    step("survey", "analytic J-profile end point (TVD, HD, N, E)",
         abs(end.tvd - tvd_ref) < 0.05 and abs(end.hd - hd_ref) < 0.05
         and abs(end.north + hd_ref * math.cos(math.radians(45))) < 0.05,
         f"TVD {end.tvd:.2f} vs {tvd_ref:.2f}, HD {end.hd:.2f} vs {hd_ref:.2f}")
    from core.api.rest_api import DrillMasterAPI
    api = DrillMasterAPI(db).get_trajectory(ids["dev"])
    api_worst = max(abs(a["tvd"] - r[1]) + abs(a["north"] - r[2]) + abs(a["east"] - r[3]) for a, r in zip(api, reference))
    step("survey", "REST get_trajectory equals the independent reference", len(api) == 61 and api_worst < 1e-6)
    from core.report_engine import EOWRReportEngine
    eowr = EOWRReportEngine(db)._collect_data(ids["dev"])["surveys"]
    eowr_worst = max(abs(s.tvd - r[1]) + abs(s.north - r[2]) + abs(s.east - r[3]) for s, r in zip(eowr, reference))
    step("survey", "EOWR surveys equal the independent reference", len(eowr) == 61 and eowr_worst < 1e-6)
    from openpyxl import load_workbook

    from core.professional_export import ProfessionalExcelExport
    xlsx = work / "pro.xlsx"
    ok_export = ProfessionalExcelExport(db).export(ids["dev"], str(xlsx))
    sheet = load_workbook(xlsx)["Survey"]
    header = [c.value for c in sheet[1]]
    cols = [header.index(h) for h in ("TVD (m)", "North (m)", "East (m)")]
    xrows = [r for r in sheet.iter_rows(min_row=2, values_only=True) if isinstance(r[cols[0]], (int, float))]
    x_worst = max(sum(abs(row[c] - ref) for c, ref in zip(cols, r[1:])) for row, r in zip(xrows, reference))
    step("survey", "Excel export equals the independent reference", ok_export and len(xrows) == 61 and x_worst < 1e-6)
    from core.trajectory_3d import scene_from_trajectory
    feet = scene_from_trajectory(trajectory, "ft")
    step("survey", "feet display is an exact conversion", abs(feet.tvd[-1] - end.tvd * FT_PER_M) < 1e-9)
    from core.well_trajectory import NO_DATA, REVIEW_REQUIRED
    qa = load_well_trajectory(db, ids["qa"])
    qa_scene = scene_from_trajectory(qa, "m")
    codes = {issue.code for issue in qa.issues}
    vert = scene_from_trajectory(load_well_trajectory(db, ids["vert"]), "m")
    step("survey", "invalid rows: review status, reasons, only valid stations drawn; no data draws nothing",
         qa.status == REVIEW_REQUIRED and list(qa_scene.md) == [0.0, 100.0] and qa.excluded_rows >= 2
         and {"CONFLICTING_DUPLICATE", "MISSING_ANGLE"} <= codes and not vert.drawable
         and load_well_trajectory(db, ids["vert"]).status == NO_DATA,
         f"SYN-QA-03 {qa.status}, excluded={qa.excluded_rows}; SYN-VERT-02 drawable={vert.drawable}")
    from PySide6.QtWidgets import QApplication

    from tabs.w6_Trajectory_Widget import SurveyDataTab
    QApplication.instance() or QApplication([])
    editor = SurveyDataTab(db)
    editor.current_well_id, editor.current_report_id = ids["dev"], None
    editor.load_data()
    well_rows = editor.survey_table.rowCount()
    editor.current_report_id = ids["reports"][6]
    editor.load_data()
    step("survey", "survey editor: well scope lists well-level rows only (W1)",
         well_rows == 0 and editor.survey_table.rowCount() == len(db.load_survey_points(report_id=ids["reports"][6])),
         f"well scope rows={well_rows}")

    db.engine.dispose()
    failed = [s for s in steps if s["status"] != "PASS"]
    print(f"\nsteps={len(steps)} passed={len(steps) - len(failed)} failed={len(failed)}")
    if args.json:
        Path(args.json).write_text(json.dumps({"ocr_engine": StandInOcr.name, "steps": steps}, indent=2,
                                              ensure_ascii=False), encoding="utf-8")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
