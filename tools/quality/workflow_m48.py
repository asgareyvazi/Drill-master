#!/usr/bin/env python3
"""M48 Operations Analysis workflows through the real W12 page, SYNTHETIC data.

A temporary FILE database (real initialisation, a real authenticated
engineer, no fabricated principal) is filled with the M48 synthetic wells
(tests/m48_synthetic.py). The real ``AnalysisWidget`` is driven through the
real ``SelectionManager`` and its own buttons. Every expected value is read
from the file with an independent ``sqlite3`` connection and plain SQL, never
from the analysis code under test.

Flows:
  performance: KPI -> evidence trend -> point -> Open Source Report (report_id)
               -> typed CSV export -> switch well (nothing stale)
  npt:         coverage states -> NPT projection NOT ASSESSED -> typed XLSX
  cost-risk:   mixed currency never summed; single-currency golden;
               rule-based risk (NOT ASSESSED with missing data, LOW with data)
  integrity:   the analysis wrote nothing (row counts) and integrity_check ok

Qt runs offscreen (``QT_QPA_PLATFORM=offscreen``); this is NOT a native
Windows desktop run.

Usage::

    QT_QPA_PLATFORM=offscreen python tools/quality/workflow_m48.py [--json OUT]

Exit status 0 only when every step passes.
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import sqlite3
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
TABLES = ("daily_reports", "drilling_parameters", "time_logs_24h", "cost_records", "safety_reports")


def _raw(path, sql, *params):
    connection = sqlite3.connect(str(path))
    try:
        return connection.execute(sql, params).fetchall()
    finally:
        connection.close()


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--json")
    args = parser.parse_args(argv)
    work = Path(tempfile.mkdtemp(prefix="drillmaster-m48-workflow-"))
    db_path = work / "synthetic.db"
    os.environ.update({"DRILLMASTER_ENV": "test", "DRILLMASTER_DATA_DIR": str(work),
                       "DRILLMASTER_DB_PATH": str(db_path), "DRILLMASTER_LOG_DIR": str(work / "logs"),
                       "DRILLMASTER_BACKUP_DIR": str(work / "backups"), "DRILLMASTER_AI_IMPORT": "0"})
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    sys.path[:0] = [str(ROOT), str(ROOT / "tests")]
    from PySide6.QtWidgets import QApplication, QMessageBox

    app = QApplication.instance() or QApplication([])
    for name in ("warning", "information", "critical", "question"):
        setattr(QMessageBox, name, staticmethod(lambda *a, **k: QMessageBox.Ok))

    from core import analysis_evidence as ev
    from core.permissions import permissions, session_user
    from core.selection_manager import SelectionManager
    from m48_synthetic import build
    from synthetic_operations import make_file_database
    import tabs.w12_Analysis as w12

    steps = []

    def step(flow, name, ok, detail=""):
        steps.append({"flow": flow, "step": len([s for s in steps if s["flow"] == flow]) + 1, "name": name,
                      "status": "PASS" if ok else "FAIL", "detail": detail})
        print(f"{flow:11s} {steps[-1]['step']:2d} {'PASS' if ok else 'FAIL'}  {name}  {detail}")

    def save_as(path):
        w12.QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (str(path), ""))

    db = make_file_database(db_path)
    engineer = session_user(db.authenticate_user("engineer", "Synthetic-Eng-1!"))
    permissions.set_user(engineer)
    db.enforce_authorization(permissions)
    _db, ids = build(db)
    counts_before = {t: _raw(db_path, f"SELECT COUNT(*) FROM {t}")[0][0] for t in TABLES}
    started = time.time()

    sel = SelectionManager()
    sel.clear()
    tab = w12.AnalysisWidget(db_manager=db)
    tab.show()

    def card(key):
        return next(c for c in tab.kpi_cards if c.metric_key == key).metric_value

    def run(kind):
        tab.analytics_type.blockSignals(True)
        tab.analytics_type.setCurrentText(kind)
        tab.analytics_type.blockSignals(False)
        tab.run_advanced_analysis()
        return tab.results_text.toPlainText()

    def evidence(metric):
        tab.evidence_metric.setCurrentIndex(tab.evidence_metric.findData(metric))
        tab.evidence_level.setCurrentIndex(tab.evidence_level.findData(ev.SCOPE_WELL))
        tab.evidence_load_btn.click()
        return tab._evidence_rows

    # ------------------------------------------------------------- performance
    wc, wa = ids["C"], ids["A"]
    sel.select_well(wc, {"name": "SYN-C"})
    (sql_avg,) = _raw(db_path, "SELECT AVG(avg_rop) FROM drilling_parameters WHERE well_id=? "
                               "AND avg_rop IS NOT NULL AND avg_rop >= 0", wc)[0]
    step("performance", "Average ROP card equals an independent SQL average, with its unit",
         card("avg_rop").value == round(sql_avg, 2) and card("avg_rop").unit == "m/hr",
         f"card={card('avg_rop').text()} sql={sql_avg}")
    rows = evidence("rop")
    sql_rop = [r[0] for r in _raw(db_path, "SELECT avg_rop FROM drilling_parameters WHERE well_id=? "
                                           "ORDER BY report_date", wc)]
    step("performance", "evidence trend: one valid observation per date, from the canonical source",
         [r.value for r in rows] == sql_rop and all(r.state == ev.STATE_VALUE for r in rows),
         f"rows={len(rows)}")
    (target,) = _raw(db_path, "SELECT id FROM daily_reports WHERE well_id=? AND report_date='2026-02-03'", wc)[0]
    tab.select_evidence_row(2)
    enabled = tab.evidence_open_btn.isEnabled()
    tab.evidence_open_btn.click()
    step("performance", "point -> Open Source Report selects exactly that report_id",
         enabled and sel.current_report_id == target and tab.current_report_id == target,
         f"report_id={sel.current_report_id}")
    csv_path = work / "flow_performance.csv"
    save_as(csv_path)
    tab.evidence_csv_btn.click()
    lines = list(csv.reader(csv_path.open(encoding="utf-8"))) if csv_path.exists() else []
    header = next((i for i, r in enumerate(lines) if r and r[0] == "metric"), None)
    body = lines[header + 1:] if header is not None else []
    meta = {r[0]: r[1] for r in lines[:header or 0] if len(r) > 1}
    step("performance", "typed CSV: values equal SQL, metadata carries scope, unit and null policy",
         [float(r[2]) for r in body] == sql_rop and meta.get("# unit") == "m/hr"
         and "# null_policy" in meta and meta.get("# well_id") == str(wc), f"rows={len(body)}")
    sel.select_well(wa, {"name": "SYN-A"})
    a_reports = {r[0] for r in _raw(db_path, "SELECT id FROM daily_reports WHERE well_id=?", wa)}
    step("performance", "switch well: report scope cleared, evidence and KPIs belong to the new well",
         tab.current_report_id is None and {r.report_id for r in tab._evidence_rows} <= a_reports
         and card("current_depth").value == 300.0, f"depth={card('current_depth').text()}")
    # SYN-B carries a legacy DrillingParameters row (report_id NULL) on the date
    # of a report-linked row. Independent SQL encoding of the rule: linked rows,
    # plus legacy rows only on dates without any linked row.
    wb = ids["B"]
    sel.select_well(wb, {"name": "SYN-B"})
    (rule_avg,) = _raw(db_path, """
        SELECT AVG(p.avg_rop) FROM drilling_parameters p
        WHERE p.well_id=? AND p.avg_rop IS NOT NULL AND p.avg_rop >= 0
          AND (p.report_id IS NOT NULL OR NOT EXISTS (
                SELECT 1 FROM drilling_parameters q WHERE q.well_id = p.well_id
                  AND q.report_date = p.report_date AND q.report_id IS NOT NULL))""", wb)[0]
    (naive_avg,) = _raw(db_path, "SELECT AVG(avg_rop) FROM drilling_parameters WHERE well_id=? "
                                 "AND avg_rop >= 0", wb)[0]
    step("performance", "legacy row on a linked date is never double-counted (card equals rule SQL)",
         card("avg_rop").value == round(rule_avg, 2) and round(naive_avg, 2) != round(rule_avg, 2),
         f"card={card('avg_rop').text()} rule_sql={rule_avg} naive_sql={naive_avg}")

    # ------------------------------------------------------------- npt
    sel.select_well(wa, {"name": "SYN-A"})      # each flow selects its own well
    coverage_text = tab.npt_coverage_label.text()
    step("npt", "coverage label names the partial and the missing day",
         "partial 1" in coverage_text and "missing 1" in coverage_text, coverage_text)
    forecast = run("NPT Forecasting")
    step("npt", "NPT projection NOT ASSESSED on 4 coverage-valid dates (2 valid zeros counted)",
         "NOT ASSESSED" in forecast and "(2 with a valid zero" in forecast)
    rows = evidence("npt")
    states = [r.state for r in rows]
    step("npt", "evidence states keep partial / missing distinct from zero",
         states == ["value", "value", "partial", "missing", "value", "value"], ",".join(states))
    xlsx_path = work / "flow_npt.xlsx"
    save_as(xlsx_path)
    tab.evidence_xlsx_btn.click()
    from openpyxl import load_workbook
    sheet = load_workbook(xlsx_path)["evidence"] if xlsx_path.exists() else None
    got = [(r[1].value, r[2].value, r[4].value) for r in sheet.iter_rows(min_row=2)] if sheet else []
    (sql_npt,) = _raw(db_path, "SELECT SUM(t.duration) FROM time_logs_24h t JOIN daily_reports d "
                               "ON t.report_id=d.id WHERE d.well_id=? AND d.report_date='2026-01-02' "
                               "AND t.is_npt=1", wa)[0]
    step("npt", "typed XLSX: valid zero numeric, partial/missing empty, NPT equals SQL",
         len(got) == 6 and got[0][1] == 0 and got[1][1] == sql_npt and got[2][1:] == (None, "partial")
         and got[3][1:] == (None, "missing"), f"npt_day2={got[1][1] if len(got) > 1 else None} sql={sql_npt}")

    # ------------------------------------------------------------- cost / risk
    currencies = _raw(db_path, "SELECT COUNT(DISTINCT currency) FROM cost_records WHERE well_id=?", wa)[0][0]
    cost_a = run("Cost Analysis")
    step("cost-risk", "mixed currencies are never summed", currencies == 2
         and "totals NOT aggregated" in cost_a and "Actual total NOT ASSESSED" in cost_a and "$" not in cost_a)
    risk_a = run("Risk Assessment")
    step("cost-risk", "missing safety data and incomplete coverage -> risk NOT ASSESSED (never LOW)",
         "Overall level: NOT ASSESSED" in risk_a)
    sel.select_well(wc, {"name": "SYN-C"})
    (actual, planned) = _raw(db_path, "SELECT SUM(actual_cost), SUM(planned_cost) FROM cost_records "
                                      "WHERE well_id=?", wc)[0]
    cost_c = run("Cost Analysis")
    step("cost-risk", "single-currency actual and estimate equal SQL totals",
         f"USD {actual:,.2f}" in cost_c and f"USD {planned:,.2f}" in cost_c, f"actual={actual}")
    risk_c = run("Risk Assessment")
    step("cost-risk", "assessable well -> rule-based level with criteria and rules shown",
         "Overall level: LOW" in risk_c and "rule:" in risk_c and "NOT a probability" in risk_c)
    sel.select_well(None)
    step("cost-risk", "deselect clears the analysis text and every KPI",
         tab.results_text.toPlainText() == "No well selected"
         and all(c.metric_value.state == ev.STATE_NO_SELECTION for c in tab.kpi_cards))

    # ------------------------------------------------------------- integrity
    counts_after = {t: _raw(db_path, f"SELECT COUNT(*) FROM {t}")[0][0] for t in TABLES}
    step("integrity", "the analysis page wrote nothing (row counts unchanged)",
         counts_after == counts_before, json.dumps(counts_after))
    step("integrity", "PRAGMA integrity_check ok", _raw(db_path, "PRAGMA integrity_check") == [("ok",)])

    tab.cleanup()
    sel.clear()
    db.close()
    app.processEvents()
    failed = [s for s in steps if s["status"] != "PASS"]
    print(f"\nsteps={len(steps)} passed={len(steps) - len(failed)} failed={len(failed)} "
          f"seconds={time.time() - started:.1f}")
    if args.json:
        Path(args.json).write_text(json.dumps({"environment": "Qt offscreen, temporary file database",
                                               "steps": steps}, indent=2, ensure_ascii=False), encoding="utf-8")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
