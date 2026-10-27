"""Evidence layer for the Operations Analysis page (W12) — Qt-free.

One place that states, for every analysis metric, what a number means and
when it is *not* a number:

* value states (``MetricValue``): a real value (a real ``0.0`` included) is
  kept apart from missing data, partial coverage, ambiguous ownership,
  invalid input, a read error, running and NOT ASSESSED;
* validity of a stored observation (``classify_number``): ``None`` is missing,
  NaN/inf/non-numeric and negative physical quantities are invalid — each
  exclusion carries its reason;
* the ROP daily series (one observation per calendar date, ownership via
  ``report_id``/``wellbore_id``), the NPT daily coverage series (positive NPT,
  valid zero with full coverage and insufficient coverage stay distinct), a
  transparent linear projection with a rolling-origin backtest, the
  rule-based risk assessment, and the typed evidence export.

Nothing here writes to the database, and nothing here invents a value: when a
definition cannot be evaluated the result is NOT ASSESSED with a reason.
"""
from __future__ import annotations

import csv
import hashlib
import math
from dataclasses import dataclass, field, asdict
from datetime import date, datetime, timezone
from typing import Iterable, Optional

# ---------------------------------------------------------------------------
# Value states
# ---------------------------------------------------------------------------

STATE_VALUE = "value"                # a real value; 0.0 is a real value
STATE_PARTIAL = "partial"            # computed from incomplete coverage
STATE_MISSING = "missing"            # no source data
STATE_AMBIGUOUS = "ambiguous"        # ownership/linkage cannot be proven
STATE_INVALID = "invalid"            # source data present but not usable
STATE_ERROR = "error"                # reading/processing failed
STATE_RUNNING = "running"            # being computed
STATE_NOT_ASSESSED = "not_assessed"  # definition cannot be evaluated
STATE_NO_SELECTION = "no_selection"  # no valid selection

STATE_TEXT = {
    STATE_MISSING: "No data",
    STATE_AMBIGUOUS: "Ambiguous",
    STATE_INVALID: "Invalid input",
    STATE_ERROR: "Read error",
    STATE_RUNNING: "Loading…",
    STATE_NOT_ASSESSED: "NOT ASSESSED",
    STATE_NO_SELECTION: "—",
}


@dataclass(frozen=True)
class MetricValue:
    """A displayed metric: its state, value, unit and the reason behind it."""

    state: str
    value: Optional[float] = None
    unit: str = ""
    reason: str = ""
    digits: int = 1

    @property
    def has_value(self) -> bool:
        return self.state in (STATE_VALUE, STATE_PARTIAL) and self.value is not None

    def number_text(self) -> str:
        """The number alone, or the state word when there is no number."""
        if not self.has_value:
            return STATE_TEXT.get(self.state, "—")
        text = f"{self.value:.{self.digits}f}"
        return text + " (partial)" if self.state == STATE_PARTIAL else text

    def text(self) -> str:
        """The number with its unit; a state word never carries a unit."""
        if not self.has_value:
            return self.number_text()
        return f"{self.number_text()} {self.unit}".rstrip()


def metric(value, unit: str = "", digits: int = 1, *,
           missing_reason: str = "no source data",
           allow_negative: bool = True) -> MetricValue:
    """Wrap a computed scalar: ``None`` -> missing, non-finite -> invalid."""
    state, number, reason = classify_number(value, allow_negative=allow_negative)
    if state == STATE_MISSING:
        return MetricValue(STATE_MISSING, None, unit, missing_reason, digits)
    if state == STATE_INVALID:
        return MetricValue(STATE_INVALID, None, unit, reason, digits)
    return MetricValue(STATE_VALUE, number, unit, "", digits)


def classify_number(raw, *, allow_negative: bool = False):
    """Return ``(state, value, reason)`` for one stored observation.

    ``None`` is missing (never 0). Booleans, non-numeric text, NaN and
    infinities are invalid. A negative value is invalid unless the quantity
    may be negative (``allow_negative``). A real ``0.0`` is a valid value.
    """
    if raw is None:
        return STATE_MISSING, None, "not recorded"
    if isinstance(raw, bool):
        return STATE_INVALID, None, "boolean is not a measurement"
    try:
        number = float(raw)
    except (TypeError, ValueError):
        return STATE_INVALID, None, "not a number"
    if not math.isfinite(number):
        return STATE_INVALID, None, "non-finite value (NaN/inf)"
    if number < 0 and not allow_negative:
        return STATE_INVALID, None, "negative value"
    return STATE_VALUE, number, ""


# ---------------------------------------------------------------------------
# Metric definitions (tooltips, export metadata, report)
# ---------------------------------------------------------------------------

KPI_DEFINITIONS = {
    "current_depth": ("Current depth", "m",
                      "Maximum known DailyReport.depth_2400 (measured depth at 24:00) in the "
                      "scope. Reports without a depth are skipped, never read as 0 m."),
    "rig_days": ("Rig days", "days",
                 "Number of distinct DailyReport.report_date values in the scope."),
    "avg_rop": ("Average ROP", "m/hr",
                "Arithmetic mean of valid DrillingParameters.avg_rop rows in the scope "
                "(finite and >= 0). A recorded 0.0 counts; missing, NaN, infinite and "
                "negative values are excluded."),
    "best_rop": ("Best ROP", "m/hr",
                 "Largest valid DrillingParameters.avg_rop row in the scope."),
    "weighted_rop": ("Footage-weighted ROP", "m/hr",
                     "Sum of footage / sum of on-bottom hours over rows with a complete valid "
                     "pair (BitPerformanceEngine.weighted_rop)."),
    "avg_wob": ("Average WOB", "klb",
                "Mean of per-row midpoints (wob_min + wob_max) / 2; rows lacking a bound "
                "or holding an invalid bound are excluded."),
    "avg_rpm": ("Average RPM", "rpm",
                "Mean of per-row midpoints (rpm_min + rpm_max) / 2; rows lacking a bound "
                "are excluded."),
    "avg_torque": ("Average torque", "klb.ft",
                   "Mean of per-row midpoints (torque_min + torque_max) / 2; rows lacking a "
                   "bound are excluded."),
    "total_npt": ("Total NPT", "hrs",
                  "Sum of TimeLog24H.duration where is_npt is True. Unknown when an NPT row "
                  "has no valid duration or a row is not classified as NPT/productive."),
    "npt_percent": ("NPT", "% of recorded time",
                    "NPT hours / recorded hours x 100 (operational_time.summarize_time_logs)."),
    "productive_percent": ("Productive time", "% of recorded time",
                           "100 - NPT % over recorded time; unknown when NPT % is unknown."),
    "npt_daily_avg": ("Daily average NPT", "h/day",
                      "Total NPT hours / number of distinct report dates in the scope."),
    "daily_progress": ("Depth change, last day", "m",
                       "depth_2400 of the last report minus the previous report when both "
                       "belong to one wellbore on consecutive dates (canonical service)."),
}


def kpi_tooltip(key: str) -> str:
    title, unit, definition = KPI_DEFINITIONS[key]
    return f"{title} [{unit}]\n{definition}"


# ---------------------------------------------------------------------------
# ROP
# ---------------------------------------------------------------------------

def valid_rop_values(values: Iterable) -> list:
    """The one validity rule for stored ROP: finite and >= 0 (0.0 kept)."""
    out = []
    for raw in values:
        state, number, _ = classify_number(raw)
        if state == STATE_VALUE:
            out.append(number)
    return out


def summarize_rop(values: Iterable) -> dict:
    """Average and best ROP over valid rows plus the exclusion counts."""
    values = list(values)
    valid, missing, invalid = [], 0, 0
    for raw in values:
        state, number, _ = classify_number(raw)
        if state == STATE_VALUE:
            valid.append(number)
        elif state == STATE_MISSING:
            missing += 1
        else:
            invalid += 1
    return {
        "average": (sum(valid) / len(valid)) if valid else None,
        "best": max(valid) if valid else None,
        "valid_count": len(valid), "missing_count": missing, "invalid_count": invalid,
    }


def midpoint_series(pairs: Iterable) -> list:
    """Midpoints of (min, max) pairs where both bounds are valid numbers >= 0."""
    out = []
    for low, high in pairs:
        s_low, v_low, _ = classify_number(low)
        s_high, v_high, _ = classify_number(high)
        if s_low == STATE_VALUE and s_high == STATE_VALUE:
            out.append((v_low + v_high) / 2.0)
    return out


@dataclass
class Observation:
    """One evidence row: a value on a date with its provenance and state."""

    metric: str
    date: Optional[date]
    value: Optional[float]
    unit: str
    state: str
    reason: str = ""
    report_id: Optional[int] = None
    wellbore_id: Optional[int] = None
    section_id: Optional[int] = None
    source: str = ""

    @property
    def drilldown_report_id(self) -> Optional[int]:
        """A report can be opened only through an explicit, unambiguous report_id."""
        if self.report_id is None or self.state == STATE_AMBIGUOUS:
            return None
        return self.report_id

    def drilldown_reason(self) -> str:
        if self.state == STATE_AMBIGUOUS:
            return "Ambiguous ownership: several sources share this date"
        if self.report_id is None:
            return "No report_id link (legacy row); the source report cannot be proven"
        return ""


LEGACY_DUPLICATE_REASON = ("legacy row without report_id duplicates a report-linked row on "
                           "this date; the report-linked data is used")


def rop_daily_series(rows: Iterable[dict], *, source_name: str = "DrillingParameters.avg_rop"):
    """One ROP observation per calendar date, with every exclusion explained.

    ``rows`` are dicts with ``date``, ``rop``, ``report_id``, ``wellbore_id``,
    ``section_id``, ``hours`` (on-bottom, optional) and ``source_id``.

    * a missing/invalid ROP row is excluded with its reason;
    * one valid row on a date is that date's observation;
    * several valid rows on one date are combined only when they belong to the
      same wellbore and every row has valid on-bottom hours > 0: the value is
      the hours-weighted mean of the stated rates. Otherwise the date is
      AMBIGUOUS and excluded — same-date records are never counted as
      independent observations;
    * a legacy row without ``report_id`` is a fallback only when no
      report-linked row exists on its date (the rule of
      ``operations_intelligence.canonical_params_clause``); otherwise it is
      excluded as AMBIGUOUS with ``LEGACY_DUPLICATE_REASON`` and the date is
      decided by its linked rows alone.
    """
    by_date: dict = {}
    excluded = []
    rows = list(rows)
    linked_dates = {row.get("date") for row in rows if row.get("report_id") is not None}
    for row in rows:
        if row.get("report_id") is None and row.get("date") in linked_dates:
            excluded.append(Observation(
                "rop", row.get("date"), None, "m/hr", STATE_AMBIGUOUS, LEGACY_DUPLICATE_REASON,
                None, None, None, f"{source_name}#{row.get('source_id')}"))
            continue
        state, number, reason = classify_number(row.get("rop"))
        if row.get("date") is None:
            state, reason = STATE_INVALID, "no report date"
        if state != STATE_VALUE:
            excluded.append(Observation(
                "rop", row.get("date"), None, "m/hr", state, reason,
                row.get("report_id"), row.get("wellbore_id"), row.get("section_id"),
                f"{source_name}#{row.get('source_id')}"))
            continue
        by_date.setdefault(row["date"], []).append(dict(row, rop=number))

    series = []
    for day in sorted(by_date):
        group = by_date[day]
        if len(group) == 1:
            row = group[0]
            series.append(Observation(
                "rop", day, row["rop"], "m/hr", STATE_VALUE, "",
                row.get("report_id"), row.get("wellbore_id"), row.get("section_id"),
                f"{source_name}#{row.get('source_id')}"))
            continue
        bores = {row.get("wellbore_id") for row in group}
        hours = [classify_number(row.get("hours")) for row in group]
        weighable = len(bores) == 1 and all(s == STATE_VALUE and h > 0 for s, h, _ in hours)
        reports = {row.get("report_id") for row in group}
        ids = ",".join(str(row.get("source_id")) for row in group)
        if not weighable:
            reason = ("rows from several wellbores share this date" if len(bores) > 1 else
                      "several ROP rows on this date without on-bottom hours to weight them")
            excluded.append(Observation(
                "rop", day, None, "m/hr", STATE_AMBIGUOUS, reason,
                None, None, None, f"{source_name}#{ids}"))
            continue
        total_h = sum(h for _, h, _ in hours)
        value = sum(row["rop"] * h for row, (_, h, _) in zip(group, hours)) / total_h
        report_id = next(iter(reports)) if len(reports) == 1 else None
        series.append(Observation(
            "rop", day, value, "m/hr", STATE_VALUE,
            f"hours-weighted mean of {len(group)} rows",
            report_id, next(iter(bores)), group[0].get("section_id") if report_id else None,
            f"{source_name}#{ids}"))
    return series, excluded


# ---------------------------------------------------------------------------
# NPT daily coverage
# ---------------------------------------------------------------------------

FULL_DAY_HOURS = 24.0
# Same threshold DataQualityService uses for "good" 24 h coverage.
COVERAGE_SUFFICIENT_PERCENT = 95.0
# Rounding tolerance above 24 h (6 min); more than that is over-recorded time.
OVER_RECORDED_TOLERANCE_HOURS = 0.1
UNCLASSIFIED_NPT = "Unclassified (no NPT code)"


@dataclass
class NptDay:
    date: date
    state: str
    npt_hours: Optional[float]
    recorded_hours: Optional[float]
    coverage_percent: Optional[float]
    reason: str
    report_ids: list = field(default_factory=list)
    wellbore_id: Optional[int] = None
    section_id: Optional[int] = None
    categories: dict = field(default_factory=dict)

    def observation(self) -> Observation:
        report_id = self.report_ids[0] if len(self.report_ids) == 1 else None
        return Observation(
            "npt", self.date, self.npt_hours if self.state == STATE_VALUE else None,
            "hrs", self.state, self.reason, report_id, self.wellbore_id, self.section_id,
            "TimeLog24H via DailyReport#" + ",".join(str(i) for i in self.report_ids))


def npt_daily_series(reports: Iterable[dict], logs: Iterable[dict]):
    """Per report date: NPT hours with an explicit coverage state.

    ``reports``: dicts ``report_id``, ``date``, ``wellbore_id``, ``section_id``.
    ``logs``: dicts ``report_id``, ``duration``, ``is_npt``, ``code``.

    * several reports on one date in the scope -> AMBIGUOUS (their time cannot
      be combined into one 24 h day without proof);
    * no time-log rows -> MISSING;
    * a row without a valid duration (missing, NaN, negative) or without an
      NPT/productive classification -> INVALID (negative durations are never
      summed);
    * more than 24 h (+ rounding tolerance) -> INVALID (over-recorded);
    * coverage below 95 % of 24 h -> PARTIAL (not a valid zero);
    * otherwise VALUE: NPT hours, where 0.0 is a valid zero with full coverage.
    """
    from core.operational_time import summarize_time_logs

    logs_by_report: dict = {}
    for log in logs:
        logs_by_report.setdefault(log.get("report_id"), []).append(log)
    by_date: dict = {}
    for rep in reports:
        by_date.setdefault(rep["date"], []).append(rep)

    days = []
    for day in sorted(d for d in by_date if d is not None):
        reps = by_date[day]
        ids = [r["report_id"] for r in reps]
        bore = reps[0].get("wellbore_id") if len(reps) == 1 else None
        section = reps[0].get("section_id") if len(reps) == 1 else None
        if len(reps) > 1:
            days.append(NptDay(day, STATE_AMBIGUOUS, None, None, None,
                               f"{len(reps)} reports share this date in the scope",
                               ids, bore, section))
            continue
        rows = logs_by_report.get(ids[0], [])
        if not rows:
            days.append(NptDay(day, STATE_MISSING, None, None, None,
                               "no time-log rows", ids, bore, section))
            continue

        class _Row:
            def __init__(self, d):
                self.duration, self.is_npt = d.get("duration"), d.get("is_npt")

        summary = summarize_time_logs([_Row(r) for r in rows])
        categories: dict = {}
        for r in rows:
            state, hours, _ = classify_number(r.get("duration"))
            if r.get("is_npt") is True and state == STATE_VALUE:
                code = (r.get("code") or "").strip() or UNCLASSIFIED_NPT
                categories[code] = categories.get(code, 0.0) + hours
        if summary["unknown_duration_count"]:
            days.append(NptDay(day, STATE_INVALID, None, None, None,
                               f"{summary['unknown_duration_count']} row(s) without a valid "
                               "duration (missing, non-finite or negative)", ids, bore, section))
            continue
        if summary["npt_hours"] is None:
            days.append(NptDay(day, STATE_INVALID, None, summary["total_hours"], None,
                               "row(s) not classified as NPT or productive", ids, bore, section))
            continue
        recorded = summary["total_hours"]
        coverage = recorded / FULL_DAY_HOURS * 100.0
        if recorded > FULL_DAY_HOURS + OVER_RECORDED_TOLERANCE_HOURS:
            days.append(NptDay(day, STATE_INVALID, None, recorded, coverage,
                               f"{recorded:.2f} h recorded exceeds 24 h (overlap suspected)",
                               ids, bore, section, categories))
            continue
        if coverage < COVERAGE_SUFFICIENT_PERCENT:
            days.append(NptDay(day, STATE_PARTIAL, summary["npt_hours"], recorded, coverage,
                               f"coverage {coverage:.1f}% < {COVERAGE_SUFFICIENT_PERCENT:.0f}% "
                               "of 24 h", ids, bore, section, categories))
            continue
        days.append(NptDay(day, STATE_VALUE, summary["npt_hours"], recorded, coverage,
                           "" if summary["npt_hours"] else "valid zero: full coverage, no NPT",
                           ids, bore, section, categories))
    return days


def npt_coverage_summary(days: list) -> dict:
    counts: dict = {}
    for d in days:
        counts[d.state] = counts.get(d.state, 0) + 1
    covered = counts.get(STATE_VALUE, 0)
    return {"days": len(days), "covered_days": covered, "by_state": counts,
            "complete": bool(days) and covered == len(days)}


# ---------------------------------------------------------------------------
# Projection (linear trend, not a validated model)
# ---------------------------------------------------------------------------

MIN_PROJECTION_POINTS = 5
BACKTEST_MIN_TRAIN = 3
PROJECTION_MODEL = "ordinary least-squares straight line of value against calendar day"
PROJECTION_ASSUMPTIONS = (
    "the recent linear trend continues unchanged over the horizon",
    "one observation per calendar date; excluded dates are not zero",
    "values are clipped at 0 because the quantity cannot be negative",
)
PROJECTION_LIMITS = (
    "a projection, not a validated forecast; no confidence interval is claimed",
    "no causal explanation is assessed (bit wear, formation, crew, equipment)",
    "sensitive to outliers and to changes of section, bit or wellbore",
)


def least_squares(xs: list, ys: list):
    """Closed-form OLS slope and intercept; ``None`` for a degenerate x set."""
    n = len(xs)
    mean_x, mean_y = sum(xs) / n, sum(ys) / n
    sxx = sum((x - mean_x) ** 2 for x in xs)
    if sxx == 0:
        return None
    slope = sum((x - mean_x) * (y - mean_y) for x, y in zip(xs, ys)) / sxx
    return slope, mean_y - slope * mean_x


def linear_projection(points: list, horizon: int, *, minimum_points: int = MIN_PROJECTION_POINTS):
    """Project a dated series forward with full data-sufficiency disclosure.

    ``points`` are ``(date, value)`` pairs with unique dates. Returns a dict
    with ``status`` ``"assessed"`` or ``"not_assessed"`` and every quantity
    a reader needs to judge it: count, distinct dates, span, model,
    assumptions, limits, the fitted line, the clipped projection and a
    rolling-origin one-step backtest against the training-mean baseline.
    """
    points = sorted(points, key=lambda p: p[0])
    dates = [p[0] for p in points]
    result = {
        "status": "not_assessed", "reason": "", "n": len(points),
        "distinct_dates": len(set(dates)), "minimum_points": minimum_points,
        "span_days": (dates[-1] - dates[0]).days if points else 0,
        "first_date": dates[0] if points else None, "last_date": dates[-1] if points else None,
        "model": PROJECTION_MODEL, "assumptions": list(PROJECTION_ASSUMPTIONS),
        "limits": list(PROJECTION_LIMITS), "slope": None, "intercept": None,
        "fitted": [], "projection": [], "backtest": None,
    }
    if len(set(dates)) != len(dates):
        result["reason"] = "dates are not unique; aggregate to one value per date first"
        return result
    if len(points) < minimum_points:
        result["reason"] = (f"{len(points)} valid dated observation(s); at least "
                            f"{minimum_points} distinct dates are required")
        return result
    xs = [(d - dates[0]).days + 1 for d in dates]
    ys = [float(v) for _, v in points]
    fit = least_squares(xs, ys)
    if fit is None:
        result["reason"] = "all observations on one day; no trend can be fitted"
        return result
    slope, intercept = fit
    result.update(status="assessed", slope=slope, intercept=intercept)
    result["x"] = xs
    result["fitted"] = [slope * x + intercept for x in xs]
    for step in range(1, horizon + 1):
        x = xs[-1] + step
        raw = slope * x + intercept
        result["projection"].append({"day": x, "value": max(0.0, raw), "clipped": raw < 0})
    result["backtest"] = rolling_origin_backtest(xs, ys)
    return result


def rolling_origin_backtest(xs: list, ys: list, *, minimum_train: int = BACKTEST_MIN_TRAIN):
    """One-step-ahead errors of the line vs the training-mean baseline."""
    model_err, base_err = [], []
    for k in range(minimum_train, len(xs)):
        fit = least_squares(xs[:k], ys[:k])
        if fit is None:
            continue
        predicted = max(0.0, fit[0] * xs[k] + fit[1])
        baseline = sum(ys[:k]) / k
        model_err.append(abs(predicted - ys[k]))
        base_err.append(abs(baseline - ys[k]))
    if not model_err:
        return None
    mae_model = sum(model_err) / len(model_err)
    mae_base = sum(base_err) / len(base_err)
    return {"origins": len(model_err), "mae_model": mae_model, "mae_baseline": mae_base,
            "baseline": "mean of the training values", "model_better": mae_model < mae_base}


def describe_projection(result: dict, unit: str, title: str) -> str:
    lines = [title, "=" * 40, "Projection (linear trend) - NOT a validated forecast."]
    lines.append(f"Valid dated observations: {result['n']} (distinct dates "
                 f"{result['distinct_dates']}, minimum {result['minimum_points']})")
    if result.get("first_date"):
        lines.append(f"Period: {result['first_date']} to {result['last_date']} "
                     f"({result['span_days']} days)")
    if result["status"] != "assessed":
        lines.append(f"Result: NOT ASSESSED - {result['reason']}")
        return "\n".join(lines)
    lines.append(f"Model: {result['model']}")
    lines.append(f"Trend slope: {result['slope']:.3f} {unit} per day")
    for item in result["projection"]:
        clip = " (clipped at 0)" if item["clipped"] else ""
        lines.append(f"  Day {item['day']}: {item['value']:.2f} {unit}{clip}")
    bt = result.get("backtest")
    if bt:
        lines.append(f"Rolling-origin backtest ({bt['origins']} one-step origins): "
                     f"MAE line {bt['mae_model']:.2f} vs baseline {bt['mae_baseline']:.2f} {unit} "
                     f"({bt['baseline']}); line {'beats' if bt['model_better'] else 'does NOT beat'} "
                     "the baseline")
    lines.append("Assumptions: " + "; ".join(result["assumptions"]))
    lines.append("Limits: " + "; ".join(result["limits"]))
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Rule-based risk (not a probability)
# ---------------------------------------------------------------------------

RISK_RULES = {
    "NPT Risk": "NPT % of recorded time: > 30 -> 9; > 20 -> 7; > 10 -> 5; otherwise 3",
    "Well Control": "PROXY derived from NPT % (> 15 -> 6, otherwise 3); no well-control data "
                    "is read, so it never counts toward the overall level",
    "Safety": "Days without LTI: > 90 -> 2; > 30 -> 5; otherwise 8",
    "Equipment": "no data source in this assessment",
    "Weather": "no data source in this assessment",
}
RISK_LEVEL_RULE = "score >= 7 HIGH; score >= 4 MEDIUM; otherwise LOW"
REQUIRED_RISK_CRITERIA = ("NPT Risk", "Safety")


def risk_scores(npt_pct, days_without_lti, has_safety_report):
    """Per-category rule scores (1-10) or ``None`` when the input is unknown."""
    scores = {}
    if npt_pct is None:
        scores["NPT Risk"] = None
    elif npt_pct > 30:
        scores["NPT Risk"] = 9
    elif npt_pct > 20:
        scores["NPT Risk"] = 7
    elif npt_pct > 10:
        scores["NPT Risk"] = 5
    else:
        scores["NPT Risk"] = 3
    scores["Equipment"] = None
    scores["Weather"] = None
    scores["Well Control"] = None if npt_pct is None else (6 if npt_pct > 15 else 3)
    if not has_safety_report or days_without_lti is None:
        scores["Safety"] = None
    else:
        d = days_without_lti
        scores["Safety"] = 2 if d > 90 else 5 if d > 30 else 8
    return scores


def risk_level(score):
    if score is None:
        return None
    return "HIGH" if score >= 7 else "MEDIUM" if score >= 4 else "LOW"


def risk_assessment(npt_pct, npt_coverage: dict, days_without_lti):
    """Criteria, their evidence state and the overall rule-based level.

    The NPT criterion is assessed only when every report date in the scope has
    full time coverage (``npt_coverage['complete']``). The overall level is the
    WORST level among the required criteria (NPT, Safety) and is reported only
    when all of them are assessed — missing data can never yield LOW.
    """
    npt_ok = npt_pct is not None and bool(npt_coverage.get("complete"))
    scores = risk_scores(npt_pct if npt_ok else None, days_without_lti,
                         days_without_lti is not None)
    criteria = []
    for name, score in scores.items():
        if name == "NPT Risk" and not npt_ok:
            reason = ("NPT % unknown" if npt_pct is None else
                      f"time coverage incomplete: {npt_coverage.get('covered_days', 0)} of "
                      f"{npt_coverage.get('days', 0)} report dates fully covered")
        elif name == "Safety" and score is None:
            reason = "no unambiguous days-without-LTI observation"
        elif score is None:
            reason = RISK_RULES[name]
        else:
            reason = ""
        criteria.append({
            "name": name, "score": score, "level": risk_level(score),
            "assessed": score is not None, "required": name in REQUIRED_RISK_CRITERIA,
            "proxy": name == "Well Control", "rule": RISK_RULES[name], "reason": reason,
        })
    required = [c for c in criteria if c["required"]]
    missing = [c["name"] for c in required if not c["assessed"]]
    if missing:
        overall = {"assessed": False, "level": None, "score": None,
                   "reason": "required criteria not assessed: " + ", ".join(missing)}
    else:
        worst = max(c["score"] for c in required)
        overall = {"assessed": True, "level": risk_level(worst), "score": worst,
                   "reason": "worst level among required criteria (" +
                             ", ".join(REQUIRED_RISK_CRITERIA) + ")"}
    return {"criteria": criteria, "overall": overall, "level_rule": RISK_LEVEL_RULE}


# ---------------------------------------------------------------------------
# Typed export of evidence rows
# ---------------------------------------------------------------------------

EXPORT_COLUMNS = ("metric", "date", "value", "unit", "state", "reason",
                  "report_id", "wellbore_id", "section_id", "source")
NULL_POLICY = ("An empty value cell means the value is not available; the state and reason "
               "columns say why. A recorded zero is written as 0.")
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _safe_text(value):
    """Neutralise spreadsheet formula prefixes in free text (CSV/XLSX)."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def export_metadata(scope_label: str, metric_key: str, unit: str, definition: str,
                    date_range: str, extra: Optional[dict] = None) -> dict:
    meta = {
        "generated_utc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "scope": scope_label, "metric": metric_key, "unit": unit,
        "definition": definition, "date_range": date_range, "null_policy": NULL_POLICY,
        "source": "DrillMaster operational database (read-only query)",
    }
    meta.update(extra or {})
    return meta


def _row_values(row: Observation) -> list:
    data = asdict(row)
    out = []
    for column in EXPORT_COLUMNS:
        value = data[column]
        if isinstance(value, date):
            value = value.isoformat()
        out.append(_safe_text(value))
    return out


def export_csv(path: str, rows: Iterable[Observation], metadata: dict) -> int:
    rows = list(rows)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        for key, value in metadata.items():
            writer.writerow([f"# {key}", _safe_text(str(value))])
        writer.writerow(EXPORT_COLUMNS)
        for row in rows:
            writer.writerow(["" if v is None else v for v in _row_values(row)])
    return len(rows)


def export_xlsx(path: str, rows: Iterable[Observation], metadata: dict) -> int:
    from openpyxl import Workbook

    rows = list(rows)
    book = Workbook()
    sheet = book.active
    sheet.title = "evidence"
    sheet.append(list(EXPORT_COLUMNS))
    for row in rows:
        sheet.append(_row_values(row))  # None stays an empty cell; numbers stay numbers
    meta = book.create_sheet("metadata")
    for key, value in metadata.items():
        meta.append([key, _safe_text(str(value))])
    book.save(path)
    return len(rows)


def filter_range(rows: Iterable[Observation], start: Optional[date], end: Optional[date]):
    out = []
    for row in rows:
        if row.date is None:
            continue
        if start is not None and row.date < start:
            continue
        if end is not None and row.date > end:
            continue
        out.append(row)
    return out


# ---------------------------------------------------------------------------
# Scope and read-only loaders
# ---------------------------------------------------------------------------

SCOPE_WELL = "well"        # whole well, or the selected wellbore when one is set
SCOPE_SECTION = "section"  # the selected section (within the wellbore when set)
SCOPE_REPORT = "report"    # the selected daily report


@dataclass(frozen=True)
class AnalysisScope:
    """Every dimension that changes an analysis result (also the cache key)."""

    well_id: int
    wellbore_id: Optional[int] = None
    section_id: Optional[int] = None
    report_id: Optional[int] = None
    level: str = SCOPE_WELL

    def problem(self) -> str:
        """Why this scope cannot be evaluated, or an empty string."""
        if not self.well_id:
            return "no well selected"
        if self.level == SCOPE_SECTION and not self.section_id:
            return "no section selected"
        if self.level == SCOPE_REPORT and not self.report_id:
            return "no daily report selected"
        return ""

    def key(self) -> str:
        return (f"well={self.well_id}|bore={self.wellbore_id}|section={self.section_id}"
                f"|report={self.report_id}|level={self.level}")

    def label(self) -> str:
        base = f"Well #{self.well_id}" + (f" / Wellbore #{self.wellbore_id}"
                                          if self.wellbore_id else " (whole well)")
        if self.level == SCOPE_SECTION:
            return f"{base} / Section #{self.section_id}"
        if self.level == SCOPE_REPORT:
            return f"{base} / Report #{self.report_id}"
        return base


def scoped_reports_query(session, scope: AnalysisScope):
    from core.database import DailyReport

    q = session.query(DailyReport).filter(DailyReport.well_id == scope.well_id)
    if scope.wellbore_id:
        q = q.filter(DailyReport.wellbore_id == scope.wellbore_id)
    if scope.level == SCOPE_SECTION:
        q = q.filter(DailyReport.section_id == scope.section_id)
    elif scope.level == SCOPE_REPORT:
        q = q.filter(DailyReport.id == scope.report_id)
    return q


def load_rop_rows(session, scope: AnalysisScope) -> list:
    """DrillingParameters rows of the scope as plain dicts (no ORM leaves).

    A wellbore/section/report scope resolves ownership through
    ``report_id -> DailyReport`` only; legacy rows without ``report_id`` are
    included solely in the whole-well view, where they cannot be attributed to
    the wrong bore, and they never offer a drill-down.
    """
    from core.database import DailyReport, DrillingParameters

    narrowed = bool(scope.wellbore_id) or scope.level != SCOPE_WELL
    q = (session.query(DrillingParameters, DailyReport)
         .outerjoin(DailyReport, DrillingParameters.report_id == DailyReport.id)
         .filter(DrillingParameters.well_id == scope.well_id))
    if narrowed:
        q = q.filter(DailyReport.id.in_(
            scoped_reports_query(session, scope).with_entities(DailyReport.id)))
    rows = []
    for p, rep in q.order_by(DrillingParameters.report_date, DrillingParameters.id).all():
        rows.append({
            "source_id": p.id, "date": p.report_date, "rop": p.avg_rop,
            "hours": p.hours_on_bottom, "report_id": p.report_id,
            "wellbore_id": rep.wellbore_id if rep is not None else None,
            "section_id": rep.section_id if rep is not None else None,
        })
    return rows


def load_npt_inputs(session, scope: AnalysisScope):
    """Reports and time-log rows of the scope as plain dicts."""
    from core.database import DailyReport, TimeLog24H

    reports = [{"report_id": r.id, "date": r.report_date, "wellbore_id": r.wellbore_id,
                "section_id": r.section_id}
               for r in scoped_reports_query(session, scope).order_by(DailyReport.report_date).all()]
    ids = [r["report_id"] for r in reports]
    logs = []
    if ids:
        for log in session.query(TimeLog24H).filter(TimeLog24H.report_id.in_(ids)).all():
            logs.append({"report_id": log.report_id, "duration": log.duration,
                         "is_npt": log.is_npt, "code": log.main_code})
    return reports, logs


def load_depth_observations(session, scope: AnalysisScope) -> list:
    """Reported depth at 24:00 per report, with its state."""
    from core.database import DailyReport

    out = []
    for r in scoped_reports_query(session, scope).order_by(DailyReport.report_date, DailyReport.id).all():
        state, value, reason = classify_number(r.depth_2400)
        out.append(Observation("depth", r.report_date, value, "m", state, reason,
                               r.id, r.wellbore_id, r.section_id, f"DailyReport.depth_2400#{r.id}"))
    by_date: dict = {}
    for row in out:
        by_date.setdefault(row.date, []).append(row)
    for rows in by_date.values():
        if len(rows) > 1:
            for row in rows:
                if row.state == STATE_VALUE:
                    row.state, row.reason = STATE_AMBIGUOUS, f"{len(rows)} reports share this date"
    return out


def evidence_for_metric(session, scope: AnalysisScope, metric_key: str) -> list:
    """Every evidence row (valid and excluded) for one metric, date-sorted."""
    if metric_key == "rop":
        series, excluded = rop_daily_series(load_rop_rows(session, scope))
        rows = series + excluded
    elif metric_key == "npt":
        rows = [d.observation() for d in npt_daily_series(*load_npt_inputs(session, scope))]
    elif metric_key == "depth":
        rows = load_depth_observations(session, scope)
    else:
        raise ValueError(f"unknown evidence metric: {metric_key!r}")
    return sorted(rows, key=lambda r: (r.date or date.min, r.source))


EVIDENCE_METRICS = {
    "rop": ("ROP", "m/hr", KPI_DEFINITIONS["avg_rop"][2] +
            " One observation per date; same-date rows are hours-weighted or ambiguous."),
    "npt": ("NPT per report day", "hrs",
            "NPT hours per report date with time coverage; 0 is a valid zero only with "
            ">= 95% of 24 h recorded."),
    "depth": ("Depth at 24:00", "m", "DailyReport.depth_2400 per report."),
}


# Bumped whenever the digest input changes, so an old revision never equals a new one.
DATA_REVISION_VERSION = 2


def _rows_digest(query) -> str:
    """Deterministic digest of every column of every row of ``query``.

    The rows are read in primary-key order and each value is serialised with
    ``repr`` (exact for floats, dates, booleans and None).  Any change to any
    stored value - including a text-only change such as ``main_code`` or
    ``currency`` that leaves counts, ids, sums and timestamps unchanged -
    changes the digest, whatever path wrote it (ORM, bulk update, raw SQL).
    """
    digest = hashlib.blake2b(digest_size=16)
    count = 0
    for row in query:
        digest.update(repr(tuple(row)).encode("utf-8", "surrogatepass"))
        digest.update(b"\x1e")
        count += 1
    return f"{count}:{digest.hexdigest()}"


def data_revision(session, well_id: int) -> tuple:
    """Content fingerprint of every input of the well's analysis (cache key part).

    M49 replaces the M48 summary (row counts, max ids, ``updated_at``, value
    sums), which missed edits that change none of those - e.g. only an NPT
    code (``TimeLog24H`` has no ``updated_at``) or only a currency code.
    Every table the Analysis page reads is digested in full for the well:
    the well and its wellbores and sections, daily reports, drilling
    parameters, mud reports, 24-h time logs, costs and safety reports.
    """
    from core.database import (CostRecord, DailyReport, DrillingParameters, MudReport, SafetyReport, Section,
                               TimeLog24H, Well, Wellbore)

    def columns(model):
        return [model.__table__.c[name] for name in model.__table__.c.keys()]

    def own(model):
        return session.query(*columns(model)).filter(model.well_id == well_id).order_by(model.id)

    parts = [str(DATA_REVISION_VERSION),
             _rows_digest(session.query(*columns(Well)).filter(Well.id == well_id))]
    parts += [_rows_digest(own(model)) for model in
              (Wellbore, Section, DailyReport, DrillingParameters, MudReport, CostRecord, SafetyReport)]
    parts.append(_rows_digest(
        session.query(*columns(TimeLog24H)).join(DailyReport, TimeLog24H.report_id == DailyReport.id)
        .filter(DailyReport.well_id == well_id).order_by(TimeLog24H.id)))
    return tuple(parts)


def _cell(value, *, for_csv: bool):
    """A typed export cell: numbers stay numbers, None stays empty, text is neutralised."""
    from datetime import time as _time
    if value is None:
        return "" if for_csv else None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return value if math.isfinite(value) else ""
    if isinstance(value, _time):
        return value.strftime("%H:%M")
    if isinstance(value, date):
        return value.isoformat() if for_csv else value
    return _safe_text(str(value))


def export_tables_xlsx(path: str, tables: dict, metadata: dict) -> int:
    """``tables``: name -> (columns, rows). One sheet each plus a metadata sheet."""
    from openpyxl import Workbook

    book = Workbook()
    book.remove(book.active)
    total = 0
    for name, (columns, rows) in tables.items():
        sheet = book.create_sheet(name[:31])
        sheet.append(list(columns))
        for row in rows:
            sheet.append([_cell(v, for_csv=False) for v in row])
            total += 1
    meta = book.create_sheet("metadata")
    for key, value in metadata.items():
        meta.append([key, _safe_text(str(value))])
    book.save(path)
    return total


def export_tables_csv(base_path: str, tables: dict, metadata: dict) -> list:
    """One CSV per table (``<base>_<name>.csv``), each starting with metadata lines."""
    base = base_path[:-4] if base_path.lower().endswith(".csv") else base_path
    written = []
    for name, (columns, rows) in tables.items():
        path = f"{base}_{name}.csv"
        with open(path, "w", newline="", encoding="utf-8") as fh:
            writer = csv.writer(fh)
            for key, value in dict(metadata, table=name).items():
                writer.writerow([f"# {key}", _safe_text(str(value))])
            writer.writerow(list(columns))
            for row in rows:
                writer.writerow([_cell(v, for_csv=True) for v in row])
        written.append(path)
    return written
