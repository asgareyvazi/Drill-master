"""Centre-to-centre trajectory comparison and clearance SCREENING (M52, Part 4 / group A).

What it does
------------
* Positions between survey stations follow the minimum-curvature circular arc
  exactly (closed form below), not a straight chord and not the nearest station.
  Station positions themselves come from the canonical engine through
  :func:`core.well_trajectory.compose_well_trajectory`; this module adds no new
  station formula, only the interpolation along the same arc.
* Closest approach is continuous: both paths are sampled (all stations plus a
  regular step) and every candidate is refined by golden-section search on the
  exact arc, so the minimum is not limited to station pairs or a common MD.
* Separation profile: for each reference MD sample, the 3D distance to the
  closest point anywhere on the offset path (and that offset MD).
* Plan vs actual: for each actual station, the closest point on the plan and
  the miss decomposed in the plan's normal plane into high/low (along the
  high-side vector) and right/left (along tangent x high side).

Contract (see ``CONTRACT``): metres, degrees; local north/east relative to each
well's own wellhead; vertical = TVD below that well's rotary table.  Comparing
two wells needs both wellheads (``northing``/``easting``) in ONE projected CRS
in metres - the database has no CRS field, so the caller must confirm it - and
both rotary-table elevations (``rte_msl``) to put both wells on depth below MSL.
Any missing item makes the result NOT ASSESSED; nothing is guessed or zeroed.

NOT included and NOT claimed: survey error models or ISCWSA uncertainty,
separation factors, grid convergence, magnetic corrections, casing/hole sizes.
Every result is labelled ``LABEL``.  It is a screening aid, never an approval.

Interpolation formula (circular arc, unit tangents t1, t2, dogleg b, length L,
fraction u of L):  t(u) = [sin((1-u)b) t1 + sin(ub) t2] / sin b, integrated:
p(u) = p1 + L / (b sin b) * {[cos((1-u)b) - cos b] t1 + [1 - cos(ub)] t2}.
At u = 1 this is the minimum-curvature step L/2 * RF * (t1 + t2), RF = 2/b tan(b/2);
for b -> 0 it is the straight line p1 + uL t1.
"""
from __future__ import annotations

import csv
import io
import math
from dataclasses import dataclass, field
from typing import Iterable, Optional, Sequence

from core.well_trajectory import compose_well_trajectory

METHOD_VERSION = "M52 centre-to-centre screening v1"
LABEL = "SCREENING - centre-to-centre, no ISCWSA uncertainty, NOT APPROVED"
ASSESSED = "ASSESSED"
NOT_ASSESSED = "NOT ASSESSED"
DEFAULT_STEP_M = 10.0
NEAR_VERTICAL_DEG = 0.5
_GOLDEN = (math.sqrt(5.0) - 1.0) / 2.0
FT_TO_M = 0.3048

CONTRACT = {
    "method": METHOD_VERSION,
    "label": LABEL,
    "units": {"length": "m", "angle": "deg"},
    "frame": "north/east relative to the reference wellhead; down = depth below MSL for well-to-well, "
             "TVD below rotary table for plan vs actual of one well",
    "crs": "both wellheads in one projected CRS in metres, confirmed by the user (not stored in the database)",
    "signs": "north +, east +, down +; high side + (up), right + (clockwise of travel seen from above)",
    "station_positions": "canonical TrajectoryCalculator (minimum curvature) via compose_well_trajectory",
    "interpolation": "exact minimum-curvature circular arc",
    "closest_approach": "sampled then golden-section refined on both paths",
    "uncertainty": "none - centre-to-centre only; no ISCWSA, no separation factor",
    "validity": "stations with MD, inclination 0-180 and azimuth; paths compared only inside their MD ranges",
    "precision": "refinement tolerance 1e-6 m in MD; results reported to 0.01 m",
}


def _tangent(inc_deg: float, azi_deg: float):
    inc, azi = math.radians(inc_deg), math.radians(azi_deg)
    return (math.sin(inc) * math.cos(azi), math.sin(inc) * math.sin(azi), math.cos(inc))


def _high_side(inc_deg: float, azi_deg: float):
    inc, azi = math.radians(inc_deg), math.radians(azi_deg)
    return (math.cos(inc) * math.cos(azi), math.cos(inc) * math.sin(azi), -math.sin(inc))


def _cross(a, b):
    return (a[1] * b[2] - a[2] * b[1], a[2] * b[0] - a[0] * b[2], a[0] * b[1] - a[1] * b[0])


def _dot(a, b):
    return a[0] * b[0] + a[1] * b[1] + a[2] * b[2]


def _sub(a, b):
    return (a[0] - b[0], a[1] - b[1], a[2] - b[2])


def _norm(a):
    return math.sqrt(_dot(a, a))


@dataclass(frozen=True)
class Station:
    md: float
    inc: float
    azi: float
    north: float
    east: float
    down: float


@dataclass(frozen=True)
class WellPath:
    """Stations in one frame; positions between them follow the minimum-curvature arc."""

    name: str
    stations: tuple

    @property
    def md_range(self):
        return self.stations[0].md, self.stations[-1].md

    def _segment(self, md: float) -> int:
        lo, hi = 0, len(self.stations) - 1
        while hi - lo > 1:
            mid = (lo + hi) // 2
            if self.stations[mid].md <= md:
                lo = mid
            else:
                hi = mid
        return lo

    def state(self, md: float):
        """(position, unit tangent) at ``md``; None outside the surveyed MD range."""
        first, last = self.md_range
        if not (first - 1e-9 <= md <= last + 1e-9):
            return None
        md = min(max(md, first), last)
        if len(self.stations) == 1:
            s = self.stations[0]
            return (s.north, s.east, s.down), _tangent(s.inc, s.azi)
        i = self._segment(md)
        a, b = self.stations[i], self.stations[min(i + 1, len(self.stations) - 1)]
        t1, t2 = _tangent(a.inc, a.azi), _tangent(b.inc, b.azi)
        length = b.md - a.md
        p1 = (a.north, a.east, a.down)
        if length <= 0:
            return p1, t1
        u = (md - a.md) / length
        beta = math.acos(max(-1.0, min(1.0, _dot(t1, t2))))
        if math.degrees(beta) >= ANTIPARALLEL_DEG:
            raise ValueError(f"Dogleg of ~180 deg between MD {a.md} and {b.md}: minimum-curvature arc undefined")
        if beta < 1e-9:
            s = u * length
            return (p1[0] + s * t1[0], p1[1] + s * t1[1], p1[2] + s * t1[2]), t1
        sb = math.sin(beta)
        k = length / (beta * sb)
        c1 = math.cos((1.0 - u) * beta) - math.cos(beta)
        c2 = 1.0 - math.cos(u * beta)
        position = tuple(p1[j] + k * (c1 * t1[j] + c2 * t2[j]) for j in range(3))
        w1, w2 = math.sin((1.0 - u) * beta) / sb, math.sin(u * beta) / sb
        tangent = tuple(w1 * t1[j] + w2 * t2[j] for j in range(3))
        n = _norm(tangent)
        return position, tuple(c / n for c in tangent)

    def position(self, md: float):
        state = self.state(md)
        return None if state is None else state[0]

    def angles(self, md: float):
        """Inclination and azimuth (deg) of the arc tangent at ``md``."""
        state = self.state(md)
        if state is None:
            return None
        t = state[1]
        inc = math.degrees(math.acos(max(-1.0, min(1.0, t[2]))))
        azi = math.degrees(math.atan2(t[1], t[0])) % 360.0
        return inc, azi

    def sample_mds(self, step: float) -> list:
        first, last = self.md_range
        values = {s.md for s in self.stations}
        if step > 0 and last > first:
            n = int(math.floor((last - first) / step))
            values.update(first + k * step for k in range(n + 1))
        return sorted(values)


ANTIPARALLEL_DEG = 179.9


def degenerate_segments(path: WellPath) -> list:
    """MDs of segments whose tangents are (nearly) opposite: the arc is undefined there."""
    bad = []
    for a, b in zip(path.stations, path.stations[1:]):
        cos = max(-1.0, min(1.0, _dot(_tangent(a.inc, a.azi), _tangent(b.inc, b.azi))))
        if math.degrees(math.acos(cos)) >= ANTIPARALLEL_DEG:
            bad.append((a.md, b.md))
    return bad


def path_from_trajectory(trajectory, *, name: str = "", north0: float = 0.0, east0: float = 0.0,
                         down0: float = 0.0) -> Optional[WellPath]:
    """WellPath from a composed :class:`core.well_trajectory.WellTrajectory` shifted by an origin."""
    if trajectory is None or not trajectory.drawable:
        return None
    stations = tuple(Station(s.md, s.inc, s.azi, s.north + north0, s.east + east0, s.tvd + down0)
                     for s in trajectory.stations)
    return WellPath(name=name, stations=stations)


def _golden_min(f, lo: float, hi: float, tol: float = 1e-6):
    if hi - lo <= tol:
        x = (lo + hi) / 2.0
        return x, f(x)
    c, d = hi - _GOLDEN * (hi - lo), lo + _GOLDEN * (hi - lo)
    fc, fd = f(c), f(d)
    while hi - lo > tol:
        if fc <= fd:
            hi, d, fd = d, c, fc
            c = hi - _GOLDEN * (hi - lo)
            fc = f(c)
        else:
            lo, c, fc = c, d, fd
            d = lo + _GOLDEN * (hi - lo)
            fd = f(d)
    x = (lo + hi) / 2.0
    return x, f(x)


def closest_point(path: WellPath, point, *, step: float = DEFAULT_STEP_M):
    """(md, distance) of the point of ``path`` closest to ``point`` (continuous)."""
    mds = path.sample_mds(step)
    distances = [_norm(_sub(path.position(md), point)) for md in mds]
    best = min(range(len(mds)), key=distances.__getitem__)
    candidates = {best}
    for i in range(1, len(mds) - 1):                     # every local minimum, not only the global sample
        if distances[i] <= distances[i - 1] and distances[i] <= distances[i + 1]:
            candidates.add(i)
    result = (mds[best], distances[best])
    for i in candidates:
        lo, hi = mds[max(i - 1, 0)], mds[min(i + 1, len(mds) - 1)]
        md, d = _golden_min(lambda m: _norm(_sub(path.position(m), point)), lo, hi)
        if d < result[1]:
            result = (md, d)
    return result


@dataclass
class Comparison:
    status: str
    reasons: list = field(default_factory=list)
    label: str = LABEL
    method: str = METHOD_VERSION
    rows: list = field(default_factory=list)
    closest: Optional[dict] = None
    kind: str = ""
    warnings: list = field(default_factory=list)

    def to_json(self) -> dict:
        return {"status": self.status, "reasons": list(self.reasons), "label": self.label,
                "method": self.method, "kind": self.kind, "closest": self.closest, "rows": list(self.rows),
                "warnings": list(self.warnings), "contract": CONTRACT}


def separation_profile(reference: WellPath, offset: WellPath, *, step: float = DEFAULT_STEP_M) -> Comparison:
    """Distance from every reference MD sample to the closest point of the offset path."""
    rows = []
    for md in reference.sample_mds(step):
        point = reference.position(md)
        off_md, distance = closest_point(offset, point, step=step)
        rows.append({"reference_md": round(md, 3), "offset_md": round(off_md, 3),
                     "distance_m": round(distance, 2)})
    best = min(range(len(rows)), key=lambda i: rows[i]["distance_m"])
    mds = reference.sample_mds(step)
    lo, hi = mds[max(best - 1, 0)], mds[min(best + 1, len(mds) - 1)]
    ref_md, distance = _golden_min(lambda m: closest_point(offset, reference.position(m), step=step)[1], lo, hi)
    off_md, distance = closest_point(offset, reference.position(ref_md), step=step)
    if distance > rows[best]["distance_m"] + 0.005:       # refinement must never report worse than a sample
        ref_md, off_md, distance = mds[best], rows[best]["offset_md"], rows[best]["distance_m"]
    closest = {"reference_md": round(ref_md, 3), "offset_md": round(off_md, 3), "distance_m": round(distance, 2)}
    return Comparison(status=ASSESSED, rows=rows, closest=closest, kind="well_to_well")


def deviation_from_plan(actual: WellPath, plan: WellPath, *, step: float = DEFAULT_STEP_M) -> Comparison:
    """Miss of every actual station from the plan, in the plan's normal plane."""
    rows, warnings = [], []
    plan_first, plan_last = plan.md_range
    for station in actual.stations:
        point = (station.north, station.east, station.down)
        plan_md, distance = closest_point(plan, point, step=step)
        position, tangent = plan.state(plan_md)
        inc, azi = plan.angles(plan_md)
        high = _high_side(inc, azi)
        right = _cross(tangent, high)
        miss = _sub(point, position)
        row = {"actual_md": round(station.md, 3), "plan_md": round(plan_md, 3), "distance_m": round(distance, 2),
               "high_m": round(_dot(miss, high), 2), "right_m": round(_dot(miss, right), 2),
               "along_m": round(_dot(miss, tangent), 2)}
        if inc < NEAR_VERTICAL_DEG:
            row["note"] = "plan near vertical: high side follows the plan azimuth"
        if plan_md <= plan_first + 1e-6 or plan_md >= plan_last - 1e-6:
            row["note"] = "closest plan point is a plan end: actual station outside the plan MD range"
        rows.append(row)
    if any("note" in r for r in rows):
        warnings.append("Some rows carry notes (near-vertical plan or outside the plan MD range)")
    worst = max(rows, key=lambda r: r["distance_m"]) if rows else None
    return Comparison(status=ASSESSED, rows=rows, closest=worst, kind="plan_vs_actual", warnings=warnings)


def _missing(value) -> bool:
    try:
        return value is None or not math.isfinite(float(value))
    except (TypeError, ValueError):  # silent-ok: non-numeric is missing - reported as a NOT ASSESSED reason
        return True


def compare_wells(reference_traj, reference_well: dict, offset_traj, offset_well: dict, *,
                  same_crs_confirmed: bool, step: float = DEFAULT_STEP_M) -> Comparison:
    """Well-to-well screening in one frame, or NOT ASSESSED with every reason listed.

    ``*_well`` mappings carry ``name``, ``northing``, ``easting`` and ``rte_msl`` (metres).
    """
    reasons = []
    if not same_crs_confirmed:
        reasons.append("Wellhead coordinates are not confirmed to share one projected CRS in metres")
    for label, well in (("reference", reference_well), ("offset", offset_well)):
        for key in ("northing", "easting", "rte_msl"):
            if _missing(well.get(key)):
                reasons.append(f"{label} well: {key} is missing")
    ref_path = path_from_trajectory(reference_traj, name=str(reference_well.get("name") or "reference"))
    off_path = path_from_trajectory(offset_traj, name=str(offset_well.get("name") or "offset"))
    if ref_path is None:
        reasons.append("reference well: fewer than two usable survey stations")
    if off_path is None:
        reasons.append("offset well: fewer than two usable survey stations")
    reasons.extend(_degenerate_reasons(ref_path, "reference well") + _degenerate_reasons(off_path, "offset well"))
    if reasons:
        return Comparison(status=NOT_ASSESSED, reasons=reasons, kind="well_to_well")
    n0, e0 = float(reference_well["northing"]), float(reference_well["easting"])
    ref_path = path_from_trajectory(reference_traj, name=ref_path.name,
                                    down0=-float(reference_well["rte_msl"]))
    off_path = path_from_trajectory(offset_traj, name=off_path.name,
                                    north0=float(offset_well["northing"]) - n0,
                                    east0=float(offset_well["easting"]) - e0,
                                    down0=-float(offset_well["rte_msl"]))
    result = separation_profile(ref_path, off_path, step=step)
    result.warnings.extend(_trajectory_warnings(reference_traj, "reference") +
                           _trajectory_warnings(offset_traj, "offset"))
    return result


def compare_plan(actual_traj, plan_traj, *, step: float = DEFAULT_STEP_M) -> Comparison:
    """Plan vs actual of ONE well (same wellhead and rotary table)."""
    reasons = []
    actual = path_from_trajectory(actual_traj, name="actual")
    plan = path_from_trajectory(plan_traj, name="plan")
    if actual is None:
        reasons.append("actual: fewer than two usable survey stations")
    if plan is None:
        reasons.append("plan: fewer than two usable plan stations")
    reasons.extend(_degenerate_reasons(actual, "actual") + _degenerate_reasons(plan, "plan"))
    if reasons:
        return Comparison(status=NOT_ASSESSED, reasons=reasons, kind="plan_vs_actual")
    result = deviation_from_plan(actual, plan, step=step)
    result.warnings.extend(_trajectory_warnings(actual_traj, "actual") + _trajectory_warnings(plan_traj, "plan"))
    return result


def _degenerate_reasons(path, label: str) -> list:
    if path is None:
        return []
    return [f"{label}: ~180 deg dogleg between MD {a:g} and {b:g} m - arc undefined, add a station"
            for a, b in degenerate_segments(path)]


def _trajectory_warnings(trajectory, label: str) -> list:
    return [f"{label}: {issue.message}" for issue in trajectory.issues if issue.severity in ("warning", "error")]


# ---- plan import ----------------------------------------------------------------------------------
_ALIASES = {"md": ("md", "measured depth", "measureddepth", "depth"),
            "inc": ("inc", "inclination", "incl"),
            "azi": ("azi", "azimuth", "az", "azim")}


class PlanImportError(ValueError):
    """The plan table cannot be read as MD / inclination / azimuth."""


def _column_map(header: Sequence[str]) -> dict:
    found = {}
    for index, raw in enumerate(header):
        key = str(raw or "").strip().lower()
        key = key.split("(")[0].split("[")[0].strip()
        for target, names in _ALIASES.items():
            if key in names and target not in found:
                found[target] = index
    missing = [k for k in ("md", "inc", "azi") if k not in found]
    if missing:
        raise PlanImportError("Missing plan column(s): " + ", ".join(missing))
    return found


def plan_rows_from_table(rows: Iterable[Sequence], *, length_unit: str = "m") -> list:
    """Rows -> [{md, inc, azi}] in metres/degrees; the first row is the header.

    ``length_unit`` is the MD unit of the file ("m" or "ft") - stated, never guessed.
    """
    if length_unit not in ("m", "ft"):
        raise PlanImportError(f"Unsupported MD unit: {length_unit}")
    factor = FT_TO_M if length_unit == "ft" else 1.0
    iterator = iter(rows)
    try:
        header = next(iterator)
    except StopIteration:
        raise PlanImportError("The plan file is empty") from None
    columns = _column_map(header)
    out = []
    for number, row in enumerate(iterator, start=2):
        if row is None or all(str(c).strip() == "" for c in row if c is not None):
            continue
        try:
            values = {k: float(str(row[i]).strip()) for k, i in columns.items()}
        except (IndexError, TypeError, ValueError):
            raise PlanImportError(f"Row {number}: MD, inclination and azimuth must be numbers") from None
        if not all(math.isfinite(v) for v in values.values()):
            raise PlanImportError(f"Row {number}: non-finite value")
        values["md"] *= factor
        out.append(values)
    if len(out) < 2:
        raise PlanImportError("A plan needs at least two stations")
    return out


def read_plan_file(path: str, *, length_unit: str = "m") -> list:
    """CSV or XLSX (first sheet) -> plan rows; no network, no external service."""
    lower = str(path).lower()
    if lower.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = [list(r) for r in workbook.worksheets[0].iter_rows(values_only=True)]
        finally:
            workbook.close()
        return plan_rows_from_table(rows, length_unit=length_unit)
    if lower.endswith((".csv", ".txt")):
        with open(path, "r", encoding="utf-8-sig", newline="") as handle:
            text = handle.read()
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t") if text.strip() else csv.excel
        return plan_rows_from_table(list(csv.reader(io.StringIO(text), dialect)), length_unit=length_unit)
    raise PlanImportError("Plan files must be .csv, .txt or .xlsx")


def compose_plan(rows: list):
    """Plan stations through the same canonical engine as the actual surveys."""
    return compose_well_trajectory(rows)


def calculation_record(comparison: Comparison, *, inputs: dict) -> dict:
    """Fields for a NEW ``TrajectoryCalculation`` row (results are never overwritten)."""
    return {"method": METHOD_VERSION[:50], "description": LABEL,
            "parameters_json": {"contract": CONTRACT, "inputs": inputs},
            "results_json": comparison.to_json()}


# ---- database (small reads; one insert) ------------------------------------------------------------
def list_wells(db_manager) -> list:
    """[(id, name)] of every well, by name."""
    from core.database import Well
    with db_manager.session_scope() as session:
        return [(int(i), str(n or f"Well {i}")) for i, n in session.query(Well.id, Well.name).order_by(Well.name)]


def well_reference(db_manager, well_id: int) -> dict:
    """Wellhead and datum fields of one well as plain values (no ORM object leaves the session)."""
    from core.database import Well
    with db_manager.session_scope() as session:
        row = session.query(Well.name, Well.northing, Well.easting, Well.rte_msl).filter(Well.id == well_id).first()
    if row is None:
        return {"name": None, "northing": None, "easting": None, "rte_msl": None}
    return {"name": row[0], "northing": row[1], "easting": row[2], "rte_msl": row[3]}


def save_comparison(db_manager, well_id: int, comparison: Comparison, *, inputs: dict,
                    user_id: Optional[int] = None) -> int:
    """Insert a NEW TrajectoryCalculation; earlier results are never updated or replaced.

    Writes go through the session, so the data-layer authorization (M43) applies.
    """
    from datetime import date

    from core.database import TrajectoryCalculation
    record = calculation_record(comparison, inputs=inputs)
    with db_manager.session_scope() as session:
        calc = TrajectoryCalculation(
            well_id=well_id, method=record["method"], calculation_date=date.today(),
            parameters_json=record["parameters_json"], results_json=record["results_json"],
            description=record["description"], calculated_by=user_id)
        session.add(calc)
        session.flush()
        return int(calc.id)
