"""Lossless common intermediate representation for document imports.

The IR is deliberately source-oriented.  Excel, MinerU, and any future
extractor may use different mechanics, but the downstream mapper receives the
same primitives and the same provenance contract.  This module does not guess
canonical fields and does not coerce values.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any, Optional


IR_VERSION = 2


@dataclass(frozen=True)
class SourceLocation:
    """A source position that is safe to carry through mapping and review."""

    file: str = ""
    sheet: Optional[str] = None
    page: Optional[int] = None
    row: Optional[int] = None
    column: Optional[int | str] = None
    cell: Optional[str] = None
    coordinates: Optional[tuple[float, ...]] = None
    table: Optional[str] = None

    def to_dict(self) -> dict[str, Any]:
        return {
            "file": self.file,
            "sheet": self.sheet,
            "page": self.page,
            "table": self.table,
            "row": self.row,
            "column": self.column,
            "cell": self.cell,
            "coordinates": list(self.coordinates) if self.coordinates else None,
        }

    @classmethod
    def from_dict(cls, payload: Optional[dict[str, Any]]) -> "SourceLocation":
        payload = payload or {}
        coordinates = payload.get("coordinates")
        return cls(
            file=str(payload.get("file", "") or ""),
            sheet=payload.get("sheet"),
            page=payload.get("page"),
            table=payload.get("table"),
            row=payload.get("row"),
            column=payload.get("column"),
            cell=payload.get("cell"),
            coordinates=tuple(coordinates) if coordinates else None,
        )


@dataclass
class RawCell:
    """One source cell/token with both source and downstream state.

    ``value`` remains the source value for backwards compatibility.  The
    explicit original/normalized fields make it impossible for a normalizer
    or review exporter to silently overwrite the source token.
    """

    value: Any
    location: SourceLocation
    # Keep the original positional contract: formula/hidden/merged/style were
    # the first fields after location in IR version 1.
    formula: Optional[str] = None
    hidden: bool = False
    merged: bool = False
    style: Optional[str] = None
    original_value: Any = None
    normalized_value: Any = None
    original_unit: Optional[str] = None
    normalized_unit: Optional[str] = None
    table: Optional[str] = None
    section_title: Optional[str] = None
    extraction_method: str = "unknown"
    confidence: Optional[float] = None
    validation_state: str = "unvalidated"
    review_state: str = "unreviewed"
    merge_anchor: Optional[str] = None

    def __post_init__(self) -> None:
        # Source values are never invented.  Mirroring the source into the
        # explicit field is metadata, not a transformation.
        if self.original_value is None and self.value is not None:
            self.original_value = self.value

    def to_dict(self) -> dict[str, Any]:
        return {
            "value": self.value,
            "original_value": self.original_value,
            "normalized_value": self.normalized_value,
            "original_unit": self.original_unit,
            "normalized_unit": self.normalized_unit,
            "table": self.table,
            "section_title": self.section_title,
            "extraction_method": self.extraction_method,
            "confidence": self.confidence,
            "validation_state": self.validation_state,
            "review_state": self.review_state,
            "formula": self.formula,
            "hidden": self.hidden,
            "merged": self.merged,
            "merge_anchor": self.merge_anchor,
            "style": self.style,
            "location": self.location.to_dict(),
        }

    @classmethod
    def from_dict(cls, payload: dict[str, Any]) -> "RawCell":
        return cls(
            value=payload.get("value"),
            original_value=payload.get("original_value"),
            normalized_value=payload.get("normalized_value"),
            original_unit=payload.get("original_unit"),
            normalized_unit=payload.get("normalized_unit"),
            table=payload.get("table"),
            section_title=payload.get("section_title"),
            extraction_method=str(payload.get("extraction_method", "unknown") or "unknown"),
            confidence=payload.get("confidence"),
            validation_state=str(payload.get("validation_state", "unvalidated") or "unvalidated"),
            review_state=str(payload.get("review_state", "unreviewed") or "unreviewed"),
            formula=payload.get("formula"),
            hidden=bool(payload.get("hidden", False)),
            merged=bool(payload.get("merged", False)),
            merge_anchor=payload.get("merge_anchor"),
            style=payload.get("style"),
            location=SourceLocation.from_dict(payload.get("location")),
        )


@dataclass
class RawTable:
    headers: list[RawCell] = field(default_factory=list)
    rows: list[list[RawCell]] = field(default_factory=list)
    name: str = ""
    location: Optional[SourceLocation] = None
    classification: str = "unknown"
    section_title: Optional[str] = None
    units: list[Optional[str]] = field(default_factory=list)
    extraction_method: str = "unknown"
    confidence: Optional[float] = None
    validation_state: str = "unvalidated"
    review_state: str = "unreviewed"

    def to_dict(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "location": self.location.to_dict() if self.location else None,
            "classification": self.classification,
            "section_title": self.section_title,
            "units": list(self.units),
            "extraction_method": self.extraction_method,
            "confidence": self.confidence,
            "validation_state": self.validation_state,
            "review_state": self.review_state,
            "headers": [cell.to_dict() for cell in self.headers],
            "rows": [[cell.to_dict() for cell in row] for row in self.rows],
        }

    @classmethod
    def from_dict(cls, payload: dict[str, Any]) -> "RawTable":
        return cls(
            headers=[RawCell.from_dict(cell) for cell in payload.get("headers", [])],
            rows=[
                [RawCell.from_dict(cell) for cell in row]
                for row in payload.get("rows", [])
            ],
            name=str(payload.get("name", "") or ""),
            location=SourceLocation.from_dict(payload["location"]) if payload.get("location") else None,
            classification=str(payload.get("classification", "unknown") or "unknown"),
            section_title=payload.get("section_title"),
            units=list(payload.get("units", []) or []),
            extraction_method=str(payload.get("extraction_method", "unknown") or "unknown"),
            confidence=payload.get("confidence"),
            validation_state=str(payload.get("validation_state", "unvalidated") or "unvalidated"),
            review_state=str(payload.get("review_state", "unreviewed") or "unreviewed"),
        )


@dataclass
class RawDocument:
    source_file: str
    tables: list[RawTable] = field(default_factory=list)
    cells: list[RawCell] = field(default_factory=list)
    text_blocks: list[tuple[str, SourceLocation]] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=lambda: {"ir_version": IR_VERSION})
    section_titles: list[tuple[str, SourceLocation]] = field(default_factory=list)

    def __post_init__(self) -> None:
        self.metadata.setdefault("ir_version", IR_VERSION)

    @property
    def page_count(self) -> int:
        pages = set(self.metadata.get("page_numbers", []) or [])
        pages.update(cell.location.page for cell in self.cells if cell.location.page is not None)
        pages.update(location.page for _text, location in self.text_blocks if location.page is not None)
        return len(pages)

    @property
    def sheet_count(self) -> int:
        sheets = {cell.location.sheet for cell in self.cells if cell.location.sheet}
        sheets.update(location.sheet for _text, location in self.text_blocks if location.sheet)
        return len(sheets)

    def to_dict(self, *, include_cells: bool = False) -> dict[str, Any]:
        """Serialize a summary by default, or the complete lossless IR.

        Existing callers receive the small summary payload.  Acceptance tests,
        audit exports, and persistence diagnostics use ``include_cells=True``
        to retain every table/header/cell/text provenance field.
        """
        payload: dict[str, Any] = {
            "ir_version": self.metadata.get("ir_version", IR_VERSION),
            "source_file": self.source_file,
            "tables": len(self.tables),
            "cells": len(self.cells),
            "text_blocks": len(self.text_blocks),
            "section_titles": len(self.section_titles),
            "metadata": dict(self.metadata),
        }
        if include_cells:
            payload["tables_detail"] = [table.to_dict() for table in self.tables]
            payload["cell_values"] = [cell.to_dict() for cell in self.cells]
            payload["text_block_values"] = [
                {"text": text, "location": location.to_dict()}
                for text, location in self.text_blocks
            ]
            payload["section_title_values"] = [
                {"text": text, "location": location.to_dict()}
                for text, location in self.section_titles
            ]
        return payload

    @classmethod
    def from_dict(cls, payload: dict[str, Any]) -> "RawDocument":
        """Deserialize both summary and complete IR payloads safely."""
        tables = [RawTable.from_dict(item) for item in payload.get("tables_detail", [])]
        cells = [RawCell.from_dict(item) for item in payload.get("cell_values", [])]
        text_blocks = [
            (str(item.get("text", "")), SourceLocation.from_dict(item.get("location")))
            for item in payload.get("text_block_values", [])
        ]
        section_titles = [
            (str(item.get("text", "")), SourceLocation.from_dict(item.get("location")))
            for item in payload.get("section_title_values", [])
        ]
        metadata = dict(payload.get("metadata", {}) or {})
        metadata.setdefault("ir_version", payload.get("ir_version", IR_VERSION))
        return cls(
            source_file=str(payload.get("source_file", "") or ""),
            tables=tables,
            cells=cells,
            text_blocks=text_blocks,
            section_titles=section_titles,
            metadata=metadata,
        )


def _mineru_location(raw_source: str, provenance: Any, *, table: Optional[str] = None,
                     row: Optional[int] = None, column: Optional[int | str] = None,
                     cell: Optional[str] = None) -> SourceLocation:
    return SourceLocation(
        file=raw_source,
        sheet=getattr(provenance, "source_sheet", None),
        page=getattr(provenance, "source_page", None),
        table=table,
        row=row if row is not None else getattr(provenance, "source_row", None),
        column=column if column is not None else getattr(provenance, "source_column", None),
        cell=cell,
        coordinates=getattr(provenance, "bounding_box", None),
    )


def raw_document_from_mineru(document: Any) -> RawDocument:
    """Adapt a MinerU document without importing MinerU itself."""
    raw = RawDocument(
        source_file=str(getattr(document, "source_file", "")),
        metadata={
            "ir_version": IR_VERSION,
            "backend": getattr(document, "backend", ""),
            "method": getattr(document, "method", ""),
            "engine": "MinerU",
            "page_numbers": [
                getattr(page, "number", None)
                for page in (getattr(document, "pages", []) or [])
                if getattr(page, "number", None) is not None
            ],
            "output_dir": getattr(document, "output_dir", None),
            "raw_files": list(getattr(document, "raw_files", []) or []),
        },
    )

    for heading in getattr(document, "headings", []) or []:
        provenance = getattr(heading, "provenance", None)
        location = _mineru_location(raw.source_file, provenance)
        raw.section_titles.append((str(getattr(heading, "text", "")), location))

    for block in getattr(document, "text_blocks", []) or []:
        provenance = getattr(block, "provenance", None)
        location = _mineru_location(raw.source_file, provenance)
        raw.text_blocks.append((str(getattr(block, "text", "")), location))

    for table_index, table in enumerate(getattr(document, "tables", []) or [], 1):
        provenance = getattr(table, "provenance", None)
        table_name = str(getattr(table, "name", "") or f"table_{table_index}")
        method = str(getattr(provenance, "extraction_method", "mineru") or "mineru")
        confidence = getattr(provenance, "confidence", None)
        base = _mineru_location(raw.source_file, provenance, table=table_name)
        section_title = table_name or None

        headers: list[RawCell] = []
        for index, header in enumerate(getattr(table, "headers", []) or [], 1):
            location = _mineru_location(
                raw.source_file, provenance, table=table_name, column=index,
                cell=f"page {base.page}:column {index}" if base.page is not None else f"column {index}",
            )
            headers.append(RawCell(
                value=header,
                location=location,
                original_value=header,
                table=table_name,
                section_title=section_title,
                extraction_method=method,
                confidence=confidence,
            ))

        rows: list[list[RawCell]] = []
        for row_number, row in enumerate(getattr(table, "rows", []) or [], 1):
            row_cells: list[RawCell] = []
            for index, value in enumerate(row, 1):
                location = _mineru_location(
                    raw.source_file, provenance, table=table_name,
                    row=row_number, column=index,
                    cell=(
                        f"page {base.page}:row {row_number}:column {index}"
                        if base.page is not None
                        else f"row {row_number}:column {index}"
                    ),
                )
                row_cells.append(RawCell(
                    value=value,
                    location=location,
                    original_value=value,
                    table=table_name,
                    section_title=section_title,
                    extraction_method=method,
                    confidence=confidence,
                ))
            rows.append(row_cells)
            raw.cells.extend(row_cells)

        raw.tables.append(RawTable(
            headers=headers,
            rows=rows,
            name=table_name,
            location=base,
            classification="table",
            section_title=section_title,
            units=[None] * len(headers),
            extraction_method=method,
            confidence=confidence,
        ))
        raw.cells.extend(headers)

    return raw


def raw_document_from_workbook(workbook: Any, *, max_cells: int = 250_000) -> RawDocument:
    """Adapt populated Excel cells, formulas, hidden state, and table layout.

    The adapter is the only workbook walk used by ``ExcelIntelligence``.  The
    mapper therefore consumes this IR rather than rereading openpyxl cells.
    """
    source_file = str(getattr(workbook, "filename", ""))
    raw = RawDocument(
        source_file=source_file,
        metadata={"ir_version": IR_VERSION, "engine": "Excel"},
    )
    for worksheet in getattr(workbook, "worksheets", []) or []:
        hidden_rows = getattr(worksheet, "row_dimensions", {})
        hidden_columns = getattr(worksheet, "column_dimensions", {})
        merged_ranges = getattr(worksheet, "merged_cells", None)
        merged_cells: dict[tuple[int, int], str] = {}
        if merged_ranges is not None:
            for merged_range in merged_ranges.ranges:
                anchor = worksheet.cell(merged_range.min_row, merged_range.min_col).coordinate
                for row in range(merged_range.min_row, merged_range.max_row + 1):
                    for col in range(merged_range.min_col, merged_range.max_col + 1):
                        merged_cells[(row, col)] = anchor

        sheet_cells: list[RawCell] = []
        for row in worksheet.iter_rows():
            for cell in row:
                anchor = merged_cells.get((cell.row, cell.column))
                # Keep blank merged cells because the merge anchor is required
                # to reproduce label/value lookup without touching openpyxl.
                if cell.value is None and anchor is None:
                    continue
                coordinate = getattr(cell, "coordinate", None)
                location = SourceLocation(
                    file=source_file,
                    sheet=worksheet.title,
                    cell=coordinate,
                    row=cell.row,
                    column=cell.column,
                )
                row_dimension = hidden_rows.get(cell.row)
                column_dimension = hidden_columns.get(getattr(cell, "column_letter", ""))
                raw_cell = RawCell(
                    value=cell.value,
                    location=location,
                    original_value=cell.value,
                    table=worksheet.title,
                    extraction_method="excel-cell",
                    confidence=None,
                    formula=cell.value if isinstance(cell.value, str) and cell.value.startswith("=") else None,
                    hidden=bool(
                        getattr(row_dimension, "hidden", False)
                        or getattr(column_dimension, "hidden", False)
                    ),
                    merged=anchor is not None,
                    merge_anchor=anchor,
                    style=getattr(cell, "style_id", None),
                )
                sheet_cells.append(raw_cell)
                if len(raw.cells) + len(sheet_cells) >= max_cells:
                    raw.metadata["truncated"] = True
                    break
            if len(raw.cells) + len(sheet_cells) >= max_cells:
                break
            # Do not rely on raw.cells while building the per-sheet list when
            # a worksheet contains only blank merged cells.
        raw.cells.extend(sheet_cells[: max(0, max_cells - len(raw.cells))])
        if len(raw.cells) >= max_cells:
            break

        # A worksheet is a source region even when it has no formal Excel
        # table.  Headers remain empty unless the file declares a real table;
        # no first-row guessing is performed.
        sheet_location = SourceLocation(file=source_file, sheet=worksheet.title, table=worksheet.title)
        raw.tables.append(RawTable(
            name=worksheet.title,
            location=sheet_location,
            classification="worksheet",
            section_title=None,
            extraction_method="excel-cell",
        ))

        # Preserve formal Excel table headers and row membership where present.
        for defined_table in getattr(getattr(worksheet, "tables", None), "values", lambda: [])():
            try:
                from openpyxl.utils.cell import range_boundaries
                min_col, min_row, max_col, max_row = range_boundaries(defined_table.ref)
            except Exception:
                continue
            by_coord = {(cell.location.row, cell.location.column): cell for cell in sheet_cells}
            headers = [
                by_coord[(min_row, column)]
                for column in range(min_col, max_col + 1)
                if (min_row, column) in by_coord
            ]
            rows = [
                [
                    by_coord[(row_number, column)]
                    for column in range(min_col, max_col + 1)
                    if (row_number, column) in by_coord
                ]
                for row_number in range(min_row + 1, max_row + 1)
            ]
            raw.tables.append(RawTable(
                headers=headers,
                rows=rows,
                name=str(getattr(defined_table, "displayName", "") or defined_table.name),
                location=SourceLocation(
                    file=source_file, sheet=worksheet.title,
                    table=str(getattr(defined_table, "displayName", "") or defined_table.name),
                    row=min_row, column=min_col,
                ),
                classification="table",
                section_title=None,
                units=[None] * len(headers),
                extraction_method="excel-table",
            ))

    raw.metadata["sheet_names"] = [worksheet.title for worksheet in getattr(workbook, "worksheets", []) or []]
    raw.metadata["sheet_count"] = len(raw.metadata["sheet_names"])
    return raw


__all__ = [
    "IR_VERSION", "SourceLocation", "RawCell", "RawTable", "RawDocument",
    "raw_document_from_mineru", "raw_document_from_workbook",
]
